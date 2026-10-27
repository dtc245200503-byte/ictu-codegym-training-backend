from datetime import date, time, timedelta
from io import BytesIO

from openpyxl import load_workbook
from app.core.security import hash_password
from sqlalchemy import event
from app.models.attendance import AttendanceRecord, AttendanceSession
from app.models.course import Course
from app.models.course_lesson import CourseLesson
from app.models.role import Role
from app.models.study_program import StudyProgram
from app.models.training_class import ClassEnrollment, TrainingClass
from app.models.user import User


def make_user(db, role_code, suffix):
    role = db.query(Role).filter_by(code=role_code).one_or_none()
    if role is None:
        role = Role(code=role_code, name=role_code.title())
        db.add(role)
        db.flush()
    user = User(
        username=f"{role_code}.{suffix}",
        email=f"{role_code}.{suffix}@example.edu.vn",
        full_name=f"{role_code.title()} {suffix}",
        hashed_password=hash_password("CorrectPassword123!"),
        roles=[role],
    )
    db.add(user)
    db.commit()
    db.refresh(user)
    return user


def login_headers(client, user):
    response = client.post(
        "/api/auth/login",
        json={"email": user.email, "password": "CorrectPassword123!"},
    )
    assert response.status_code == 200
    return {"Authorization": "Bearer " + response.json()["access_token"]}


def make_class(db, teacher, code):
    course = Course(
        code=f"{code}-COURSE",
        name=f"{code} Course",
        department="IT",
        credits=3,
        session_count=3,
    )
    db.add(course)
    db.flush()
    lesson = CourseLesson(
        course_id=course.id,
        sequence=1,
        title="Lesson 1",
        objectives="Learn the topic",
    )
    db.add(lesson)
    training_class = TrainingClass(
        code=code,
        name=code,
        course_id=course.id,
        teacher_id=teacher.id,
        room="SUMMARY-ROOM",
        term="2026-Fall",
        day_of_week=date.today().weekday(),
        start_time=time(9, 0),
        end_time=time(11, 0),
        capacity=20,
        status="active",
    )
    db.add(training_class)
    db.commit()
    db.refresh(training_class)
    return training_class


def make_session(db, training_class, session_date):
    lesson = (
        db.query(CourseLesson)
        .filter_by(course_id=training_class.course_id)
        .order_by(CourseLesson.sequence)
        .first()
    )
    session = AttendanceSession(
        class_id=training_class.id,
        session_date=session_date,
        start_time=time(9, 0),
        end_time=time(11, 0),
        teacher_id=training_class.teacher_id,
        room=training_class.room,
        lesson_id=lesson.id,
    )
    db.add(session)
    db.commit()
    db.refresh(session)
    return session


def test_class_attendance_summary_returns_matrix_and_aggregates(client, test_db):
    teacher = make_user(test_db, "teacher", "attendance_summary_api")
    manager = make_user(test_db, "manager", "attendance_summary_api")
    student = make_user(test_db, "student", "attendance_summary_api")
    training_class = make_class(test_db, teacher, code="SUMMARY-API")
    program = StudyProgram(code="SUMMARY-PROGRAM", name="Summary Program", total_credits=120)
    test_db.add(program)
    test_db.flush()
    training_class.program_id = program.id
    enrollment = ClassEnrollment(
        class_id=training_class.id,
        student_id=student.id,
        enrollment_status="active",
    )
    test_db.add(enrollment)

    statuses = [
        "present",
        "late",
        "excused_absent",
        "unexcused_absent",
        "absent",
    ]
    sessions = [
        make_session(
            test_db,
            training_class,
            session_date=date.today() - timedelta(days=6 - index),
        )
        for index in range(7)
    ]
    sessions[6].is_cancelled = True
    for session, attendance_status in zip(sessions[:5], statuses):
        test_db.add(
            AttendanceRecord(
                session_id=session.id,
                student_id=student.id,
                marked_by=teacher.id,
                status=attendance_status,
            )
        )
    # The legacy enrollment-status endpoint persists this sentinel in a
    # per-session record when attendance is cleared.
    test_db.add(
        AttendanceRecord(
            session_id=sessions[5].id,
            student_id=student.id,
            marked_by=teacher.id,
            status="not_recorded",
        )
    )
    test_db.commit()

    response = client.get(
        f"/api/classes/{training_class.id}/attendance/summary",
        headers=login_headers(client, teacher),
    )

    assert response.status_code == 200
    data = response.json()
    assert data["class"] == {
        "id": training_class.id,
        "code": training_class.code,
        "name": training_class.name,
    }
    assert [session["id"] for session in data["sessions"]] == [
        session.id for session in sessions[:6]
    ]
    row = data["students"][0]
    cells = row["attendance_by_session"]
    assert cells[str(sessions[0].id)] == {"eligible": True, "status": "present"}
    assert cells[str(sessions[1].id)] == {"eligible": True, "status": "late"}
    assert cells[str(sessions[2].id)]["status"] == "excused_absent"
    assert cells[str(sessions[3].id)]["status"] == "unexcused_absent"
    assert cells[str(sessions[4].id)]["status"] == "absent"
    assert cells[str(sessions[5].id)] == {"eligible": True, "status": None}
    assert str(sessions[6].id) not in cells
    assert row["total_sessions"] == 6
    assert row["attended_count"] == 2
    assert row["present_count"] == 1
    assert row["late_count"] == 1
    assert row["excused_absent_count"] == 1
    assert row["unexcused_absent_count"] == 1
    assert row["legacy_absent_count"] == 1
    assert row["unrecorded_count"] == 1
    assert row["absence_threshold_percent"] == 20.0
    assert row["absence_rate_percent"] == 50.0
    assert row["at_risk"] is True

    manager_headers = login_headers(client, manager)
    teacher_headers = login_headers(client, teacher)
    denied_policy = client.put(f"/api/study-programs/{program.id}/attendance-policy",
        headers=teacher_headers, json={"absence_threshold_percent": 40})
    assert denied_policy.status_code == 403
    invalid_policy = client.put(f"/api/study-programs/{program.id}/attendance-policy",
        headers=manager_headers, json={"absence_threshold_percent": 100.01})
    assert invalid_policy.status_code == 422
    configured_policy = client.put(f"/api/study-programs/{program.id}/attendance-policy",
        headers=manager_headers, json={"absence_threshold_percent": 50})
    assert configured_policy.status_code == 200
    changed_summary = client.get(f"/api/classes/{training_class.id}/attendance/summary", headers=teacher_headers)
    assert changed_summary.json()["students"][0]["at_risk"] is True
    assert changed_summary.json()["students"][0]["warning_status"] == "at_threshold"
    assert changed_summary.json()["students"][0]["allowed_absences"] == 3
    assert row["absent_count"] == 3
    assert row["attendance_rate"] == 33.3


def test_class_attendance_summary_returns_null_rate_without_eligible_sessions(
    client, test_db
):
    teacher = make_user(test_db, "teacher", "attendance_summary_empty")
    student = make_user(test_db, "student", "attendance_summary_empty")
    training_class = make_class(test_db, teacher, code="SUMMARY-EMPTY")
    test_db.add(
        ClassEnrollment(
            class_id=training_class.id,
            student_id=student.id,
            enrollment_status="active",
        )
    )
    future_session = make_session(
        test_db,
        training_class,
        session_date=date.today() + timedelta(days=1),
    )

    response = client.get(
        f"/api/classes/{training_class.id}/attendance/summary",
        headers=login_headers(client, teacher),
    )

    assert response.status_code == 200
    data = response.json()
    assert data["sessions"] == []
    assert len(data["students"]) == 1
    assert str(future_session.id) not in data["students"][0]["attendance_by_session"]
    assert data["students"][0]["attendance_rate"] is None


def test_class_attendance_export_preserves_matrix_and_totals_and_is_manager_only(
    client, test_db
):
    manager = make_user(test_db, "manager", "attendance_export")
    teacher = make_user(test_db, "teacher", "attendance_export")
    student = make_user(test_db, "student", "attendance_export")
    training_class = make_class(test_db, teacher, code="SUMMARY-EXPORT")
    test_db.add(
        ClassEnrollment(
            class_id=training_class.id,
            student_id=student.id,
            enrollment_status="active",
        )
    )
    test_db.commit()
    session = make_session(
        test_db,
        training_class,
        session_date=date.today() - timedelta(days=1),
    )
    test_db.add(
        AttendanceRecord(
            session_id=session.id,
            student_id=student.id,
            marked_by=teacher.id,
            status="excused_absent",
        )
    )
    test_db.commit()

    response = client.get(
        f"/api/classes/{training_class.id}/attendance/export",
        headers=login_headers(client, manager),
    )

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "SUMMARY-EXPORT-attendance-" in response.headers["content-disposition"]
    sheet = load_workbook(BytesIO(response.content), read_only=True).active
    rows = list(sheet.iter_rows(values_only=True))
    headers = rows[3]
    student_row = rows[4]
    assert headers[:3] == ("Mã học viên", "Họ tên", f"{session.session_date} 09:00–11:00")
    assert "Vắng có phép" in headers
    assert "Tỷ lệ chuyên cần (%)" in headers
    assert student_row[0] == student.student_code or student_row[0] == student.username
    assert student_row[1] == student.full_name
    assert student_row[2] == "Vắng có phép"
    assert student_row[3] == 0  # Có mặt
    assert student_row[5] == 1  # Vắng có phép
    assert student_row[-2] == 1  # Tổng vắng
    assert student_row[-1] == 0.0

    denied = client.get(
        f"/api/classes/{training_class.id}/attendance/export",
        headers=login_headers(client, teacher),
    )
    assert denied.status_code == 403


def test_class_attendance_summary_search_filter_and_server_pagination(client, test_db):
    teacher = make_user(test_db, "teacher", "summary_page")
    low = make_user(test_db, "student", "summary_page_low")
    high = make_user(test_db, "student", "summary_page_high")
    unrecorded = make_user(test_db, "student", "summary_page_unrecorded")
    low.full_name = "An Nguyen"
    high.full_name = "Binh Tran"
    unrecorded.full_name = "Chi Le"
    training_class = make_class(test_db, teacher, code="SUMMARY-PAGE")
    test_db.add_all(
        [
            ClassEnrollment(class_id=training_class.id, student_id=user.id)
            for user in (low, high, unrecorded)
        ]
    )
    test_db.commit()
    session = make_session(
        test_db,
        training_class,
        session_date=date.today() - timedelta(days=1),
    )
    test_db.add_all(
        [
            AttendanceRecord(
                session_id=session.id,
                student_id=low.id,
                marked_by=teacher.id,
                status="unexcused_absent",
            ),
            AttendanceRecord(
                session_id=session.id,
                student_id=high.id,
                marked_by=teacher.id,
                status="present",
            ),
        ]
    )
    test_db.commit()
    headers = login_headers(client, teacher)

    first_page = client.get(
        f"/api/classes/{training_class.id}/attendance/summary?page=1&page_size=1&sort=attendance_asc",
        headers=headers,
    )
    assert first_page.status_code == 200
    assert first_page.json()["students"][0]["id"] == low.id

    second_page = client.get(
        f"/api/classes/{training_class.id}/attendance/summary?page=2&page_size=1&sort=attendance_asc",
        headers=headers,
    )
    assert second_page.status_code == 200
    page_data = second_page.json()
    assert page_data["pagination"] == {
        "page": 2, "page_size": 1, "total": 3,
        "pages": 3, "has_previous": True, "has_next": True,
    }
    assert page_data["total_students"] == 3
    assert page_data["totals"]["absent_count"] == 1
    assert page_data["students"][0]["id"] == unrecorded.id

    descending = client.get(
        f"/api/classes/{training_class.id}/attendance/summary?page=1&page_size=1&sort=attendance_desc",
        headers=headers,
    )
    assert descending.status_code == 200
    assert descending.json()["students"][0]["id"] == high.id

    filtered = client.get(
        f"/api/classes/{training_class.id}/attendance/summary?q=An%20Nguyen&status=unexcused_absent&page=1&page_size=20",
        headers=headers,
    )
    assert filtered.status_code == 200
    filtered_data = filtered.json()
    assert filtered_data["pagination"]["total"] == 1
    assert [row["id"] for row in filtered_data["students"]] == [low.id]


def test_class_attendance_summary_large_roster_pages_after_global_rate_sort(client, test_db):
    teacher = make_user(test_db, "teacher", "summary_large")
    students = [make_user(test_db, "student", f"summary_large_{i:02d}") for i in range(25)]
    training_class = make_class(test_db, teacher, code="SUMMARY-LARGE")
    test_db.add_all([ClassEnrollment(class_id=training_class.id, student_id=s.id) for s in students])
    test_db.commit()
    session = make_session(test_db, training_class, session_date=date.today() - timedelta(days=2))
    test_db.add_all([
        AttendanceRecord(session_id=session.id, student_id=student.id,
            marked_by=teacher.id, status="present" if index >= 10 else "unexcused_absent")
        for index, student in enumerate(students)
    ])
    test_db.commit()
    headers = login_headers(client, teacher)

    page_one = client.get(f"/api/classes/{training_class.id}/attendance/summary?page=1&page_size=20",
        headers=headers)
    page_two = client.get(f"/api/classes/{training_class.id}/attendance/summary?page=2&page_size=20",
        headers=headers)
    assert page_one.status_code == page_two.status_code == 200
    first_ids = [row["id"] for row in page_one.json()["students"]]
    second_ids = [row["id"] for row in page_two.json()["students"]]
    assert len(first_ids) == 20 and len(second_ids) == 5
    assert not set(first_ids).intersection(second_ids)
    assert page_one.json()["pagination"]["total"] == 25
    assert page_one.json()["totals"]["absent_count"] == 10
    assert page_one.json()["totals"]["attended_count"] == 15
    assert all(row["attendance_rate"] == 0 for row in page_one.json()["students"][:10])
    assert all(row["attendance_rate"] == 100 for row in page_one.json()["students"][10:])


def test_class_attendance_summary_enforces_role_class_access_and_not_found(
    client, test_db
):
    admin = make_user(test_db, "admin", "attendance_summary_admin")
    manager = make_user(test_db, "manager", "attendance_summary_manager")
    teacher = make_user(test_db, "teacher", "attendance_summary_owner")
    other_teacher = make_user(test_db, "teacher", "attendance_summary_other")
    assigned_ta = make_user(test_db, "ta", "attendance_summary_assigned")
    other_ta = make_user(test_db, "ta", "attendance_summary_unassigned")
    student = make_user(test_db, "student", "attendance_summary_reader")
    training_class = make_class(test_db, teacher, code="SUMMARY-ACCESS")
    training_class.assistants.append(assigned_ta)
    test_db.commit()

    admin_allowed = client.get(
        f"/api/classes/{training_class.id}/attendance/summary",
        headers=login_headers(client, admin),
    )
    manager_allowed = client.get(
        f"/api/classes/{training_class.id}/attendance/summary",
        headers=login_headers(client, manager),
    )
    teacher_allowed = client.get(
        f"/api/classes/{training_class.id}/attendance/summary",
        headers=login_headers(client, teacher),
    )
    ta_allowed = client.get(
        f"/api/classes/{training_class.id}/attendance/summary",
        headers=login_headers(client, assigned_ta),
    )
    forbidden = client.get(
        f"/api/classes/{training_class.id}/attendance/summary",
        headers=login_headers(client, other_teacher),
    )
    ta_forbidden = client.get(
        f"/api/classes/{training_class.id}/attendance/summary",
        headers=login_headers(client, other_ta),
    )
    role_forbidden = client.get(
        f"/api/classes/{training_class.id}/attendance/summary",
        headers=login_headers(client, student),
    )
    missing = client.get(
        "/api/classes/999999/attendance/summary",
        headers=login_headers(client, teacher),
    )

    assert admin_allowed.status_code == 200
    assert manager_allowed.status_code == 200
    assert teacher_allowed.status_code == 200
    assert ta_allowed.status_code == 200
    assert forbidden.status_code == 403
    assert ta_forbidden.status_code == 403
    assert role_forbidden.status_code == 403
    assert missing.status_code == 404


def test_class_attendance_summary_select_count_does_not_scale_with_matrix_size(
    client, test_db
):
    teacher = make_user(test_db, "teacher", "attendance_summary_queries")
    small_class = make_class(test_db, teacher, code="SUMMARY-SMALL")
    large_class = make_class(test_db, teacher, code="SUMMARY-LARGE")

    for index, training_class in enumerate((small_class, large_class)):
        student_count = 1 if index == 0 else 5
        session_count = 1 if index == 0 else 5
        students = [
            make_user(test_db, "student", f"attendance_summary_q{index}_{n}")
            for n in range(student_count)
        ]
        test_db.add_all(
            [
                ClassEnrollment(
                    class_id=training_class.id,
                    student_id=student.id,
                    enrollment_status="active",
                )
                for student in students
            ]
        )
        for day_index in range(session_count):
            make_session(
                test_db,
                training_class,
                date.today() - timedelta(days=day_index + 1),
            )
        test_db.commit()

    headers = login_headers(client, teacher)
    select_statements = []

    def count_selects(_conn, _cursor, statement, _parameters, _context, _many):
        if statement.lstrip().upper().startswith("SELECT"):
            select_statements.append(statement)

    bind = test_db.get_bind()
    event.listen(bind, "before_cursor_execute", count_selects)
    try:
        small_response = client.get(
            f"/api/classes/{small_class.id}/attendance/summary", headers=headers
        )
        small_query_count = len(select_statements)
        select_statements.clear()
        large_response = client.get(
            f"/api/classes/{large_class.id}/attendance/summary", headers=headers
        )
        large_query_count = len(select_statements)
    finally:
        event.remove(bind, "before_cursor_execute", count_selects)

    assert small_response.status_code == 200
    assert large_response.status_code == 200
    assert large_query_count == small_query_count
