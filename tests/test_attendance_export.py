from datetime import date
from io import BytesIO
from urllib.parse import unquote

import pytest
from openpyxl import load_workbook

from app.attendance_export import MIME_TYPE
from app.crud import attendance_summary as crud
from app.models.attendance import SessionAttendance, StudentClass
from app.models.role_permission import RolePermission
from app.models.training_program import TrainingClass
from app.models.user import User
from test_attendance_summary import ctx


def export(ctx, **kwargs):
    return ctx[0]("/classes/1/export", actor=1, **kwargs)


def sheet(response):
    assert response.status_code == 200, response.text if response.status_code != 200 else ""
    assert response.headers["content-type"] == MIME_TYPE
    return load_workbook(BytesIO(response.content)).active


def test_excel_preserves_matrix_totals_numeric_percent_and_colors(ctx):
    matrix = ctx[0](actor=1).json()
    response = export(ctx)
    ws = sheet(response)
    assert ws.max_row == 5 + len(matrix["students"])
    assert ws.max_column == 3 + len(matrix["sessions"])
    assert [ws.cell(5, col).value for col in range(1, 4)] == ["Học viên", "Tổng vắng", "Chuyên cần"]
    assert [ws.cell(6, col).value for col in range(4, 11)] == ["CM", "M", "CP", "KP", "Huỷ", "?", "Sắp"]
    assert ws["B6"].value == 2 and ws["B6"].data_type == "n"
    assert ws["C6"].value == .5 and ws["C6"].number_format == "0.00%"
    assert ws["C7"].value == "—"
    assert "1 vắng có phép" in ws["B6"].comment.text
    assert "2/4" in ws["C6"].comment.text
    assert ws["D6"].fill.fgColor.rgb != ws["G6"].fill.fgColor.rgb
    for index, session in enumerate(matrix["sessions"], 4):
        assert session["topic"] in ws.cell(5, index).value
        assert session["subject_name"] in ws.cell(5, index).value
    assert ws.freeze_panes == "D6"
    assert ws.auto_filter.ref == "A5:J8"
    assert "Chú giải" in ws.parent.sheetnames
    assert response.headers["cache-control"] == "no-store"


def test_export_includes_students_beyond_first_page_and_selected_sort(ctx):
    with ctx[1]() as db:
        for id in range(100, 158):
            db.add(User(user_id=id, username=f"HV{id}", full_name=f"Z {id}", email=f"z{id}@example.com", password="unused", role_id=3))
        db.flush()
        db.add_all([StudentClass(student_id=id, class_id=1, enrolled_date=date(2026, 10, 1)) for id in range(100, 158)])
        db.add(SessionAttendance(session_id=1, student_id=6, student_name="N6", student_code="HV6", status="unexcused_absence"))
        db.commit()
    assert len(ctx[0](actor=1).json()["students"]) == 50
    ws = sheet(export(ctx, sort="rate_asc"))
    assert ws.max_row == 66  # 61 students, including the 11 beyond page one.
    assert "HV6" in ws["A6"].value
    assert "HV4" in ws["A7"].value
    assert "HV157" in ws.cell(ws.max_row, 1).value


def test_filename_uses_class_code_and_vietnam_report_date(ctx, monkeypatch):
    with ctx[1]() as db:
        db.get(TrainingClass, 1).class_code = "LỚP/60:01"
        db.commit()
    monkeypatch.setattr(crud, "today", lambda: date(2026, 10, 21))
    response = export(ctx)
    assert "Chuyen-can_LỚP_60_01_2026-10-21.xlsx" in unquote(response.headers["content-disposition"])
    assert "21/10/2026" in sheet(response)["A3"].value
    assert ctx[0](actor=1).json()["class_code"] == "LỚP/60:01"


def test_literal_formula_like_names_and_control_characters_are_safe(ctx):
    with ctx[1]() as db:
        db.get(User, 4).full_name = '=HYPERLINK("https://example.invalid")\x01'
        db.get(User, 4).username = "+HV4"
        db.commit()
    ws = sheet(export(ctx))
    name = next(cell for cell in ws["A"] if cell.value and "HYPERLINK" in str(cell.value))
    assert name.data_type == "s"
    assert name.value.startswith("=HYPERLINK")
    assert "\x01" not in name.value
    assert all(cell.data_type != "f" for row in ws for cell in row)


@pytest.mark.parametrize("actor,code", [(None, 401), (2, 403), (4, 403), (7, 403), (8, 403)])
def test_export_requires_manager_role_even_with_forged_token_claim(ctx, actor, code):
    assert ctx[0]("/classes/1/export", actor=actor).status_code == code


def test_revoked_permission_denies_export_and_ui_capability(ctx):
    assert ctx[0](actor=1).json()["can_export"] is True
    assert ctx[0](actor=2).json()["can_export"] is False
    with ctx[1]() as db:
        db.query(RolePermission).filter_by(role_id=1).delete()
        db.commit()
    assert export(ctx).status_code == 403


def test_empty_class_exports_valid_headers_and_missing_class_rejected(ctx):
    response = ctx[0]("/classes/3/export", actor=1)
    ws = sheet(response)
    assert (ws.max_row, ws.max_column) == (5, 3)
    assert "LOP-3_2026-10-18.xlsx" in response.headers["content-disposition"]
    assert ctx[0]("/classes/999/export", actor=1).status_code == 404
    assert export(ctx, sort="invalid").status_code == 422
