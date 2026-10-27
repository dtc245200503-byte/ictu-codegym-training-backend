from datetime import date, datetime, time, timedelta, timezone
from io import BytesIO

from openpyxl import Workbook, load_workbook
import pytest
from sqlalchemy import select, text
from sqlalchemy.exc import IntegrityError
from app.core.security import hash_password
from app.core.security import verify_password
from app.core.config import settings
from app.models.audit_log import AuditLog
from app.models.auth_session import AuthSession
from app.models.attendance import AttendanceSession
from app.models.course import Course
from app.models.classroom import Classroom
from app.models.role import Role
from app.models.notification import Notification, NotificationOutbox
from app.models.study_program import ProgramFeeSchedule, StudyProgram
from app.models.tuition import (
    TuitionDiscount,
    TuitionInstallment,
    TuitionPayment,
    TuitionPaymentReversal,
    TuitionPlan,
)
from app.models.training_class import (
    Assignment,
    AssignmentGradeVersion,
    AssignmentRecipient,
    AssignmentSubmission,
    ClassEnrollment,
    RubricTemplate,
    TrainingClass,
)
from app.models.user import User
from app.services.attendance_summary import HO_CHI_MINH
from app.api.routes import management as management_routes
from app.main import app as fastapi_app


def create_user(db, *, username, email, full_name, role, student_code=None):
    assigned_role = db.query(Role).filter(Role.code == role).one()
    user = User(
        username=username,
        email=email,
        full_name=full_name,
        student_code=student_code,
        hashed_password=hash_password("CorrectPassword123!"),
        roles=[assigned_role],
    )
    db.add(user)
    db.commit()
    db.refresh(user)
    return user


def create_role(db, code, name):
    role = Role(code=code, name=name)
    db.add(role)
    db.commit()
    return role


def get_admin(client, db):
    for code, name in (
        ("admin", "Admin"),
        ("manager", "Training Manager"),
        ("admissions", "Admissions"),
        ("accountant", "Accountant"),
        ("teacher", "Instructor"),
        ("ta", "TA"),
        ("student", "Student"),
        ("guest", "Guest"),
    ):
        create_role(db, code, name)
    user = create_user(
        db,
        username="admin01",
        email="admin@example.edu.vn",
        full_name="System Admin",
        role="admin",
    )
    response = client.post(
        "/api/auth/login",
        json={"email": user.email, "password": "CorrectPassword123!"},
    )
    assert response.status_code == 200
    return response.json()["access_token"]


def test_account_create_and_role_filtered_list(client, test_db):
    admin_token = get_admin(client, test_db)
    headers = {"Authorization": f"Bearer {admin_token}"}

    response = client.post(
        "/api/accounts",
        headers=headers,
        json={
            "username": "sv0001",
            "email": "sv0001@example.edu.vn",
            "full_name": "Nguyen Van A",
            "role": "student",
            "student_code": "SV0001",
            "phone": "09012340001",
        },
    )
    assert response.status_code == 201
    account = response.json()
    assert account["credential_delivery"] == "manual"
    temporary_password = account["temporary_password"]
    assert len(temporary_password) >= 12
    assert temporary_password != "123456"
    assert account["roles"] == ["student"]
    created_user = test_db.query(User).filter(User.username == "sv0001").one()
    assert created_user.must_change_password is True
    assert created_user.hashed_password.startswith("$argon2id$")
    assert verify_password(temporary_password, created_user.hashed_password)
    first_login = client.post(
        "/api/auth/login",
        json={"identifier": "sv0001@example.edu.vn", "password": temporary_password},
    )
    assert first_login.status_code == 200
    assert first_login.json()["user"]["must_change_password"] is True
    blocked = client.get(
        "/api/dashboard/summary",
        headers={"Authorization": "Bearer " + first_login.json()["access_token"]},
    )
    assert blocked.status_code == 403

    response = client.get(
        "/api/accounts?role=student&q=SV0001&page=1&page_size=10",
        headers=headers,
    )
    assert response.status_code == 200
    assert response.json()["total"] == 1
    assert response.json()["items"][0]["student_code"] == "SV0001"
    by_phone = client.get("/api/accounts?role=student&q=09012340001&status=active",
        headers=headers)
    assert by_phone.status_code == 200
    assert by_phone.json()["page_size"] == 20
    assert [item["id"] for item in by_phone.json()["items"]] == [created_user.id]
    created_user.is_locked = True
    test_db.commit()
    locked = client.get("/api/accounts?role=student&status=locked", headers=headers)
    assert [item["id"] for item in locked.json()["items"]] == [created_user.id]
    created_user.is_locked = False
    created_user.is_active = False
    test_db.commit()
    inactive = client.get("/api/accounts?role=student&status=inactive", headers=headers)
    assert [item["id"] for item in inactive.json()["items"]] == [created_user.id]

    duplicate = client.post("/api/accounts", headers=headers, json={
        "username": "different.username", "email": "sv0001@example.edu.vn",
        "full_name": "Duplicate Email", "role": "student"})
    assert duplicate.status_code == 409
    assert "email" in duplicate.json()["detail"].lower()


def test_account_creation_sends_activation_credentials_through_configured_transport(
    client, test_db, monkeypatch
):
    token = get_admin(client, test_db)
    for setting, value in (("MAIL_USERNAME", "smtp-user"), ("MAIL_PASSWORD", "smtp-pass"),
        ("MAIL_FROM", "no-reply@example.test"), ("MAIL_SERVER", "smtp.example.test")):
        monkeypatch.setattr(management_routes.settings, setting, value)
    calls = []

    async def fake_send(email, full_name, username, password):
        calls.append((email, full_name, username, password))

    monkeypatch.setattr(management_routes, "send_account_credentials_email", fake_send)
    response = client.post("/api/accounts", headers={"Authorization": f"Bearer {token}"},
        json={"username": "activated.teacher", "email": "activated@example.edu.vn",
            "full_name": "Activated Teacher", "role": "teacher",
            "password": "TempPass123!"})
    assert response.status_code == 201
    assert response.json()["credential_delivery"] == "email"
    assert "temporary_password" not in response.json()
    assert calls == [("activated@example.edu.vn", "Activated Teacher",
        "activated.teacher", "TempPass123!")]


def test_course_list_uses_stable_database_pages_with_filter_metadata(client, test_db):
    token = get_admin(client, test_db)
    test_db.add_all(Course(code=f"PAGE-COURSE-{index:02}",
        name=f"Page Course {index:02}", department="Page Department", credits=3)
        for index in range(25))
    test_db.commit()
    headers = {"Authorization": f"Bearer {token}"}
    first = client.get("/api/courses?q=Page%20Course&department=Page%20Department",
        headers=headers)
    second = client.get("/api/courses?q=Page%20Course&department=Page%20Department&page=2&page_size=20",
        headers=headers)
    assert first.status_code == second.status_code == 200
    assert len(first.json()) == 20
    assert len(second.json()) == 5
    assert first.headers["X-Total-Count"] == "25"
    assert first.headers["X-Total-Pages"] == "2"
    assert first.headers["X-Current-Page"] == "1"
    assert first.headers["X-Has-Previous"] == "false"
    assert first.headers["X-Has-Next"] == "true"
    assert second.headers["X-Has-Previous"] == "true"
    assert second.headers["X-Has-Next"] == "false"
    ids = [item["id"] for item in first.json() + second.json()]
    assert len(ids) == len(set(ids)) == 25
    assert client.get("/api/courses?page_size=101", headers=headers).status_code == 422


def test_class_and_roster_lists_page_in_database_and_preserve_scope(client, test_db):
    token = get_admin(client, test_db)
    teacher = create_user(test_db, username="paged.teacher", email="paged.teacher@example.edu.vn",
        full_name="Paged Teacher", role="teacher")
    course = Course(code="PAGED-CLASS-COURSE", name="Paged Class Course", department="IT", credits=3)
    test_db.add(course)
    test_db.flush()
    classes = [TrainingClass(code=f"PAGED-CLASS-{i:02}", name=f"Paged Class {i:02}",
        course_id=course.id, teacher_id=teacher.id, room="Paged Room", term="2026-Fall",
        day_of_week=i % 7, start_time=time(9), end_time=time(11), capacity=30, status="active")
        for i in range(25)]
    test_db.add_all(classes)
    test_db.commit()
    headers = {"Authorization": f"Bearer {token}"}
    first = client.get(f"/api/classes?course_id={course.id}&status=active", headers=headers)
    second = client.get(f"/api/classes?course_id={course.id}&status=active&page=2&page_size=20",
        headers=headers)
    assert len(first.json()) == 20 and len(second.json()) == 5
    assert first.headers["X-Total-Count"] == "25"
    assert first.headers["X-Has-Next"] == "true"
    assert second.headers["X-Has-Previous"] == "true"
    class_ids = [item["id"] for item in first.json() + second.json()]
    assert len(class_ids) == len(set(class_ids)) == 25

    roster_class = classes[0]
    students = [create_user(test_db, username=f"paged-student-{i:02}",
        email=f"paged-student-{i:02}@example.edu.vn", full_name=f"Paged Student {i:02}",
        role="student") for i in range(25)]
    test_db.add_all(ClassEnrollment(class_id=roster_class.id, student_id=student.id)
        for student in students)
    test_db.commit()
    roster_first = client.get(f"/api/classes/{roster_class.id}/students", headers=headers)
    roster_second = client.get(f"/api/classes/{roster_class.id}/students?page=2&page_size=20",
        headers=headers)
    assert len(roster_first.json()) == 20 and len(roster_second.json()) == 5
    assert roster_first.headers["X-Total-Count"] == "25"
    roster_ids = [item["id"] for item in roster_first.json() + roster_second.json()]
    assert len(roster_ids) == len(set(roster_ids)) == 25


def test_account_list_respects_admissions_scope_without_restricting_admin_roles(client, test_db):
    get_admin(client, test_db)
    admissions_role = test_db.query(Role).filter(Role.code == "admissions").one()
    elevated_user = create_user(
        test_db,
        username="admin.admissions",
        email="admin.admissions@example.edu.vn",
        full_name="Admin and Admissions",
        role="admin",
    )
    elevated_user.roles.append(admissions_role)
    test_db.commit()
    elevated_login = client.post(
        "/api/auth/login",
        json={"email": elevated_user.email, "password": "CorrectPassword123!"},
    )
    elevated_headers = {
        "Authorization": "Bearer " + elevated_login.json()["access_token"]
    }

    for query in (
        "/api/accounts?page=1&page_size=10",
        "/api/accounts?role=teacher&page=1&page_size=10",
        "/api/accounts?role=ta&page=1&page_size=10",
    ):
        assert client.get(query, headers=elevated_headers).status_code == 200

    admissions_user = create_user(
        test_db,
        username="admissions.only",
        email="admissions.only@example.edu.vn",
        full_name="Admissions Only",
        role="admissions",
    )
    admissions_login = client.post(
        "/api/auth/login",
        json={"email": admissions_user.email, "password": "CorrectPassword123!"},
    )
    admissions_headers = {
        "Authorization": "Bearer " + admissions_login.json()["access_token"]
    }

    assert client.get(
        "/api/accounts?role=student&page=1&page_size=10", headers=admissions_headers
    ).status_code == 200
    assert client.get(
        "/api/accounts?page=1&page_size=10", headers=admissions_headers
    ).status_code == 403
    assert client.get(
        "/api/accounts?role=teacher&page=1&page_size=10", headers=admissions_headers
    ).status_code == 403


def test_account_create_uses_admin_selected_initial_password(client, test_db):
    admin_token = get_admin(client, test_db)
    response = client.post(
        "/api/accounts",
        headers={"Authorization": "Bearer " + admin_token},
        json={
            "username": "teacher.custom",
            "email": "teacher.custom@example.edu.vn",
            "full_name": "Custom Password Teacher",
            "role": "teacher",
            "password": "TeacherTemp123!",
        },
    )

    assert response.status_code == 201
    assert response.json()["temporary_password"] == "TeacherTemp123!"
    user = test_db.query(User).filter(User.username == "teacher.custom").one()
    assert user.must_change_password is True
    assert verify_password("TeacherTemp123!", user.hashed_password)


def test_account_create_and_role_update_support_multiple_roles(client, test_db):
    admin_token = get_admin(client, test_db)
    headers = {"Authorization": f"Bearer {admin_token}"}
    created = client.post(
        "/api/accounts",
        headers=headers,
        json={
            "username": "teacher.admissions",
            "email": "teacher.admissions@example.edu.vn",
            "full_name": "Teacher and Advisor",
            "roles": ["teacher", "admissions"],
        },
    )
    assert created.status_code == 201, created.text
    assert set(created.json()["roles"]) == {"teacher", "admissions"}
    created_login = client.post("/api/auth/login", json={
        "email": "teacher.admissions@example.edu.vn",
        "password": created.json()["temporary_password"],
    })
    assert created_login.status_code == 200
    created_headers = {"Authorization": f"Bearer {created_login.json()['access_token']}"}

    updated = client.put(
        f"/api/accounts/{created.json()['id']}/role",
        headers=headers,
        json={"roles": ["teacher", "ta", "admissions"]},
    )
    assert updated.status_code == 200
    assert set(updated.json()["roles"]) == {"teacher", "ta", "admissions"}
    next_request = client.get("/api/me", headers=created_headers)
    assert next_request.status_code == 200
    assert set(next_request.json()["roles"]) == {"teacher", "ta", "admissions"}
    role_change = (
        test_db.query(AuditLog)
        .filter_by(action="account.role_changed", target_user_id=created.json()["id"])
        .one()
    )
    assert role_change.actor_id is not None
    assert role_change.created_at is not None
    assert role_change.details == "before=admissions,teacher;after=admissions,ta,teacher"

    unchanged = client.put(
        f"/api/accounts/{created.json()['id']}/role",
        headers=headers,
        json={"roles": ["teacher", "ta", "admissions"]},
    )
    assert unchanged.status_code == 200
    assert (
        test_db.query(AuditLog)
        .filter_by(action="account.role_changed", target_user_id=created.json()["id"])
        .count()
        == 1
    )

    admin_user = test_db.query(User).filter_by(email="admin@example.edu.vn").one()
    self_revoke = client.put(f"/api/accounts/{admin_user.id}/role", headers=headers,
        json={"roles": ["manager"]})
    assert self_revoke.status_code == 400
    assert "admin" in test_db.get(User, admin_user.id).roles[0].code

    invalid = client.put(
        f"/api/accounts/{created.json()['id']}/role",
        headers=headers,
        json={"roles": ["teacher", "unknown"]},
    )
    assert invalid.status_code == 422


def test_account_lock_and_unlock_require_audit_reason(client, test_db):
    admin_token = get_admin(client, test_db)
    student = create_user(
        test_db,
        username="lock.student",
        email="lock.student@example.edu.vn",
        full_name="Lock Student",
        role="student",
    )
    headers = {"Authorization": f"Bearer {admin_token}"}
    assert client.post(f"/api/accounts/{student.id}/lock", headers=headers, json={}).status_code == 422

    locked = client.post(
        f"/api/accounts/{student.id}/lock",
        headers=headers,
        json={"reason": "Tạm khóa để xác minh hồ sơ."},
    )
    assert locked.status_code == 200
    assert locked.json()["is_locked"] is True
    unlocked = client.post(
        f"/api/accounts/{student.id}/unlock",
        headers=headers,
        json={"reason": "Đã xác minh thông tin."},
    )
    assert unlocked.status_code == 200
    assert unlocked.json()["is_locked"] is False


def test_locking_staff_revokes_sessions_rejects_login_and_reports_handover_classes(client, test_db):
    admin_token = get_admin(client, test_db)
    teacher = create_user(test_db, username="handover.teacher",
        email="handover.teacher@example.edu.vn", full_name="Handover Teacher",
        role="teacher")
    course = Course(code="HANDOVER101", name="Handover Course", department="IT", credits=3)
    test_db.add(course)
    test_db.flush()
    training_class = TrainingClass(code="HANDOVER-01", name="Handover Class",
        course_id=course.id, teacher_id=teacher.id, room="Room 1", term="2026-Fall",
        day_of_week=0, start_time=time(9), end_time=time(11), capacity=20, status="active")
    test_db.add(training_class)
    test_db.commit()
    teacher_login = client.post("/api/auth/login", json={
        "email": teacher.email, "password": "CorrectPassword123!"})
    assert teacher_login.status_code == 200
    admin_headers = {"Authorization": f"Bearer {admin_token}"}
    locked = client.post(f"/api/accounts/{teacher.id}/lock", headers=admin_headers,
        json={"reason": "Staff departure"})
    assert locked.status_code == 200
    assert locked.json()["handover_classes"] == [training_class.code]
    assert test_db.query(AuthSession).filter_by(user_id=teacher.id).one().revoked_at is not None
    teacher_headers = {"Authorization": f"Bearer {teacher_login.json()['access_token']}"}
    assert client.get("/api/me", headers=teacher_headers).status_code == 401
    assert client.post("/api/auth/refresh", json={
        "refresh_token": teacher_login.json()["refresh_token"]}).status_code == 401
    assert client.post("/api/auth/login", json={
        "email": teacher.email, "password": "CorrectPassword123!"}).status_code == 403
    audit = test_db.query(AuditLog).filter_by(action="account.locked",
        target_user_id=teacher.id).one()
    assert training_class.code in audit.details


def test_course_catalog_stores_outcomes_and_reuses_weighted_course_across_programs(client, test_db):
    admin_token = get_admin(client, test_db)
    headers = {"Authorization": f"Bearer {admin_token}"}
    course_response = client.post("/api/courses", headers=headers, json={
        "code": "REUSE101", "name": "Reusable Course", "department": "IT",
        "credits": 3, "session_count": 20,
        "learning_outcomes": "Design and explain a normalized relational schema.",
    })
    assert course_response.status_code == 201
    course = course_response.json()
    assert course["session_count"] == 20
    assert course["learning_outcomes"].startswith("Design and explain")
    assert client.post("/api/courses", headers=headers, json={
        "code": "REUSE101", "name": "Duplicate", "department": "IT", "credits": 3,
    }).status_code == 409

    programs = []
    for number in (1, 2):
        created_program = client.post("/api/study-programs", headers=headers, json={
            "code": f"REUSE-P{number}", "name": f"Reuse Program {number}", "total_credits": 120,
        })
        assert created_program.status_code == 201
        programs.append(created_program.json())
    for program, weight in zip(programs, (30, 45)):
        linked = client.put(f"/api/study-programs/{program['id']}/courses", headers=headers, json={
            "courses": [{"course_id": course["id"], "weight": weight}],
        })
        assert linked.status_code == 200
        syllabus = client.get(f"/api/study-programs/{program['id']}/subjects", headers=headers)
        assert syllabus.json()[0]["weight"] == f"{weight}.00"
        assert syllabus.json()[0]["learning_outcomes"] == course["learning_outcomes"]

    teacher = create_user(test_db, username="reuse.teacher", email="reuse.teacher@example.edu.vn",
        full_name="Reuse Teacher", role="teacher")
    retained_class = TrainingClass(code="REUSE101-A", name="Reusable Class", course_id=course["id"],
        program_id=programs[0]["id"], teacher_id=teacher.id, room="R1", term="2026-Fall",
        day_of_week=1, start_time=time(9), end_time=time(11), capacity=20, status="active")
    test_db.add(retained_class)
    test_db.commit()
    assert client.delete(f"/api/courses/{course['id']}", headers=headers).status_code == 409
    denied_student = create_user(test_db, username="reuse.student", email="reuse.student@example.edu.vn",
        full_name="Reuse Student", role="student")
    denied_token = client.post("/api/auth/login", json={
        "email": denied_student.email, "password": "CorrectPassword123!",
    }).json()["access_token"]
    assert client.post("/api/courses", headers={"Authorization": f"Bearer {denied_token}"}, json={
        "code": "NO-COURSE", "name": "Denied", "department": "IT", "credits": 1,
    }).status_code == 403


def test_course_syllabus_sessions_can_be_managed(client, test_db):
    admin_token = get_admin(client, test_db)
    headers = {"Authorization": f"Bearer {admin_token}"}
    course = Course(code="SYL101", name="Syllabus Course", department="IT", credits=3)
    test_db.add(course)
    test_db.commit()
    dependent = client.post(
        "/api/courses",
        headers=headers,
        json={
            "code": "SYL102",
            "name": "Advanced Syllabus",
            "department": "IT",
            "credits": 3,
            "prerequisite_id": course.id,
        },
    )
    assert dependent.status_code == 201
    assert dependent.json()["prerequisite_id"] == course.id

    second = client.post(
        f"/api/courses/{course.id}/lessons",
        headers=headers,
        json={"sequence": 2, "title": "Databases", "objectives": "Model relational data", "description": "Relational design"},
    )
    first = client.post(
        f"/api/courses/{course.id}/lessons",
        headers=headers,
        json={"sequence": 1, "title": "Introduction", "objectives": "Understand course foundations"},
    )
    assert second.status_code == first.status_code == 201
    beyond_limit = client.post(
        f"/api/courses/{course.id}/lessons",
        headers=headers,
        json={
            "sequence": 17,
            "title": "Beyond syllabus",
            "objectives": "Should be rejected",
        },
    )
    assert beyond_limit.status_code == 422
    cannot_shrink = client.patch(
        f"/api/courses/{course.id}",
        headers=headers,
        json={"session_count": 1},
    )
    assert cannot_shrink.status_code == 409
    duplicate = client.post(
        f"/api/courses/{course.id}/lessons",
        headers=headers,
        json={"sequence": 1, "title": "Duplicate sequence", "objectives": "Duplicate lesson"},
    )
    assert duplicate.status_code == 409
    listed = client.get(f"/api/courses/{course.id}/lessons", headers=headers)
    assert [lesson["sequence"] for lesson in listed.json()] == [1, 2]
    edited = client.patch(
        f"/api/courses/{course.id}/lessons/{second.json()['id']}",
        headers=headers,
        json={"title": "SQL and databases"},
    )
    assert edited.status_code == 200
    assert edited.json()["title"] == "SQL and databases"
    clone = client.post(
        f"/api/courses/{dependent.json()['id']}/lessons/clone",
        headers=headers,
        json={"source_course_id": course.id},
    )
    assert clone.status_code == 201
    assert [lesson["objectives"] for lesson in clone.json()] == [
        "Understand course foundations",
        "Model relational data",
    ]
    assert client.delete(
        f"/api/courses/{course.id}/lessons/{first.json()['id']}",
        headers=headers,
    ).status_code == 204


def test_new_admin_created_account_can_login_with_email_or_username(client, test_db):
    admin_token = get_admin(client, test_db)
    response = client.post(
        "/api/accounts",
        headers={"Authorization": f"Bearer {admin_token}"},
        json={
            "username": "new.admin",
            "email": "new.admin@example.edu.vn",
            "full_name": "New Administrator",
            "role": "admin",
        },
    )

    assert response.status_code == 201
    created = response.json()
    account = test_db.query(User).filter(User.username == "new.admin").one()
    assert account.is_active is True
    assert account.must_change_password is True
    assert account.hashed_password != created["temporary_password"]
    assert verify_password(created["temporary_password"], account.hashed_password)

    for identifier in (created["email"], created["username"]):
        login = client.post(
            "/api/auth/login",
            json={
                "identifier": identifier,
                "password": created["temporary_password"],
            },
        )
        assert login.status_code == 200
        assert login.json()["access_token"]
        assert login.json()["refresh_token"]
        assert login.json()["user"]["must_change_password"] is True


def test_admin_must_change_password_before_accessing_other_apis(client, test_db):
    admin_token = get_admin(client, test_db)
    response = client.post(
        "/api/accounts",
        headers={"Authorization": f"Bearer {admin_token}"},
        json={
            "username": "first.login.admin",
            "email": "first.login.admin@example.edu.vn",
            "full_name": "First Login Admin",
            "role": "admin",
        },
    )
    created = response.json()
    login = client.post(
        "/api/auth/login",
        json={
            "identifier": created["username"],
            "password": created["temporary_password"],
        },
    )
    assert login.status_code == 200
    access_token = login.json()["access_token"]
    headers = {"Authorization": f"Bearer {access_token}"}
    assert client.get("/api/me", headers=headers).json()["must_change_password"] is True
    blocked = client.get("/api/dashboard/summary", headers=headers)
    assert blocked.status_code == 403
    assert blocked.json()["detail"]["code"] == "password_change_required"

    changed = client.post(
        "/api/auth/change-password",
        headers=headers,
        json={
            "current_password": created["temporary_password"],
            "new_password": "NewAdminPassword456!",
        },
    )
    assert changed.status_code == 200
    assert client.get("/api/me", headers=headers).json()["must_change_password"] is False
    assert client.get("/api/dashboard/summary", headers=headers).status_code == 200


def test_non_admin_cannot_manage_accounts(client, test_db):
    create_role(test_db, "student", "Học viên")
    student = create_user(
        test_db,
        username="student01",
        email="student@example.edu.vn",
        full_name="Student",
        role="student",
    )
    token = client.post(
        "/api/auth/login",
        json={"email": student.email, "password": "CorrectPassword123!"},
    ).json()["access_token"]

    response = client.get(
        "/api/accounts", headers={"Authorization": f"Bearer {token}"}
    )
    assert response.status_code == 403


def test_excel_import_skips_invalid_rows_and_creates_valid_accounts(client, test_db):
    admin_token = get_admin(client, test_db)
    headers = {
        "Authorization": f"Bearer {admin_token}",
        "Content-Type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    }
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["username", "email", "full_name", "role", "student_code"])
    sheet.append(["sv0001", "duplicate@example.edu.vn", "Student One", "student", "SV0001"])
    sheet.append(["sv0002", "duplicate@example.edu.vn", "Student Two", "student", "SV0002"])
    invalid_file = BytesIO()
    workbook.save(invalid_file)

    preview = client.post(
        "/api/accounts/import?preview=true",
        headers=headers,
        content=invalid_file.getvalue(),
    )
    assert preview.status_code == 200
    assert preview.json()["valid_count"] == 1
    assert preview.json()["invalid_count"] == 1
    assert preview.json()["rows"][1]["errors"] == ["email đã tồn tại"]
    assert test_db.query(User).count() == 1

    response = client.post(
        "/api/accounts/import",
        headers=headers,
        content=invalid_file.getvalue(),
    )
    assert response.status_code == 200
    assert response.json()["created_count"] == 1
    assert response.json()["skipped_count"] == 1
    assert response.json()["row_errors"][0]["row"] == 3
    assert test_db.query(User).count() == 2

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["username", "email", "full_name", "role", "student_code"])
    sheet.append(["sv0003", "sv0003@example.edu.vn", "Student Three", "student", "SV0003"])
    valid_file = BytesIO()
    workbook.save(valid_file)

    response = client.post(
        "/api/accounts/import",
        headers=headers,
        content=valid_file.getvalue(),
    )
    assert response.status_code == 200
    assert response.json()["created_count"] == 1
    assert response.json()["accounts"][0]["credential_delivery"] == "manual"
    temporary_password = response.json()["accounts"][0]["temporary_password"]
    assert len(temporary_password) >= 12
    assert temporary_password != "123456"
    imported_user = test_db.query(User).filter(User.username == "sv0003").one()
    assert imported_user.must_change_password is True
    assert imported_user.hashed_password.startswith("$argon2id$")
    assert verify_password(temporary_password, imported_user.hashed_password)


def test_account_import_template_and_file_boundaries_are_admin_only(client, test_db):
    admin_token = get_admin(client, test_db)
    admin_headers = {"Authorization": f"Bearer {admin_token}"}
    template = client.get("/api/accounts/template", headers=admin_headers)
    assert template.status_code == 200
    workbook = load_workbook(BytesIO(template.content), read_only=True, data_only=True)
    assert next(workbook.active.iter_rows(values_only=True)) == (
        "username", "email", "full_name", "role", "student_code", "phone", "password"
    )
    workbook.close()

    student = create_user(
        test_db, username="import_student", email="import.student@example.edu.vn",
        full_name="Import Student", role="student",
    )
    student_token = client.post("/api/auth/login", json={
        "email": student.email, "password": "CorrectPassword123!",
    }).json()["access_token"]
    student_headers = {"Authorization": f"Bearer {student_token}"}
    assert client.post("/api/accounts/import", headers=student_headers, content=b"bad").status_code == 403
    assert client.get("/api/accounts/template", headers=student_headers).status_code == 403

    assert client.post("/api/accounts/import", headers=admin_headers, content=b"not an xlsx").status_code == 400
    assert client.post("/api/accounts/import", headers=admin_headers, content=b"x" * (10 * 1024 * 1024 + 1)).status_code == 413

    invalid_workbook = Workbook()
    invalid_workbook.active.append(["username", "email"])
    invalid_bytes = BytesIO()
    invalid_workbook.save(invalid_bytes)
    missing_headers = client.post("/api/accounts/import", headers=admin_headers, content=invalid_bytes.getvalue())
    assert missing_headers.status_code == 422
    assert test_db.query(User).filter(User.username == "import_student").count() == 1


def test_study_program_catalog_create_versioned_fee_deactivate_and_delete_guards(client, test_db):
    admin_token = get_admin(client, test_db)
    admin_headers = {"Authorization": f"Bearer {admin_token}"}
    created = client.post("/api/study-programs", headers=admin_headers, json={
        "code": "catalog-01", "name": "Catalog Program", "total_credits": 120,
        "duration_semesters": 8, "standard_tuition_vnd": 12000000,
    })
    assert created.status_code == 201, created.text
    program = created.json()
    assert program["code"] == "CATALOG-01"
    assert program["is_active"] is True
    assert program["standard_tuition_vnd"] == "12000000"
    duplicate = client.post("/api/study-programs", headers=admin_headers, json={
        "code": "CATALOG-01", "name": "Duplicate Program", "total_credits": 100,
    })
    assert duplicate.status_code == 409
    student = create_user(test_db, username="catalog_student", email="catalog.student@example.edu.vn",
        full_name="Catalog Student", role="student")
    student_token = client.post("/api/auth/login", json={
        "email": student.email, "password": "CorrectPassword123!",
    }).json()["access_token"]
    denied = client.post("/api/study-programs", headers={"Authorization": f"Bearer {student_token}"}, json={
        "code": "STUDENT-CATALOG", "name": "Student Cannot Create", "total_credits": 100,
    })
    assert denied.status_code == 403

    preserved_fee = ProgramFeeSchedule(program_id=program["id"], version=1,
        amount_vnd=10000000, valid_from=date.today() - timedelta(days=30), created_by=1)
    test_db.add(preserved_fee)
    test_db.commit()

    updated = client.patch(f"/api/study-programs/{program['id']}", headers=admin_headers, json={
        "standard_tuition_vnd": 15000000, "description": "Current catalog details",
    })
    assert updated.status_code == 200
    assert updated.json()["standard_tuition_vnd"] == "15000000"
    schedules = test_db.query(ProgramFeeSchedule).filter_by(program_id=program["id"]).all()
    assert len(schedules) == 1 and str(schedules[0].amount_vnd) == "10000000"

    manager_user = test_db.query(User).filter(User.username == "admin01").one()
    retained_course = Course(code="CATALOG-COURSE", name="Catalog Course", department="IT", credits=3)
    test_db.add(retained_course)
    test_db.flush()
    program_model = test_db.get(StudyProgram, program["id"])
    program_model.courses.append(retained_course)
    retained_class = TrainingClass(code="CATALOG-CLASS", name="Catalog Class",
        course_id=retained_course.id, program_id=program["id"], teacher_id=manager_user.id,
        room="R1", term="2026-Fall", day_of_week=1, start_time=time(9), end_time=time(11),
        capacity=20, status="active")
    test_db.add(retained_class)
    test_db.commit()
    assert client.delete(f"/api/study-programs/{program['id']}", headers=admin_headers).status_code == 409

    disabled = client.patch(f"/api/study-programs/{program['id']}", headers=admin_headers,
        json={"is_active": False})
    assert disabled.status_code == 200 and disabled.json()["is_active"] is False
    inactive_class = client.post("/api/classes", headers=admin_headers, json={
        "code": "INACTIVE-PROGRAM-CLASS", "name": "Should be rejected",
        "course_id": retained_course.id, "program_id": program["id"],
        "teacher_id": manager_user.id, "room": "R1", "term": "2026-Fall",
        "day_of_week": 1, "start_time": "09:00", "end_time": "11:00", "capacity": 20,
    })
    assert inactive_class.status_code == 409
    catalog = client.get("/api/study-programs", headers=admin_headers)
    assert next(item for item in catalog.json() if item["id"] == program["id"])["standard_tuition_vnd"] == "15000000"
    assert client.delete(f"/api/study-programs/{program['id']}", headers=admin_headers).status_code == 409

    manager = create_user(test_db, username="catalog_manager", email="catalog.manager@example.edu.vn",
        full_name="Catalog Manager", role="manager")
    manager_token = client.post("/api/auth/login", json={
        "email": manager.email, "password": "CorrectPassword123!",
    }).json()["access_token"]
    assert client.patch(f"/api/study-programs/{program['id']}",
        headers={"Authorization": f"Bearer {manager_token}"}, json={"name": "Manager Edited"}).status_code == 200


def test_change_password_revokes_refresh_session(client, test_db):
    create_role(test_db, "student", "Học viên")
    student = create_user(
        test_db,
        username="student01",
        email="student@example.edu.vn",
        full_name="Student",
        role="student",
    )
    login = client.post(
        "/api/auth/login",
        json={"email": student.email, "password": "CorrectPassword123!"},
    ).json()

    response = client.post(
        "/api/auth/change-password",
        headers={"Authorization": f"Bearer {login['access_token']}"},
        json={
            "current_password": "CorrectPassword123!",
            "new_password": "NewPassword456!",
        },
    )
    assert response.status_code == 200
    assert client.post(
        "/api/auth/refresh",
        json={"refresh_token": login["refresh_token"]},
    ).status_code == 401
    assert client.post(
        "/api/auth/login",
        json={"email": student.email, "password": "NewPassword456!"},
    ).status_code == 200


def test_change_password_requires_current_and_policy_compliant_new_password(client, test_db):
    create_role(test_db, "student", "Student")
    student = create_user(test_db, username="password-policy-student",
        email="password-policy@example.edu.vn", full_name="Password Policy",
        role="student")
    login = client.post("/api/auth/login", json={"email": student.email,
        "password": "CorrectPassword123!"}).json()
    headers = {"Authorization": f"Bearer {login['access_token']}"}
    wrong_current = client.post("/api/auth/change-password", headers=headers,
        json={"current_password": "WrongPassword123!", "new_password": "NewPassword456!"})
    assert wrong_current.status_code == 400
    same_password = client.post("/api/auth/change-password", headers=headers,
        json={"current_password": "CorrectPassword123!", "new_password": "CorrectPassword123!"})
    assert same_password.status_code == 400
    for invalid in ("Ab1!xxx", "Password!!!", "12345678!"):
        response = client.post("/api/auth/change-password", headers=headers,
            json={"current_password": "CorrectPassword123!", "new_password": invalid})
        assert response.status_code == 422


def test_class_enrollment_rejects_student_schedule_conflict(client, test_db):
    admin_token = get_admin(client, test_db)
    teacher = create_user(
        test_db,
        username="teacher01",
        email="teacher@example.edu.vn",
        full_name="Teacher",
        role="teacher",
    )
    second_teacher = create_user(
        test_db,
        username="teacher02",
        email="teacher02@example.edu.vn",
        full_name="Second Teacher",
        role="teacher",
    )
    student = create_user(
        test_db,
        username="student01",
        email="student@example.edu.vn",
        full_name="Student",
        role="student",
    )
    course_a = Course(
        code="CS101", name="Programming", department="Computer Science", credits=3
    )
    course_b = Course(
        code="CS102", name="Algorithms", department="Computer Science", credits=3
    )
    test_db.add_all([course_a, course_b])
    test_db.commit()

    class_a = TrainingClass(
        code="CS101-01",
        name="Programming 01",
        course_id=course_a.id,
        teacher_id=teacher.id,
        room="A101",
        term="2026-1",
        day_of_week=0,
        start_time=time(9, 0),
        end_time=time(11, 0),
        capacity=30,
    )
    test_db.add(class_a)
    test_db.commit()
    test_db.add(ClassEnrollment(class_id=class_a.id, student_id=student.id))
    test_db.commit()

    response = client.post(
        "/api/classes",
        headers={"Authorization": f"Bearer {admin_token}"},
        json={
            "code": "CS102-01",
            "name": "Algorithms 01",
            "course_id": course_b.id,
            "teacher_id": second_teacher.id,
            "room": "A102",
            "term": "2026-1",
            "day_of_week": 0,
            "start_time": "10:30",
            "end_time": "12:00",
            "capacity": 30,
        },
    )
    assert response.status_code == 201
    response = client.post(
        f"/api/classes/{response.json()['id']}/students",
        headers={"Authorization": f"Bearer {admin_token}"},
        json={"student_ids": [student.id]},
    )
    assert response.status_code == 409
    assert "trùng" in response.json()["detail"]


def test_classroom_catalog_program_class_creation_and_capacity_rules(client, test_db):
    admin_token = get_admin(client, test_db)
    headers = {"Authorization": "Bearer " + admin_token}
    insecure_room = client.post("/api/rooms", headers=headers, json={
        "code": "INSECURE-ROOM", "capacity": 20, "online_url": "http://meet.example.edu.vn/room",
    })
    assert insecure_room.status_code == 422
    teacher = create_user(
        test_db,
        username="room.teacher",
        email="room.teacher@example.edu.vn",
        full_name="Room Teacher",
        role="teacher",
    )
    course = Course(
        code="ROOM101",
        name="Room Catalog",
        department="Computer Science",
        credits=3,
    )
    test_db.add(course)
    test_db.flush()
    second_course = Course(code="ROOM102", name="Room Catalog II", department="Computer Science", credits=3)
    test_db.add(second_course)
    test_db.flush()
    program = StudyProgram(
        code="ROOM-PROG",
        name="Room Program",
        total_credits=120,
        courses=[course, second_course],
    )
    test_db.add(program)
    test_db.commit()

    room_response = client.post(
        "/api/rooms",
        headers=headers,
        json={
            "code": "LAB-9",
            "capacity": 30,
            "online_url": "https://meet.example.edu.vn/lab-9",
        },
    )
    assert room_response.status_code == 201
    room = room_response.json()
    assert room["capacity"] == 30
    assert client.get("/api/rooms", headers=headers).json()[0]["code"] == "LAB-9"

    class_payload = {
        "code": "ROOM101-01",
        "name": "Room Catalog Section 01",
        "course_id": course.id,
        "program_id": program.id,
        "planned_start_date": "2026-03-01",
        "delivery_mode": "online",
        "teacher_id": teacher.id,
        "room": "LAB-9",
        "term": "2026-2",
        "day_of_week": 2,
        "start_time": "13:00",
        "end_time": "15:00",
        "capacity": 31,
    }
    oversized = client.post("/api/classes", headers=headers, json=class_payload)
    assert oversized.status_code == 422

    class_payload["capacity"] = 20
    created = client.post("/api/classes", headers=headers, json=class_payload)
    assert created.status_code == 201
    assert created.json()["room"] == "LAB-9"
    assert created.json()["program_id"] == program.id
    assert created.json()["planned_start_date"] == "2026-03-01"
    assert created.json()["delivery_mode"] == "online"
    assert [item["code"] for item in created.json()["program_curriculum_snapshot"]] == [course.code, second_course.code]
    assert len(client.get("/api/classes?term=2026-2&q=ROOM101", headers=headers).json()) == 1
    assert len(client.get(f"/api/classes?program_id={program.id}", headers=headers).json()) == 1
    assert client.get("/api/classes?term=2026-1", headers=headers).json() == []
    test_db.add(AttendanceSession(class_id=created.json()["id"], session_date=date(2026, 2, 3),
        start_time=time(13), end_time=time(15)))
    test_db.commit()
    date_filtered = client.get("/api/classes?start_date=2026-02-01&end_date=2026-02-28", headers=headers)
    assert date_filtered.status_code == 200 and len(date_filtered.json()) == 1
    assert date_filtered.json()[0]["session_progress"] == "1/1"
    assert client.get("/api/classes?start_date=2026-03-01&end_date=2026-02-28", headers=headers).status_code == 422

    students = [
        create_user(
            test_db,
            username=f"bulk.student{index}",
            email=f"bulk.student{index}@example.edu.vn",
            full_name=f"Bulk Student {index}",
            role="student",
        )
        for index in (1, 2)
    ]
    enrolled = client.post(
        f"/api/classes/{created.json()['id']}/students",
        headers=headers,
        json={"student_ids": [student.id for student in students]},
    )
    assert enrolled.status_code == 201
    assert enrolled.json()["enrolled_count"] == 2
    roster = client.get(f"/api/classes/{created.json()['id']}/students", headers=headers)
    assert roster.status_code == 200
    assert all(item["enrolled_by_id"] is not None for item in roster.json())
    assert all(item["enrolled_at"] for item in roster.json())
    duplicate_request = client.post(f"/api/classes/{created.json()['id']}/students",
        headers=headers, json={"student_ids": [students[0].id, students[0].id]})
    assert duplicate_request.status_code == 422

    too_small = client.patch(
        f"/api/rooms/{room['id']}",
        headers=headers,
        json={"capacity": 10},
    )
    assert too_small.status_code == 409
    renamed = client.patch(
        f"/api/rooms/{room['id']}",
        headers=headers,
        json={"code": "LAB-10"},
    )
    assert renamed.status_code == 200
    assert test_db.query(TrainingClass).filter_by(code="ROOM101-01").one().room == "LAB-10"

    wrong_program = client.post(
        "/api/classes",
        headers=headers,
        json={
            **class_payload,
            "code": "ROOM101-02",
            "program_id": program.id + 1,
            "room": "LAB-10",
        },
    )
    assert wrong_program.status_code == 422
    missing_room = client.post(
        "/api/classes",
        headers=headers,
        json={
            **class_payload,
            "code": "ROOM101-03",
            "program_id": program.id,
            "room": "UNLISTED",
            "day_of_week": 4,
        },
    )
    assert missing_room.status_code == 422


def test_program_curriculum_persists_course_order_and_prerequisite(client, test_db):
    admin_token = get_admin(client, test_db)
    headers = {"Authorization": "Bearer " + admin_token}
    first = Course(code="CUR101", name="Foundations", department="IT", credits=3)
    second = Course(code="CUR201", name="Advanced", department="IT", credits=3)
    program = StudyProgram(
        code="CUR-PROG",
        name="Curriculum Program",
        total_credits=120,
        courses=[first, second],
    )
    test_db.add(program)
    test_db.commit()

    response = client.put(
        f"/api/study-programs/{program.id}/courses",
        headers=headers,
        json={
            "courses": [
                {"course_id": first.id},
                {"course_id": second.id, "prerequisite_course_id": first.id},
            ]
        },
    )
    assert response.status_code == 200
    result = client.get("/api/study-programs", headers=headers)
    program_payload = next(item for item in result.json() if item["id"] == program.id)
    assert [item["id"] for item in program_payload["courses"]] == [first.id, second.id]
    assert program_payload["courses"][1]["prerequisite_course_id"] == first.id

    curriculum = client.get(
        f"/api/study-programs/{program.id}/subjects",
        headers=headers,
    )
    assert curriculum.status_code == 200
    assert [item["subject_id"] for item in curriculum.json()] == [first.id, second.id]
    assert curriculum.json()[1]["order_index"] == 2
    assert curriculum.json()[1]["prerequisite_subject_id"] == first.id

    third = Course(code="CUR301", name="Capstone", department="IT", credits=3)
    test_db.add(third)
    test_db.commit()
    added = client.post(
        f"/api/study-programs/{program.id}/subjects",
        headers=headers,
        json={"subject_id": third.id},
    )
    assert added.status_code == 201
    assert added.json()["order_index"] == 3
    curriculum_page = client.get(
        f"/api/study-programs/{program.id}/subjects?q=CUR&sort=code&page=2&page_size=1",
        headers=headers,
    )
    assert curriculum_page.status_code == 200
    assert [item["code"] for item in curriculum_page.json()] == ["CUR201"]
    assert curriculum_page.headers["x-total-count"] == "3"
    embedded_course_page = client.get(
        "/api/study-programs?course_q=CUR&course_sort=code&course_page=2&course_page_size=1",
        headers=headers,
    )
    embedded_program = next(item for item in embedded_course_page.json() if item["id"] == program.id)
    assert [item["code"] for item in embedded_program["courses"]] == ["CUR201"]
    assert embedded_program["courses_pagination"]["total"] == 3
    capstone_search = client.get(
        f"/api/study-programs/{program.id}/subjects?q=Capstone", headers=headers
    )
    assert [item["id"] for item in capstone_search.json()] == [third.id]
    duplicate = client.post(
        f"/api/study-programs/{program.id}/subjects",
        headers=headers,
        json={"subject_id": third.id},
    )
    assert duplicate.status_code == 409

    updated = client.put(
        f"/api/study-programs/{program.id}/subjects",
        headers=headers,
        json={
            "subjects": [
                {"subject_id": first.id},
                {"subject_id": second.id, "prerequisite_subject_id": first.id},
                {"subject_id": third.id, "prerequisite_subject_id": second.id},
            ]
        },
    )
    assert updated.status_code == 200
    removed = client.delete(
        f"/api/study-programs/{program.id}/subjects/{second.id}",
        headers=headers,
    )
    assert removed.status_code == 204
    after_remove = client.get(
        f"/api/study-programs/{program.id}/subjects",
        headers=headers,
    )
    assert [item["subject_id"] for item in after_remove.json()] == [first.id, third.id]
    assert [item["order_index"] for item in after_remove.json()] == [1, 2]
    assert after_remove.json()[1]["prerequisite_subject_id"] is None

    invalid_order = client.put(
        f"/api/study-programs/{program.id}/courses",
        headers=headers,
        json={
            "courses": [
                {"course_id": second.id, "prerequisite_course_id": first.id},
                {"course_id": first.id},
            ]
        },
    )
    assert invalid_order.status_code == 422


def test_bulk_enrollment_checks_physical_room_capacity(client, test_db):
    admin_token = get_admin(client, test_db)
    headers = {"Authorization": f"Bearer {admin_token}"}
    teacher = create_user(
        test_db, username="capacity.teacher", email="capacity.teacher@example.edu.vn",
        full_name="Capacity Teacher", role="teacher",
    )
    students = [
        create_user(
            test_db,
            username=f"capacity.student.{index}",
            email=f"capacity.student.{index}@example.edu.vn",
            full_name=f"Capacity Student {index}",
            role="student",
        )
        for index in range(2)
    ]
    course = Course(code="CAP101", name="Capacity", department="IT", credits=3)
    test_db.add(course)
    test_db.commit()
    room = client.post(
        "/api/rooms",
        headers=headers,
        json={"code": "CAP-ROOM", "capacity": 5},
    )
    assert room.status_code == 201
    created = client.post(
        "/api/classes",
        headers=headers,
        json={
            "code": "CAP101-01",
            "name": "Capacity Test",
            "course_id": course.id,
            "teacher_id": teacher.id,
            "room": "CAP-ROOM",
            "term": "2026-1",
            "day_of_week": 1,
            "start_time": "09:00",
            "end_time": "11:00",
            "capacity": 5,
        },
    )
    assert created.status_code == 201
    test_db.query(Classroom).filter_by(code="CAP-ROOM").one().capacity = 1
    test_db.commit()

    response = client.post(
        f"/api/classes/{created.json()['id']}/students",
        headers=headers,
        json={"student_ids": [student.id for student in students]},
    )
    assert response.status_code == 409
    assert "phòng học" in response.json()["detail"]


def test_class_schedule_changes_notify_students_and_roster_status_is_exportable(
    client, test_db
):
    admin_token = get_admin(client, test_db)
    teacher = create_user(
        test_db,
        username="teacher01",
        email="teacher@example.edu.vn",
        full_name="Teacher",
        role="teacher",
    )
    student = create_user(
        test_db,
        username="student01",
        email="student@example.edu.vn",
        full_name="Student",
        role="student",
        student_code="SV0001",
    )
    course = Course(
        code="CS201", name="Database", department="Computer Science", credits=3
    )
    test_db.add(course)
    test_db.commit()
    created = client.post(
        "/api/classes",
        headers={"Authorization": f"Bearer {admin_token}"},
        json={
            "code": "CS201-01",
            "name": "Database 01",
            "course_id": course.id,
            "teacher_id": teacher.id,
            "room": "B201",
            "term": "2026-1",
            "day_of_week": 1,
            "start_time": "09:00",
            "end_time": "11:00",
            "capacity": 30,
        },
    )
    assert created.status_code == 201
    class_id = created.json()["id"]
    enrolled = client.post(
        f"/api/classes/{class_id}/students",
        headers={"Authorization": f"Bearer {admin_token}"},
        json={"student_ids": [student.id]},
    )
    assert enrolled.status_code == 201
    tuition_response = client.patch(
        f"/api/classes/{class_id}/students/{student.id}/status",
        headers={"Authorization": f"Bearer {admin_token}"},
        json={"tuition_status": "paid"},
    )
    assert tuition_response.status_code == 200
    tuition_audit = (
        test_db.query(AuditLog)
        .filter_by(
            action="enrollment.tuition_status_changed",
            target_user_id=student.id,
        )
        .one()
    )
    assert tuition_audit.actor_id is not None
    assert tuition_audit.created_at is not None
    assert tuition_audit.details == f"class_id={class_id};before=unpaid;after=paid"

    test_db.connection().exec_driver_sql(
        """
        CREATE TRIGGER reject_tuition_audit
        BEFORE INSERT ON audit_logs
        WHEN NEW.action = 'enrollment.tuition_status_changed'
        BEGIN
            SELECT RAISE(ABORT, 'simulated audit persistence failure');
        END;
        """
    )
    with pytest.raises(IntegrityError):
        client.patch(
            f"/api/classes/{class_id}/students/{student.id}/status",
            headers={"Authorization": f"Bearer {admin_token}"},
            json={"tuition_status": "unpaid"},
        )
    test_db.rollback()
    persisted_enrollment = test_db.scalar(
        select(ClassEnrollment).where(
            ClassEnrollment.class_id == class_id,
            ClassEnrollment.student_id == student.id,
        )
    )
    assert persisted_enrollment.tuition_status == "paid"
    assert (
        test_db.query(AuditLog)
        .filter_by(
            action="enrollment.tuition_status_changed",
            target_user_id=student.id,
        )
        .count()
        == 1
    )

    teacher_token = client.post(
        "/api/auth/login",
        json={"email": teacher.email, "password": "CorrectPassword123!"},
    ).json()["access_token"]
    denied_tuition = client.patch(
        f"/api/classes/{class_id}/students/{student.id}/status",
        headers={"Authorization": f"Bearer {teacher_token}"},
        json={"tuition_status": "unpaid"},
    )
    assert denied_tuition.status_code == 403
    assert (
        test_db.query(AuditLog)
        .filter_by(
            action="enrollment.tuition_status_changed",
            target_user_id=student.id,
        )
        .count()
        == 1
    )
    status_response = client.patch(
        f"/api/classes/{class_id}/students/{student.id}/status",
        headers={"Authorization": f"Bearer {teacher_token}"},
        json={"attendance_status": "present"},
    )
    assert status_response.status_code == 200
    assert status_response.json()["attendance_status"] == "present"

    response = client.patch(
        f"/api/classes/{class_id}",
        headers={"Authorization": f"Bearer {admin_token}"},
        json={"room": "B205", "start_time": "10:00", "end_time": "12:00"},
    )
    assert response.status_code == 200
    assert response.json()["room"] == "B205"

    student_token = client.post(
        "/api/auth/login",
        json={"email": student.email, "password": "CorrectPassword123!"},
    ).json()["access_token"]
    notifications = client.get(
        "/api/notifications",
        headers={"Authorization": f"Bearer {student_token}"},
    )
    assert notifications.status_code == 200
    assert notifications.json()["unread_count"] == 1
    notification_id = notifications.json()["items"][0]["id"]
    assert client.post(
        f"/api/notifications/{notification_id}/read",
        headers={"Authorization": f"Bearer {student_token}"},
    ).json()["is_read"] is True

    roster = client.get(
        f"/api/classes/{class_id}/students",
        headers={"Authorization": f"Bearer {admin_token}"},
    )
    assert roster.status_code == 200
    assert roster.json()[0]["tuition_status"] == "paid"
    assert roster.json()[0]["attendance_status"] == "present"
    exported = client.get(
        f"/api/classes/{class_id}/students/export",
        headers={"Authorization": f"Bearer {admin_token}"},
    )
    assert exported.status_code == 200
    assert exported.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


def test_all_roles_ta_attendance_and_class_assignments(client, test_db):
    admin_token = get_admin(client, test_db)
    teacher = create_user(
        test_db,
        username="teacher01",
        email="teacher@example.edu.vn",
        full_name="Instructor",
        role="teacher",
    )
    assistant = create_user(
        test_db,
        username="ta01",
        email="ta@example.edu.vn",
        full_name="Teaching Assistant",
        role="ta",
    )
    second_assistant = create_user(
        test_db,
        username="ta02",
        email="ta02@example.edu.vn",
        full_name="Second Teaching Assistant",
        role="ta",
    )
    student = create_user(
        test_db,
        username="student01",
        email="student@example.edu.vn",
        full_name="Student",
        role="student",
    )
    course = Course(
        code="CS301", name="Web Development", department="Computer Science", credits=3
    )
    test_db.add(course)
    test_db.commit()

    created = client.post(
        "/api/classes",
        headers={"Authorization": f"Bearer {admin_token}"},
        json={
            "code": "CS301-01",
            "name": "Web Development 01",
            "course_id": course.id,
            "teacher_id": teacher.id,
            "assistant_ids": [assistant.id, second_assistant.id],
            "room": "C301",
            "term": "2026-1",
            "day_of_week": 2,
            "start_time": "09:00",
            "end_time": "11:00",
            "capacity": 30,
        },
    )
    assert created.status_code == 201
    class_id = created.json()["id"]
    assert created.json()["assistant_name"] == assistant.full_name
    assert {member["id"] for member in created.json()["assistants"]} == {
        assistant.id,
        second_assistant.id,
    }
    activated = client.patch(
        f"/api/classes/{class_id}",
        headers={"Authorization": f"Bearer {admin_token}"},
        json={"status": "active"},
    )
    assert activated.status_code == 200
    assert client.post(
        f"/api/classes/{class_id}/students",
        headers={"Authorization": f"Bearer {admin_token}"},
        json={"student_ids": [student.id]},
    ).status_code == 201

    ta_token = client.post(
        "/api/auth/login",
        json={"email": assistant.email, "password": "CorrectPassword123!"},
    ).json()["access_token"]
    ta_headers = {"Authorization": f"Bearer {ta_token}"}
    assert [item["id"] for item in client.get("/api/classes", headers=ta_headers).json()] == [
        class_id
    ]
    second_ta_token = client.post(
        "/api/auth/login",
        json={"email": second_assistant.email, "password": "CorrectPassword123!"},
    ).json()["access_token"]
    assert [item["id"] for item in client.get(
        "/api/classes", headers={"Authorization": f"Bearer {second_ta_token}"}
    ).json()] == [class_id]
    attendance = client.patch(
        f"/api/classes/{class_id}/students/{student.id}/status",
        headers=ta_headers,
        json={"attendance_status": "present"},
    )
    assert attendance.status_code == 200
    assert attendance.json()["attendance_status"] == "present"

    assignment = client.post(
        f"/api/classes/{class_id}/assignments",
        headers=ta_headers,
        json={
            "title": "Build a small website",
            "description": "Submit the source code.",
            "due_at": "2026-12-01T23:59:00",
        },
    )
    assert assignment.status_code == 201
    assert assignment.json()["class_id"] == class_id
    secondary_ta_assignments = client.get(
        f"/api/assignments?class_id={class_id}&page=1&page_size=20",
        headers={"Authorization": f"Bearer {second_ta_token}"},
    )
    assert secondary_ta_assignments.status_code == 200
    assert [item["title"] for item in secondary_ta_assignments.json()["items"]] == [
        "Build a small website"
    ]
    assignment_notifications = (
        test_db.query(Notification).filter_by(user_id=student.id).all()
    )
    assert len(assignment_notifications) == 1
    assert assignment_notifications[0].class_id == class_id
    assert assignment_notifications[0].title == "Bài tập mới"
    assert assignment_notifications[0].message == (
        "Build a small website đã được giao cho lớp CS301-01."
    )
    past_deadline = client.post(
        f"/api/classes/{class_id}/assignments",
        headers=ta_headers,
        json={
            "title": "Assignment with past deadline",
            "due_at": "2000-01-01T00:00:00",
        },
    )
    assert past_deadline.status_code == 422
    assert (
        test_db.query(Assignment).filter_by(class_id=class_id).count() == 1
    )
    assert test_db.query(Notification).filter_by(user_id=student.id).count() == 1

    test_db.connection().exec_driver_sql(
        """
        CREATE TRIGGER reject_assignment_notification
        BEFORE INSERT ON notifications
        WHEN NEW.title = 'Bài tập mới'
        BEGIN
            SELECT RAISE(ABORT, 'simulated assignment notification failure');
        END;
        """
    )
    with pytest.raises(IntegrityError):
        client.post(
            f"/api/classes/{class_id}/assignments",
            headers=ta_headers,
            json={"title": "Assignment notification rollback"},
        )
    test_db.rollback()
    assert test_db.query(Assignment).filter_by(class_id=class_id).count() == 1
    assert test_db.query(Notification).filter_by(user_id=student.id).count() == 1

    student_token = client.post(
        "/api/auth/login",
        json={"email": student.email, "password": "CorrectPassword123!"},
    ).json()["access_token"]
    student_headers = {"Authorization": f"Bearer {student_token}"}
    visible_assignments = client.get("/api/assignments", headers=student_headers)
    assert visible_assignments.status_code == 200
    assert visible_assignments.json()[0]["title"] == "Build a small website"
    student_notifications = client.get("/api/notifications", headers=student_headers)
    assert student_notifications.status_code == 200
    assert student_notifications.json()["items"][0]["class_id"] == class_id
    assert client.post(
        f"/api/classes/{class_id}/assignments",
        headers=student_headers,
        json={"title": "Unauthorized"},
    ).status_code == 403

    roles = client.get("/api/roles", headers={"Authorization": f"Bearer {admin_token}"})
    assert roles.status_code == 200
    assert {role["code"] for role in roles.json()} == {
        "admin",
        "manager",
        "admissions",
        "accountant",
        "teacher",
        "ta",
        "student",
        "guest",
    }
    completed = client.patch(
        f"/api/classes/{class_id}",
        headers={"Authorization": f"Bearer {admin_token}"},
        json={"status": "completed"},
    )
    assert completed.status_code == 200
    invalid_transition = client.patch(
        f"/api/classes/{class_id}",
        headers={"Authorization": f"Bearer {admin_token}"},
        json={"status": "active"},
    )
    assert invalid_transition.status_code == 409


def test_notification_cursor_pagination_and_mark_all_are_user_scoped(client, test_db):
    admin_token = get_admin(client, test_db)
    admin = test_db.query(User).filter_by(username="admin01").one()
    other = create_user(
        test_db,
        username="student_notifications",
        email="student.notifications@example.edu.vn",
        full_name="Notification Student",
        role="student",
    )
    now = datetime.now(timezone.utc).replace(tzinfo=None)
    own_items = [
        Notification(
            user_id=admin.id,
            title=f"Notice {index}",
            message=f"Message {index}",
            created_at=now - timedelta(minutes=index),
        )
        for index in range(3)
    ]
    other_item = Notification(
        user_id=other.id,
        title="Private notice",
        message="Only for the student",
        created_at=now,
    )
    test_db.add_all([*own_items, other_item])
    test_db.commit()

    first_page = client.get("/api/notifications?limit=2", headers={"Authorization": f"Bearer {admin_token}"})
    assert first_page.status_code == 200
    first = first_page.json()
    assert first["unread_count"] == 3
    assert first["has_more"] is True
    assert [item["title"] for item in first["items"]] == ["Notice 0", "Notice 1"]

    cursor = first["next_cursor"]
    second_page = client.get(
        "/api/notifications",
        headers={"Authorization": f"Bearer {admin_token}"},
        params={
            "limit": 2,
            "cursor_created_at": cursor["created_at"],
            "cursor_id": cursor["id"],
        },
    )
    assert second_page.status_code == 200
    second = second_page.json()
    assert second["has_more"] is False
    assert [item["title"] for item in second["items"]] == ["Notice 2"]
    assert not ({item["id"] for item in first["items"]} & {item["id"] for item in second["items"]})

    marked = client.post(
        "/api/notifications/read-all",
        headers={"Authorization": f"Bearer {admin_token}"},
    )
    assert marked.status_code == 200
    assert marked.json() == {"updated_count": 3, "unread_count": 0}
    assert test_db.query(Notification).filter_by(user_id=other.id, read_at=None).count() == 1

    invalid_cursor = client.get(
        "/api/notifications?cursor_id=1",
        headers={"Authorization": f"Bearer {admin_token}"},
    )
    assert invalid_cursor.status_code == 422


def test_program_fee_schedules_are_decimal_versioned_and_non_overlapping(client, test_db):
    admin_token = get_admin(client, test_db)
    admin_headers = {"Authorization": f"Bearer {admin_token}"}
    accountant = create_user(
        test_db,
        username="accountant_fee_schedule",
        email="accountant.fee.schedule@example.edu.vn",
        full_name="Fee Schedule Accountant",
        role="accountant",
    )
    program = StudyProgram(
        code="FEE-PROGRAM",
        name="Fee Schedule Program",
        total_credits=60,
        duration_semesters=4,
    )
    test_db.add(program)
    test_db.commit()

    path = f"/api/study-programs/{program.id}/fee-schedules"
    created = client.post(
        path,
        headers=admin_headers,
        json={
            "amount_vnd": "12500000",
            "valid_from": "2026-01-01",
            "valid_to": "2026-06-30",
        },
    )
    assert created.status_code == 201
    assert created.json()["amount_vnd"] == "12500000"
    assert created.json()["currency"] == "VND"
    assert created.json()["version"] == 1

    second_version = client.post(
        path,
        headers=admin_headers,
        json={
            "amount_vnd": "13000000",
            "valid_from": "2026-07-01",
        },
    )
    assert second_version.status_code == 201
    assert second_version.json()["version"] == 2

    overlap = client.post(
        path,
        headers=admin_headers,
        json={
            "amount_vnd": "14000000",
            "valid_from": "2026-06-30",
        },
    )
    assert overlap.status_code == 409
    invalid_range = client.post(
        path,
        headers=admin_headers,
        json={
            "amount_vnd": "14000000",
            "valid_from": "2026-09-01",
            "valid_to": "2026-08-31",
        },
    )
    assert invalid_range.status_code == 422
    fractional_vnd = client.post(
        path,
        headers=admin_headers,
        json={
            "amount_vnd": "100.5",
            "valid_from": "2027-01-01",
        },
    )
    assert fractional_vnd.status_code == 422

    effective = client.get(
        path + "?effective_on=2026-04-01", headers=admin_headers
    )
    assert effective.status_code == 200
    assert [item["id"] for item in effective.json()] == [created.json()["id"]]
    assert test_db.query(ProgramFeeSchedule).count() == 2

    accountant_token = client.post(
        "/api/auth/login",
        json={"email": accountant.email, "password": "CorrectPassword123!"},
    ).json()["access_token"]
    accountant_headers = {"Authorization": f"Bearer {accountant_token}"}
    assert client.get(path, headers=accountant_headers).status_code == 200
    accountant_program = StudyProgram(
        code="FEE-ACCOUNTANT", name="Accountant Fee Program", total_credits=30
    )
    test_db.add(accountant_program)
    test_db.commit()
    accountant_created = client.post(
        f"/api/study-programs/{accountant_program.id}/fee-schedules",
        headers=accountant_headers,
        json={"amount_vnd": "1000000", "valid_from": "2028-01-01"},
    )
    assert accountant_created.status_code == 201
    assert accountant_created.json()["created_by"] == accountant.id
    teacher = create_user(test_db, username="fee_schedule_teacher",
        email="fee.schedule.teacher@example.edu.vn", full_name="Fee Schedule Teacher", role="teacher")
    teacher_token = client.post("/api/auth/login", json={
        "email": teacher.email, "password": "CorrectPassword123!",
    }).json()["access_token"]
    teacher_headers = {"Authorization": f"Bearer {teacher_token}"}
    assert client.get(path, headers=teacher_headers).status_code == 403


def test_tuition_plan_batch_is_exact_and_idempotent(client, test_db):
    admin_token = get_admin(client, test_db)
    accountant = create_user(test_db, username="accountant_tuition_batch",
        email="accountant.tuition.batch@example.edu.vn", full_name="Tuition Batch Accountant", role="accountant")
    accountant_token = client.post("/api/auth/login", json={
        "email": accountant.email, "password": "CorrectPassword123!",
    }).json()["access_token"]
    accountant_headers = {"Authorization": f"Bearer {accountant_token}"}
    teacher = create_user(
        test_db,
        username="teacher_tuition_batch",
        email="teacher.tuition.batch@example.edu.vn",
        full_name="Tuition Teacher",
        role="teacher",
    )
    students = [
        create_user(
            test_db,
            username=f"tuition_student_{index}",
            email=f"tuition.student.{index}@example.edu.vn",
            full_name=f"Tuition Student {index}",
            role="student",
        )
        for index in range(2)
    ]
    course = Course(
        code="TUITION101", name="Tuition Course", department="IT", credits=3
    )
    program = StudyProgram(
        code="TUITION-PROGRAM",
        name="Tuition Program",
        total_credits=60,
        duration_semesters=4,
    )
    program.courses.append(course)
    test_db.add_all([course, program])
    test_db.commit()
    schedule = ProgramFeeSchedule(
        program_id=program.id,
        version=1,
        amount_vnd=12_500_000,
        valid_from=date(2026, 1, 1),
        valid_to=date(2026, 12, 31),
    )
    training_class = TrainingClass(
        code="TUITION101-01",
        name="Tuition Class",
        course_id=course.id,
        teacher_id=teacher.id,
        room="Room 1",
        term="2026-Fall",
        day_of_week=0,
        start_time=time(9, 0),
        end_time=time(11, 0),
        capacity=20,
        status="active",
    )
    test_db.add_all([schedule, training_class])
    test_db.flush()
    active_enrollment = ClassEnrollment(
        class_id=training_class.id,
        student_id=students[0].id,
        enrollment_status="active",
    )
    withdrawn_enrollment = ClassEnrollment(
        class_id=training_class.id,
        student_id=students[1].id,
        enrollment_status="withdrawn",
        effective_date=date(2026, 1, 1),
    )
    test_db.add_all([active_enrollment, withdrawn_enrollment])
    test_db.commit()

    payload = {
        "class_id": training_class.id,
        "fee_schedule_id": schedule.id,
        "applied_on": "2026-06-01",
        "installments": [
            {"amount_vnd": "5000000", "due_date": "2026-06-15"},
            {"amount_vnd": "7500000", "due_date": "2026-08-15"},
        ],
    }
    path = "/api/tuition-plans/batch"
    mismatch = {**payload, "installments": [
        {"amount_vnd": "5000000", "due_date": "2026-06-15"}
    ]}
    assert client.post(path, headers=accountant_headers, json=mismatch).status_code == 422
    assert test_db.query(TuitionPlan).count() == 0

    created = client.post(path, headers=accountant_headers, json=payload)
    assert created.status_code == 201
    assert created.json()["created_count"] == 1
    assert created.json()["skipped_count"] == 0
    assert created.json()["plans"][0]["enrollment_id"] == active_enrollment.id
    assert test_db.query(TuitionPlan).count() == 1
    assert test_db.query(TuitionInstallment).count() == 2
    retry = client.post(path, headers=accountant_headers, json=payload)
    assert retry.status_code == 201
    assert retry.json()["created_count"] == 0
    assert retry.json()["skipped_count"] == 1
    assert test_db.query(TuitionPlan).count() == 1
    assert test_db.query(TuitionInstallment).count() == 2
    changed_retry = {
        **payload,
        "installments": [
            {"amount_vnd": "6000000", "due_date": "2026-06-15"},
            {"amount_vnd": "6500000", "due_date": "2026-08-15"},
        ],
    }
    assert client.post(
        path, headers=accountant_headers, json=changed_retry
    ).status_code == 409
    assert test_db.query(TuitionInstallment).count() == 2
    teacher_token = client.post("/api/auth/login", json={
        "email": teacher.email, "password": "CorrectPassword123!",
    }).json()["access_token"]
    assert client.post(path, headers={"Authorization": f"Bearer {teacher_token}"},
        json=payload).status_code == 403


def test_tuition_payments_are_partial_idempotent_and_reversible(client, test_db):
    admin_token = get_admin(client, test_db)
    accountant = create_user(
        test_db,
        username="accountant_payment",
        email="accountant.payment@example.edu.vn",
        full_name="Payment Accountant",
        role="accountant",
    )
    teacher = create_user(
        test_db,
        username="teacher_payment",
        email="teacher.payment@example.edu.vn",
        full_name="Payment Teacher",
        role="teacher",
    )
    student = create_user(
        test_db,
        username="student_payment",
        email="student.payment@example.edu.vn",
        full_name="Payment Student",
        role="student",
        student_code="SV-PAY-01",
    )
    other_student = create_user(
        test_db,
        username="student_other_payment",
        email="student.other.payment@example.edu.vn",
        full_name="Other Payment Student",
        role="student",
        student_code="SV-PAY-02",
    )
    course = Course(code="PAY101", name="Payment Course", department="IT", credits=3)
    program = StudyProgram(
        code="PAY-PROGRAM", name="Payment Program", total_credits=60
    )
    program.courses.append(course)
    test_db.add_all([course, program])
    test_db.commit()
    fee_schedule = ProgramFeeSchedule(
        program_id=program.id,
        version=1,
        amount_vnd=12_500_000,
        valid_from=date(2026, 1, 1),
        valid_to=date(2026, 12, 31),
    )
    training_class = TrainingClass(
        code="PAY101-01",
        name="Payment Class",
        course_id=course.id,
        program_id=program.id,
        teacher_id=teacher.id,
        room="Room 1",
        term="2026-Fall",
        day_of_week=1,
        start_time=time(9, 0),
        end_time=time(11, 0),
        capacity=20,
        status="active",
    )
    test_db.add_all([fee_schedule, training_class])
    test_db.flush()
    enrollment = ClassEnrollment(
        class_id=training_class.id,
        student_id=student.id,
        enrollment_status="active",
    )
    test_db.add(enrollment)
    test_db.flush()
    plan = TuitionPlan(
        enrollment_id=enrollment.id,
        fee_schedule_id=fee_schedule.id,
        total_due_vnd=12_500_000,
        applied_on=date(2026, 6, 1),
    )
    test_db.add(plan)
    test_db.flush()
    installment = TuitionInstallment(
        plan_id=plan.id,
        installment_number=1,
        amount_due_vnd=12_500_000,
        due_date=date(2026, 12, 15),
    )
    test_db.add(installment)
    test_db.commit()

    accountant_token = client.post(
        "/api/auth/login",
        json={"email": accountant.email, "password": "CorrectPassword123!"},
    ).json()["access_token"]
    accountant_headers = {"Authorization": f"Bearer {accountant_token}"}
    manager = create_user(
        test_db,
        username="manager_payment_discount",
        email="manager.payment.discount@example.edu.vn",
        full_name="Training Manager",
        role="manager",
    )
    manager_token = client.post(
        "/api/auth/login",
        json={"email": manager.email, "password": "CorrectPassword123!"},
    ).json()["access_token"]
    manager_headers = {"Authorization": f"Bearer {manager_token}"}
    discount_path = f"/api/tuition-plans/{plan.id}/discounts"
    assert client.post(
        discount_path,
        headers=accountant_headers,
        json={"kind": "percent", "value": "101", "reason": "Invalid range"},
    ).status_code == 422
    pending_discount_response = client.post(
        discount_path,
        headers=accountant_headers,
        json={"kind": "percent", "value": "40", "reason": "Scholarship approval"},
    )
    assert pending_discount_response.status_code == 201
    pending_discount = pending_discount_response.json()
    assert pending_discount["amount_vnd"] == "5000000"
    assert pending_discount["status"] == "pending"
    discount_approve_path = f"/api/tuition-discounts/{pending_discount['id']}/approve"
    assert client.post(discount_approve_path, headers=accountant_headers).status_code == 403
    teacher_token = client.post(
        "/api/auth/login",
        json={"email": teacher.email, "password": "CorrectPassword123!"},
    ).json()["access_token"]
    teacher_headers = {"Authorization": f"Bearer {teacher_token}"}
    assert client.post(discount_approve_path, headers=teacher_headers).status_code == 403
    approved_discount = client.post(discount_approve_path, headers=manager_headers)
    assert approved_discount.status_code == 200
    assert approved_discount.json()["status"] == "approved"
    assert approved_discount.json()["reviewed_by"] == manager.id
    assert client.post(discount_approve_path, headers=manager_headers).status_code == 200
    assert test_db.query(TuitionDiscount).filter_by(plan_id=plan.id, status="approved").count() == 1
    automatic_fixed_discount = client.post(
        discount_path,
        headers=accountant_headers,
        json={"kind": "fixed", "value": "500000", "reason": "Fixed early payment award"},
    )
    assert automatic_fixed_discount.status_code == 201
    assert automatic_fixed_discount.json()["status"] == "approved"
    assert automatic_fixed_discount.json()["requested_by"] == accountant.id
    rejected_discount = client.post(
        discount_path,
        headers=accountant_headers,
        json={"kind": "fixed", "value": "4000000", "reason": "Rejected scholarship"},
    )
    assert rejected_discount.status_code == 201
    reject_path = f"/api/tuition-discounts/{rejected_discount.json()['id']}/reject"
    assert client.post(reject_path, headers=manager_headers).json()["status"] == "rejected"
    assert client.post(
        f"/api/tuition-discounts/{rejected_discount.json()['id']}/approve",
        headers=manager_headers,
    ).status_code == 409
    assert client.post(
        discount_path,
        headers=accountant_headers,
        json={"kind": "fixed", "value": "8000000", "reason": "Exceeds open balance"},
    ).status_code == 409
    payment_path = f"/api/tuition-installments/{installment.id}/payments"
    payload = {
        "amount_vnd": "5000000",
        "method": "cash",
        "paid_on": "2026-10-26",
        "idempotency_key": "cash-counter-2026-0001",
        "note": "Thu đợt một",
    }
    assert client.post(payment_path, json=payload).status_code == 401
    invalid_amount = client.post(
        payment_path,
        headers=accountant_headers,
        json={**payload, "amount_vnd": "0", "idempotency_key": "invalid-amount-0001"},
    )
    assert invalid_amount.status_code == 422
    assert client.post(
        "/api/tuition-installments/999999/payments",
        headers=accountant_headers,
        json=payload,
    ).status_code == 404
    test_db.execute(text(
        """CREATE TRIGGER reject_tuition_payment_audit
        BEFORE INSERT ON audit_logs
        WHEN NEW.action = 'tuition.payment_recorded'
        BEGIN SELECT RAISE(ABORT, 'reject tuition audit'); END"""
    ))
    test_db.commit()
    atomicity_failure = client.post(
        payment_path,
        headers=accountant_headers,
        json={**payload, "idempotency_key": "atomicity-check-0001"},
    )
    assert atomicity_failure.status_code == 409
    assert test_db.query(TuitionPayment).count() == 0
    test_db.execute(text("DROP TRIGGER reject_tuition_payment_audit"))
    test_db.commit()
    first = client.post(payment_path, headers=accountant_headers, json=payload)
    assert first.status_code == 201
    assert first.json()["receipt_code"].startswith("ICTU-")
    assert first.json()["collected_by"] == accountant.id
    assert first.json()["amount_vnd"] == "5000000"
    assert first.json()["installment_remaining_vnd"] == "2000000"
    assert first.json()["installment_status"] == "partial"
    assert first.json()["collected_by"] == accountant.id

    replay = client.post(payment_path, headers=accountant_headers, json=payload)
    assert replay.status_code == 200
    assert replay.json()["id"] == first.json()["id"]
    assert test_db.query(TuitionPayment).count() == 1
    key_collision = client.post(
        payment_path,
        headers=accountant_headers,
        json={**payload, "amount_vnd": "4000000"},
    )
    assert key_collision.status_code == 409
    overpayment = client.post(
        payment_path,
        headers=accountant_headers,
        json={
            **payload,
            "amount_vnd": "7500001",
            "idempotency_key": "cash-counter-2026-0002",
        },
    )
    assert overpayment.status_code == 409
    assert test_db.query(TuitionPayment).count() == 1

    teacher_token = client.post(
        "/api/auth/login",
        json={"email": teacher.email, "password": "CorrectPassword123!"},
    ).json()["access_token"]
    assert client.post(
        payment_path,
        headers={"Authorization": f"Bearer {teacher_token}"},
        json=payload,
    ).status_code == 403

    second = client.post(
        payment_path,
        headers=accountant_headers,
        json={
            **payload,
            "amount_vnd": "2000000",
            "method": "bank_transfer",
            "idempotency_key": "bank-transfer-2026-01",
            "note": None,
        },
    )
    assert second.status_code == 201
    assert second.json()["receipt_code"] != first.json()["receipt_code"]
    assert second.json()["installment_remaining_vnd"] == "0"
    assert second.json()["installment_status"] == "paid"

    pdf = client.get(
        f"/api/tuition-payments/{first.json()['id']}/receipt.pdf",
        headers=accountant_headers,
    )
    assert pdf.status_code == 200
    assert pdf.headers["content-type"] == "application/pdf"
    assert pdf.content.startswith(b"%PDF-1.4")
    assert first.json()["receipt_code"].encode("ascii") in pdf.content
    assert b"Payment Student" in pdf.content
    assert b"Amount: 5000000 VND" in pdf.content
    assert b"Paid on: 2026-10-26" in pdf.content

    student_token = client.post(
        "/api/auth/login",
        json={"email": student.email, "password": "CorrectPassword123!"},
    ).json()["access_token"]
    other_student_token = client.post(
        "/api/auth/login",
        json={"email": other_student.email, "password": "CorrectPassword123!"},
    ).json()["access_token"]
    own_tuition = client.get(
        "/api/students/me/tuition",
        headers={"Authorization": f"Bearer {student_token}"},
    )
    assert own_tuition.status_code == 200
    assert len(own_tuition.json()["plans"]) == 1
    assert own_tuition.json()["page"] == 1 and own_tuition.json()["total"] == 1
    assert own_tuition.json()["pages"] == 1 and own_tuition.json()["has_next"] is False
    assert own_tuition.json()["plans"][0]["total_due_vnd"] == "7000000"
    assert own_tuition.json()["plans"][0]["installments"][0]["status"] == "paid"
    payment_history = client.get("/api/tuition-payments?page=1&page_size=1&sort=paid_on_asc",
        headers=accountant_headers)
    assert payment_history.status_code == 200
    assert len(payment_history.json()["items"]) == 1
    assert payment_history.json()["total"] == 2
    assert payment_history.headers["x-has-next"] == "true"
    own_payment_history = client.get("/api/tuition-payments?page_size=100",
        headers={"Authorization": f"Bearer {student_token}"})
    assert own_payment_history.status_code == 200
    assert all(item["student_id"] == student.id for item in own_payment_history.json()["items"])
    other_tuition = client.get(
        "/api/students/me/tuition",
        headers={"Authorization": f"Bearer {other_student_token}"},
    )
    assert other_tuition.status_code == 200
    assert other_tuition.json()["plans"] == []
    other_payment_history = client.get("/api/tuition-payments",
        headers={"Authorization": f"Bearer {other_student_token}"})
    assert other_payment_history.status_code == 200
    assert other_payment_history.json()["items"] == []
    assert client.get("/api/tuition-payments?paid_from=2026-12-01&paid_to=2026-01-01",
        headers=accountant_headers).status_code == 422
    debts = client.get("/api/tuition-debts?status=paid", headers=accountant_headers)
    assert debts.status_code == 200
    assert debts.json()["total"] == 1
    assert debts.json()["total_due_vnd"] == "7000000"
    assert debts.json()["total_paid_vnd"] == "7000000"
    assert debts.json()["total_remaining_vnd"] == "0"
    debt_export = client.get("/api/tuition-debts/export.xlsx", headers=accountant_headers)
    assert debt_export.status_code == 200
    assert debt_export.content.startswith(b"PK")
    debt_sheet = load_workbook(BytesIO(debt_export.content), data_only=True).active
    assert debt_sheet.max_row == 3
    assert debt_sheet.cell(2, 1).value == "SV-PAY-01"
    assert debt_sheet.cell(2, 6).value == "7000000"
    assert debt_sheet.cell(2, 7).value == "7000000"
    assert debt_sheet.cell(2, 8).value == "0"
    assert debt_sheet.cell(3, 1).value == "TOTAL"
    assert float(debt_sheet.cell(3, 6).value) == 7000000
    assert float(debt_sheet.cell(3, 7).value) == 7000000
    assert client.get("/api/tuition-debts", headers={"Authorization": f"Bearer {student_token}"}).status_code == 403
    assert client.get(
        f"/api/tuition-payments/{first.json()['id']}",
        headers={"Authorization": f"Bearer {student_token}"},
    ).status_code == 200
    assert client.get(
        f"/api/tuition-payments/{first.json()['id']}",
        headers={"Authorization": f"Bearer {other_student_token}"},
    ).status_code == 404
    assert client.get(
        f"/api/tuition-payments/{first.json()['id']}/receipt.pdf",
        headers={"Authorization": f"Bearer {other_student_token}"},
    ).status_code == 404
    assert client.get(
        "/api/tuition-payments/999999",
        headers=accountant_headers,
    ).status_code == 404

    reverse_path = f"/api/tuition-payments/{first.json()['id']}/reverse"
    reverse_payload = {
        "idempotency_key": "reverse-payment-2026-0001",
        "reason": "Thu nhầm, hoàn lại",
    }
    reversal = client.post(
        reverse_path, headers=accountant_headers, json=reverse_payload
    )
    assert reversal.status_code == 201
    assert reversal.json()["amount_vnd"] == "5000000"
    assert reversal.json()["payment"]["reversed"] is True
    assert reversal.json()["payment"]["installment_remaining_vnd"] == "5000000"
    revenue = client.get("/api/tuition-revenue", headers=accountant_headers)
    assert revenue.status_code == 200
    assert revenue.json()["net_revenue_vnd"] == "2000000"
    assert revenue.json()["by_class"][0]["class_id"] == training_class.id
    assert revenue.json()["by_class"][0]["net_revenue_vnd"] == "2000000"
    assert revenue.json()["by_program"][0]["program_id"] == program.id
    assert revenue.json()["by_program"][0]["net_revenue_vnd"] == "2000000"
    assert revenue.json()["by_program"][0]["total_due_vnd"] == "7000000"
    assert revenue.json()["by_program"][0]["paid_vnd"] == "2000000"
    assert revenue.json()["by_program"][0]["remaining_vnd"] == "5000000"
    assert revenue.json()["payment_count"] == 2
    assert revenue.json()["reversal_count"] == 1
    date_range = client.get(
        "/api/tuition-revenue?from_date=2026-10-26&to_date=2026-10-26",
        headers=accountant_headers,
    )
    assert date_range.status_code == 200
    assert date_range.json()["from_date"] == "2026-10-26"
    assert date_range.json()["to_date"] == "2026-10-26"
    assert date_range.json()["net_revenue_vnd"] == "7000000"
    assert date_range.json()["monthly"] == [{"month": "2026-10", "net_revenue_vnd": "7000000"}]
    assert client.get(
        "/api/tuition-revenue?from_date=2026-10-27&to_date=2026-10-26",
        headers=accountant_headers,
    ).status_code == 422
    revenue_export = client.get("/api/tuition-revenue/export.xlsx", headers=accountant_headers)
    revenue_workbook = load_workbook(BytesIO(revenue_export.content), data_only=True)
    assert "By class" in revenue_workbook.sheetnames
    assert "By program" in revenue_workbook.sheetnames
    assert revenue_workbook["By class"].max_row == 2
    assert revenue_export.status_code == 200
    assert revenue_export.content.startswith(b"PK")
    revenue_sheet = load_workbook(BytesIO(revenue_export.content), data_only=True).active
    assert revenue_sheet.max_row == 3
    assert revenue_sheet.cell(2, 2).value == "2000000"
    assert revenue_sheet.cell(3, 2).value == "2000000"
    assert client.get("/api/tuition-revenue", headers={"Authorization": f"Bearer {student_token}"}).status_code == 403
    reversal_replay = client.post(
        reverse_path, headers=accountant_headers, json=reverse_payload
    )
    assert reversal_replay.status_code == 200
    duplicate_reverse = client.post(
        reverse_path,
        headers=accountant_headers,
        json={"idempotency_key": "reverse-payment-2026-0002", "reason": "Lý do khác"},
    )
    assert duplicate_reverse.status_code == 409
    assert test_db.query(TuitionPayment).count() == 2
    assert test_db.query(TuitionPaymentReversal).count() == 1
    tuition_audit = client.get("/api/audit-logs?entity_type=tuition", headers={"Authorization": f"Bearer {admin_token}"})
    assert tuition_audit.status_code == 200
    tuition_actions = {item["action"] for item in tuition_audit.json()["items"]}
    assert {
        "tuition.discount_requested",
        "tuition.discount_approved",
        "tuition.discount_rejected",
        "tuition.payment_recorded",
        "tuition.payment_reversed",
    } <= tuition_actions


def test_assignment_list_search_pagination_and_student_class_scope(client, test_db):
    admin_token = get_admin(client, test_db)
    admin = test_db.query(User).filter_by(username="admin01").one()
    teacher = create_user(
        test_db,
        username="teacher_assignment_list",
        email="teacher.assignment.list@example.edu.vn",
        full_name="Assignment List Teacher",
        role="teacher",
    )
    student = create_user(
        test_db,
        username="student_assignment_list",
        email="student.assignment.list@example.edu.vn",
        full_name="Assignment List Student",
        role="student",
    )

    def make_training_class(code):
        course = Course(
            code=f"{code}-COURSE",
            name=f"{code} Course",
            department="IT",
            credits=3,
            session_count=3,
        )
        test_db.add(course)
        test_db.flush()
        training_class = TrainingClass(
            code=code,
            name=f"{code} Class",
            course_id=course.id,
            teacher_id=teacher.id,
            room="Room 1",
            term="2026-Fall",
            day_of_week=0,
            start_time=time(9, 0),
            end_time=time(11, 0),
            capacity=20,
            status="active",
        )
        test_db.add(training_class)
        test_db.flush()
        return training_class

    visible_class = make_training_class("ASG-VISIBLE")
    hidden_class = make_training_class("ASG-HIDDEN")
    test_db.add(
        ClassEnrollment(class_id=visible_class.id, student_id=student.id)
    )
    test_db.add_all(
        [
            Assignment(
                class_id=visible_class.id,
                created_by=admin.id,
                title="Alpha intro",
                due_at=datetime(2026, 11, 2, 9, 0),
            ),
            Assignment(
                class_id=visible_class.id,
                created_by=admin.id,
                title="Alpha advanced",
                due_at=datetime(2026, 11, 1, 9, 0),
            ),
            Assignment(
                class_id=visible_class.id,
                created_by=admin.id,
                title="Alpha no deadline",
                due_at=None,
            ),
            Assignment(
                class_id=hidden_class.id,
                created_by=admin.id,
                title="Hidden class task",
                due_at=datetime(2026, 11, 3, 9, 0),
            ),
        ]
    )
    test_db.commit()

    headers = {"Authorization": f"Bearer {admin_token}"}
    legacy = client.get("/api/assignments", headers=headers)
    assert legacy.status_code == 200
    assert isinstance(legacy.json(), list)
    assert len(legacy.json()) == 4
    assert all(item["max_score"] is None for item in legacy.json())

    first_page = client.get(
        "/api/assignments?q=alpha&page=1&page_size=1", headers=headers
    )
    assert first_page.status_code == 200
    assert first_page.json()["total"] == 3
    assert first_page.json()["pages"] == 3
    assert [item["title"] for item in first_page.json()["items"]] == ["Alpha advanced"]

    second_page = client.get(
        "/api/assignments?q=alpha&page=2&page_size=1", headers=headers
    )
    assert second_page.status_code == 200
    assert [item["title"] for item in second_page.json()["items"]] == ["Alpha intro"]

    third_page = client.get(
        "/api/assignments?q=alpha&page=3&page_size=1", headers=headers
    )
    assert third_page.status_code == 200
    assert [item["title"] for item in third_page.json()["items"]] == [
        "Alpha no deadline"
    ]

    class_filtered = client.get(
        f"/api/assignments?class_id={visible_class.id}&page=1&page_size=20",
        headers=headers,
    )
    assert class_filtered.status_code == 200
    assert class_filtered.json()["total"] == 3
    assert {item["class_id"] for item in class_filtered.json()["items"]} == {
        visible_class.id
    }

    combined_filters = client.get(
        f"/api/assignments?q=advanced&class_id={visible_class.id}&page=1&page_size=20",
        headers=headers,
    )
    assert combined_filters.status_code == 200
    assert combined_filters.json()["total"] == 1
    assert [item["title"] for item in combined_filters.json()["items"]] == [
        "Alpha advanced"
    ]

    student_token = client.post(
        "/api/auth/login",
        json={"email": student.email, "password": "CorrectPassword123!"},
    ).json()["access_token"]
    student_list = client.get(
        "/api/assignments?page=1&page_size=20",
        headers={"Authorization": f"Bearer {student_token}"},
    )
    assert student_list.status_code == 200
    assert [item["title"] for item in student_list.json()["items"]] == [
        "Alpha advanced",
        "Alpha intro",
        "Alpha no deadline",
    ]
    student_items = student_list.json()["items"]
    assert all(item["submission_status"] == "not_submitted" for item in student_items)
    assert all(item["is_overdue"] is False for item in student_items)
    assert all(
        item["remaining_seconds"] is None
        for item in student_items
        if item["due_at"] is None
    )
    assert all(
        item["remaining_seconds"] > 0
        for item in student_items
        if item["due_at"] is not None
    )
    student_hidden_class = client.get(
        f"/api/assignments?class_id={hidden_class.id}&page=1&page_size=20",
        headers={"Authorization": f"Bearer {student_token}"},
    )
    assert student_hidden_class.status_code == 200
    assert student_hidden_class.json()["total"] == 0

    invalid_page_size = client.get(
        "/api/assignments?page=1&page_size=101", headers=headers
    )
    assert invalid_page_size.status_code == 422


def test_assignment_draft_publish_snapshots_roster_and_notifies_once(client, test_db, tmp_path, monkeypatch):
    admin_token = get_admin(client, test_db)
    teacher = create_user(
        test_db,
        username="teacher_draft_assignment",
        email="teacher.draft.assignment@example.edu.vn",
        full_name="Draft Assignment Teacher",
        role="teacher",
    )
    unassigned_ta = create_user(
        test_db,
        username="ta_unassigned_publish",
        email="ta.unassigned.publish@example.edu.vn",
        full_name="Unassigned TA",
        role="ta",
    )
    student = create_user(
        test_db,
        username="student_draft_assignment",
        email="student.draft.assignment@example.edu.vn",
        full_name="Draft Assignment Student",
        role="student",
    )
    unsubmitted_student = create_user(
        test_db,
        username="student_no_submission",
        email="student.no.submission@example.edu.vn",
        full_name="Student Without Submission",
        role="student",
    )
    outsider_student = create_user(
        test_db,
        username="student_outside_assignment_roster",
        email="student.outside.assignment.roster@example.edu.vn",
        full_name="Outside Assignment Roster",
        role="student",
    )
    course = Course(
        code="DRAFT101",
        name="Draft Assignment Course",
        department="IT",
        credits=3,
        session_count=3,
    )
    test_db.add(course)
    test_db.flush()
    training_class = TrainingClass(
        code="DRAFT101-01",
        name="Draft Assignment Class",
        course_id=course.id,
        teacher_id=teacher.id,
        room="Room 1",
        term="2026-Fall",
        day_of_week=0,
        start_time=time(9, 0),
        end_time=time(11, 0),
        capacity=20,
        status="active",
    )
    test_db.add(training_class)
    test_db.flush()
    student_class = ClassEnrollment(class_id=training_class.id, student_id=student.id)
    unsubmitted_class = ClassEnrollment(
        class_id=training_class.id, student_id=unsubmitted_student.id
    )
    session = AttendanceSession(
        class_id=training_class.id,
        session_date=date.today() + timedelta(days=7),
        start_time=time(9, 0),
        end_time=time(11, 0),
        teacher_id=teacher.id,
        room="Room 1",
    )
    test_db.add_all([student_class, unsubmitted_class, session])
    test_db.commit()

    teacher_token = client.post(
        "/api/auth/login",
        json={"email": teacher.email, "password": "CorrectPassword123!"},
    ).json()["access_token"]
    teacher_headers = {"Authorization": f"Bearer {teacher_token}"}
    student_token = client.post(
        "/api/auth/login",
        json={"email": student.email, "password": "CorrectPassword123!"},
    ).json()["access_token"]
    student_headers = {"Authorization": f"Bearer {student_token}"}
    outsider_token = client.post(
        "/api/auth/login",
        json={"email": outsider_student.email, "password": "CorrectPassword123!"},
    ).json()["access_token"]
    outsider_headers = {"Authorization": f"Bearer {outsider_token}"}

    mismatched_session = client.post(
        f"/api/classes/{training_class.id}/assignments",
        headers=teacher_headers,
        json={"title": "Wrong session", "status": "draft", "session_id": session.id + 100},
    )
    assert mismatched_session.status_code == 422

    for invalid_score in ("0", "100.001"):
        invalid_max_score = client.post(
            f"/api/classes/{training_class.id}/assignments",
            headers=teacher_headers,
            json={
                "title": "Invalid score",
                "status": "draft",
                "max_score": invalid_score,
            },
        )
        assert invalid_max_score.status_code == 422

    draft = client.post(
        f"/api/classes/{training_class.id}/assignments",
        headers=teacher_headers,
        json={
            "title": "Draft task",
            "description": "Draft description",
            "due_at": (datetime.now(timezone.utc) + timedelta(days=2)).isoformat(),
            "max_score": "100.00",
            "status": "draft",
            "session_id": session.id,
        },
    )
    assert draft.status_code == 201
    assert draft.json()["status"] == "draft"
    assert draft.json()["max_score"] == "100.00"
    assert draft.json()["session_id"] == session.id
    assert client.get("/api/assignments", headers=student_headers).json() == []
    assert test_db.query(Notification).filter_by(user_id=student.id).count() == 0
    monkeypatch.setattr(settings, "PRIVATE_FILE_STORAGE_DIR", str(tmp_path / "private"))
    instruction_file = client.post(
        f"/api/assignments/{draft.json()['id']}/files",
        headers={**teacher_headers, "X-File-Name": "instructions.pdf", "Content-Type": "application/pdf"},
        content=b"%PDF-1.7\nassignment instructions",
    )
    assert instruction_file.status_code == 201
    assert "storage_key" not in instruction_file.json()
    assert client.get(
        f"/api/assignments/{draft.json()['id']}/files", headers=student_headers
    ).status_code == 404

    invalid_rubric = client.post(
        "/api/rubric-templates",
        headers=teacher_headers,
        json={"name": "Invalid", "criteria": [
            {"title": "A", "weight": "40"}, {"title": "B", "weight": "50"}
        ]},
    )
    assert invalid_rubric.status_code == 422
    test_db.connection().exec_driver_sql(
        """
        CREATE TRIGGER reject_rubric_audit
        BEFORE INSERT ON audit_logs
        WHEN NEW.action = 'rubric.template_created'
        BEGIN
            SELECT RAISE(ABORT, 'simulated rubric audit failure');
        END;
        """
    )
    with pytest.raises(IntegrityError):
        client.post(
            "/api/rubric-templates",
            headers=teacher_headers,
            json={"name": "Must roll back", "criteria": [{"title": "A", "weight": "100"}]},
        )
    test_db.rollback()
    assert test_db.query(RubricTemplate).filter_by(name="Must roll back").count() == 0
    test_db.connection().exec_driver_sql("DROP TRIGGER reject_rubric_audit")
    test_db.commit()
    rubric_response = client.post(
        "/api/rubric-templates",
        headers=teacher_headers,
        json={
            "name": "Assignment rubric",
            "criteria": [
                {"title": "Correctness", "weight": "50.00"},
                {"title": "Explanation", "weight": "50.00"},
            ],
        },
    )
    assert rubric_response.status_code == 201
    rubric = rubric_response.json()
    manager = create_user(
        test_db,
        username="manager_rubric_scope",
        email="manager.rubric.scope@example.edu.vn",
        full_name="Rubric Scope Manager",
        role="manager",
    )
    manager_token = client.post(
        "/api/auth/login",
        json={"email": manager.email, "password": "CorrectPassword123!"},
    ).json()["access_token"]
    manager_headers = {"Authorization": f"Bearer {manager_token}"}
    assert client.get("/api/rubric-templates", headers=manager_headers).json() == []
    assert client.put(
        f"/api/assignments/{draft.json()['id']}/rubric?template_id={rubric['id']}",
        headers=manager_headers,
    ).status_code == 404
    attached_rubric = client.put(
        f"/api/assignments/{draft.json()['id']}/rubric?template_id={rubric['id']}",
        headers=teacher_headers,
    )
    assert attached_rubric.status_code == 200
    assert attached_rubric.json()["template_version"] == 1
    changed_template = client.put(
        f"/api/rubric-templates/{rubric['id']}",
        headers=teacher_headers,
        json={
            "name": "Changed rubric",
            "criteria": [
                {"title": "New criterion A", "weight": "50"},
                {"title": "New criterion B", "weight": "50"},
            ],
        },
    )
    assert changed_template.status_code == 200
    rubric_audit = client.get(
        "/api/audit-logs?entity_type=rubric",
        headers={"Authorization": f"Bearer {admin_token}"},
    )
    assert [item["action"] for item in rubric_audit.json()["items"]] == [
        "rubric.template_updated",
        "rubric.template_created",
    ]
    assert "criteria_count=2" in rubric_audit.json()["items"][0]["details"]
    immutable_snapshot = client.get(
        f"/api/assignments/{draft.json()['id']}/rubric", headers=teacher_headers
    )
    assert immutable_snapshot.json()["criteria"][0]["title"] == "Correctness"
    unassigned_token = client.post(
        "/api/auth/login",
        json={"email": unassigned_ta.email, "password": "CorrectPassword123!"},
    ).json()["access_token"]
    assert client.get("/api/rubric-templates", headers={
        "Authorization": f"Bearer {unassigned_token}"
    }).json() == []
    denied_snapshot = client.put(
        f"/api/assignments/{draft.json()['id']}/rubric?template_id={rubric['id']}",
        headers={"Authorization": f"Bearer {unassigned_token}"},
    )
    assert denied_snapshot.status_code == 403
    denied_publish = client.post(
        f"/api/assignments/{draft.json()['id']}/publish",
        headers={"Authorization": f"Bearer {unassigned_token}"},
    )
    assert denied_publish.status_code == 403
    assert client.get(
        f"/api/assignments/{draft.json()['id']}/grading-queue",
        headers={"Authorization": f"Bearer {unassigned_token}"},
    ).status_code == 404
    assert client.get(
        f"/api/assignments/{draft.json()['id']}/submissions/progress",
        headers={"Authorization": f"Bearer {unassigned_token}"},
    ).status_code == 403

    published = client.post(
        f"/api/assignments/{draft.json()['id']}/publish",
        headers=teacher_headers,
    )
    assert published.status_code == 200
    assert published.json()["status"] == "published"
    assert published.json()["max_score"] == "100.00"
    assert client.get(
        f"/api/assignments/{draft.json()['id']}/rubric", headers=student_headers
    ).status_code == 403
    student_assignments = client.get(
        "/api/assignments", headers=student_headers
    ).json()
    assert [item["title"] for item in student_assignments] == ["Draft task"]
    assert student_assignments[0]["max_score"] == "100.00"
    assert student_assignments[0]["attachments"][0]["filename"] == "instructions.pdf"
    assert client.get(
        student_assignments[0]["attachments"][0]["download_url"], headers=student_headers
    ).content.startswith(b"%PDF-")
    assert client.get(
        f"/api/assignments/{draft.json()['id']}/files", headers=outsider_headers
    ).status_code == 404
    assert client.get(
        instruction_file.json()["download_url"], headers=outsider_headers
    ).status_code == 404
    assert test_db.query(Notification).filter_by(user_id=student.id).count() == 1
    assert test_db.query(AssignmentRecipient).filter_by(
        assignment_id=draft.json()["id"], student_id=student.id
    ).count() == 1

    submission_path = f"/api/assignments/{draft.json()['id']}/submissions"
    test_db.connection().exec_driver_sql(
        """
        CREATE TRIGGER reject_submission_audit
        BEFORE INSERT ON audit_logs
        WHEN NEW.action = 'assignment.submitted'
        BEGIN
            SELECT RAISE(ABORT, 'simulated submission audit failure');
        END;
        """
    )
    failed_submission = client.post(
        submission_path,
        headers=student_headers,
        json={"repository_url": "https://github.com/ictu/rollback-project"},
    )
    assert failed_submission.status_code == 409
    test_db.rollback()
    assert test_db.query(AssignmentSubmission).filter_by(
        assignment_id=draft.json()["id"], student_id=student.id
    ).count() == 0
    test_db.connection().exec_driver_sql("DROP TRIGGER reject_submission_audit")
    test_db.commit()

    first_submission = client.post(
        submission_path,
        headers=student_headers,
        json={"repository_url": "https://github.com/ictu/student-project"},
    )
    assert first_submission.status_code == 201
    assert first_submission.json()["version_number"] == 1
    assert first_submission.json()["repository_url"] == "https://github.com/ictu/student-project"
    submission_audit = client.get(
        "/api/audit-logs?entity_type=submission",
        headers={"Authorization": f"Bearer {admin_token}"},
    )
    assert submission_audit.status_code == 200
    assert submission_audit.json()["total"] == 1
    assert submission_audit.json()["items"][0]["action"] == "assignment.submitted"
    assert "repository_url" not in (submission_audit.json()["items"][0]["details"] or "")
    insecure_submission = client.post(
        submission_path,
        headers=student_headers,
        json={"repository_url": "http://github.com/ictu/student-project"},
    )
    assert insecure_submission.status_code == 422
    second_submission = client.post(
        submission_path,
        headers=student_headers,
        json={"repository_url": "https://github.com/ictu/student-project-v2"},
    )
    assert second_submission.status_code == 201
    assert second_submission.json()["version_number"] == 2
    student_assignment_list = client.get(
        "/api/assignments?page=1&page_size=20", headers=student_headers
    )
    assert student_assignment_list.status_code == 200
    listed_assignment = next(
        item
        for item in student_assignment_list.json()["items"]
        if item["id"] == draft.json()["id"]
    )
    assert listed_assignment["submission_status"] == "submitted"
    assert listed_assignment["submission_version"] == 2
    grading_queue_response = client.get(
        f"/api/assignments/{draft.json()['id']}/grading-queue",
        headers=teacher_headers,
    )
    assert grading_queue_response.status_code == 200
    assert grading_queue_response.json()["total_ungraded"] == 1
    assert grading_queue_response.json()["items"][0]["submission_id"] == second_submission.json()["id"]
    assert grading_queue_response.json()["next_submission_id"] == second_submission.json()["id"]
    assert grading_queue_response.json()["items"][0]["repository_url"] == (
        "https://github.com/ictu/student-project-v2"
    )
    assert test_db.query(AssignmentSubmission).filter_by(
        assignment_id=draft.json()["id"], student_id=student.id
    ).count() == 2
    assert [item["version_number"] for item in client.get(
        submission_path, headers=student_headers
    ).json()] == [1, 2]
    staff_roster = client.get(submission_path, headers=teacher_headers).json()
    assert [(item["student_id"], item["status"], item["version_count"]) for item in staff_roster] == [
        (student.id, "not_graded", 2),
        (unsubmitted_student.id, "not_submitted", 0),
    ]
    ungraded_filter = client.get(
        submission_path + "?status=not_graded", headers=teacher_headers
    )
    assert ungraded_filter.status_code == 200
    assert [item["student_id"] for item in ungraded_filter.json()] == [student.id]
    assert client.get(
        submission_path + "?status=not-a-status", headers=teacher_headers
    ).status_code == 422
    not_submitted_filter = client.get(
        submission_path + "?status=not_submitted", headers=teacher_headers
    )
    assert not_submitted_filter.status_code == 200
    assert [item["student_id"] for item in not_submitted_filter.json()] == [unsubmitted_student.id]
    progress = client.get(
        submission_path + "/progress", headers=teacher_headers
    )
    assert progress.status_code == 200
    assert progress.json() == {
        "assignment_id": draft.json()["id"],
        "total_recipients": 2,
        "submitted_count": 1,
        "not_submitted_count": 1,
        "graded_count": 0,
        "not_graded_count": 1,
        "grading_progress_percent": 0.0,
    }
    grade_path = f"/api/assignments/{draft.json()['id']}/students/{student.id}/grades"
    invalid_grade = client.post(
        grade_path,
        headers=teacher_headers,
        json={"submission_id": second_submission.json()["id"], "score": "100.001"},
    )
    assert invalid_grade.status_code == 422
    over_max_grade = client.post(
        grade_path,
        headers=teacher_headers,
        json={"submission_id": second_submission.json()["id"], "score": "100.01"},
    )
    assert over_max_grade.status_code == 422
    missing_criterion = client.post(
        grade_path,
        headers=teacher_headers,
        json={
            "submission_id": second_submission.json()["id"],
            "rubric_scores": [{"criterion_id": 1, "score_percent": "90"}],
        },
    )
    assert missing_criterion.status_code == 422
    mismatched_total = client.post(
        grade_path,
        headers=teacher_headers,
        json={
            "submission_id": second_submission.json()["id"],
            "score": "91",
            "rubric_scores": [
                {"criterion_id": 1, "score_percent": "90"},
                {"criterion_id": 2, "score_percent": "90"},
            ],
        },
    )
    assert mismatched_total.status_code == 422
    grade_draft = client.post(
        grade_path,
        headers=teacher_headers,
        json={
            "submission_id": second_submission.json()["id"],
            "rubric_scores": [
                {"criterion_id": 1, "score_percent": "90.51", "comment": "Accurate"},
                {"criterion_id": 2, "score_percent": "90.00"},
            ],
            "feedback": "Good work",
            "internal_comment": "Teacher-only note",
        },
    )
    assert grade_draft.status_code == 201
    assert grade_draft.json()["status"] == "draft"
    assert grade_draft.json()["score"] == "90.26"
    assert len(grade_draft.json()["rubric_scores"]) == 2
    assert client.get(
        f"/api/assignments/{draft.json()['id']}/grades", headers=student_headers
    ).json() == []
    denied_grade = client.post(
        grade_path,
        headers={"Authorization": f"Bearer {unassigned_token}"},
        json={"submission_id": second_submission.json()["id"], "score": "90"},
    )
    assert denied_grade.status_code == 403
    published_grade = client.post(
        f"/api/assignments/{draft.json()['id']}/grades/{grade_draft.json()['id']}/publish",
        headers=teacher_headers,
    )
    assert published_grade.status_code == 200
    assert published_grade.json()["status"] == "published"
    visible_grades = client.get(
        f"/api/assignments/{draft.json()['id']}/grades", headers=student_headers
    )
    assert visible_grades.status_code == 200
    assert visible_grades.json()[0]["score"] == "90.26"
    assert visible_grades.json()[0]["feedback"] == "Good work"
    assert "internal_comment" not in visible_grades.json()[0]
    assert len(visible_grades.json()[0]["rubric_scores"]) == 2
    assert visible_grades.json()[0]["rubric_scores"][0]["title"] == "Correctness"
    student_assignment_list = client.get(
        "/api/assignments?page=1&page_size=20", headers=student_headers
    )
    listed_assignment = next(
        item
        for item in student_assignment_list.json()["items"]
        if item["id"] == draft.json()["id"]
    )
    assert listed_assignment["submission_status"] == "graded"
    locked_submission = client.post(
        submission_path,
        headers=student_headers,
        json={"repository_url": "https://github.com/ictu/student-project-after-grade"},
    )
    assert locked_submission.status_code == 409
    assert test_db.query(AssignmentSubmission).filter_by(
        assignment_id=draft.json()["id"], student_id=student.id
    ).count() == 2
    assert test_db.query(Notification).filter_by(user_id=student.id).count() == 2
    graded_roster = client.get(submission_path, headers=teacher_headers).json()
    graded_row = next(item for item in graded_roster if item["student_id"] == student.id)
    assert graded_row["status"] == "graded"
    assert graded_row["published_grade"]["score"] == "90.26"
    graded_filter = client.get(
        submission_path + "?status=graded", headers=teacher_headers
    )
    assert [item["student_id"] for item in graded_filter.json()] == [student.id]
    assert client.get(
        f"/api/assignments/{draft.json()['id']}/grading-queue",
        headers=teacher_headers,
    ).json()["total_ungraded"] == 0
    assert test_db.query(NotificationOutbox).filter(NotificationOutbox.dedupe_key.like("notification:%")).count() == 3
    assert client.get(
        submission_path + "/progress", headers=teacher_headers
    ).json()["graded_count"] == 1
    grade_audit = client.get(
        "/api/audit-logs?entity_type=grade",
        headers={"Authorization": f"Bearer {admin_token}"},
    )
    assert grade_audit.status_code == 200
    assert [item["action"] for item in grade_audit.json()["items"]] == [
        "grade.published",
        "grade.draft_saved",
    ]
    second_grade_draft = client.post(
        grade_path,
        headers=teacher_headers,
        json={
            "submission_id": second_submission.json()["id"],
            "score": "88",
            "rubric_scores": [
                {"criterion_id": 1, "score_percent": "88"},
                {"criterion_id": 2, "score_percent": "88"},
            ],
        },
    )
    assert second_grade_draft.status_code == 201
    test_db.connection().exec_driver_sql(
        """
        CREATE TRIGGER reject_grade_notification
        BEFORE INSERT ON notifications
        WHEN NEW.title = 'Điểm bài tập đã được công bố'
        BEGIN
            SELECT RAISE(ABORT, 'simulated grade notification failure');
        END;
        """
    )
    with pytest.raises(IntegrityError):
        client.post(
            f"/api/assignments/{draft.json()['id']}/grades/{second_grade_draft.json()['id']}/publish",
            headers=teacher_headers,
        )
    test_db.rollback()
    failed_grade = test_db.get(AssignmentGradeVersion, second_grade_draft.json()["id"])
    assert failed_grade.status == "draft"
    assert test_db.query(AuditLog).filter_by(action="grade.published").count() == 1
    test_db.connection().exec_driver_sql("DROP TRIGGER reject_grade_notification")
    test_db.commit()
    submitted_assignment = test_db.get(Assignment, draft.json()["id"])
    submitted_assignment.due_at = datetime.now(timezone.utc).replace(tzinfo=None) - timedelta(days=1)
    test_db.commit()
    late_attempt = client.post(
        submission_path,
        headers=student_headers,
        json={"repository_url": "https://github.com/ictu/student-project-late"},
    )
    assert late_attempt.status_code == 409
    assert test_db.query(AssignmentSubmission).filter_by(
        assignment_id=draft.json()["id"], student_id=student.id
    ).count() == 2

    late_enrolled_student = create_user(
        test_db,
        username="student_after_publish",
        email="student.after.publish@example.edu.vn",
        full_name="Student After Publish",
        role="student",
    )
    test_db.add(
        ClassEnrollment(class_id=training_class.id, student_id=late_enrolled_student.id)
    )
    test_db.commit()
    late_enrolled_token = client.post(
        "/api/auth/login",
        json={"email": late_enrolled_student.email, "password": "CorrectPassword123!"},
    ).json()["access_token"]
    assert client.get(
        "/api/assignments",
        headers={"Authorization": f"Bearer {late_enrolled_token}"},
    ).json() == []
    late_enrolled_submission = client.post(
        submission_path,
        headers={"Authorization": f"Bearer {late_enrolled_token}"},
        json={"repository_url": "https://github.com/ictu/late-project"},
    )
    assert late_enrolled_submission.status_code == 403

    repeated_publish = client.post(
        f"/api/assignments/{draft.json()['id']}/publish",
        headers=teacher_headers,
    )
    assert repeated_publish.status_code == 200
    assert test_db.query(Notification).filter_by(user_id=student.id).count() == 2

    expired_draft = client.post(
        f"/api/classes/{training_class.id}/assignments",
        headers=teacher_headers,
        json={
            "title": "Expired draft",
            "due_at": "2000-01-01T00:00:00",
            "status": "draft",
        },
    )
    assert expired_draft.status_code == 201
    expired_publish = client.post(
        f"/api/assignments/{expired_draft.json()['id']}/publish",
        headers=teacher_headers,
    )
    assert expired_publish.status_code == 422
    test_db.refresh(test_db.get(Assignment, expired_draft.json()["id"]))
    assert test_db.get(Assignment, expired_draft.json()["id"]).status == "draft"
    assert test_db.query(Notification).filter_by(user_id=student.id).count() == 2

    redo_path = f"/api/assignments/{draft.json()['id']}/students/{student.id}/redo-requests"
    redo_payload = {
        "feedback": "Please correct the missing acceptance tests.",
        "due_at": (datetime.now(timezone.utc) + timedelta(days=2)).isoformat(),
    }
    denied_redo = client.post(redo_path, headers={"Authorization": f"Bearer {unassigned_token}"}, json=redo_payload)
    assert denied_redo.status_code == 403
    redo = client.post(redo_path, headers=teacher_headers, json=redo_payload)
    assert redo.status_code == 201
    assert redo.json()["request_number"] == 1
    assert redo.json()["requested_by"] == teacher.id
    assert client.post(redo_path, headers=teacher_headers, json=redo_payload).status_code == 409
    assert client.get(f"/api/assignments/{draft.json()['id']}/rework-requests", headers=student_headers).status_code == 200
    needs_redo = client.get("/api/assignments", headers=student_headers).json()[0]
    assert needs_redo["submission_status"] == "needs_redo"
    resubmitted = client.post(submission_path, headers=student_headers, json={
        "repository_url": "https://github.com/ictu/student-project-redo",
    })
    assert resubmitted.status_code == 201 and resubmitted.json()["version_number"] == 3
    resubmitted_assignment = client.get("/api/assignments", headers=student_headers).json()[0]
    assert resubmitted_assignment["submission_status"] == "submitted"
    assert test_db.query(AssignmentSubmission).filter_by(
        assignment_id=draft.json()["id"], student_id=student.id
    ).count() == 3
    completed_rework = client.get(f"/api/assignments/{draft.json()['id']}/rework-requests", headers=student_headers)
    assert completed_rework.json()[0]["status"] == "completed"
    assert completed_rework.json()[0]["completed_submission_id"] == resubmitted.json()["id"]
    monkeypatch.setattr(settings, "PRIVATE_FILE_STORAGE_DIR", str(tmp_path / "private"))
    file_path = f"/api/assignments/{draft.json()['id']}/submissions/{resubmitted.json()['id']}/files"
    uploaded = client.post(file_path, headers={**student_headers, "X-File-Name": "solution.pdf",
        "Content-Type": "application/pdf"}, content=b"%PDF-1.7\nprivate submission")
    assert uploaded.status_code == 201
    assert uploaded.json()["size_bytes"] == len(b"%PDF-1.7\nprivate submission")
    assert "storage_key" not in uploaded.json()
    second_file = client.post(file_path, headers={**student_headers, "X-File-Name": "notes.txt",
        "Content-Type": "text/plain"}, content=b"submission notes")
    assert second_file.status_code == 201
    assert len(client.get(file_path, headers=student_headers).json()) == 2
    file_review_queue = client.get(
        f"/api/assignments/{draft.json()['id']}/grading-queue", headers=teacher_headers
    )
    assert file_review_queue.status_code == 200
    reviewed_submission = next(
        item for item in file_review_queue.json()["items"]
        if item["submission_id"] == resubmitted.json()["id"]
    )
    assert {item["filename"] for item in reviewed_submission["files"]} == {
        "solution.pdf", "notes.txt"
    }
    assert all("download_url" in item for item in reviewed_submission["files"])
    too_large = client.post(file_path, headers={**student_headers, "X-File-Name": "large.pdf",
        "Content-Type": "application/pdf"}, content=b"%PDF-1.7" + b"x" * (20 * 1024 * 1024))
    assert too_large.status_code == 413
    denied_file_list = client.get(file_path, headers={"Authorization": f"Bearer {unassigned_token}"})
    assert denied_file_list.status_code == 403
    downloaded = client.get(uploaded.json()["download_url"], headers=student_headers)
    assert downloaded.status_code == 200 and downloaded.content.startswith(b"%PDF-")
    denied_download = client.get(uploaded.json()["download_url"], headers={"Authorization": f"Bearer {unassigned_token}"})
    assert denied_download.status_code == 403
    bad_signature = client.post(file_path, headers={**student_headers, "X-File-Name": "fake.pdf",
        "Content-Type": "application/pdf"}, content=b"not a pdf")
    assert bad_signature.status_code == 415
    path_traversal = client.post(file_path, headers={**student_headers, "X-File-Name": "../escape.pdf",
        "Content-Type": "application/pdf"}, content=b"%PDF-1.7")
    assert path_traversal.status_code == 422
    submitted_assignment = test_db.get(Assignment, draft.json()["id"])
    submitted_assignment.late_submission_policy = "ALLOW_LATE"
    test_db.commit()
    redo_grade = client.post(grade_path, headers=teacher_headers, json={
        "submission_id": resubmitted.json()["id"],
        "rubric_scores": [
            {"criterion_id": 1, "score_percent": "85"},
            {"criterion_id": 2, "score_percent": "85"},
        ],
    })
    assert redo_grade.status_code == 201
    assert client.post(f"/api/assignments/{draft.json()['id']}/grades/{redo_grade.json()['id']}/publish",
        headers=teacher_headers).status_code == 200
    reopened = client.post(f"/api/assignments/{draft.json()['id']}/students/{student.id}/reopen",
        headers=teacher_headers)
    assert reopened.status_code == 201 and reopened.json()["request_type"] == "reopen"
    reopened_submission = client.post(submission_path, headers=student_headers, json={
        "repository_url": "https://github.com/ictu/student-project-reopened",
    })
    assert reopened_submission.status_code == 201 and reopened_submission.json()["version_number"] == 4
    rework_history = client.get(f"/api/assignments/{draft.json()['id']}/rework-requests",
        headers=student_headers).json()
    assert rework_history[1]["request_type"] == "reopen"
    assert rework_history[1]["status"] == "completed"


def test_audit_log_list_is_admin_only_and_filters_supported_sensitive_events(client, test_db):
    admin_token = get_admin(client, test_db)
    admin = test_db.query(User).filter_by(username="admin01").one()
    teacher = create_user(
        test_db,
        username="teacher_audit_filter",
        email="teacher.audit.filter@example.edu.vn",
        full_name="Audit Filter Teacher",
        role="teacher",
    )
    base_time = datetime(2026, 10, 22, 9, 0)
    test_db.add_all(
        [
            AuditLog(
                actor_id=admin.id,
                target_user_id=teacher.id,
                action="account.role_changed",
                details="before=teacher;after=manager",
                created_at=base_time,
            ),
            AuditLog(
                actor_id=admin.id,
                target_user_id=teacher.id,
                action="enrollment.tuition_status_changed",
                details="class_id=7;before=unpaid;after=paid",
                created_at=base_time + timedelta(minutes=1),
            ),
            AuditLog(
                actor_id=admin.id,
                target_user_id=teacher.id,
                action="account.password_changed",
                details="must not be exposed by sprint audit endpoint",
                created_at=base_time + timedelta(minutes=2),
            ),
            AuditLog(
                actor_id=admin.id,
                target_user_id=teacher.id,
                action="account.role_changed",
                details="outside local calendar day",
                created_at=base_time + timedelta(hours=8, minutes=30),
            ),
        ]
    )
    test_db.commit()

    headers = {"Authorization": f"Bearer {admin_token}"}
    response = client.get(
        "/api/audit-logs",
        headers=headers,
        params={
            "actor_id": admin.id,
            "entity_type": "tuition",
            "from_date": "2026-10-22T00:00:00Z",
            "to_date": "2026-10-23T00:00:00Z",
        },
    )
    assert response.status_code == 200
    assert response.json()["total"] == 1
    assert response.json()["items"][0]["action"] == "enrollment.tuition_status_changed"
    assert response.json()["items"][0]["details"] == "class_id=7;before=unpaid;after=paid"

    local_day = client.get(
        "/api/audit-logs",
        headers=headers,
        params={
            "from_date": "2026-10-22T00:00:00+07:00",
            "to_date": "2026-10-22T23:59:59.999999+07:00",
        },
    )
    assert local_day.status_code == 200
    assert local_day.json()["total"] == 2

    invalid_range = client.get(
        "/api/audit-logs?from_date=2026-10-23T00:00:00Z&to_date=2026-10-22T00:00:00Z",
        headers=headers,
    )
    assert invalid_range.status_code == 422

    student = create_user(
        test_db,
        username="student_audit_denied",
        email="student.audit.denied@example.edu.vn",
        full_name="Audit Denied Student",
        role="student",
    )
    student_token = client.post(
        "/api/auth/login",
        json={"email": student.email, "password": "CorrectPassword123!"},
    ).json()["access_token"]
    assert client.get(
        "/api/audit-logs",
        headers={"Authorization": f"Bearer {student_token}"},
    ).status_code == 403


def test_tuition_debt_report_filters_overdue_and_exports_days_and_totals(client, test_db):
    get_admin(client, test_db)
    accountant = create_user(test_db, username="overdue_accountant",
        email="overdue.accountant@example.edu.vn", full_name="Overdue Accountant", role="accountant")
    teacher = create_user(test_db, username="overdue_teacher",
        email="overdue.teacher@example.edu.vn", full_name="Overdue Teacher", role="teacher")
    student = create_user(test_db, username="overdue_student", email="overdue.student@example.edu.vn",
        full_name="Overdue Student", role="student", student_code="SV-OVERDUE")
    course = Course(code="OVERDUE101", name="Overdue Course", department="IT", credits=3)
    program = StudyProgram(code="OVERDUE-PROGRAM", name="Overdue Program", total_credits=60)
    program.courses.append(course)
    test_db.add_all([course, program])
    test_db.flush()
    schedule = ProgramFeeSchedule(program_id=program.id, version=1, amount_vnd=500000,
        valid_from=date.today() - timedelta(days=30))
    training_class = TrainingClass(code="OVERDUE101-A", name="Overdue Class", course_id=course.id,
        program_id=program.id, teacher_id=teacher.id, room="R1", term="2026-Fall",
        day_of_week=1, start_time=time(9), end_time=time(11), capacity=10, status="active")
    test_db.add_all([schedule, training_class])
    test_db.flush()
    enrollment = ClassEnrollment(class_id=training_class.id, student_id=student.id)
    test_db.add(enrollment)
    test_db.flush()
    plan = TuitionPlan(enrollment_id=enrollment.id, fee_schedule_id=schedule.id,
        total_due_vnd=500000, applied_on=date.today() - timedelta(days=30))
    test_db.add(plan)
    test_db.flush()
    test_db.add_all(TuitionInstallment(plan_id=plan.id, installment_number=number,
        amount_due_vnd=20000, due_date=date.today() - timedelta(days=7))
        for number in range(1, 26))
    test_db.commit()

    login = client.post("/api/auth/login", json={
        "email": accountant.email, "password": "CorrectPassword123!",
    })
    headers = {"Authorization": f"Bearer {login.json()['access_token']}"}
    response = client.get("/api/tuition-debts?status=overdue&page=1&page_size=20", headers=headers)
    assert response.status_code == 200
    assert response.json()["total"] == 25
    assert len(response.json()["items"]) == 20
    assert response.json()["pages"] == 2 and response.json()["has_next"] is True
    assert response.json()["total_due_vnd"] == "500000"
    assert response.json()["total_remaining_vnd"] == "500000"
    row = response.json()["items"][0]
    assert row["status"] == "overdue"
    assert row["overdue_days"] == 7
    assert row["student_code"] == "SV-OVERDUE"
    second_page = client.get("/api/tuition-debts?status=overdue&page=2&page_size=20",
        headers=headers)
    assert len(second_page.json()["items"]) == 5
    assert second_page.json()["has_previous"] is True and second_page.json()["has_next"] is False
    all_page_ids = [item["installment_id"] for item in response.json()["items"]
        + second_page.json()["items"]]
    assert len(all_page_ids) == len(set(all_page_ids)) == 25
    assert client.get(f"/api/tuition-debts?class_id={training_class.id}&student_id={student.id}",
        headers=headers).json()["total"] == 25
    export = client.get("/api/tuition-debts/export.xlsx?status=overdue", headers=headers)
    assert export.status_code == 200
    sheet = load_workbook(BytesIO(export.content), data_only=True).active
    assert sheet.cell(2, 1).value == "SV-OVERDUE"
    assert sheet.cell(2, 9).value == "overdue"
    assert sheet.cell(2, 10).value == 7
    assert sheet.cell(27, 1).value == "TOTAL"
    assert sheet.cell(27, 8).value == 500000


def test_assignment_late_policy_accept_reject_and_unconfigured(client, test_db, monkeypatch):
    get_admin(client, test_db)
    teacher = create_user(test_db, username="late-policy-teacher",
        email="late-policy-teacher@example.edu.vn", full_name="Late Policy Teacher", role="teacher")
    student = create_user(test_db, username="late-policy-student",
        email="late-policy-student@example.edu.vn", full_name="Late Policy Student", role="student")
    course = Course(code="LATE101", name="Late Policy Course", department="IT", credits=3)
    test_db.add(course)
    test_db.flush()
    training_class = TrainingClass(code="LATE101-01", name="Late Policy Class",
        course_id=course.id, teacher_id=teacher.id, room="Room 1", term="2026-Fall",
        day_of_week=1, start_time=time(9), end_time=time(11), capacity=20, status="active")
    test_db.add(training_class)
    test_db.flush()
    test_db.add(ClassEnrollment(class_id=training_class.id, student_id=student.id))
    due_at = datetime.now(HO_CHI_MINH).replace(tzinfo=None) - timedelta(hours=2, minutes=5)
    assignments = [
        Assignment(class_id=training_class.id, created_by=teacher.id, title="Late accepted",
                   due_at=due_at, status="published", late_submission_policy="ALLOW_LATE"),
        Assignment(class_id=training_class.id, created_by=teacher.id, title="Late rejected",
                   due_at=due_at, status="published", late_submission_policy="BLOCK_LATE"),
        Assignment(class_id=training_class.id, created_by=teacher.id, title="Late unconfigured",
                   due_at=due_at, status="published"),
    ]
    test_db.add_all(assignments)
    test_db.commit()
    student_token = client.post("/api/auth/login", json={
        "email": student.email, "password": "CorrectPassword123!"}).json()["access_token"]
    student_headers = {"Authorization": f"Bearer {student_token}"}
    accepted = client.post(f"/api/assignments/{assignments[0].id}/submissions",
        headers=student_headers, json={"repository_url": "https://example.test/late"})
    assert accepted.status_code == 201
    assert accepted.json()["is_late"] is True
    assert accepted.json()["late_hours"] == 3
    student_assignments = client.get("/api/assignments", headers=student_headers)
    late_item = next(item for item in student_assignments.json()
                     if item["id"] == assignments[0].id)
    assert late_item["is_late"] is True and late_item["late_hours"] == 3
    teacher_token = client.post("/api/auth/login", json={
        "email": teacher.email, "password": "CorrectPassword123!"}).json()["access_token"]
    staff_rows = client.get(f"/api/assignments/{assignments[0].id}/submissions?status=late",
        headers={"Authorization": f"Bearer {teacher_token}"})
    assert staff_rows.status_code == 200 and staff_rows.json()[0]["late_hours"] == 3
    assert staff_rows.headers["x-total-count"] == "1"
    assert staff_rows.headers["x-current-page"] == "1"
    student_history_page = client.get(
        f"/api/assignments/{assignments[0].id}/submissions?page=1&page_size=1",
        headers=student_headers,
    )
    assert student_history_page.status_code == 200
    assert student_history_page.headers["x-total-count"] == "1"
    assert client.post(f"/api/assignments/{assignments[1].id}/submissions",
        headers=student_headers, json={"repository_url": "https://example.test/rejected"}).status_code == 409
    assert client.post(f"/api/assignments/{assignments[2].id}/submissions",
        headers=student_headers, json={"repository_url": "https://example.test/unconfigured"}).status_code == 409

    exact_utc = datetime(2026, 4, 1, 12, 0, tzinfo=timezone.utc)
    exact_assignment = Assignment(class_id=training_class.id, created_by=teacher.id,
        title="Exact deadline", due_at=exact_utc.astimezone(HO_CHI_MINH).replace(tzinfo=None),
        status="published", late_submission_policy="BLOCK_LATE")
    test_db.add(exact_assignment)
    test_db.commit()
    class FrozenDateTime(datetime):
        @classmethod
        def now(cls, tz=None):
            return exact_utc if tz is not None else exact_utc.replace(tzinfo=None)
    monkeypatch.setattr(management_routes, "datetime", FrozenDateTime)
    exact_submission = client.post(f"/api/assignments/{exact_assignment.id}/submissions",
        headers=student_headers, json={"repository_url": "https://example.test/exact"})
    assert exact_submission.status_code == 201
    assert exact_submission.json()["is_late"] is False
    assert exact_submission.json()["late_hours"] is None


def test_rbac_default_deny_and_finance_academic_role_boundaries(client, test_db):
    expected_roles = {"admin", "manager", "admissions", "accountant", "teacher", "ta", "student", "guest"}
    test_db.add_all([Role(code=role, name=role.title()) for role in expected_roles])
    test_db.commit()
    users = {
        role: create_user(test_db, username=f"rbac_matrix_{role}",
            email=f"rbac.matrix.{role}@example.edu.vn", full_name=f"RBAC {role}", role=role)
        for role in expected_roles - {"guest"}
    }
    guest_role = test_db.query(Role).filter_by(code="guest").one()
    guest = User(username="rbac_matrix_guest", email="rbac.matrix.guest@example.edu.vn",
        full_name="RBAC guest", hashed_password=hash_password("CorrectPassword123!"), roles=[guest_role])
    test_db.add(guest)
    test_db.commit()

    headers = {}
    for role, user in (*users.items(), ("guest", guest)):
        token = client.post("/api/auth/login", json={
            "email": user.email, "password": "CorrectPassword123!"}).json()["access_token"]
        headers[role] = {"Authorization": f"Bearer {token}"}

    assert {item["code"] for item in client.get("/api/roles", headers=headers["admin"]).json()} == expected_roles
    assert client.get("/api/roles", headers=headers["manager"]).status_code == 403

    for role in ("admin", "accountant"):
        assert client.get("/api/tuition-debts", headers=headers[role]).status_code == 200
    for role in expected_roles - {"admin", "accountant"}:
        denied = client.get("/api/tuition-debts", headers=headers[role])
        assert denied.status_code == 403
        assert denied.json()["detail"] == "Bạn không có quyền thực hiện thao tác này."

    invalid_policy = {
        "program_id": 999999, "course_id": 999999, "pass_score": "60",
        "components": [{"key": "exam", "name": "Exam", "weight": "100"}],
    }
    for role in expected_roles - {"admin", "manager"}:
        assert client.post("/api/grade-policies", headers=headers[role], json=invalid_policy).status_code == 403
    assert client.post("/api/grade-policies", headers=headers["manager"], json=invalid_policy).status_code == 404


def test_openapi_declares_auth_for_every_nonpublic_api_operation():
    public_operations = {
        ("POST", "/api/auth/login"),
        ("POST", "/api/auth/refresh"),
        ("POST", "/api/auth/logout"),
        ("POST", "/api/auth/forgot-password"),
        ("POST", "/api/auth/reset-password"),
        ("POST", "/api/leads"),
        ("GET", "/"),
        ("GET", "/health"),
    }
    schema = fastapi_app.openapi()
    unsecured = {
        (method.upper(), path)
        for path, operations in schema["paths"].items()
        for method, operation in operations.items()
        if method.lower() in {"get", "post", "put", "patch", "delete"}
        and not operation.get("security")
    }
    assert unsecured == public_operations
