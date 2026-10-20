"""Excel representation of the same matrix used by the attendance summary."""
from io import BytesIO
import re

from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

MIME_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
STATES = {
    "present": ("CM", "Có mặt", "DFF5E9"),
    "late": ("M", "Đi muộn", "FFF2C9"),
    "excused_absence": ("CP", "Vắng có phép", "E6EDFF"),
    "unexcused_absence": ("KP", "Vắng không phép", "FFE3DE"),
    "unmarked": ("?", "Chưa điểm danh", "FFFFFF"),
    "not_enrolled": ("—", "Không thuộc lớp tại ngày học", "F1F4F8"),
    "upcoming": ("Sắp", "Buổi tương lai", "F1F4F8"),
    "cancelled": ("Huỷ", "Buổi đã huỷ", "E8EAF0"),
}
MEMBERSHIPS = {
    "active": "Đang học", "former": "Đã rời lớp / bảo lưu",
    "upcoming": "Chưa vào lớp", "historical": "Hồ sơ lịch sử",
}


def safe_text(value):
    return ILLEGAL_CHARACTERS_RE.sub("", str(value))


def text_cell(sheet, row, column, value):
    cell = sheet.cell(row, column, safe_text(value))
    # User-controlled codes/names/topics must remain literal text, including '='.
    cell.data_type = "s"
    return cell


def filename(matrix):
    code = re.sub(r'[<>:"/\\|?*\x00-\x1f\x7f]', "_", matrix["class_code"]).strip(" .")[:80]
    return f"Chuyen-can_{code or 'LOP-' + str(matrix['class_id'])}_{matrix['today'].isoformat()}.xlsx"


def workbook_bytes(matrix):
    columns, rows = 3 + len(matrix["sessions"]), 5 + len(matrix["students"])
    if columns > 16384 or rows > 1048576:
        raise HTTPException(422, "Bảng chuyên cần vượt giới hạn số dòng hoặc cột của Excel.")
    book = Workbook()
    sheet = book.active
    sheet.title = "Chuyên cần"
    last_column = get_column_letter(columns)
    for row, value in enumerate([
        "BẢNG CHUYÊN CẦN",
        f"{matrix['class_code']} · {matrix['class_name']}",
        f"Ngày xuất: {matrix['today'].strftime('%d/%m/%Y')} (giờ Việt Nam) · {matrix['total']} học viên",
        "Tỷ lệ = (Có mặt + Đi muộn) / số buổi có kết quả. Chú giải tại trang “Chú giải”.",
    ], 1):
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=columns)
        cell = text_cell(sheet, row, 1, value)
        cell.font = Font(name="Calibri", size=16 if row == 1 else 11, bold=row <= 2, color="173B68")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.row_dimensions[row].height = 30
    headers = ["Học viên", "Tổng vắng", "Chuyên cần"]
    for index, session in enumerate(matrix["sessions"], 1):
        suffix = "\nĐã huỷ" if session["status"] == "cancelled" else "\nChưa đến ngày" if session["upcoming"] else ""
        headers.append(f"Buổi {index}\n{session['session_date'].strftime('%d/%m/%Y')} {session['start_time']}\n"
                       f"{session['subject_name']}\n{session['topic']}{suffix}")
    for column, header in enumerate(headers, 1):
        cell = text_cell(sheet, 5, column, header)
        cell.fill = PatternFill("solid", fgColor="173B68")
        cell.font = Font(name="Calibri", bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.column_dimensions[get_column_letter(column)].width = 23 if column > 3 else 17
    sheet.column_dimensions["A"].width = 34
    sheet.row_dimensions[5].height = 100
    edge = Side(style="thin", color="DCE5EF")
    for row, student in enumerate(matrix["students"], 6):
        text_cell(sheet, row, 1, f"{student['student_name']}\n{student['student_code']}\n{MEMBERSHIPS[student['membership']]}")
        sheet.cell(row, 2, student["absent_sessions"]).comment = Comment(
            f"{student['excused_absences']} vắng có phép · {student['unexcused_absences']} vắng không phép", "Hệ thống đào tạo")
        if student["attendance_rate"] is None:
            text_cell(sheet, row, 3, "—")
        else:
            sheet.cell(row, 3, student["attendance_rate"] / 100).number_format = "0.00%"
        sheet.cell(row, 3).comment = Comment(
            f"{student['attended_sessions']}/{student['recorded_sessions']} buổi có kết quả", "Hệ thống đào tạo")
        for column, state in enumerate(student["cells"], 4):
            short, label, color = STATES[state]
            cell = text_cell(sheet, row, column, short)
            cell.fill = PatternFill("solid", fgColor=color)
            cell.comment = Comment(label, "Hệ thống đào tạo")
        for cell in sheet[row]:
            cell.font = Font(name="Calibri", size=11, color="263B52")
            cell.alignment = Alignment(horizontal="left" if cell.column == 1 else "center", vertical="center", wrap_text=True)
            cell.border = Border(left=edge, right=edge, top=edge, bottom=edge)
        sheet.row_dimensions[row].height = 56
    sheet.freeze_panes = "D6"
    sheet.auto_filter.ref = f"A5:{last_column}{max(5, rows)}"
    sheet.print_title_rows = "1:5"
    sheet.print_title_cols = "A:C"
    sheet.print_area = f"A1:{last_column}{max(5, rows)}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.sheet_view.showGridLines = False
    legend = book.create_sheet("Chú giải")
    for row, (short, label, color) in enumerate(STATES.values(), 1):
        text_cell(legend, row, 1, short).fill = PatternFill("solid", fgColor=color)
        text_cell(legend, row, 2, label)
    text_cell(legend, 10, 2, "Tổng vắng = vắng có phép + vắng không phép.")
    text_cell(legend, 11, 2, "Buổi huỷ, tương lai, chưa điểm danh và thời gian không thuộc lớp không tính vào tỷ lệ.")
    text_cell(legend, 12, 2, "Chưa có kết quả: —. Các nhận xét ô tổng chứa số CP/KP và số buổi có mặt/có kết quả.")
    legend.column_dimensions["A"].width = 12
    legend.column_dimensions["B"].width = 100
    output = BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()
