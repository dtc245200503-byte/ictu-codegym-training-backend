from datetime import datetime, timezone
from io import BytesIO

from fastapi import APIRouter, Depends, HTTPException, Query, Response
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from sqlalchemy import func, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.api.dependencies import get_current_user, require_roles
from app.api.pagination import page_metadata, set_pagination_headers
from app.db.session import get_db
from app.models.survey import (
    AnonymousSurveyResponse, ClassSurvey, SurveyQuestion, SurveyRespondentReceipt,
)
from app.models.training_class import ClassEnrollment, TrainingClass
from app.models.user import User
from app.schemas.academic import ClassSurveyCreate, SurveyResponseCreate

router = APIRouter(tags=["Course surveys"])


def _survey_question(question: SurveyQuestion) -> dict:
    return {"id": question.id, "position": question.position, "prompt": question.prompt,
            "question_type": question.question_type, "is_required": question.is_required,
            "scale_min": question.scale_min, "scale_max": question.scale_max}


def _get_survey(db: Session, survey_id: int) -> ClassSurvey:
    survey = db.get(ClassSurvey, survey_id)
    if survey is None:
        raise HTTPException(status_code=404, detail="Survey not found.")
    return survey


def _report(db: Session, survey: ClassSurvey) -> dict:
    questions = db.scalars(select(SurveyQuestion).where(
        SurveyQuestion.survey_id == survey.id
    ).order_by(SurveyQuestion.position)).all()
    responses = db.scalars(select(AnonymousSurveyResponse).where(
        AnonymousSurveyResponse.survey_id == survey.id
    )).all()
    aggregate = []
    text_answers = []
    for question in questions:
        values = [row.answers.get(str(question.id)) for row in responses
                  if str(question.id) in row.answers]
        if question.question_type == "rating":
            numeric = [value for value in values if isinstance(value, int) and not isinstance(value, bool)]
            aggregate.append({"question_id": question.id, "prompt": question.prompt,
                "question_type": "rating", "response_count": len(numeric),
                "average": round(sum(numeric) / len(numeric), 2) if numeric else None})
        else:
            aggregate.append({"question_id": question.id, "prompt": question.prompt,
                "question_type": "text", "response_count": len(values), "average": None})
            text_answers.extend({"question_id": question.id, "prompt": question.prompt,
                                 "answer": value} for value in values if isinstance(value, str))
    training_class = db.get(TrainingClass, survey.class_id)
    return {"survey_id": survey.id, "title": survey.title, "class_id": survey.class_id,
            "class_name": training_class.name if training_class else None,
            "teacher_id": training_class.teacher_id if training_class else None,
            "response_count": len(responses), "questions": aggregate,
            "text_answers": text_answers}


@router.post("/surveys", status_code=201)
def create_survey(payload: ClassSurveyCreate, db: Session = Depends(get_db),
                  user: User = Depends(require_roles("admin", "manager"))):
    training_class = db.get(TrainingClass, payload.class_id)
    if training_class is None:
        raise HTTPException(status_code=404, detail="Class not found.")
    if db.scalar(select(ClassSurvey.id).where(ClassSurvey.class_id == payload.class_id)) is not None:
        raise HTTPException(status_code=409, detail="This class already has a survey.")
    deadline = (payload.deadline.replace(tzinfo=timezone.utc) if payload.deadline.tzinfo is None
                else payload.deadline.astimezone(timezone.utc)).replace(tzinfo=None)
    now_utc = datetime.now(timezone.utc).replace(tzinfo=None)
    if deadline <= now_utc:
        raise HTTPException(status_code=422, detail="Survey deadline must be in the future.")
    survey = ClassSurvey(class_id=payload.class_id, title=payload.title,
        description=payload.description, deadline=deadline,
        status="open", created_by=user.id)
    db.add(survey)
    db.flush()
    for position, question in enumerate(payload.questions, start=1):
        db.add(SurveyQuestion(survey_id=survey.id, position=position,
            prompt=question.prompt, question_type=question.question_type,
            is_required=question.is_required,
            scale_min=question.scale_min if question.question_type == "rating" else None,
            scale_max=question.scale_max if question.question_type == "rating" else None))
    try:
        db.commit()
    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(status_code=409, detail="This class already has a survey.") from exc
    db.refresh(survey)
    return {"id": survey.id, "class_id": survey.class_id, "title": survey.title,
            "deadline": survey.deadline.isoformat(), "status": survey.status,
            "questions": [_survey_question(item) for item in db.scalars(
                select(SurveyQuestion).where(SurveyQuestion.survey_id == survey.id)
                .order_by(SurveyQuestion.position)).all()]}


@router.get("/students/me/surveys")
def my_surveys(response: Response, page: int = Query(default=1, ge=1),
               page_size: int = Query(default=20, ge=1, le=100),
               db: Session = Depends(get_db), user: User = Depends(require_roles("student"))):
    now = datetime.now(timezone.utc).replace(tzinfo=None)
    query = select(ClassSurvey, TrainingClass).join(
        TrainingClass, ClassSurvey.class_id == TrainingClass.id
    ).join(ClassEnrollment, ClassEnrollment.class_id == TrainingClass.id).where(
        ClassEnrollment.student_id == user.id,
        ClassEnrollment.enrollment_status == "active",
        ClassSurvey.status == "open", ClassSurvey.deadline > now,
    )
    total = db.scalar(select(func.count()).select_from(
        query.with_only_columns(ClassSurvey.id).order_by(None).distinct().subquery()
    )) or 0
    rows = db.execute(query.order_by(ClassSurvey.deadline, ClassSurvey.id)
        .offset((page - 1) * page_size).limit(page_size)).all()
    set_pagination_headers(response, total=total, page=page, page_size=page_size)
    result = []
    for survey, training_class in rows:
        receipt = db.scalar(select(SurveyRespondentReceipt.id).where(
            SurveyRespondentReceipt.survey_id == survey.id,
            SurveyRespondentReceipt.student_id == user.id,
        ))
        result.append({"survey_id": survey.id, "title": survey.title,
            "class_id": training_class.id, "class_name": training_class.name,
            "deadline": survey.deadline.isoformat(), "responded": receipt is not None,
            "questions": [_survey_question(item) for item in db.scalars(select(SurveyQuestion)
                .where(SurveyQuestion.survey_id == survey.id)
                .order_by(SurveyQuestion.position)).all()]})
    return result


@router.post("/surveys/{survey_id}/responses", status_code=201)
def submit_survey_response(survey_id: int, payload: SurveyResponseCreate,
                           db: Session = Depends(get_db),
                           user: User = Depends(require_roles("student"))):
    survey = _get_survey(db, survey_id)
    if survey.status != "open" or survey.deadline <= datetime.now(timezone.utc).replace(tzinfo=None):
        raise HTTPException(status_code=409, detail="Survey is closed.")
    enrollment = db.scalar(select(ClassEnrollment.id).where(
        ClassEnrollment.class_id == survey.class_id,
        ClassEnrollment.student_id == user.id,
        ClassEnrollment.enrollment_status == "active",
    ))
    if enrollment is None:
        raise HTTPException(status_code=403, detail="Active class enrollment required.")
    questions = db.scalars(select(SurveyQuestion).where(
        SurveyQuestion.survey_id == survey.id
    )).all()
    by_id = {str(question.id): question for question in questions}
    answers = {str(question_id): value for question_id, value in payload.answers.items()}
    unknown = set(answers) - set(by_id)
    missing = {key for key, question in by_id.items()
               if question.is_required and key not in answers}
    if unknown or missing:
        raise HTTPException(status_code=422, detail={"unknown_question_ids": sorted(unknown),
            "missing_question_ids": sorted(missing)})
    for key, value in answers.items():
        question = by_id[key]
        if question.question_type == "rating":
            if isinstance(value, bool) or not isinstance(value, int) or not question.scale_min <= value <= question.scale_max:
                raise HTTPException(status_code=422, detail=f"Invalid rating for question {key}.")
        elif not isinstance(value, str) or len(value) > 5000:
            raise HTTPException(status_code=422, detail=f"Invalid text answer for question {key}.")
    if db.scalar(select(SurveyRespondentReceipt.id).where(
        SurveyRespondentReceipt.survey_id == survey.id,
        SurveyRespondentReceipt.student_id == user.id,
    )) is not None:
        raise HTTPException(status_code=409, detail="A response was already submitted.")
    db.add(AnonymousSurveyResponse(survey_id=survey.id, answers=answers))
    db.add(SurveyRespondentReceipt(survey_id=survey.id, student_id=user.id))
    try:
        db.commit()
    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(status_code=409, detail="A response was already submitted.") from exc
    return {"survey_id": survey.id, "submitted": True,
            "message": "Response stored without respondent identity."}


@router.get("/surveys/{survey_id}/respondents")
def survey_respondents(survey_id: int, page: int = Query(default=1, ge=1),
                       page_size: int = Query(default=20, ge=1, le=100),
                       db: Session = Depends(get_db),
                       user: User = Depends(require_roles("manager"))):
    survey = _get_survey(db, survey_id)
    roster_query = select(User).join(ClassEnrollment,
        ClassEnrollment.student_id == User.id).where(
        ClassEnrollment.class_id == survey.class_id,
        ClassEnrollment.enrollment_status == "active")
    total = db.scalar(select(func.count()).select_from(
        roster_query.with_only_columns(User.id).order_by(None).subquery()
    )) or 0
    roster = db.scalars(roster_query.order_by(User.id)
        .offset((page - 1) * page_size).limit(page_size)).all()
    roster_ids = [student.id for student in roster]
    submitted_ids = set(db.scalars(select(SurveyRespondentReceipt.student_id).where(
        SurveyRespondentReceipt.survey_id == survey_id,
        SurveyRespondentReceipt.student_id.in_(roster_ids or [-1]))).all())
    return {"survey_id": survey_id,
            **page_metadata(total=total, page=page, page_size=page_size),
            "responded_student_ids": sorted(submitted_ids),
            "pending_student_ids": [student.id for student in roster if student.id not in submitted_ids]}


@router.get("/surveys/{survey_id}/results")
def survey_results(survey_id: int, db: Session = Depends(get_db),
                   user: User = Depends(require_roles("admin", "manager"))):
    return _report(db, _get_survey(db, survey_id))


@router.get("/survey-reports/teachers/{teacher_id}")
def teacher_survey_comparison(teacher_id: int, response: Response,
                              page: int = Query(default=1, ge=1),
                              page_size: int = Query(default=20, ge=1, le=100),
                              db: Session = Depends(get_db),
                              user: User = Depends(require_roles("admin", "manager"))):
    query = select(ClassSurvey).join(TrainingClass,
        ClassSurvey.class_id == TrainingClass.id).where(
        TrainingClass.teacher_id == teacher_id)
    total = db.scalar(select(func.count()).select_from(query.order_by(None).subquery())) or 0
    surveys = db.scalars(query.order_by(ClassSurvey.class_id, ClassSurvey.id)
        .offset((page - 1) * page_size).limit(page_size)).all()
    reports = [_report(db, survey) for survey in surveys]
    return {"teacher_id": teacher_id, **page_metadata(total=total, page=page, page_size=page_size),
            "classes": reports}


@router.get("/survey-reports/teachers/{teacher_id}/export.xlsx")
def export_teacher_survey_comparison(teacher_id: int, db: Session = Depends(get_db),
                                     user: User = Depends(require_roles("admin", "manager"))):
    surveys = db.scalars(select(ClassSurvey).join(TrainingClass,
        ClassSurvey.class_id == TrainingClass.id).where(
        TrainingClass.teacher_id == teacher_id).order_by(ClassSurvey.class_id, ClassSurvey.id)).all()
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Class comparison"
    summary.append(["Class", "Survey", "Question", "Responses", "Average rating"])
    comments = workbook.create_sheet("Anonymous comments")
    comments.append(["Class", "Survey", "Question", "Answer"])
    for survey in surveys:
        report = _report(db, survey)
        for question in report["questions"]:
            summary.append([report["class_name"], report["title"], question["prompt"],
                            question["response_count"], question["average"]])
        for answer in report["text_answers"]:
            comments.append([report["class_name"], report["title"],
                             answer["prompt"], answer["answer"]])
    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return StreamingResponse(stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="teacher-{teacher_id}-survey-comparison.xlsx"'})


@router.get("/surveys/{survey_id}/results.xlsx")
def export_survey_results(survey_id: int, db: Session = Depends(get_db),
                          user: User = Depends(require_roles("admin", "manager"))):
    report = _report(db, _get_survey(db, survey_id))
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Survey summary"
    sheet.append(["Question", "Type", "Responses", "Average"])
    for item in report["questions"]:
        sheet.append([item["prompt"], item["question_type"], item["response_count"], item["average"]])
    text_sheet = workbook.create_sheet("Anonymous comments")
    text_sheet.append(["Question", "Answer"])
    for item in report["text_answers"]:
        text_sheet.append([item["prompt"], item["answer"]])
    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return StreamingResponse(stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="survey-{survey_id}-results.xlsx"'})
