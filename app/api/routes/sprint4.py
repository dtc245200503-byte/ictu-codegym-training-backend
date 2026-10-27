import base64
import binascii
from datetime import date, datetime, time, timedelta, timezone
from io import BytesIO
from typing import Any
from urllib.parse import quote

from fastapi import APIRouter, Depends, HTTPException, Query, Response, status
from openpyxl import Workbook
from sqlalchemy import String, and_, case, cast, func, or_, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session, joinedload, load_only, selectinload

from app.api.dependencies import get_current_user, require_roles
from app.api.pagination import page_metadata, set_pagination_headers
from app.db.session import get_db
from app.models.audit_log import AuditLog
from app.models.attendance import (
    AbsenceRequest,
    AttendanceChangeLog,
    AttendanceRecord,
    AttendanceRiskState,
    AttendanceScheduleHoliday,
    AttendanceSession,
)
from app.models.course import Course
from app.models.course_lesson import CourseLesson
from app.models.notification import Notification
from app.models.student_enrollment_history import StudentEnrollmentHistory
from app.models.study_program import StudyProgram
from app.models.training_class import ClassEnrollment, TrainingClass
from app.services.attendance_summary import HO_CHI_MINH, summarize_class_attendance
from app.services.notification_outbox import enqueue_notification_email
from app.services.attendance_risk import attendance_warning, refresh_attendance_risk_states
from app.models.user import User
from app.schemas.sprint4 import (
    AbsenceRequestCreate,
    AbsenceRequestDecision,
    AttendanceBulkUpdate,
    EnrollmentChangeCreate,
    ScheduleGenerateRequest,
    ScheduleHolidayCreate,
    ScheduleSessionUpdate,
)

router = APIRouter(tags=["Sprint 4 schedule and attendance"])
MANAGER_ROLES = ("admin", "manager")
STAFF_ROLES = ("admin", "manager", "teacher", "ta")
DAY_NAMES = ("Thứ Hai", "Thứ Ba", "Thứ Tư", "Thứ Năm", "Thứ Sáu", "Thứ Bảy", "Chủ Nhật")


def _has_class_staff_access(user: User, training_class: TrainingClass) -> bool:
    roles = {role.code for role in user.roles}
    if roles.intersection(MANAGER_ROLES):
        return True
    if "teacher" in roles and training_class.teacher_id == user.id:
        return True
    return "ta" in roles and (
        training_class.assistant_id == user.id
        or any(assistant.id == user.id for assistant in training_class.assistants)
    )


def _assert_class_staff_access(
    user: User,
    training_class: TrainingClass,
    session: AttendanceSession | None = None,
) -> None:
    if _has_class_staff_access(user, training_class):
        return
    if (
        session is not None
        and "teacher" in {role.code for role in user.roles}
        and _session_teacher(session) == user.id
    ):
        return
    raise HTTPException(status_code=403, detail="Bạn không được phân công vào lớp này.")


def _effective_enrollments(
    db: Session, training_class: TrainingClass, session_date: date
) -> list[ClassEnrollment]:
    enrollments = db.scalars(
        select(ClassEnrollment)
        .options(joinedload(ClassEnrollment.student))
        .where(ClassEnrollment.class_id == training_class.id,
            _effective_enrollment_predicate(session_date))
        .order_by(ClassEnrollment.id)
    ).all()
    return enrollments


def _effective_enrollment_predicate(session_date: date):
    return or_(
        and_(ClassEnrollment.enrollment_status == "active",
            or_(ClassEnrollment.effective_date.is_(None),
                ClassEnrollment.effective_date <= session_date)),
        and_(ClassEnrollment.enrollment_status != "active",
            ClassEnrollment.effective_date > session_date),
    )


def _session_teacher(session: AttendanceSession) -> int:
    return session.teacher_id or session.training_class.teacher_id


def _session_room(session: AttendanceSession) -> str:
    return session.room or session.training_class.room


def _session_payload(session: AttendanceSession) -> dict[str, Any]:
    training_class = session.training_class
    lesson = session.lesson
    teacher_id = _session_teacher(session)
    return {
        "id": session.id,
        "class_id": session.class_id,
        "class_code": training_class.code,
        "class_name": training_class.name,
        "course_id": training_class.course_id,
        "course_code": training_class.course.code,
        "course_name": training_class.course.name,
        "session_date": session.session_date.isoformat(),
        "weekday": DAY_NAMES[session.session_date.weekday()],
        "start_time": session.start_time.strftime("%H:%M"),
        "end_time": session.end_time.strftime("%H:%M"),
        "teacher_id": teacher_id,
        "teacher_name": (
            session.teacher.full_name
            if session.teacher_id is not None and session.teacher is not None
            else training_class.teacher.full_name
        ),
        "room": _session_room(session),
        "lesson_id": session.lesson_id,
        "lesson_sequence": lesson.sequence if lesson else None,
        "lesson_title": lesson.title if lesson else None,
        "is_cancelled": session.is_cancelled,
        "materials_open_early": session.materials_open_early,
        "cancellation_reason": session.cancellation_reason,
        "conflict_override": session.conflict_override,
        "has_attendance": bool(session.records),
    }


def _load_class(db: Session, class_id: int) -> TrainingClass:
    training_class = db.scalar(
        select(TrainingClass)
        .options(
            joinedload(TrainingClass.course),
            joinedload(TrainingClass.teacher),
            selectinload(TrainingClass.assistants),
        )
        .where(TrainingClass.id == class_id)
    )
    if training_class is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy lớp học.")
    return training_class


def _candidate_conflicts(
    db: Session,
    training_class: TrainingClass,
    session_date: date,
    start_time: time,
    end_time: time,
    teacher_id: int,
    room: str,
    exclude_session_id: int | None = None,
) -> list[dict[str, Any]]:
    conflicts: list[dict[str, Any]] = []
    explicit_sessions = db.scalars(
        select(AttendanceSession)
        .options(
            joinedload(AttendanceSession.training_class).joinedload(TrainingClass.course),
            joinedload(AttendanceSession.training_class).joinedload(TrainingClass.teacher),
            joinedload(AttendanceSession.teacher),
            joinedload(AttendanceSession.lesson),
        )
        .join(TrainingClass, AttendanceSession.class_id == TrainingClass.id)
        .where(
            TrainingClass.id != training_class.id,
            AttendanceSession.session_date == session_date,
            AttendanceSession.id != (exclude_session_id or -1),
            AttendanceSession.is_cancelled.is_(False),
            AttendanceSession.start_time < end_time,
            AttendanceSession.end_time > start_time,
        )
    ).unique().all()
    explicit_class_ids = {session.class_id for session in explicit_sessions}
    for session in explicit_sessions:
        same_teacher = _session_teacher(session) == teacher_id
        same_room = _session_room(session).casefold() == room.casefold()
        if same_teacher or same_room:
            conflicts.append(
                {
                    "date": session_date.isoformat(),
                    "class_id": session.class_id,
                    "class_code": session.training_class.code,
                    "session_id": session.id,
                    "start_time": session.start_time.strftime("%H:%M"),
                    "end_time": session.end_time.strftime("%H:%M"),
                    "teacher_conflict": same_teacher,
                    "room_conflict": same_room,
                }
            )

    recurring_classes = db.scalars(
        select(TrainingClass)
        .options(joinedload(TrainingClass.teacher))
        .where(
            TrainingClass.id != training_class.id,
            TrainingClass.status.in_(("preparing", "active")),
            TrainingClass.day_of_week == session_date.weekday(),
            TrainingClass.start_time < end_time,
            TrainingClass.end_time > start_time,
        )
    ).all()
    for other_class in recurring_classes:
        if other_class.id in explicit_class_ids:
            continue
        same_teacher = other_class.teacher_id == teacher_id
        same_room = other_class.room.casefold() == room.casefold()
        if same_teacher or same_room:
            conflicts.append(
                {
                    "date": session_date.isoformat(),
                    "class_id": other_class.id,
                    "class_code": other_class.code,
                    "session_id": None,
                    "start_time": other_class.start_time.strftime("%H:%M"),
                    "end_time": other_class.end_time.strftime("%H:%M"),
                    "teacher_conflict": same_teacher,
                    "room_conflict": same_room,
                }
            )
    return conflicts


def _generate_preview(
    db: Session, training_class: TrainingClass, request: ScheduleGenerateRequest
) -> list[dict[str, Any]]:
    if training_class.status not in {"preparing", "active"}:
        raise HTTPException(status_code=409, detail="Lớp không ở trạng thái được xếp lịch.")
    teacher_id = request.teacher_id or training_class.teacher_id
    room = request.room or training_class.room
    start_time = request.start_time or training_class.start_time
    end_time = request.end_time or training_class.end_time
    if end_time <= start_time:
        raise HTTPException(status_code=422, detail="Giờ kết thúc phải sau giờ bắt đầu.")
    teacher = db.get(User, teacher_id)
    if teacher is None or "teacher" not in {role.code for role in teacher.roles}:
        raise HTTPException(status_code=422, detail="Giảng viên được phân công không hợp lệ.")

    lessons = db.scalars(
        select(CourseLesson)
        .where(CourseLesson.course_id == training_class.course_id)
        .order_by(CourseLesson.sequence)
    ).all()
    holidays = {
        holiday.holiday_date: holiday.name
        for holiday in db.scalars(
            select(AttendanceScheduleHoliday).where(
                AttendanceScheduleHoliday.holiday_date >= request.start_date
            )
        ).all()
    }

    sessions = []
    current = request.start_date
    examined_days = 0
    while len(sessions) < request.session_count and examined_days < 3650:
        examined_days += 1
        if current.weekday() not in request.weekdays:
            current += timedelta(days=1)
            continue
        if current in holidays:
            current += timedelta(days=1)
            continue
        if db.scalar(
            select(AttendanceSession.id).where(
                AttendanceSession.class_id == training_class.id,
                AttendanceSession.session_date == current,
            )
        ):
            raise HTTPException(
                status_code=409,
                detail=f"Lớp đã có buổi học ngày {current.isoformat()}.",
            )
        lesson = lessons[len(sessions)] if len(sessions) < len(lessons) else None
        conflicts = _candidate_conflicts(
            db,
            training_class,
            current,
            start_time,
            end_time,
            teacher_id,
            room,
        )
        sessions.append(
            {
                "session_date": current.isoformat(),
                "weekday": DAY_NAMES[current.weekday()],
                "start_time": start_time.strftime("%H:%M"),
                "end_time": end_time.strftime("%H:%M"),
                "teacher_id": teacher_id,
                "teacher_name": teacher.full_name,
                "room": room,
                "lesson_id": lesson.id if lesson else None,
                "lesson_sequence": lesson.sequence if lesson else len(sessions) + 1,
                "lesson_title": lesson.title if lesson else f"Buổi {len(sessions) + 1}: Chuyên đề & Thực hành",
                "conflicts": conflicts,
            }
        )
        current += timedelta(days=1)
    if len(sessions) != request.session_count:
        raise HTTPException(status_code=422, detail="Không thể sinh đủ số buổi yêu cầu.")
    return sessions


def _notify_schedule_change(
    db: Session,
    training_class: TrainingClass,
    user_ids: set[int],
    title: str,
    message: str,
    teacher_id: int | None = None,
) -> None:
    recipients = user_ids | {training_class.teacher_id}
    if teacher_id is not None:
        recipients.add(teacher_id)
    for user_id in recipients:
        notification = Notification(
            user_id=user_id,
            class_id=training_class.id,
            title=title,
            message=message,
        )
        db.add(notification)
        enqueue_notification_email(db, notification, db.get(User, user_id))


@router.get("/schedule/holidays")
def list_schedule_holidays(
    response: Response,
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=20, ge=1, le=100),
    start_date: date | None = None,
    end_date: date | None = None,
    _user: User = Depends(require_roles(*MANAGER_ROLES)),
    db: Session = Depends(get_db),
) -> list[dict]:
    query = select(AttendanceScheduleHoliday)
    if start_date:
        query = query.where(AttendanceScheduleHoliday.holiday_date >= start_date)
    if end_date:
        query = query.where(AttendanceScheduleHoliday.holiday_date <= end_date)
    total = db.scalar(select(func.count()).select_from(query.subquery())) or 0
    rows = db.scalars(query.order_by(AttendanceScheduleHoliday.holiday_date,
        AttendanceScheduleHoliday.id).offset((page - 1) * page_size).limit(page_size)).all()
    set_pagination_headers(response, total=total, page=page, page_size=page_size)
    return [
        {"id": item.id, "holiday_date": item.holiday_date.isoformat(), "name": item.name}
        for item in rows
    ]


@router.post("/schedule/holidays", status_code=status.HTTP_201_CREATED)
def create_schedule_holiday(
    request: ScheduleHolidayCreate,
    user: User = Depends(require_roles(*MANAGER_ROLES)),
    db: Session = Depends(get_db),
) -> dict:
    if db.scalar(
        select(AttendanceScheduleHoliday.id).where(
            AttendanceScheduleHoliday.holiday_date == request.holiday_date
        )
    ):
        raise HTTPException(status_code=409, detail="Ngày nghỉ đã được khai báo.")
    holiday = AttendanceScheduleHoliday(
        holiday_date=request.holiday_date,
        name=request.name,
        created_by=user.id,
    )
    db.add(holiday)
    db.commit()
    db.refresh(holiday)
    return {"id": holiday.id, "holiday_date": holiday.holiday_date.isoformat(), "name": holiday.name}


@router.delete("/schedule/holidays/{holiday_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_schedule_holiday(
    holiday_id: int,
    _user: User = Depends(require_roles(*MANAGER_ROLES)),
    db: Session = Depends(get_db),
) -> None:
    holiday = db.get(AttendanceScheduleHoliday, holiday_id)
    if holiday is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy ngày nghỉ.")
    db.delete(holiday)
    db.commit()


@router.post("/classes/{class_id}/sessions/preview")
def preview_class_sessions(
    class_id: int,
    request: ScheduleGenerateRequest,
    _manager: User = Depends(require_roles(*MANAGER_ROLES)),
    db: Session = Depends(get_db),
) -> dict:
    training_class = _load_class(db, class_id)
    sessions = _generate_preview(db, training_class, request)
    conflicts = [conflict for session in sessions for conflict in session["conflicts"]]
    return {
        "class_id": class_id,
        "session_count": len(sessions),
        "conflict_count": len(conflicts),
        "sessions": sessions,
    }


@router.post("/classes/{class_id}/sessions", status_code=status.HTTP_201_CREATED)
def create_class_sessions(
    class_id: int,
    request: ScheduleGenerateRequest,
    manager: User = Depends(require_roles(*MANAGER_ROLES)),
    db: Session = Depends(get_db),
) -> dict:
    training_class = _load_class(db, class_id)
    sessions = _generate_preview(db, training_class, request)
    conflicts = [conflict for session in sessions for conflict in session["conflicts"]]
    if conflicts and not request.allow_conflicts:
        raise HTTPException(
            status_code=409,
            detail={
                "message": "Có xung đột giảng viên hoặc phòng học.",
                "conflicts": conflicts,
                "override_required": True,
            },
        )
    records = [
        AttendanceSession(
            class_id=class_id,
            session_date=date.fromisoformat(item["session_date"]),
            start_time=time.fromisoformat(item["start_time"]),
            end_time=time.fromisoformat(item["end_time"]),
            teacher_id=item["teacher_id"],
            room=item["room"],
            lesson_id=item["lesson_id"],
            conflict_override=bool(item["conflicts"]),
        )
        for item in sessions
    ]
    db.add_all(records)
    try:
        db.flush()
        for item in records:
            enrolled_student_ids = {
                enrollment.student_id
                for enrollment in _effective_enrollments(
                    db, training_class, item.session_date
                )
            }
            _notify_schedule_change(
                db,
                training_class,
                enrolled_student_ids,
                "Thời khóa biểu mới",
                f"{training_class.code}: buổi học ngày {item.session_date.isoformat()} đã được xếp lịch.",
            )
        db.add(
            AuditLog(
                actor_id=manager.id,
                action="schedule.generated",
                details=f"class_id={class_id}; sessions={len(records)}; overridden={bool(conflicts)}",
            )
        )
        db.commit()
    except IntegrityError as error:
        db.rollback()
        raise HTTPException(status_code=409, detail="Lịch có buổi trùng ngày; chưa lưu buổi nào.") from error
    return {
        "class_id": class_id,
        "created_count": len(records),
        "sessions": [
            {
                "id": item.id,
                "session_date": item.session_date.isoformat(),
                "lesson_id": item.lesson_id,
                "conflict_override": item.conflict_override,
            }
            for item in records
        ],
    }


def _schedule_query(db: Session):
    return select(AttendanceSession).options(
        joinedload(AttendanceSession.training_class).joinedload(TrainingClass.course),
        joinedload(AttendanceSession.training_class).joinedload(TrainingClass.teacher),
        joinedload(AttendanceSession.teacher),
        joinedload(AttendanceSession.lesson),
        selectinload(AttendanceSession.records),
    )


@router.get("/schedule/sessions")
def list_schedule_sessions(
    start_date: date,
    end_date: date,
    response: Response,
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=20, ge=1, le=100),
    class_id: int | None = None,
    teacher_id: int | None = None,
    room: str | None = Query(default=None, max_length=80),
    user: User = Depends(require_roles("admin", "manager", "admissions", "accountant", "teacher", "ta", "student")),
    db: Session = Depends(get_db),
) -> list[dict]:
    if end_date < start_date or (end_date - start_date).days > 370:
        raise HTTPException(status_code=422, detail="Khoảng ngày không hợp lệ (tối đa 371 ngày).")
    query = _schedule_query(db).where(
        AttendanceSession.session_date.between(start_date, end_date)
    )
    roles = {role.code for role in user.roles}
    if class_id:
        query = query.where(AttendanceSession.class_id == class_id)
    if roles.intersection({"teacher", "ta"}):
        query = query.join(TrainingClass, AttendanceSession.class_id == TrainingClass.id).where(
            or_(
                TrainingClass.teacher_id == user.id,
                TrainingClass.assistant_id == user.id,
                TrainingClass.assistants.any(User.id == user.id),
                AttendanceSession.teacher_id == user.id,
            )
        )
    elif "student" in roles and not roles.intersection(MANAGER_ROLES):
        query = query.join(TrainingClass, AttendanceSession.class_id == TrainingClass.id).join(
            ClassEnrollment, ClassEnrollment.class_id == TrainingClass.id
        ).where(
            ClassEnrollment.student_id == user.id,
            or_(
                and_(
                    ClassEnrollment.enrollment_status == "active",
                    or_(
                        ClassEnrollment.effective_date.is_(None),
                        ClassEnrollment.effective_date <= AttendanceSession.session_date,
                    ),
                ),
                and_(
                    ClassEnrollment.enrollment_status != "active",
                    ClassEnrollment.effective_date > AttendanceSession.session_date,
                ),
            ),
        )
    if teacher_id:
        query = query.where(
            or_(
                AttendanceSession.teacher_id == teacher_id,
                and_(
                    AttendanceSession.teacher_id.is_(None),
                    AttendanceSession.class_id.in_(
                        select(TrainingClass.id).where(TrainingClass.teacher_id == teacher_id)
                    ),
                ),
            )
        )
    if room:
        query = query.where(
            or_(
                AttendanceSession.room == room,
                and_(
                    AttendanceSession.room.is_(None),
                    AttendanceSession.class_id.in_(
                        select(TrainingClass.id).where(TrainingClass.room == room)
                    ),
                ),
            )
        )
    total = db.scalar(select(func.count()).select_from(
        query.with_only_columns(AttendanceSession.id).order_by(None).distinct().subquery()
    )) or 0
    sessions = db.scalars(query.order_by(AttendanceSession.session_date,
        AttendanceSession.start_time, AttendanceSession.id)
        .offset((page - 1) * page_size).limit(page_size)).unique().all()
    set_pagination_headers(response, total=total, page=page, page_size=page_size)
    return [_session_payload(session) for session in sessions]


@router.get("/schedule/conflicts")
def list_schedule_conflicts(
    start_date: date | None = None,
    end_date: date | None = None,
    _manager: User = Depends(require_roles(*MANAGER_ROLES)),
    db: Session = Depends(get_db),
) -> dict:
    today = date.today()
    actual_start = start_date or (today - timedelta(days=30))
    actual_end = end_date or (today + timedelta(days=60))
    if actual_end < actual_start:
        raise HTTPException(status_code=422, detail="end_date must be on or after start_date.")
    if (actual_end - actual_start).days > 90:
        raise HTTPException(
            status_code=422,
            detail="Conflict checks are limited to a 90-day window; narrow the requested date range.",
        )

    query = _schedule_query(db).where(
        AttendanceSession.session_date.between(actual_start, actual_end),
        AttendanceSession.is_cancelled.is_(False),
    )
    sessions = db.scalars(
        query.order_by(AttendanceSession.session_date, AttendanceSession.start_time)
    ).unique().all()

    conflicts: list[dict[str, Any]] = []
    sessions_by_date: dict[date, list[AttendanceSession]] = {}
    for s in sessions:
        sessions_by_date.setdefault(s.session_date, []).append(s)

    for session_date, day_sessions in sessions_by_date.items():
        n = len(day_sessions)
        for i in range(n):
            s1 = day_sessions[i]
            t1 = _session_teacher(s1)
            r1 = _session_room(s1).strip()
            for j in range(i + 1, n):
                s2 = day_sessions[j]
                if s1.class_id == s2.class_id:
                    continue
                if s1.start_time < s2.end_time and s1.end_time > s2.start_time:
                    t2 = _session_teacher(s2)
                    r2 = _session_room(s2).strip()
                    teacher_conflict = (t1 == t2)
                    room_conflict = bool(r1 and r2 and r1.casefold() == r2.casefold())
                    if teacher_conflict or room_conflict:
                        conflict_type = (
                            "both" if (teacher_conflict and room_conflict)
                            else ("teacher" if teacher_conflict else "room")
                        )
                        overridden = bool(s1.conflict_override or s2.conflict_override)
                        teacher_name = (
                            s1.teacher.full_name
                            if s1.teacher
                            else s1.training_class.teacher.full_name
                        )
                        entity_name = (
                            teacher_name if conflict_type == "teacher"
                            else (r1 if conflict_type == "room" else f"{teacher_name} & {r1}")
                        )
                        conflicts.append({
                            "id": f"{s1.id}-{s2.id}",
                            "date": session_date.isoformat(),
                            "weekday": DAY_NAMES[session_date.weekday()],
                            "conflict_type": conflict_type,
                            "teacher_conflict": teacher_conflict,
                            "room_conflict": room_conflict,
                            "conflict_entity": entity_name,
                            "status": "overridden" if overridden else "unresolved",
                            "severity": "medium" if overridden else "high",
                            "overlap_time": f"{max(s1.start_time, s2.start_time).strftime('%H:%M')} – {min(s1.end_time, s2.end_time).strftime('%H:%M')}",
                            "session_1": _session_payload(s1),
                            "session_2": _session_payload(s2),
                        })

    return {
        "start_date": actual_start.isoformat(),
        "end_date": actual_end.isoformat(),
        "total_conflicts": len(conflicts),
        "unresolved_count": sum(1 for c in conflicts if c["status"] == "unresolved"),
        "overridden_count": sum(1 for c in conflicts if c["status"] == "overridden"),
        "teacher_conflicts": sum(1 for c in conflicts if c["teacher_conflict"]),
        "room_conflicts": sum(1 for c in conflicts if c["room_conflict"]),
        "conflicts": conflicts,
    }


@router.get("/me/schedule")
def personal_schedule(
    response: Response,
    start_date: date = Query(default_factory=date.today),
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=20, ge=1, le=100),
    end_date: date | None = None,
    class_id: int | None = Query(default=None, ge=1),
    user: User = Depends(require_roles("student")),
    db: Session = Depends(get_db),
) -> list[dict]:
    actual_end = end_date or start_date + timedelta(days=6)
    query = _schedule_query(db).join(
        ClassEnrollment, ClassEnrollment.class_id == AttendanceSession.class_id
    ).where(
        ClassEnrollment.student_id == user.id,
        or_(
            and_(
                ClassEnrollment.enrollment_status == "active",
                or_(
                    ClassEnrollment.effective_date.is_(None),
                    ClassEnrollment.effective_date <= AttendanceSession.session_date,
                ),
            ),
            and_(
                ClassEnrollment.enrollment_status != "active",
                ClassEnrollment.effective_date > AttendanceSession.session_date,
            ),
        ),
        AttendanceSession.session_date.between(start_date, actual_end),
    )
    if class_id is not None:
        query = query.where(AttendanceSession.class_id == class_id)
    total = db.scalar(select(func.count()).select_from(
        query.with_only_columns(AttendanceSession.id).order_by(None).distinct().subquery()
    )) or 0
    sessions = db.scalars(query.order_by(AttendanceSession.session_date,
        AttendanceSession.start_time, AttendanceSession.id)
        .offset((page - 1) * page_size).limit(page_size)).unique().all()
    set_pagination_headers(response, total=total, page=page, page_size=page_size)
    return [_session_payload(session) for session in sessions]


@router.get("/classes/{class_id}/sessions")
def list_class_sessions(
    class_id: int,
    response: Response,
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=20, ge=1, le=100),
    start_date: date | None = None,
    end_date: date | None = None,
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
) -> list[dict]:
    training_class = _load_class(db, class_id)
    roles = {role.code for role in user.roles}
    session_teacher_only = False
    if not roles.intersection(MANAGER_ROLES + ("admissions", "accountant")):
        if "student" in roles:
            if not db.scalar(
                select(ClassEnrollment.id).where(
                    ClassEnrollment.class_id == class_id,
                    ClassEnrollment.student_id == user.id,
                    ClassEnrollment.enrollment_status == "active",
                )
            ):
                raise HTTPException(status_code=403, detail="Bạn không tham gia lớp này.")
        elif "teacher" in roles and not _has_class_staff_access(user, training_class):
            session_teacher_only = True
        else:
            _assert_class_staff_access(user, training_class)
    query = _schedule_query(db).where(AttendanceSession.class_id == class_id)
    if session_teacher_only:
        query = query.where(AttendanceSession.teacher_id == user.id)
    if start_date:
        query = query.where(AttendanceSession.session_date >= start_date)
    if end_date:
        query = query.where(AttendanceSession.session_date <= end_date)
    total = db.scalar(select(func.count()).select_from(
        query.with_only_columns(AttendanceSession.id).order_by(None).distinct().subquery()
    )) or 0
    sessions = db.scalars(query.order_by(AttendanceSession.session_date,
        AttendanceSession.start_time, AttendanceSession.id)
        .offset((page - 1) * page_size).limit(page_size)).unique().all()
    set_pagination_headers(response, total=total, page=page, page_size=page_size)
    return [
        _session_payload(item)
        for item in sessions
    ]


@router.get("/students/{student_id}/enrollments")
def list_student_enrollments(
    student_id: int,
    response: Response,
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=20, ge=1, le=100),
    _manager: User = Depends(require_roles(*MANAGER_ROLES)),
    db: Session = Depends(get_db),
) -> list[dict]:
    query = select(ClassEnrollment).where(
        ClassEnrollment.student_id == student_id,
        ClassEnrollment.enrollment_status == "active",
    )
    total = db.scalar(select(func.count()).select_from(query.subquery())) or 0
    enrollments = db.scalars(
        query
        .options(
            joinedload(ClassEnrollment.training_class)
            .joinedload(TrainingClass.course)
        )
        .order_by(ClassEnrollment.id)
        .offset((page - 1) * page_size).limit(page_size)
    ).all()
    set_pagination_headers(response, total=total, page=page, page_size=page_size)
    return [
        {
            "class_id": enrollment.class_id,
            "class_code": enrollment.training_class.code,
            "class_name": enrollment.training_class.name,
            "course_name": enrollment.training_class.course.name,
        }
        for enrollment in enrollments
    ]


@router.get("/schedule/sessions/{session_id}")
def get_schedule_session(
    session_id: int,
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
) -> dict:
    session = db.scalar(
        _schedule_query(db).where(AttendanceSession.id == session_id)
    )
    if session is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy buổi học.")
    roles = {role.code for role in user.roles}
    if not roles.intersection(MANAGER_ROLES + ("admissions", "accountant")):
        if "student" in roles:
            enrolled = db.scalar(
                select(ClassEnrollment.id).where(
                    ClassEnrollment.class_id == session.class_id,
                    ClassEnrollment.student_id == user.id,
                )
            )
            if not enrolled:
                raise HTTPException(status_code=403, detail="Bạn không tham gia lớp này.")
        elif not _has_class_staff_access(user, session.training_class) and session.teacher_id != user.id:
            raise HTTPException(status_code=403, detail="Bạn không được phân công vào lớp này.")
    return _session_payload(session)


@router.patch("/schedule/sessions/{session_id}")
def update_schedule_session(
    session_id: int,
    request: ScheduleSessionUpdate,
    manager: User = Depends(require_roles(*MANAGER_ROLES)),
    db: Session = Depends(get_db),
) -> dict:
    session = db.scalar(
        _schedule_query(db).where(AttendanceSession.id == session_id)
    )
    if session is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy buổi học.")
    if session.records:
        raise HTTPException(status_code=409, detail="Buổi học đã điểm danh nên không thể chỉnh sửa.")
    updates = request.model_dump(exclude_unset=True)
    allow_conflicts = updates.pop("allow_conflicts", False)
    if not updates:
        raise HTTPException(status_code=400, detail="Chưa có thay đổi cho buổi học.")
    if updates.get("is_cancelled") is True and not (updates.get("cancellation_reason") or "").strip():
        raise HTTPException(status_code=422, detail="Hủy buổi học bắt buộc nhập lý do.")
    if updates.get("is_cancelled") is False:
        updates["cancellation_reason"] = None
    next_date = updates.get("session_date", session.session_date)
    next_start = updates.get("start_time", session.start_time)
    next_end = updates.get("end_time", session.end_time)
    next_teacher_id = updates.get("teacher_id", _session_teacher(session))
    next_room = updates.get("room", _session_room(session))
    if next_end <= next_start:
        raise HTTPException(status_code=422, detail="Giờ kết thúc phải sau giờ bắt đầu.")
    if "teacher_id" in updates:
        teacher = db.get(User, next_teacher_id)
        if teacher is None or "teacher" not in {role.code for role in teacher.roles}:
            raise HTTPException(status_code=422, detail="Giảng viên dạy thay không hợp lệ.")
    conflicts = [] if updates.get("is_cancelled") is True else _candidate_conflicts(
        db,
        session.training_class,
        next_date,
        next_start,
        next_end,
        next_teacher_id,
        next_room,
        exclude_session_id=session.id,
    )
    if conflicts and not allow_conflicts:
        raise HTTPException(
            status_code=409,
            detail={
                "message": "Buổi học bị trùng giảng viên hoặc phòng.",
                "conflicts": conflicts,
                "override_required": True,
            },
        )
    for field, value in updates.items():
        setattr(session, field, value)
    session.conflict_override = bool(conflicts)
    active_students = {
        enrollment.student_id
        for enrollment in _effective_enrollments(db, session.training_class, next_date)
    }
    title = "Buổi học đã bị hủy" if session.is_cancelled else "Thời khóa biểu được cập nhật"
    _notify_schedule_change(
        db,
        session.training_class,
        active_students,
        title,
        f"{session.training_class.code}, ngày {next_date.isoformat()}: "
        f"{session.cancellation_reason if session.is_cancelled else 'lịch học đã thay đổi.'}",
        teacher_id=next_teacher_id,
    )
    try:
        db.commit()
    except IntegrityError as error:
        db.rollback()
        raise HTTPException(status_code=409, detail="Lớp đã có buổi học trong ngày này.") from error
    db.refresh(session)
    return _session_payload(session)


@router.get("/attendance/sessions/{session_id}")
def get_session_attendance(
    session_id: int,
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=20, ge=1, le=100),
    user: User = Depends(require_roles(*STAFF_ROLES)),
    db: Session = Depends(get_db),
) -> dict:
    session = db.scalar(
        _schedule_query(db).where(AttendanceSession.id == session_id)
    )
    if session is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy buổi học.")
    _assert_class_staff_access(user, session.training_class, session)
    enrollment_query = select(ClassEnrollment).where(
        ClassEnrollment.class_id == session.class_id,
        _effective_enrollment_predicate(session.session_date),
    )
    total = db.scalar(select(func.count()).select_from(enrollment_query.subquery())) or 0
    enrollments = db.scalars(enrollment_query.options(
        joinedload(ClassEnrollment.student)
    ).order_by(ClassEnrollment.id).offset((page - 1) * page_size).limit(page_size)).all()
    records_by_student = {record.student_id: record for record in session.records}
    return {
        "session": _session_payload(session),
        **page_metadata(total=total, page=page, page_size=page_size),
        "students": [
            {
                "id": enrollment.student_id,
                "student_code": enrollment.student.student_code or enrollment.student.username,
                "full_name": enrollment.student.full_name,
                "status": (
                    records_by_student[enrollment.student_id].status
                    if enrollment.student_id in records_by_student
                    and records_by_student[enrollment.student_id].status != "not_recorded"
                    else None
                ),
            }
            for enrollment in enrollments
        ],
    }


@router.get("/classes/{class_id}/attendance/summary")
def get_class_attendance_summary(
    class_id: int,
    q: str = Query(default="", max_length=150),
    attendance_status: str = Query(
        default="all",
        alias="status",
        pattern="^(all|present|late|excused_absent|unexcused_absent|absent|not_recorded)$",
    ),
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=20, ge=1, le=100),
    sort: str = Query(default="attendance_asc", pattern="^(attendance_asc|attendance_desc)$"),
    user: User = Depends(require_roles(*STAFF_ROLES)),
    db: Session = Depends(get_db),
) -> dict:
    training_class = db.scalar(
        select(TrainingClass)
        .options(
            load_only(
                TrainingClass.id,
                TrainingClass.code,
                TrainingClass.name,
                TrainingClass.teacher_id,
                TrainingClass.assistant_id,
                TrainingClass.program_id,
            ),
            selectinload(TrainingClass.assistants).load_only(User.id),
        )
        .where(TrainingClass.id == class_id)
    )
    if training_class is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy lớp học.")
    _assert_class_staff_access(user, training_class)

    local_now = datetime.now(HO_CHI_MINH)
    eligible_sessions = (
        select(
            AttendanceSession.id.label("id"),
            AttendanceSession.session_date.label("session_date"),
            AttendanceSession.start_time.label("start_time"),
            AttendanceSession.end_time.label("end_time"),
        )
        .where(
            AttendanceSession.class_id == class_id,
            AttendanceSession.is_cancelled.is_(False),
            or_(
                AttendanceSession.session_date < local_now.date(),
                and_(AttendanceSession.session_date == local_now.date(),
                    AttendanceSession.end_time < local_now.time().replace(tzinfo=None)),
            ),
        )
        .subquery()
    )
    sessions = db.execute(select(eligible_sessions).order_by(
        eligible_sessions.c.session_date, eligible_sessions.c.start_time, eligible_sessions.c.id
    )).all()
    session_ids = [row.id for row in sessions]
    eligible_enrollment = or_(
        and_(ClassEnrollment.enrollment_status == "active",
            or_(ClassEnrollment.effective_date.is_(None),
                ClassEnrollment.effective_date <= eligible_sessions.c.session_date)),
        and_(ClassEnrollment.enrollment_status != "active",
            ClassEnrollment.effective_date.is_not(None),
            eligible_sessions.c.session_date < ClassEnrollment.effective_date),
    )
    present_count = func.sum(case((AttendanceRecord.status == "present", 1), else_=0))
    late_count = func.sum(case((AttendanceRecord.status == "late", 1), else_=0))
    excused_count = func.sum(case((AttendanceRecord.status == "excused_absent", 1), else_=0))
    unexcused_count = func.sum(case((AttendanceRecord.status == "unexcused_absent", 1), else_=0))
    legacy_count = func.sum(case((AttendanceRecord.status == "absent", 1), else_=0))
    enrollment_stats = (
        select(
            ClassEnrollment.id.label("enrollment_id"),
            ClassEnrollment.student_id.label("student_id"),
            ClassEnrollment.enrollment_status.label("enrollment_status"),
            ClassEnrollment.effective_date.label("effective_date"),
            User.full_name.label("full_name"),
            func.coalesce(User.student_code, User.username).label("student_code"),
            func.count(eligible_sessions.c.id).label("total_sessions"),
            func.coalesce(present_count, 0).label("present_count"),
            func.coalesce(late_count, 0).label("late_count"),
            func.coalesce(excused_count, 0).label("excused_absent_count"),
            func.coalesce(unexcused_count, 0).label("unexcused_absent_count"),
            func.coalesce(legacy_count, 0).label("legacy_absent_count"),
        )
        .join(User, User.id == ClassEnrollment.student_id)
        .outerjoin(eligible_sessions, and_(eligible_enrollment, eligible_sessions.c.id.is_not(None)))
        .outerjoin(AttendanceRecord, and_(
            AttendanceRecord.session_id == eligible_sessions.c.id,
            AttendanceRecord.student_id == ClassEnrollment.student_id,
        ))
        .where(ClassEnrollment.class_id == class_id)
        .group_by(ClassEnrollment.id, ClassEnrollment.student_id,
            ClassEnrollment.enrollment_status, ClassEnrollment.effective_date,
            User.full_name, User.student_code, User.username)
        .subquery()
    )
    included = or_(
        and_(enrollment_stats.c.enrollment_status == "active",
            or_(enrollment_stats.c.effective_date.is_(None),
                enrollment_stats.c.effective_date <= local_now.date())),
        enrollment_stats.c.total_sessions > 0,
    )
    absent_count = (enrollment_stats.c.excused_absent_count
        + enrollment_stats.c.unexcused_absent_count + enrollment_stats.c.legacy_absent_count)
    attendance_rate = case(
        (enrollment_stats.c.total_sessions > 0,
            func.round((enrollment_stats.c.present_count + enrollment_stats.c.late_count)
                * 100.0 / func.nullif(enrollment_stats.c.total_sessions, 0), 1)),
        else_=None,
    )
    filtered = select(enrollment_stats, absent_count.label("absent_count"),
        attendance_rate.label("attendance_rate")).where(included)
    if q.strip():
        search = f"%{q.strip()}%"
        filtered = filtered.where(or_(enrollment_stats.c.full_name.ilike(search),
            enrollment_stats.c.student_code.ilike(search),
            cast(enrollment_stats.c.student_id, String).ilike(search)))
    if attendance_status == "present":
        filtered = filtered.where(enrollment_stats.c.present_count > 0)
    elif attendance_status == "late":
        filtered = filtered.where(enrollment_stats.c.late_count > 0)
    elif attendance_status == "excused_absent":
        filtered = filtered.where(enrollment_stats.c.excused_absent_count > 0)
    elif attendance_status == "unexcused_absent":
        filtered = filtered.where(enrollment_stats.c.unexcused_absent_count > 0)
    elif attendance_status == "absent":
        filtered = filtered.where(enrollment_stats.c.legacy_absent_count > 0)
    elif attendance_status == "not_recorded":
        recorded = (enrollment_stats.c.present_count + enrollment_stats.c.late_count
            + enrollment_stats.c.excused_absent_count + enrollment_stats.c.unexcused_absent_count
            + enrollment_stats.c.legacy_absent_count)
        filtered = filtered.where(enrollment_stats.c.total_sessions > recorded)
    total_students = db.scalar(select(func.count()).select_from(
        select(enrollment_stats.c.student_id).where(included).subquery())) or 0
    total_filtered = db.scalar(select(func.count()).select_from(filtered.subquery())) or 0
    totals_row = db.execute(select(
        func.coalesce(func.sum(enrollment_stats.c.present_count + enrollment_stats.c.late_count), 0),
        func.coalesce(func.sum(absent_count), 0),
        func.coalesce(func.sum(enrollment_stats.c.excused_absent_count), 0),
        func.coalesce(func.sum(enrollment_stats.c.unexcused_absent_count), 0),
        func.coalesce(func.sum(enrollment_stats.c.legacy_absent_count), 0),
        func.coalesce(func.sum(enrollment_stats.c.total_sessions
            - enrollment_stats.c.present_count - enrollment_stats.c.late_count
            - enrollment_stats.c.excused_absent_count - enrollment_stats.c.unexcused_absent_count
            - enrollment_stats.c.legacy_absent_count), 0),
    ).where(included)).one()
    if sort == "attendance_desc":
        sort_rate = attendance_rate.desc()
    else:
        sort_rate = attendance_rate.asc()
    actual_page = page or 1
    actual_page_size = page_size or max(total_filtered, 1)
    page_rows = db.execute(filtered.order_by(attendance_rate.is_(None), sort_rate,
            func.lower(enrollment_stats.c.full_name), enrollment_stats.c.student_id)
        .offset((actual_page - 1) * actual_page_size).limit(actual_page_size)).all()
    page_student_ids = [row.student_id for row in page_rows]
    records = db.execute(select(AttendanceRecord.session_id, AttendanceRecord.student_id,
            AttendanceRecord.status)
        .where(AttendanceRecord.session_id.in_(session_ids or [-1]),
            AttendanceRecord.student_id.in_(page_student_ids or [-1]))).all()
    record_by_pair = {(row.session_id, row.student_id): row.status for row in records}
    session_payload = [{"id": row.id, "session_date": row.session_date.isoformat(),
        "start_time": row.start_time.strftime("%H:%M:%S"),
        "end_time": row.end_time.strftime("%H:%M:%S")} for row in sessions]
    program = db.get(StudyProgram, training_class.program_id) if training_class.program_id is not None else None
    absence_threshold = float(program.absence_threshold_percent) if program is not None else 20.0
    students = []
    for row in page_rows:
        student_cells = {}
        for session in sessions:
            active_for_session = (
                (row.enrollment_status == "active" and
                    (row.effective_date is None or row.effective_date <= session.session_date))
                or (row.enrollment_status != "active" and row.effective_date is not None
                    and session.session_date < row.effective_date)
            )
            stored_status = record_by_pair.get((session.id, row.student_id)) if active_for_session else None
            cell_status = None if stored_status in (None, "not_recorded") else stored_status
            student_cells[str(session.id)] = {"eligible": active_for_session, "status": cell_status}
        student_absent = row.absent_count
        absent_rate = round(student_absent / row.total_sessions * 100, 1) if row.total_sessions else None
        warning = attendance_warning(student_absent, row.total_sessions, absence_threshold)
        students.append({
            "id": row.student_id,
            "student_code": row.student_code,
            "full_name": row.full_name,
            "attendance_by_session": student_cells,
            "total_sessions": row.total_sessions,
            "attended_count": row.present_count + row.late_count,
            "present_count": row.present_count,
            "late_count": row.late_count,
            "excused_absent_count": row.excused_absent_count,
            "unexcused_absent_count": row.unexcused_absent_count,
            "legacy_absent_count": row.legacy_absent_count,
            "unrecorded_count": max(0, row.total_sessions - row.present_count - row.late_count
                - row.excused_absent_count - row.unexcused_absent_count - row.legacy_absent_count),
            "absent_count": student_absent,
            "attendance_rate": float(row.attendance_rate) if row.attendance_rate is not None else None,
            "absence_threshold_percent": absence_threshold,
            "absence_rate_percent": absent_rate,
            **warning,
            "at_risk": warning["warning_status"] in ("at_threshold", "over_threshold"),
        })
    response = {
        "class": {
            "id": training_class.id,
            "code": training_class.code,
            "name": training_class.name,
        },
        "sessions": session_payload,
        "total_students": total_students,
        "totals": {
            "attended_count": int(totals_row[0]),
            "absent_count": int(totals_row[1]),
            "excused_absent_count": int(totals_row[2]),
            "unexcused_absent_count": int(totals_row[3]),
            "legacy_absent_count": int(totals_row[4]),
            "unrecorded_count": int(totals_row[5]),
        },
        "students": students,
    }
    if page is not None or page_size is not None:
        response["pagination"] = page_metadata(
            total=total_filtered, page=actual_page, page_size=actual_page_size
        )
    return response


@router.get("/students/me/attendance-history")
def get_my_attendance_history(
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=20, ge=1, le=100),
    class_id: int | None = Query(default=None, ge=1),
    user: User = Depends(require_roles("student")),
    db: Session = Depends(get_db),
) -> dict[str, Any]:
    """Return only the authenticated student's eligible past attendance."""
    enrollments = db.scalars(
        select(ClassEnrollment)
        .options(
            joinedload(ClassEnrollment.training_class).joinedload(TrainingClass.course)
        )
        .where(
            ClassEnrollment.student_id == user.id,
            *([ClassEnrollment.class_id == class_id] if class_id is not None else []),
        )
        .order_by(ClassEnrollment.id)
    ).all()
    class_ids = {enrollment.class_id for enrollment in enrollments}
    if not class_ids:
        return {
            "items": [],
            "classes": [],
            "totals": {"present_count": 0, "late_count": 0, "absent_count": 0, "unrecorded_count": 0, "total_sessions": 0},
            "attendance_rate": None,
            "total": 0,
            "page": page,
            "page_size": page_size,
            "pages": 0,
        }

    sessions = db.scalars(
        select(AttendanceSession)
        .options(
            load_only(
                AttendanceSession.id,
                AttendanceSession.class_id,
                AttendanceSession.session_date,
                AttendanceSession.start_time,
                AttendanceSession.end_time,
                AttendanceSession.is_cancelled,
            )
        )
        .where(AttendanceSession.class_id.in_(class_ids))
    ).all()
    session_ids = [session.id for session in sessions]
    records = (
        db.scalars(
            select(AttendanceRecord)
            .options(
                load_only(
                    AttendanceRecord.session_id,
                    AttendanceRecord.student_id,
                    AttendanceRecord.status,
                )
            )
            .where(
                AttendanceRecord.session_id.in_(session_ids),
                AttendanceRecord.student_id == user.id,
            )
        ).all()
        if session_ids
        else []
    )
    now = datetime.now(timezone.utc)
    sessions_by_class: dict[int, list[AttendanceSession]] = {}
    session_by_id = {session.id: session for session in sessions}
    records_by_class: dict[int, list[AttendanceRecord]] = {}
    for session in sessions:
        sessions_by_class.setdefault(session.class_id, []).append(session)
    for record in records:
        session = session_by_id.get(record.session_id)
        if session is not None:
            records_by_class.setdefault(session.class_id, []).append(record)

    classes: list[dict[str, Any]] = []
    items: list[dict[str, Any]] = []
    totals = {
        "present_count": 0,
        "late_count": 0,
        "absent_count": 0,
        "unrecorded_count": 0,
        "total_sessions": 0,
    }
    program_ids = {enrollment.training_class.program_id for enrollment in enrollments
                   if enrollment.training_class.program_id is not None}
    program_thresholds = dict(db.execute(
        select(StudyProgram.id, StudyProgram.absence_threshold_percent)
        .where(StudyProgram.id.in_(program_ids))
    ).all()) if program_ids else {}
    for enrollment in enrollments:
        training_class = enrollment.training_class
        class_sessions = sessions_by_class.get(training_class.id, [])
        summary = summarize_class_attendance(
            enrollments=[enrollment],
            sessions=class_sessions,
            records=records_by_class.get(training_class.id, []),
            as_of=now,
        )
        if not summary.students:
            continue
        student = summary.students[0]
        threshold = float(program_thresholds.get(training_class.program_id, 20))
        warning = attendance_warning(student.absent_count, student.total_sessions, threshold)
        classes.append(
            {
                "class_id": training_class.id,
                "class_code": training_class.code,
                "class_name": training_class.name,
                "course_name": training_class.course.name,
                "present_count": student.present_count,
                "late_count": student.late_count,
                "absent_count": student.absent_count,
                "unrecorded_count": student.unrecorded_count,
                "total_sessions": student.total_sessions,
                "attendance_rate": student.attendance_rate,
                "absence_threshold_percent": threshold,
                **warning,
            }
        )
        totals["present_count"] += student.present_count
        totals["late_count"] += student.late_count
        totals["absent_count"] += student.absent_count
        totals["unrecorded_count"] += student.unrecorded_count
        totals["total_sessions"] += student.total_sessions
    history_query = select(AttendanceSession, AttendanceRecord, TrainingClass).join(
        TrainingClass, TrainingClass.id == AttendanceSession.class_id
    ).join(ClassEnrollment, ClassEnrollment.class_id == AttendanceSession.class_id).outerjoin(
        AttendanceRecord, (AttendanceRecord.session_id == AttendanceSession.id)
        & (AttendanceRecord.student_id == user.id)
    ).where(AttendanceSession.class_id.in_(class_ids),
        AttendanceSession.is_cancelled.is_(False),
        AttendanceSession.session_date < datetime.now(HO_CHI_MINH).date(),
        ClassEnrollment.student_id == user.id,
        _effective_enrollment_predicate(AttendanceSession.session_date))
    total = db.scalar(select(func.count()).select_from(
        history_query.with_only_columns(AttendanceSession.id).order_by(None).distinct().subquery()
    )) or 0
    page_rows = db.execute(history_query.order_by(AttendanceSession.session_date.desc(),
        AttendanceSession.start_time.desc(), AttendanceSession.id.desc())
        .offset((page - 1) * page_size).limit(page_size)).all()
    items = [{"session_id": session.id, "session_date": session.session_date.isoformat(),
        "start_time": session.start_time.strftime("%H:%M:%S"),
        "end_time": session.end_time.strftime("%H:%M:%S"),
        "status": record.status if record and record.status != "not_recorded" else None,
        "class_id": training_class.id, "class_code": training_class.code,
        "class_name": training_class.name, "course_name": training_class.course.name}
        for session, record, training_class in page_rows]
    classes.sort(key=lambda item: (item["class_code"].casefold(), item["class_id"]))
    attended_count = totals["present_count"] + totals["late_count"]
    attendance_rate = (
        round(attended_count / totals["total_sessions"] * 100, 1)
        if totals["total_sessions"]
        else None
    )
    return {
        "items": items,
        "classes": classes,
        "totals": totals,
        "attendance_rate": attendance_rate,
        "total": total,
        "page": page,
        "page_size": page_size,
        "pages": (total + page_size - 1) // page_size,
        "has_previous": page > 1 and total > 0,
        "has_next": page * page_size < total,
    }


@router.get("/classes/{class_id}/attendance/export")
def export_class_attendance(
    class_id: int,
    user: User = Depends(require_roles(*MANAGER_ROLES)),
    db: Session = Depends(get_db),
) -> Response:
    summary = get_class_attendance_summary(
        class_id=class_id,
        q="",
        attendance_status="all",
        page=None,
        page_size=None,
        sort="attendance_asc",
        user=user,
        db=db,
    )
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Chuyên cần"
    def excel_text(value: str) -> str:
        text = str(value)
        return f"'{text}" if text.lstrip().startswith(("=", "+", "-", "@")) else text

    sheet.append(
        [
            excel_text(
                f"Bảng chuyên cần — {summary['class']['code']} · {summary['class']['name']}"
            )
        ]
    )
    exported_on = datetime.now(timezone.utc).astimezone(HO_CHI_MINH).date()
    sheet.append([f"Ngày xuất: {exported_on.isoformat()}"])
    sheet.append([])

    headers = ["Mã học viên", "Họ tên"]
    headers.extend(
        f"{session['session_date']} {session['start_time'][:5]}–{session['end_time'][:5]}"
        for session in summary["sessions"]
    )
    headers.extend(
        [
            "Có mặt",
            "Đi muộn",
            "Vắng có phép",
            "Vắng không phép",
            "Vắng trạng thái cũ",
            "Chưa ghi nhận",
            "Tổng vắng",
            "Tỷ lệ chuyên cần (%)",
        ]
    )
    sheet.append(headers)

    labels = {
        "present": "Có mặt",
        "late": "Đi muộn",
        "excused_absent": "Vắng có phép",
        "unexcused_absent": "Vắng không phép",
        "absent": "Vắng trạng thái cũ",
    }
    for student in summary["students"]:
        cells = []
        for session in summary["sessions"]:
            cell = student["attendance_by_session"].get(str(session["id"]))
            if not cell or not cell["eligible"]:
                cells.append("Không áp dụng")
            elif cell["status"] is None:
                cells.append("Chưa ghi nhận")
            else:
                cells.append(labels[cell["status"]])
        sheet.append(
            [
                excel_text(student["student_code"] or f"#{student['id']}"),
                excel_text(student["full_name"]),
                *cells,
                student["present_count"],
                student["late_count"],
                student["excused_absent_count"],
                student["unexcused_absent_count"],
                student["legacy_absent_count"],
                student["unrecorded_count"],
                student["absent_count"],
                student["attendance_rate"],
            ]
        )

    last_column = sheet.cell(row=4, column=len(headers)).column_letter
    sheet.freeze_panes = "C5"
    sheet.auto_filter.ref = f"A4:{last_column}{sheet.max_row}"
    for column_cells in sheet.columns:
        column_letter = column_cells[0].column_letter
        max_length = max(len(str(cell.value or "")) for cell in column_cells)
        sheet.column_dimensions[column_letter].width = min(max(max_length + 2, 12), 32)

    output = BytesIO()
    workbook.save(output)
    safe_code = "".join(
        char if char.isalnum() or char in "-_" else "-"
        for char in summary["class"]["code"]
    ).strip("-") or f"class-{class_id}"
    filename = f"{safe_code}-attendance-{exported_on.isoformat()}.xlsx"
    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename, safe='')}"},
    )


@router.get("/students/me/attendance-risk")
def get_my_attendance_risk(
    response: Response,
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=20, ge=1, le=100),
    student: User = Depends(require_roles("student")),
    db: Session = Depends(get_db),
) -> list[dict]:
    query = select(AttendanceRiskState, TrainingClass).join(
        TrainingClass, TrainingClass.id == AttendanceRiskState.class_id
    ).where(AttendanceRiskState.student_id == student.id,
            AttendanceRiskState.is_at_risk.is_(True))
    total = db.scalar(select(func.count()).select_from(
        query.with_only_columns(AttendanceRiskState.class_id).order_by(None).subquery()
    )) or 0
    rows = db.execute(query.order_by(TrainingClass.code, TrainingClass.id)
        .offset((page - 1) * page_size).limit(page_size)).all()
    set_pagination_headers(response, total=total, page=page, page_size=page_size)
    return [{"class_id": training_class.id, "class_code": training_class.code,
             "class_name": training_class.name, "is_at_risk": state.is_at_risk,
             "absent_count": state.absent_count,
             "completed_session_count": state.completed_session_count,
             "absence_threshold_percent": str(state.absence_threshold_percent),
             "flagged_at": state.flagged_at.isoformat() if state.flagged_at else None}
            for state, training_class in rows]


@router.put("/attendance/sessions/{session_id}")
def mark_session_attendance(
    session_id: int,
    request: AttendanceBulkUpdate,
    user: User = Depends(require_roles(*STAFF_ROLES)),
    db: Session = Depends(get_db),
) -> dict:
    session = db.scalar(
        _schedule_query(db).where(AttendanceSession.id == session_id)
    )
    if session is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy buổi học.")
    _assert_class_staff_access(user, session.training_class, session)
    if session.is_cancelled:
        raise HTTPException(status_code=409, detail="Buổi học đã bị hủy.")
    roles = {role.code for role in user.roles}
    if session.session_date < date.today() and not roles.intersection(MANAGER_ROLES):
        raise HTTPException(status_code=403, detail="Chỉ Quản lý đào tạo mới được sửa điểm danh ngày quá khứ.")
    active_student_ids = {
        enrollment.student_id
        for enrollment in _effective_enrollments(db, session.training_class, session.session_date)
    }
    requested_ids = {item.student_id for item in request.records}
    if not requested_ids.issubset(active_student_ids):
        raise HTTPException(status_code=422, detail="Danh sách có học viên không thuộc buổi học.")
    now = datetime.now(timezone.utc).replace(tzinfo=None)
    records_by_student = {record.student_id: record for record in session.records}
    updated_count = 0
    for marked in request.records:
        record = records_by_student.get(marked.student_id)
        if marked.status is None:
            if record is None:
                continue
            db.delete(record)
            enrollment = db.scalar(
                select(ClassEnrollment).where(
                    ClassEnrollment.class_id == session.class_id,
                    ClassEnrollment.student_id == marked.student_id,
                )
            )
            if enrollment:
                enrollment.attendance_status = "not_recorded"
                enrollment.last_attendance_at = None
            updated_count += 1
            continue

        old_status = record.status if record else None
        if record is None:
            record = AttendanceRecord(
                session_id=session.id,
                student_id=marked.student_id,
                marked_by=user.id,
                status=marked.status,
                marked_at=now,
            )
            db.add(record)
            db.flush()
        else:
            record.status = marked.status
            record.marked_by = user.id
            record.marked_at = now
        if old_status != marked.status:
            db.add(
                AttendanceChangeLog(
                    record_id=record.id,
                    changed_by=user.id,
                    old_status=old_status,
                    new_status=marked.status,
                )
            )
        enrollment = db.scalar(
            select(ClassEnrollment).where(
                ClassEnrollment.class_id == session.class_id,
                ClassEnrollment.student_id == marked.student_id,
            )
        )
        if enrollment:
            enrollment.attendance_status = (
                "present" if marked.status == "present"
                else "late" if marked.status == "late"
                else "absent" if marked.status in ("excused_absent", "unexcused_absent")
                else "not_recorded"
            )
            enrollment.last_attendance_at = now
        updated_count += 1
    refresh_attendance_risk_states(db, session.training_class)
    db.commit()
    return {"session_id": session.id, "updated_count": updated_count}


def _absence_payload(request: AbsenceRequest) -> dict:
    return {
        "id": request.id,
        "session_id": request.session_id,
        "student_id": request.student_id,
        "reason": request.reason,
        "evidence_filename": request.evidence_filename,
        "has_evidence": request.evidence_data is not None,
        "status": request.status,
        "decision_reason": request.decision_reason,
        "decided_at": request.decided_at.isoformat() if request.decided_at else None,
        "created_at": request.created_at.isoformat(),
    }


@router.post("/attendance/sessions/{session_id}/absence-requests", status_code=status.HTTP_201_CREATED)
def create_absence_request(
    session_id: int,
    request: AbsenceRequestCreate,
    user: User = Depends(require_roles("student")),
    db: Session = Depends(get_db),
) -> dict:
    session = db.scalar(
        _schedule_query(db).where(AttendanceSession.id == session_id)
    )
    if session is None or session.is_cancelled:
        raise HTTPException(status_code=404, detail="Không tìm thấy buổi học.")
    if user.id not in {
        enrollment.student_id
        for enrollment in _effective_enrollments(
            db, session.training_class, session.session_date
        )
    }:
        raise HTTPException(status_code=403, detail="Bạn không tham gia lớp này.")
    if session.session_date < date.today():
        raise HTTPException(status_code=409, detail="Không thể xin phép cho buổi đã qua.")
    existing_pending = db.scalar(
        select(AbsenceRequest.id).where(
            AbsenceRequest.session_id == session.id,
            AbsenceRequest.student_id == user.id,
            AbsenceRequest.status == "pending",
        )
    )
    if existing_pending:
        raise HTTPException(status_code=409, detail="Đã có đơn chờ duyệt cho buổi học này.")
    evidence_data = None
    if request.evidence_base64:
        try:
            evidence_data = base64.b64decode(request.evidence_base64, validate=True)
        except (ValueError, binascii.Error) as error:
            raise HTTPException(status_code=422, detail="Tệp minh chứng không hợp lệ.") from error
        if len(evidence_data) > 10 * 1024 * 1024:
            raise HTTPException(status_code=413, detail="Tệp minh chứng vượt quá 10 MB.")
        if not request.evidence_filename or request.evidence_content_type not in {
            "application/pdf", "image/jpeg", "image/png", "image/webp",
        }:
            raise HTTPException(status_code=422, detail="Định dạng tệp minh chứng không được hỗ trợ.")
    absence = AbsenceRequest(
        session_id=session.id,
        student_id=user.id,
        reason=request.reason,
        evidence_filename=request.evidence_filename if evidence_data else None,
        evidence_content_type=request.evidence_content_type if evidence_data else None,
        evidence_data=evidence_data,
    )
    db.add(absence)
    db.commit()
    db.refresh(absence)
    return _absence_payload(absence)


@router.get("/me/absence-requests")
def list_my_absence_requests(
    response: Response,
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=20, ge=1, le=100),
    user: User = Depends(require_roles("student")),
    db: Session = Depends(get_db),
) -> list[dict]:
    query = select(AbsenceRequest).where(AbsenceRequest.student_id == user.id)
    total = db.scalar(select(func.count()).select_from(query.subquery())) or 0
    requests = db.scalars(query.order_by(AbsenceRequest.created_at.desc(), AbsenceRequest.id.desc())
        .offset((page - 1) * page_size).limit(page_size)).all()
    set_pagination_headers(response, total=total, page=page, page_size=page_size)
    return [_absence_payload(item) for item in requests]


@router.get("/classes/{class_id}/absence-requests")
def list_class_absence_requests(
    class_id: int,
    response: Response,
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=20, ge=1, le=100),
    user: User = Depends(require_roles(*STAFF_ROLES)),
    db: Session = Depends(get_db),
) -> list[dict]:
    training_class = _load_class(db, class_id)
    session_teacher_only = (
        "teacher" in {role.code for role in user.roles}
        and not _has_class_staff_access(user, training_class)
    )
    if not session_teacher_only:
        _assert_class_staff_access(user, training_class)
    query = (
        select(AbsenceRequest)
        .join(AttendanceSession, AbsenceRequest.session_id == AttendanceSession.id)
        .where(AttendanceSession.class_id == class_id)
    )
    if session_teacher_only:
        has_assignment = db.scalar(
            select(AttendanceSession.id).where(
                AttendanceSession.class_id == class_id,
                AttendanceSession.teacher_id == user.id,
            ).limit(1)
        )
        if has_assignment is None:
            raise HTTPException(
                status_code=403, detail="Bạn không được phân công vào lớp này."
            )
        query = query.where(AttendanceSession.teacher_id == user.id)
    total = db.scalar(select(func.count()).select_from(
        query.with_only_columns(AbsenceRequest.id).order_by(None).distinct().subquery()
    )) or 0
    requests = db.scalars(query.order_by(AbsenceRequest.created_at.desc(), AbsenceRequest.id.desc())
        .offset((page - 1) * page_size).limit(page_size)).all()
    set_pagination_headers(response, total=total, page=page, page_size=page_size)
    return [_absence_payload(item) for item in requests]


@router.get("/absence-requests/{request_id}/evidence")
def get_absence_evidence(
    request_id: int,
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
) -> Response:
    absence = db.get(AbsenceRequest, request_id)
    if absence is None or absence.evidence_data is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy tệp minh chứng.")
    if absence.student_id != user.id:
        session = db.get(AttendanceSession, absence.session_id)
        if session is None:
            raise HTTPException(status_code=404, detail="Không tìm thấy buổi học.")
        training_class = _load_class(db, session.class_id)
        _assert_class_staff_access(user, training_class, session)
    return Response(
        content=absence.evidence_data,
        media_type=absence.evidence_content_type or "application/octet-stream",
        headers={"Content-Disposition": f'attachment; filename="{absence.evidence_filename or "evidence"}"'},
    )


@router.patch("/absence-requests/{request_id}")
def decide_absence_request(
    request_id: int,
    decision: AbsenceRequestDecision,
    user: User = Depends(require_roles(*STAFF_ROLES)),
    db: Session = Depends(get_db),
) -> dict:
    absence = db.get(AbsenceRequest, request_id)
    if absence is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy đơn xin vắng.")
    session = db.scalar(
        _schedule_query(db).where(AttendanceSession.id == absence.session_id)
    )
    _assert_class_staff_access(user, session.training_class, session)
    if absence.status != "pending":
        raise HTTPException(status_code=409, detail="Đơn đã được xử lý.")
    now = datetime.now(timezone.utc).replace(tzinfo=None)
    absence.status = decision.status
    absence.decision_reason = decision.decision_reason
    absence.decision_by = user.id
    absence.decided_at = now
    if decision.status == "approved":
        record = db.scalar(
            select(AttendanceRecord).where(
                AttendanceRecord.session_id == session.id,
                AttendanceRecord.student_id == absence.student_id,
            )
        )
        old_status = record.status if record else None
        if record is None:
            record = AttendanceRecord(
                session_id=session.id,
                student_id=absence.student_id,
                marked_by=user.id,
                status="excused_absent",
                marked_at=now,
            )
            db.add(record)
            db.flush()
        else:
            record.status = "excused_absent"
            record.marked_by = user.id
            record.marked_at = now
        if old_status != "excused_absent":
            db.add(
                AttendanceChangeLog(
                    record_id=record.id,
                    changed_by=user.id,
                    old_status=old_status,
                    new_status="excused_absent",
                )
            )
        enrollment = db.scalar(
            select(ClassEnrollment).where(
                ClassEnrollment.class_id == session.class_id,
                ClassEnrollment.student_id == absence.student_id,
            )
        )
        if enrollment:
            enrollment.attendance_status = "absent"
            enrollment.last_attendance_at = now
    notification = Notification(
        user_id=absence.student_id,
        class_id=session.class_id,
        title="Đơn xin vắng đã được xử lý",
        message=f"{session.training_class.code} · {session.session_date}: {decision.status}. {decision.decision_reason}",
    )
    db.add(notification)
    enqueue_notification_email(db, notification, db.get(User, absence.student_id))
    db.commit()
    db.refresh(absence)
    return _absence_payload(absence)


@router.post("/student-enrollment-changes")
def process_student_enrollment_change(
    request: EnrollmentChangeCreate,
    user: User = Depends(require_roles(*MANAGER_ROLES)),
    db: Session = Depends(get_db),
) -> dict:
    old_enrollment = db.scalar(
        select(ClassEnrollment).where(
            ClassEnrollment.class_id == request.class_id,
            ClassEnrollment.student_id == request.student_id,
        )
    )
    if old_enrollment is None or old_enrollment.enrollment_status != "active":
        raise HTTPException(status_code=404, detail="Không tìm thấy ghi danh đang hoạt động.")
    student = db.get(User, request.student_id)
    if student is None or not any(role.code == "student" for role in student.roles):
        raise HTTPException(status_code=422, detail="Tài khoản không phải học viên.")
    target_class = None
    if request.action == "transfer":
        target_class = db.scalar(
            select(TrainingClass)
            .options(selectinload(TrainingClass.enrollments))
            .where(TrainingClass.id == request.target_class_id)
        )
        if (
            target_class is None
            or target_class.status not in {"preparing", "active"}
            or target_class.id == request.class_id
        ):
            raise HTTPException(status_code=422, detail="Lớp đích không hợp lệ.")
        if len(
            [item for item in target_class.enrollments if item.enrollment_status == "active"]
        ) >= target_class.capacity:
            raise HTTPException(status_code=409, detail="Lớp đích đã đủ sức chứa.")
        target_enrollment = db.scalar(
            select(ClassEnrollment).where(
                ClassEnrollment.class_id == target_class.id,
                ClassEnrollment.student_id == student.id,
            )
        )
        if target_enrollment and target_enrollment.enrollment_status == "active":
            raise HTTPException(status_code=409, detail="Học viên đã ghi danh lớp đích.")
        if target_enrollment:
            target_enrollment.enrollment_status = "active"
            target_enrollment.effective_date = request.effective_date
        else:
            db.add(
                ClassEnrollment(
                    class_id=target_class.id,
                    student_id=student.id,
                    enrollment_status="active",
                    effective_date=request.effective_date,
                )
            )
    old_enrollment.enrollment_status = {
        "suspend": "suspended",
        "transfer": "transferred",
        "withdraw": "withdrawn",
    }[request.action]
    old_enrollment.effective_date = request.effective_date
    history = StudentEnrollmentHistory(
        student_id=student.id,
        class_id=request.class_id,
        target_class_id=target_class.id if target_class else None,
        action=request.action,
        reason=request.reason,
        effective_date=request.effective_date,
        created_by=user.id,
    )
    db.add(history)
    db.add(
        Notification(
            user_id=student.id,
            class_id=request.class_id,
            title="Cập nhật tình trạng ghi danh",
            message=f"{request.action} có hiệu lực {request.effective_date.isoformat()}: {request.reason}",
        )
    )
    db.commit()
    db.refresh(history)
    return {
        "id": history.id,
        "student_id": history.student_id,
        "class_id": history.class_id,
        "target_class_id": history.target_class_id,
        "action": history.action,
        "reason": history.reason,
        "effective_date": history.effective_date.isoformat(),
    }


@router.get("/students/{student_id}/enrollment-history")
def get_student_enrollment_history(
    student_id: int,
    response: Response,
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=20, ge=1, le=100),
    _manager: User = Depends(require_roles(*MANAGER_ROLES)),
    db: Session = Depends(get_db),
) -> list[dict]:
    query = select(StudentEnrollmentHistory).where(StudentEnrollmentHistory.student_id == student_id)
    total = db.scalar(select(func.count()).select_from(query.subquery())) or 0
    history = db.scalars(query.order_by(StudentEnrollmentHistory.effective_date.desc(),
        StudentEnrollmentHistory.id.desc()).offset((page - 1) * page_size).limit(page_size)).all()
    set_pagination_headers(response, total=total, page=page, page_size=page_size)
    return [
        {
            "id": item.id,
            "class_id": item.class_id,
            "target_class_id": item.target_class_id,
            "action": item.action,
            "reason": item.reason,
            "effective_date": item.effective_date.isoformat(),
            "created_at": item.created_at.isoformat(),
        }
        for item in history
    ]
