from io import BytesIO
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal, ROUND_HALF_UP
import secrets
from math import ceil

from fastapi import APIRouter, Depends, HTTPException, Query, Request, Response, status
from fastapi.responses import JSONResponse
from fastapi.responses import FileResponse
from openpyxl import Workbook, load_workbook
from sqlalchemy import and_, case, exists, false, func, or_, select, update
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session, aliased, joinedload, selectinload

from app.api.dependencies import get_current_user, require_roles
from app.api.pagination import page_metadata, set_pagination_headers
from app.core.config import settings
from app.core.security import generate_temporary_password, hash_password, verify_password
from app.db.session import get_db
from app.models.audit_log import AuditLog
from app.models.attendance import (
    AttendanceChangeLog,
    AttendanceRecord,
    AttendanceRiskState,
    AttendanceSession,
)
from app.models.auth_session import AuthSession
from app.models.classroom import Classroom
from app.models.course import Course
from app.models.course_lesson import CourseLesson
from app.models.lead import Lead
from app.models.notification import Notification, NotificationOutbox
from app.models.announcement import Announcement, SessionLearningLink, announcement_classes
from app.models.role import Role
from app.models.study_program import ProgramFeeSchedule, StudyProgram, program_subjects
from app.models.student_profile import StudentProfile
from app.models.student_enrollment_history import StudentEnrollmentHistory
from app.models.tuition import (
    TuitionDiscount,
    TuitionInstallment,
    TuitionPayment,
    TuitionPaymentReversal,
    TuitionPlan,
)
from app.models.training_class import (
    Assignment,
    AssignmentGradeVersion,
    AssignmentReworkRequest,
    PrivateAssignmentFile,
    AssignmentRecipient,
    AssignmentSubmission,
    ClassEnrollment,
    RubricTemplate,
    TrainingClass,
)
from app.models.user import User
from app.schemas.management import (
    AccountCreate,
    AccountStatusRequest,
    AccountUpdate,
    AssignmentCreate,
    AssignmentGradeCreate,
    AssignmentSubmissionCreate,
    AssignmentReworkRequestCreate,
    ClassReminderPreferenceUpdate,
    SessionMaterialsVisibilityUpdate,
    ProgramAttendancePolicyUpdate,
    AnnouncementCreate,
    SessionLearningLinkCreate,
    RubricScoreInput,
    RubricTemplateCreate,
    ProgramFeeScheduleCreate,
    TuitionPlanBatchCreate,
    TuitionPaymentCreate,
    TuitionPaymentReversalCreate,
    TuitionDiscountCreate,
    ClassroomCreate,
    ClassroomUpdate,
    ChangePasswordRequest,
    CourseCreate,
    CourseLessonCreate,
    CourseLessonsClone,
    CourseLessonUpdate,
    CourseUpdate,
    EnrollmentRequest,
    EnrollmentStatusUpdate,
    RoleAssignmentRequest,
    StudyProgramCoursesUpdate,
    StudyProgramCreate,
    StudyProgramUpdate,
    StudyProgramSubjectCreate,
    StudyProgramSubjectsUpdate,
    StudentProfileUpdate,
    TrainingClassCreate,
    TrainingClassUpdate,
)
from app.services.attendance_summary import HO_CHI_MINH
from app.services.email_service import send_account_credentials_email
from app.services.receipt_pdf import build_receipt_pdf
from app.services.private_storage import private_file_path, save_private_file
from app.services.notification_outbox import enqueue_notification_email
from app.services.attendance_risk import refresh_attendance_risk_states

router = APIRouter(tags=["Training management"])
ADMIN_ROLES = ("admin", "manager")
ALL_ROLES = (
    "admin",
    "manager",
    "admissions",
    "accountant",
    "teacher",
    "ta",
    "student",
    "guest",
)
ROLE_NAMES = {
    "admin": "Admin",
    "manager": "Training Manager",
    "admissions": "Admissions",
    "accountant": "Accountant",
    "teacher": "Instructor",
    "ta": "TA",
    "student": "Student",
    "guest": "Guest",
}
WEEKDAYS = ("Thứ Hai", "Thứ Ba", "Thứ Tư", "Thứ Năm", "Thứ Sáu", "Thứ Bảy", "Chủ Nhật")


def _user_payload(user: User) -> dict:
    return {
        "id": user.id,
        "username": user.username,
        "student_code": user.student_code,
        "email": user.email,
        "full_name": user.full_name,
        "phone": user.phone,
        "birth_date": user.birth_date.isoformat() if user.birth_date else None,
        "address": user.address,
        "roles": [role.code for role in user.roles],
        "is_active": user.is_active,
        "is_locked": user.is_locked,
        "must_change_password": user.must_change_password,
        "created_at": user.created_at.isoformat() if user.created_at else None,
    }


def _course_payload(course: Course) -> dict:
    return {
        "id": course.id,
        "code": course.code,
        "name": course.name,
        "department": course.department,
        "credits": course.credits,
        "session_count": course.session_count,
        "description": course.description,
        "learning_outcomes": course.learning_outcomes,
        "prerequisite_id": course.prerequisite_id,
        "prerequisite_name": (
            course.prerequisite.name if course.prerequisite else None
        ),
        "is_active": course.is_active,
    }


def _class_payload(training_class: TrainingClass) -> dict:
    today = date.today()
    sessions = getattr(training_class, "sessions", [])
    total_sessions = len(sessions)
    completed_sessions = sum(
        1 for s in sessions if (bool(getattr(s, "records", None)) or s.session_date <= today)
    ) if sessions else 0
    return {
        "id": training_class.id,
        "code": training_class.code,
        "name": training_class.name,
        "course_id": training_class.course_id,
        "program_id": training_class.program_id,
        "planned_start_date": training_class.planned_start_date.isoformat() if training_class.planned_start_date else None,
        "delivery_mode": training_class.delivery_mode,
        "program_curriculum_snapshot": training_class.program_curriculum_snapshot or [],
        "course_code": training_class.course.code,
        "course_name": training_class.course.name,
        "teacher_id": training_class.teacher_id,
        "teacher_name": training_class.teacher.full_name,
        "assistant_id": training_class.assistant_id,
        "assistant_name": (
            training_class.assistant.full_name if training_class.assistant else None
        ),
        "room": training_class.room,
        "term": training_class.term,
        "day_of_week": training_class.day_of_week,
        "weekday": WEEKDAYS[training_class.day_of_week],
        "start_time": training_class.start_time.strftime("%H:%M"),
        "end_time": training_class.end_time.strftime("%H:%M"),
        "capacity": training_class.capacity,
        "student_count": sum(1 for enrollment in training_class.enrollments if enrollment.enrollment_status == "active"),
        "status": training_class.status,
        "assistants": [
            {"id": assistant.id, "full_name": assistant.full_name}
            for assistant in training_class.assistants
        ],
        "total_sessions": total_sessions,
        "completed_sessions": completed_sessions,
        "session_progress": f"{completed_sessions}/{total_sessions}" if total_sessions > 0 else "Chưa có lịch",
    }


def _record_audit(
    db: Session,
    actor: User,
    action: str,
    target_user_id: int | None = None,
    details: str | None = None,
) -> None:
    db.add(
        AuditLog(
            actor_id=actor.id,
            target_user_id=target_user_id,
            action=action,
            details=details,
        )
    )


@router.get("/audit-logs")
def list_audit_logs(
    actor_id: int | None = Query(default=None, ge=1),
    entity_type: str | None = Query(default=None, pattern="^(role|tuition|grade|submission|rubric)$"),
    from_date: datetime | None = None,
    to_date: datetime | None = None,
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=20, ge=1, le=100),
    _admin: User = Depends(require_roles("admin")),
    db: Session = Depends(get_db),
) -> dict:
    start = (
        from_date.astimezone(timezone.utc).replace(tzinfo=None)
        if from_date is not None and from_date.tzinfo is not None
        else from_date
    )
    end = (
        to_date.astimezone(timezone.utc).replace(tzinfo=None)
        if to_date is not None and to_date.tzinfo is not None
        else to_date
    )
    if start is not None and end is not None and start > end:
        raise HTTPException(status_code=422, detail="Khoảng thời gian không hợp lệ.")

    query = select(AuditLog).where(
        AuditLog.action.in_(
            (
                "account.role_changed",
                "enrollment.tuition_status_changed",
                "fee_schedule.created",
                "tuition.plan_batch_created",
                "tuition.payment_recorded",
                "tuition.payment_reversed",
                "tuition.discount_requested",
                "tuition.discount_approved",
                "tuition.discount_rejected",
                "grade.draft_saved",
                "grade.published",
                "assignment.submitted",
                "rubric.template_created",
                "rubric.template_updated",
            )
        )
    )
    if actor_id is not None:
        query = query.where(AuditLog.actor_id == actor_id)
    if entity_type == "role":
        query = query.where(AuditLog.action == "account.role_changed")
    elif entity_type == "tuition":
        query = query.where(
            AuditLog.action.in_(
                (
                    "enrollment.tuition_status_changed",
                    "fee_schedule.created",
                    "tuition.plan_batch_created",
                    "tuition.payment_recorded",
                    "tuition.payment_reversed",
                    "tuition.discount_requested",
                    "tuition.discount_approved",
                    "tuition.discount_rejected",
                )
            )
        )
    elif entity_type == "grade":
        query = query.where(
            AuditLog.action.in_(("grade.draft_saved", "grade.published"))
        )
    elif entity_type == "submission":
        query = query.where(AuditLog.action == "assignment.submitted")
    elif entity_type == "rubric":
        query = query.where(
            AuditLog.action.in_(("rubric.template_created", "rubric.template_updated"))
        )
    if start is not None:
        query = query.where(AuditLog.created_at >= start)
    if end is not None:
        query = query.where(AuditLog.created_at <= end)

    total = db.scalar(select(func.count()).select_from(query.subquery())) or 0
    records = db.scalars(
        query.order_by(AuditLog.created_at.desc(), AuditLog.id.desc())
        .offset((page - 1) * page_size)
        .limit(page_size)
    ).all()
    return {
        "items": [
            {
                "id": record.id,
                "actor_id": record.actor_id,
                "target_user_id": record.target_user_id,
                "action": record.action,
                "entity_type": (
                    "role"
                    if record.action == "account.role_changed"
                    else "tuition"
                    if record.action == "enrollment.tuition_status_changed"
                    else "grade"
                    if record.action in ("grade.draft_saved", "grade.published")
                    else "submission"
                    if record.action == "assignment.submitted"
                    else "rubric"
                ),
                "details": record.details,
                "created_at": record.created_at.isoformat(),
            }
            for record in records
        ],
        "total": total,
        "page": page,
        "page_size": page_size,
        "pages": (total + page_size - 1) // page_size,
        "has_previous": page > 1 and total > 0,
        "has_next": page * page_size < total,
    }


def _program_curriculum_snapshot(db: Session, program_id: int | None) -> list[dict] | None:
    if program_id is None:
        return None
    rows = db.execute(
        select(Course, program_subjects.c.order_index, program_subjects.c.weight,
            program_subjects.c.prerequisite_subject_id)
        .join(program_subjects, program_subjects.c.subject_id == Course.id)
        .where(program_subjects.c.program_id == program_id)
        .order_by(program_subjects.c.order_index, Course.id)
    ).all()
    return [{
        "course_id": course.id,
        "code": course.code,
        "name": course.name,
        "credits": course.credits,
        "order_index": order_index,
        "weight": str(weight),
        "prerequisite_course_id": prerequisite_id,
    } for course, order_index, weight, prerequisite_id in rows]


@router.post("/auth/change-password")
def change_password(
    request: ChangePasswordRequest,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
) -> dict:
    if not verify_password(request.current_password, current_user.hashed_password):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Mật khẩu hiện tại không chính xác.",
        )
    if verify_password(request.new_password, current_user.hashed_password):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Mật khẩu mới phải khác mật khẩu hiện tại.",
        )

    now = datetime.now(timezone.utc)
    current_user.hashed_password = hash_password(request.new_password)
    current_user.must_change_password = False
    db.query(AuthSession).filter(
        AuthSession.user_id == current_user.id,
        AuthSession.revoked_at.is_(None),
    ).update({AuthSession.revoked_at: now}, synchronize_session=False)
    db.commit()
    return {"message": "Đổi mật khẩu thành công. Vui lòng đăng nhập lại."}


@router.get("/dashboard/summary")
def dashboard_summary(
    current_user: User = Depends(require_roles(*ALL_ROLES)),
    db: Session = Depends(get_db),
) -> dict:
    roles = {role.code for role in current_user.roles}
    class_query = select(func.count(TrainingClass.id))
    if roles.intersection(("teacher", "ta")) and not roles.intersection(ADMIN_ROLES):
        class_query = class_query.where(
            or_(
                TrainingClass.teacher_id == current_user.id,
                TrainingClass.assistant_id == current_user.id,
                TrainingClass.assistants.any(User.id == current_user.id),
            )
        )
    elif "student" in roles and not roles.intersection(ADMIN_ROLES + ("teacher",)):
        class_query = class_query.join(ClassEnrollment).where(
            ClassEnrollment.student_id == current_user.id
        )
    elif "guest" in roles and not roles.intersection(ADMIN_ROLES):
        class_query = class_query.where(false())

    return {
        "students": db.scalar(
            select(func.count(func.distinct(User.id)))
            .join(User.roles)
            .where(Role.code == "student", User.is_active.is_(True))
        )
        or 0,
        "teachers": db.scalar(
            select(func.count(func.distinct(User.id)))
            .join(User.roles)
            .where(Role.code == "teacher", User.is_active.is_(True))
        )
        or 0,
        "courses": db.scalar(
            select(func.count(Course.id)).where(Course.is_active.is_(True))
        )
        or 0,
        "classes": db.scalar(class_query) or 0,
    }


@router.get("/accounts")
def list_accounts(
    q: str = Query(default="", max_length=150),
    role: str | None = Query(
        default=None,
        pattern="^(admin|manager|admissions|accountant|teacher|ta|student|guest)$",
    ),
    status: str | None = Query(default=None, pattern="^(active|inactive|locked)$"),
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=20, ge=1, le=100),
    current_user: User = Depends(require_roles(*ADMIN_ROLES, "admissions")),
    db: Session = Depends(get_db),
) -> dict:
    current_roles = {assigned_role.code for assigned_role in current_user.roles}
    if (
        "admissions" in current_roles
        and not current_roles.intersection(ADMIN_ROLES)
        and role != "student"
    ):
        raise HTTPException(
            status_code=403,
            detail="Admissions chỉ được tra cứu danh sách học viên.",
        )
    query = select(User).options(selectinload(User.roles))
    if role:
        query = query.join(User.roles).where(Role.code == role)
    if status == "active":
        query = query.where(User.is_active.is_(True), User.is_locked.is_(False))
    elif status == "inactive":
        query = query.where(User.is_active.is_(False))
    elif status == "locked":
        query = query.where(User.is_locked.is_(True))
    if q.strip():
        search = f"%{q.strip()}%"
        query = query.where(
            or_(
                User.full_name.ilike(search),
                User.email.ilike(search),
                User.username.ilike(search),
                User.student_code.ilike(search),
                User.phone.ilike(search),
            )
        )
    total = db.scalar(select(func.count()).select_from(query.order_by(None).subquery())) or 0
    users = db.scalars(
        query.order_by(User.created_at.desc(), User.id.desc())
        .offset((page - 1) * page_size)
        .limit(page_size)
    ).all()
    return {
        "items": [_user_payload(user) for user in users],
        "total": total,
        "page": page,
        "page_size": page_size,
        "pages": (total + page_size - 1) // page_size,
        "has_previous": page > 1 and total > 0,
        "has_next": page * page_size < total,
    }


@router.get("/students")
def list_student_profiles(
    q: str = Query(default="", max_length=150),
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=20, ge=1, le=100),
    _staff: User = Depends(require_roles("admin", "manager", "admissions")),
    db: Session = Depends(get_db),
) -> dict:
    query = (
        select(User, StudentProfile)
        .join(User.roles)
        .outerjoin(StudentProfile, StudentProfile.user_id == User.id)
        .where(Role.code == "student")
    )
    if q.strip():
        search = f"%{q.strip()}%"
        query = query.where(
            or_(
                User.full_name.ilike(search),
                User.email.ilike(search),
                User.username.ilike(search),
                User.student_code.ilike(search),
                User.phone.ilike(search),
                StudentProfile.study_program.ilike(search),
                StudentProfile.entry_level.ilike(search),
            )
        )
    total = db.scalar(
        select(func.count()).select_from(query.order_by(None).subquery())
    ) or 0
    rows = db.execute(
        query.order_by(User.created_at.desc(), User.id.desc())
        .offset((page - 1) * page_size)
        .limit(page_size)
    ).all()
    items = []
    for student, profile in rows:
        items.append(
            {
                **_user_payload(student),
                "study_program": profile.study_program if profile else None,
                "entry_level": profile.entry_level if profile else None,
                "graduation_status": profile.graduation_status if profile else "studying",
                "lead_id": profile.lead_id if profile else None,
                "profile_created_at": (
                    profile.created_at.isoformat() if profile else None
                ),
            }
        )
    return {
        "items": items,
        "total": total,
        "page": page,
        "page_size": page_size,
        "pages": (total + page_size - 1) // page_size,
        "has_previous": page > 1 and total > 0,
        "has_next": page * page_size < total,
    }


@router.get("/students/{student_id}")
def get_student_profile(
    student_id: int,
    _current_user: User = Depends(require_roles("admin", "manager", "admissions", "teacher", "ta", "student")),
    db: Session = Depends(get_db),
) -> dict:
    roles = {role.code for role in _current_user.roles}
    if "student" in roles and not roles.intersection({"admin", "manager", "admissions", "teacher", "ta"}):
        if _current_user.id != student_id:
            raise HTTPException(status_code=403, detail="Bạn chỉ có thể xem hồ sơ của chính mình.")
    if roles.intersection({"teacher", "ta"}) and not roles.intersection({"admin", "manager", "admissions"}):
        has_class_scope = db.scalar(select(ClassEnrollment.id).join(
            TrainingClass, TrainingClass.id == ClassEnrollment.class_id
        ).where(
            ClassEnrollment.student_id == student_id,
            or_(TrainingClass.teacher_id == _current_user.id,
                TrainingClass.assistants.any(User.id == _current_user.id)),
        ).limit(1))
        if has_class_scope is None:
            raise HTTPException(status_code=403, detail="Bạn không được phân công vào lớp của học viên này.")

    student = db.scalar(
        select(User).options(selectinload(User.roles)).where(User.id == student_id)
    )
    if student is None or "student" not in {role.code for role in student.roles}:
        raise HTTPException(status_code=404, detail="Không tìm thấy hồ sơ học viên.")

    profile = db.get(StudentProfile, student.id)

    enrollments = db.scalars(
        select(ClassEnrollment)
        .options(
            joinedload(ClassEnrollment.training_class).joinedload(TrainingClass.course),
            joinedload(ClassEnrollment.training_class).joinedload(TrainingClass.teacher),
        )
        .where(ClassEnrollment.student_id == student_id)
        .order_by(ClassEnrollment.enrolled_at.desc())
    ).all()
    enrollment_list = [
        {
            "id": enr.id,
            "class_id": enr.class_id,
            "class_code": enr.training_class.code,
            "class_name": enr.training_class.name,
            "course_id": enr.training_class.course_id,
            "course_code": enr.training_class.course.code,
            "course_name": enr.training_class.course.name,
            "teacher_name": enr.training_class.teacher.full_name,
            "room": enr.training_class.room,
            "term": enr.training_class.term,
            "enrollment_status": enr.enrollment_status,
            "tuition_status": enr.tuition_status,
            "attendance_status": enr.attendance_status,
            "enrolled_at": enr.enrolled_at.isoformat() if enr.enrolled_at else None,
            "effective_date": enr.effective_date.isoformat() if enr.effective_date else None,
        }
        for enr in enrollments
    ]

    history_rows = db.scalars(
        select(StudentEnrollmentHistory)
        .where(StudentEnrollmentHistory.student_id == student_id)
        .order_by(StudentEnrollmentHistory.created_at.desc())
    ).all()

    class_ids = set()
    for h in history_rows:
        if h.class_id:
            class_ids.add(h.class_id)
        if h.target_class_id:
            class_ids.add(h.target_class_id)

    classes_map = {}
    if class_ids:
        for c in db.scalars(select(TrainingClass).where(TrainingClass.id.in_(class_ids))).all():
            classes_map[c.id] = f"{c.code} · {c.name}"

    history_list = [
        {
            "id": h.id,
            "action": h.action,
            "reason": h.reason,
            "effective_date": h.effective_date.isoformat(),
            "class_id": h.class_id,
            "class_name": classes_map.get(h.class_id, f"Lớp #{h.class_id}"),
            "target_class_id": h.target_class_id,
            "target_class_name": classes_map.get(h.target_class_id) if h.target_class_id else None,
            "created_at": h.created_at.isoformat() if h.created_at else None,
        }
        for h in history_rows
    ]

    attendance_records = db.scalars(
        select(AttendanceRecord)
        .options(
            joinedload(AttendanceRecord.session).joinedload(AttendanceSession.training_class),
            joinedload(AttendanceRecord.session).joinedload(AttendanceSession.lesson),
        )
        .where(AttendanceRecord.student_id == student_id)
        .order_by(AttendanceRecord.marked_at.desc())
    ).all()

    total_attended = len(attendance_records)
    present_count = sum(1 for r in attendance_records if r.status == "present")
    late_count = sum(1 for r in attendance_records if r.status == "late")
    excused_count = sum(1 for r in attendance_records if r.status == "excused_absent")
    unexcused_count = sum(1 for r in attendance_records if r.status == "unexcused_absent")
    attendance_rate = (
        round(((present_count + late_count) / total_attended) * 100, 1)
        if total_attended > 0
        else 100.0
    )

    attendance_list = [
        {
            "id": r.id,
            "session_id": r.session_id,
            "session_date": r.session.session_date.isoformat(),
            "start_time": r.session.start_time.strftime("%H:%M"),
            "end_time": r.session.end_time.strftime("%H:%M"),
            "class_code": r.session.training_class.code,
            "class_name": r.session.training_class.name,
            "lesson_title": r.session.lesson.title if r.session.lesson else None,
            "status": r.status,
            "marked_at": r.marked_at.isoformat() if r.marked_at else None,
        }
        for r in attendance_records[:100]
    ]
    risk_rows = db.execute(select(AttendanceRiskState, TrainingClass).join(
        TrainingClass, TrainingClass.id == AttendanceRiskState.class_id
    ).where(AttendanceRiskState.student_id == student_id,
            AttendanceRiskState.is_at_risk.is_(True)).order_by(TrainingClass.code)).all()
    attendance_risks = [{
        "class_id": training_class.id,
        "class_code": training_class.code,
        "class_name": training_class.name,
        "absent_count": state.absent_count,
        "completed_session_count": state.completed_session_count,
        "absence_threshold_percent": str(state.absence_threshold_percent),
        "flagged_at": state.flagged_at.isoformat() if state.flagged_at else None,
    } for state, training_class in risk_rows]

    return {
        **_user_payload(student),
        "study_program": profile.study_program if profile else None,
        "entry_level": profile.entry_level if profile else None,
        "notes": profile.notes if profile else None,
        "graduation_status": profile.graduation_status if profile else "studying",
        "lead_id": profile.lead_id if profile else None,
        "profile_created_at": profile.created_at.isoformat() if profile else None,
        "enrollments": enrollment_list,
        "enrollment_history": history_list,
        "attendance_stats": {
            "total_sessions": total_attended,
            "present_count": present_count,
            "late_count": late_count,
            "excused_count": excused_count,
            "unexcused_count": unexcused_count,
            "attendance_rate": attendance_rate,
        },
        "attendance_records": attendance_list,
        "attendance_risks": attendance_risks,
    }


@router.patch("/students/{student_id}")
def update_student_profile(
    student_id: int,
    request: StudentProfileUpdate,
    _staff: User = Depends(require_roles("admin", "manager", "admissions")),
    db: Session = Depends(get_db),
) -> dict:
    student = db.scalar(
        select(User).options(selectinload(User.roles)).where(User.id == student_id)
    )
    if student is None or "student" not in {role.code for role in student.roles}:
        raise HTTPException(status_code=404, detail="Không tìm thấy hồ sơ học viên.")
    student.full_name = request.full_name.strip()
    student.phone = request.phone.strip() if request.phone else None
    student.birth_date = request.birth_date
    student.address = request.address.strip() if request.address else None
    profile = db.get(StudentProfile, student.id)
    if profile is None:
        profile = StudentProfile(user_id=student.id)
        db.add(profile)
    profile.study_program = (
        request.study_program.strip() if request.study_program else None
    )
    profile.entry_level = request.entry_level
    profile.notes = request.notes
    db.commit()
    db.refresh(student)
    return {
        **_user_payload(student),
        "study_program": profile.study_program,
        "entry_level": profile.entry_level,
        "notes": profile.notes,
        "graduation_status": profile.graduation_status,
        "lead_id": profile.lead_id,
        "profile_created_at": profile.created_at.isoformat(),
    }


def _classroom_payload(classroom: Classroom) -> dict:
    return {
        "id": classroom.id,
        "code": classroom.code,
        "capacity": classroom.capacity,
        "online_url": classroom.online_url,
    }


@router.get("/rooms")
def list_classrooms(
    response: Response,
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=20, ge=1, le=100),
    _user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
) -> list[dict]:
    query = select(Classroom)
    total = db.scalar(select(func.count()).select_from(query.subquery())) or 0
    rooms = db.scalars(query.order_by(Classroom.code, Classroom.id)
        .offset((page - 1) * page_size).limit(page_size)).all()
    set_pagination_headers(response, total=total, page=page, page_size=page_size)
    return [
        _classroom_payload(room)
        for room in rooms
    ]


@router.post("/rooms", status_code=status.HTTP_201_CREATED)
def create_classroom(
    request: ClassroomCreate,
    _admin: User = Depends(require_roles(*ADMIN_ROLES)),
    db: Session = Depends(get_db),
) -> dict:
    code = request.code.strip()
    if db.scalar(select(Classroom.id).where(Classroom.code == code)):
        raise HTTPException(status_code=409, detail="Mã phòng đã tồn tại.")
    classroom = Classroom(
        code=code,
        capacity=request.capacity,
        online_url=request.online_url.strip() if request.online_url else None,
    )
    db.add(classroom)
    db.commit()
    db.refresh(classroom)
    return _classroom_payload(classroom)


@router.patch("/rooms/{room_id}")
def update_classroom(
    room_id: int,
    request: ClassroomUpdate,
    _admin: User = Depends(require_roles(*ADMIN_ROLES)),
    db: Session = Depends(get_db),
) -> dict:
    classroom = db.get(Classroom, room_id)
    if classroom is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy phòng học.")
    updates = request.model_dump(exclude_unset=True)
    if not updates:
        raise HTTPException(status_code=400, detail="Không có dữ liệu cần cập nhật.")
    if "code" in updates and updates["code"] != classroom.code and db.scalar(
        select(Classroom.id).where(Classroom.code == updates["code"])
    ):
        raise HTTPException(status_code=409, detail="Mã phòng đã tồn tại.")
    next_capacity = updates.get("capacity", classroom.capacity)
    oversized_class = db.scalar(
        select(TrainingClass.id).where(
            TrainingClass.room == classroom.code,
            TrainingClass.capacity > next_capacity,
        )
    )
    if oversized_class:
        raise HTTPException(
            status_code=409,
            detail="Sức chứa phòng không thể nhỏ hơn sức chứa lớp học phần đang sử dụng.",
        )
    old_code = classroom.code
    for field, value in updates.items():
        setattr(classroom, field, value.strip() if field in {"code", "online_url"} and value else value)
    if classroom.code != old_code:
        db.query(TrainingClass).filter(TrainingClass.room == old_code).update(
            {TrainingClass.room: classroom.code}, synchronize_session=False
        )
    db.commit()
    db.refresh(classroom)
    return _classroom_payload(classroom)


@router.delete("/rooms/{room_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_classroom(
    room_id: int,
    _admin: User = Depends(require_roles(*ADMIN_ROLES)),
    db: Session = Depends(get_db),
) -> None:
    classroom = db.get(Classroom, room_id)
    if classroom is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy phòng học.")
    if db.scalar(select(TrainingClass.id).where(TrainingClass.room == classroom.code)):
        raise HTTPException(
            status_code=409, detail="Không thể xóa phòng đang được lớp học sử dụng."
        )
    db.delete(classroom)
    db.commit()


@router.post("/accounts", status_code=status.HTTP_201_CREATED)
async def create_account(
    request: AccountCreate,
    _admin: User = Depends(require_roles("admin")),
    db: Session = Depends(get_db),
) -> dict:
    email = str(request.email).lower()
    if db.scalar(select(User.id).where(User.email == email)):
        raise HTTPException(status_code=409, detail="Email đã được sử dụng.")
    if db.scalar(select(User.id).where(User.username == request.username)):
        raise HTTPException(status_code=409, detail="Mã định danh đã được sử dụng.")
    if request.student_code and db.scalar(
        select(User.id).where(User.student_code == request.student_code)
    ):
        raise HTTPException(status_code=409, detail="Mã học viên đã được sử dụng.")

    selected_codes = request.selected_roles()
    assigned_roles = db.scalars(select(Role).where(Role.code.in_(selected_codes))).all()
    if len(assigned_roles) != len(selected_codes):
        raise HTTPException(
            status_code=503,
            detail="Danh mục vai trò chưa được khởi tạo. Chạy migration hệ thống trước.",
        )

    temporary_password = request.password
    mail_configured = all(
        value.strip()
        for value in (
            settings.MAIL_USERNAME,
            settings.MAIL_PASSWORD,
            settings.MAIL_FROM,
            settings.MAIL_SERVER,
        )
    )
    user = User(
        username=request.username.strip(),
        email=email,
        full_name=request.full_name.strip(),
        student_code=request.student_code,
        phone=request.phone,
        hashed_password=hash_password(temporary_password),
        must_change_password=True,
        roles=assigned_roles,
    )
    db.add(user)
    try:
        db.flush()
        _record_audit(
            db,
            _admin,
            "account.created",
            target_user_id=user.id,
            details=f"roles={','.join(selected_codes)}",
        )
        if mail_configured:
            await send_account_credentials_email(
                email,
                user.full_name,
                user.username,
                temporary_password,
            )
        db.commit()
        db.refresh(user)
    except Exception as error:
        db.rollback()
        if mail_configured:
            raise HTTPException(
                status_code=502,
                detail="Không thể gửi thông tin tài khoản qua email. Tài khoản chưa được tạo.",
            ) from error
        if isinstance(error, IntegrityError):
            raise HTTPException(
                status_code=409,
                detail="Email hoặc mã định danh đã được sử dụng.",
            ) from error
        raise

    result = _user_payload(user)
    result["credential_delivery"] = "email" if mail_configured else "manual"
    if not mail_configured:
        result["temporary_password"] = temporary_password
    return result


@router.get("/accounts/template")
def account_import_template(
    _admin: User = Depends(require_roles("admin")),
) -> Response:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Accounts"
    sheet.append(
        [
            "username",
            "email",
            "full_name",
            "role",
            "student_code",
            "phone",
            "password",
        ]
    )
    sheet.append(
        [
            "sv0001",
            "student@example.edu.vn",
            "Nguyễn Văn A",
            "student",
            "SV0001",
            "",
            "",
        ]
    )
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = "A1:G2"
    output = BytesIO()
    workbook.save(output)
    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": 'attachment; filename="ictu-account-template.xlsx"'
        },
    )


@router.post("/accounts/import")
async def import_accounts(
    request: Request,
    preview: bool = Query(default=False),
    admin: User = Depends(require_roles("admin")),
    db: Session = Depends(get_db),
) -> dict:
    workbook_bytes = await request.body()
    if not workbook_bytes or len(workbook_bytes) > 10 * 1024 * 1024:
        raise HTTPException(
            status_code=413, detail="File Excel trống hoặc vượt quá 10 MB."
        )
    try:
        workbook = load_workbook(BytesIO(workbook_bytes), read_only=True, data_only=True)
    except Exception as error:
        raise HTTPException(
            status_code=400, detail="File không đúng định dạng Excel .xlsx."
        ) from error

    sheet = workbook.active
    headers = [
        str(value).strip().lower() if value is not None else ""
        for value in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
    ]
    required_headers = {"username", "email", "full_name", "role"}
    if not required_headers.issubset(headers):
        workbook.close()
        missing = ", ".join(sorted(required_headers - set(headers)))
        raise HTTPException(
            status_code=422, detail=f"Thiếu cột bắt buộc trong file: {missing}."
        )

    validated: list[tuple[int, dict[str, str]]] = []
    row_errors: list[dict] = []
    preview_rows: list[dict] = []
    file_emails: set[str] = set()
    file_usernames: set[str] = set()
    file_student_codes: set[str] = set()
    for row_number, values in enumerate(
        sheet.iter_rows(min_row=2, values_only=True), start=2
    ):
        if not any(value is not None and str(value).strip() for value in values):
            continue
        row = {
            key: str(values[index]).strip() if index < len(values) and values[index] is not None else ""
            for index, key in enumerate(headers)
        }
        errors: list[str] = []
        if len(row.get("username", "")) < 3:
            errors.append("username phải có ít nhất 3 ký tự")
        if "@" not in row.get("email", "") or "." not in row.get("email", "").split("@")[-1]:
            errors.append("email không hợp lệ")
        if not row.get("full_name"):
            errors.append("full_name không được để trống")
        if row.get("role") not in ROLE_NAMES:
            errors.append(
                "role phải là admin, manager, admissions, accountant, teacher, ta, student hoặc guest"
            )

        email = row.get("email", "").lower()
        username = row.get("username", "")
        student_code = row.get("student_code", "")
        if email in file_emails or db.scalar(select(User.id).where(User.email == email)):
            errors.append("email đã tồn tại")
        if username in file_usernames or db.scalar(
            select(User.id).where(User.username == username)
        ):
            errors.append("username đã tồn tại")
        if student_code and (
            student_code in file_student_codes
            or db.scalar(select(User.id).where(User.student_code == student_code))
        ):
            errors.append("student_code đã tồn tại")

        file_emails.add(email)
        file_usernames.add(username)
        if student_code:
            file_student_codes.add(student_code)
        if errors:
            row_errors.append({"row": row_number, "errors": errors})
        else:
            row["email"] = email
            validated.append((row_number, row))
        preview_rows.append(
            {
                "row": row_number,
                "username": username,
                "email": email,
                "full_name": row.get("full_name", ""),
                "valid": not errors,
                "errors": errors,
            }
        )
    workbook.close()

    if preview:
        return {
            "preview": True,
            "total_rows": len(preview_rows),
            "valid_count": len(validated),
            "invalid_count": len(row_errors),
            "rows": preview_rows,
        }
    if not validated:
        raise HTTPException(
            status_code=422,
            detail={
                "message": "Không có dòng hợp lệ để nhập.",
                "rows": row_errors,
            },
        )

    roles = {
        role.code: role
        for role in db.scalars(select(Role).where(Role.code.in_(ROLE_NAMES))).all()
    }
    if len(roles) != len(ROLE_NAMES):
        raise HTTPException(
            status_code=503,
            detail="Danh mục vai trò chưa được khởi tạo. Chạy migration hệ thống trước.",
        )

    created: list[tuple[User, str]] = []
    for _, row in validated:
        password = row.get("password") or generate_temporary_password()
        user = User(
            username=row["username"],
            email=row["email"],
            full_name=row["full_name"],
            student_code=row.get("student_code") or None,
            phone=row.get("phone") or None,
            hashed_password=hash_password(password),
            must_change_password=True,
            roles=[roles[row["role"]]],
        )
        db.add(user)
        created.append((user, password))
    try:
        db.flush()
        for user, _ in created:
            _record_audit(
                db,
                admin,
                "account.imported",
                target_user_id=user.id,
                details=f"role={user.roles[0].code}",
            )
        db.commit()
    except IntegrityError as error:
        db.rollback()
        raise HTTPException(
            status_code=409,
            detail="Dữ liệu tài khoản bị trùng trong lúc nhập; chưa có tài khoản nào được tạo.",
        ) from error

    mail_configured = all(
        value.strip()
        for value in (
            settings.MAIL_USERNAME,
            settings.MAIL_PASSWORD,
            settings.MAIL_FROM,
            settings.MAIL_SERVER,
        )
    )
    results = []
    for user, password in created:
        delivery = "manual"
        if mail_configured:
            try:
                await send_account_credentials_email(
                    user.email, user.full_name, user.username, password
                )
                delivery = "email"
            except Exception:
                delivery = "failed"
        result = {"id": user.id, "email": user.email, "credential_delivery": delivery}
        if delivery != "email":
            result["temporary_password"] = password
        results.append(result)
    return {
        "created_count": len(results),
        "skipped_count": len(row_errors),
        "total_rows": len(preview_rows),
        "row_errors": row_errors,
        "accounts": results,
    }


@router.patch("/accounts/{user_id}")
def update_account(
    user_id: int,
    request: AccountUpdate,
    admin: User = Depends(require_roles("admin")),
    db: Session = Depends(get_db),
) -> dict:
    user = db.get(User, user_id)
    if user is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy tài khoản.")
    updates = request.model_dump(exclude_unset=True)
    if not updates:
        raise HTTPException(status_code=400, detail="Không có dữ liệu cần cập nhật.")
    for field, value in updates.items():
        setattr(user, field, value)
    _record_audit(
        db,
        admin,
        "account.updated",
        target_user_id=user.id,
        details=",".join(sorted(updates)),
    )
    db.commit()
    db.refresh(user)
    return _user_payload(user)


@router.post("/accounts/{user_id}/lock")
def lock_account(
    user_id: int,
    request: AccountStatusRequest,
    admin: User = Depends(require_roles("admin")),
    db: Session = Depends(get_db),
) -> dict:
    user = db.get(User, user_id)
    if user is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy tài khoản.")
    if user.id == admin.id:
        raise HTTPException(status_code=400, detail="Không thể khóa tài khoản đang sử dụng.")
    user.is_locked = True
    now = datetime.now(timezone.utc).replace(tzinfo=None)
    db.query(AuthSession).filter(
        AuthSession.user_id == user.id,
        AuthSession.revoked_at.is_(None),
    ).update({AuthSession.revoked_at: now}, synchronize_session=False)
    handover_classes = db.scalars(
        select(TrainingClass.code).where(
            or_(
                TrainingClass.teacher_id == user.id,
                TrainingClass.assistant_id == user.id,
                TrainingClass.assistants.any(User.id == user.id),
            )
        )
    ).all()
    _record_audit(
        db,
        admin,
        "account.locked",
        target_user_id=user.id,
        details=(
            f"{request.reason.strip()} | "
            f"handover_classes={','.join(handover_classes) or 'none'}"
        ),
    )
    db.commit()
    return {
        "id": user.id,
        "is_locked": user.is_locked,
        "handover_classes": handover_classes,
    }


@router.post("/accounts/{user_id}/unlock")
def unlock_account(
    user_id: int,
    request: AccountStatusRequest,
    admin: User = Depends(require_roles("admin")),
    db: Session = Depends(get_db),
) -> dict:
    user = db.get(User, user_id)
    if user is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy tài khoản.")
    user.is_locked = False
    _record_audit(
        db,
        admin,
        "account.unlocked",
        target_user_id=user.id,
        details=request.reason.strip(),
    )
    db.commit()
    return {"id": user.id, "is_locked": user.is_locked}


@router.put("/accounts/{user_id}/role")
def assign_account_role(
    user_id: int,
    request: RoleAssignmentRequest,
    admin: User = Depends(require_roles("admin")),
    db: Session = Depends(get_db),
) -> dict:
    user = db.scalar(
        select(User).options(selectinload(User.roles)).where(User.id == user_id)
    )
    if user is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy tài khoản.")
    selected_codes = request.selected_roles()
    roles = db.scalars(select(Role).where(Role.code.in_(selected_codes))).all()
    roles_by_code = {role.code: role for role in roles}
    if len(roles_by_code) != len(selected_codes):
        raise HTTPException(status_code=422, detail="Có vai trò không tồn tại.")
    if user.id == admin.id and "admin" not in selected_codes:
        raise HTTPException(status_code=400, detail="Không thể gỡ quyền quản trị của chính mình.")
    previous_codes = sorted(role.code for role in user.roles)
    next_codes = sorted(selected_codes)
    user.roles = [roles_by_code[code] for code in selected_codes]
    if previous_codes != next_codes:
        _record_audit(
            db,
            admin,
            "account.role_changed",
            target_user_id=user.id,
            details=f"before={','.join(previous_codes)};after={','.join(next_codes)}",
        )
    db.commit()
    return _user_payload(user)


@router.get("/roles")
def list_roles(
    _admin: User = Depends(require_roles("admin")),
) -> list[dict]:
    return [{"code": code, "name": name} for code, name in ROLE_NAMES.items()]


@router.get("/courses")
def list_courses(
    response: Response,
    q: str = Query(default="", max_length=150),
    department: str | None = Query(default=None, max_length=100),
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=20, ge=1, le=100),
    _user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
) -> list[dict]:
    query = select(Course).options(joinedload(Course.prerequisite))
    if q.strip():
        search = f"%{q.strip()}%"
        query = query.where(or_(Course.code.ilike(search), Course.name.ilike(search)))
    if department:
        query = query.where(Course.department == department)
    total = db.scalar(select(func.count()).select_from(query.order_by(None).subquery())) or 0
    courses = db.scalars(query.order_by(Course.department, Course.code, Course.id)
        .offset((page - 1) * page_size).limit(page_size)).all()
    set_pagination_headers(response, total=total, page=page, page_size=page_size)
    return [
        _course_payload(course)
        for course in courses
    ]


def _course_lesson_payload(lesson: CourseLesson) -> dict:
    return {
        "id": lesson.id,
        "course_id": lesson.course_id,
        "sequence": lesson.sequence,
        "title": lesson.title,
        "objectives": lesson.objectives,
        "description": lesson.description,
    }


@router.get("/courses/{course_id}/lessons")
def list_course_lessons(
    course_id: int,
    response: Response,
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=20, ge=1, le=100),
    _user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
) -> list[dict]:
    if db.get(Course, course_id) is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy môn học.")
    query = select(CourseLesson).where(CourseLesson.course_id == course_id)
    total = db.scalar(select(func.count()).select_from(query.subquery())) or 0
    lessons = db.scalars(query.order_by(CourseLesson.sequence, CourseLesson.id)
        .offset((page - 1) * page_size).limit(page_size)).all()
    set_pagination_headers(response, total=total, page=page, page_size=page_size)
    return [_course_lesson_payload(lesson) for lesson in lessons]


@router.post("/courses/{course_id}/lessons", status_code=status.HTTP_201_CREATED)
def create_course_lesson(
    course_id: int,
    request: CourseLessonCreate,
    _admin: User = Depends(require_roles(*ADMIN_ROLES)),
    db: Session = Depends(get_db),
) -> dict:
    course = db.get(Course, course_id)
    if course is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy môn học.")
    if request.sequence > course.session_count:
        raise HTTPException(
            status_code=422, detail="Số buổi vượt quá số buổi khai báo của môn học."
        )
    if db.scalar(
        select(CourseLesson.id).where(
            CourseLesson.course_id == course_id,
            CourseLesson.sequence == request.sequence,
        )
    ):
        raise HTTPException(status_code=409, detail="Thứ tự buổi học đã được sử dụng.")
    lesson = CourseLesson(course_id=course_id, **request.model_dump())
    db.add(lesson)
    db.commit()
    db.refresh(lesson)
    return _course_lesson_payload(lesson)


@router.patch("/courses/{course_id}/lessons/{lesson_id}")
def update_course_lesson(
    course_id: int,
    lesson_id: int,
    request: CourseLessonUpdate,
    _admin: User = Depends(require_roles(*ADMIN_ROLES)),
    db: Session = Depends(get_db),
) -> dict:
    lesson = db.scalar(
        select(CourseLesson).where(
            CourseLesson.id == lesson_id, CourseLesson.course_id == course_id
        )
    )
    if lesson is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy buổi học.")
    course = db.get(Course, course_id)
    updates = request.model_dump(exclude_unset=True)
    next_sequence = updates.get("sequence", lesson.sequence)
    if next_sequence > course.session_count:
        raise HTTPException(
            status_code=422, detail="Số buổi vượt quá số buổi khai báo của môn học."
        )
    if db.scalar(
        select(CourseLesson.id).where(
            CourseLesson.course_id == course_id,
            CourseLesson.sequence == next_sequence,
            CourseLesson.id != lesson_id,
        )
    ):
        raise HTTPException(status_code=409, detail="Thứ tự buổi học đã được sử dụng.")
    for field, value in updates.items():
        setattr(lesson, field, value)
    db.commit()
    db.refresh(lesson)
    return _course_lesson_payload(lesson)


@router.post("/courses/{course_id}/lessons/clone", status_code=status.HTTP_201_CREATED)
def clone_course_lessons(
    course_id: int,
    request: CourseLessonsClone,
    _admin: User = Depends(require_roles(*ADMIN_ROLES)),
    db: Session = Depends(get_db),
) -> list[dict]:
    target = db.get(Course, course_id)
    source = db.get(Course, request.source_course_id)
    if target is None or source is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy môn học nguồn hoặc đích.")
    if target.id == source.id:
        raise HTTPException(status_code=422, detail="Môn nguồn và môn đích phải khác nhau.")
    if db.scalar(select(CourseLesson.id).where(CourseLesson.course_id == target.id)):
        raise HTTPException(
            status_code=409, detail="Môn đích đã có syllabus; hãy xóa nội dung trước khi nhân bản."
        )
    source_lessons = db.scalars(
        select(CourseLesson)
        .where(CourseLesson.course_id == source.id)
        .order_by(CourseLesson.sequence)
    ).all()
    if len(source_lessons) > target.session_count:
        raise HTTPException(
            status_code=422,
            detail="Syllabus nguồn có nhiều buổi hơn số buổi của môn đích.",
        )
    copied = [
        CourseLesson(
            course_id=target.id,
            sequence=lesson.sequence,
            title=lesson.title,
            objectives=lesson.objectives,
            description=lesson.description,
        )
        for lesson in source_lessons
    ]
    db.add_all(copied)
    db.commit()
    return [_course_lesson_payload(lesson) for lesson in copied]


@router.delete(
    "/courses/{course_id}/lessons/{lesson_id}", status_code=status.HTTP_204_NO_CONTENT
)
def delete_course_lesson(
    course_id: int,
    lesson_id: int,
    _admin: User = Depends(require_roles(*ADMIN_ROLES)),
    db: Session = Depends(get_db),
) -> None:
    lesson = db.scalar(
        select(CourseLesson).where(
            CourseLesson.id == lesson_id, CourseLesson.course_id == course_id
        )
    )
    if lesson is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy buổi học.")
    db.delete(lesson)
    db.commit()


@router.get("/admissions-staff")
def list_admissions_staff(
    response: Response,
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=20, ge=1, le=100),
    _user: User = Depends(require_roles("admin", "admissions")),
    db: Session = Depends(get_db),
) -> list[dict]:
    query = select(User).join(User.roles).where(
        User.is_active.is_(True), User.is_locked.is_(False), Role.code == "admissions"
    )
    total = db.scalar(select(func.count()).select_from(query.subquery())) or 0
    staff = db.scalars(
        query
        .options(selectinload(User.roles))
        .order_by(User.full_name, User.id)
        .offset((page - 1) * page_size).limit(page_size)
    ).all()
    set_pagination_headers(response, total=total, page=page, page_size=page_size)
    return [
        {"id": member.id, "full_name": member.full_name, "email": member.email}
        for member in staff
    ]


def _program_fee_schedule_payload(schedule: ProgramFeeSchedule) -> dict:
    return {
        "id": schedule.id,
        "program_id": schedule.program_id,
        "version": schedule.version,
        "amount_vnd": str(schedule.amount_vnd),
        "currency": "VND",
        "valid_from": schedule.valid_from.isoformat(),
        "valid_to": schedule.valid_to.isoformat() if schedule.valid_to else None,
        "created_by": schedule.created_by,
        "created_at": schedule.created_at.isoformat() if schedule.created_at else None,
    }


@router.put("/study-programs/{program_id}/attendance-policy")
def update_program_attendance_policy(
    program_id: int,
    request: ProgramAttendancePolicyUpdate,
    current_user: User = Depends(require_roles("admin", "manager")),
    db: Session = Depends(get_db),
) -> dict:
    program = db.get(StudyProgram, program_id)
    if program is None:
        raise HTTPException(status_code=404, detail="Study program not found.")
    program.absence_threshold_percent = request.absence_threshold_percent
    classes = db.scalars(select(TrainingClass).where(TrainingClass.program_id == program_id)).all()
    for training_class in classes:
        refresh_attendance_risk_states(db, training_class)
    _record_audit(db, current_user, "attendance.threshold.updated",
        details=f"program_id={program_id};absence_threshold_percent={request.absence_threshold_percent}")
    db.commit()
    return {"program_id": program_id, "absence_threshold_percent": str(program.absence_threshold_percent)}


@router.get("/study-programs/{program_id}/fee-schedules")
def list_program_fee_schedules(
    program_id: int,
    response: Response,
    effective_on: date | None = Query(default=None),
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=20, ge=1, le=100),
    _accountant: User = Depends(require_roles("admin", "accountant")),
    db: Session = Depends(get_db),
) -> list[dict]:
    if db.get(StudyProgram, program_id) is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy chương trình đào tạo.")
    query = select(ProgramFeeSchedule).where(
        ProgramFeeSchedule.program_id == program_id
    )
    if effective_on is not None:
        query = query.where(
            ProgramFeeSchedule.valid_from <= effective_on,
            or_(
                ProgramFeeSchedule.valid_to.is_(None),
                ProgramFeeSchedule.valid_to >= effective_on,
            ),
        )
    total = db.scalar(select(func.count()).select_from(query.subquery())) or 0
    schedules = db.scalars(query.order_by(ProgramFeeSchedule.valid_from,
        ProgramFeeSchedule.version, ProgramFeeSchedule.id)
        .offset((page - 1) * page_size).limit(page_size)).all()
    set_pagination_headers(response, total=total, page=page, page_size=page_size)
    return [_program_fee_schedule_payload(item) for item in schedules]


@router.post(
    "/study-programs/{program_id}/fee-schedules",
    status_code=status.HTTP_201_CREATED,
)
def create_program_fee_schedule(
    program_id: int,
    request: ProgramFeeScheduleCreate,
    current_user: User = Depends(require_roles("admin", "accountant")),
    db: Session = Depends(get_db),
) -> dict:
    program = db.scalar(
        select(StudyProgram).where(StudyProgram.id == program_id).with_for_update()
    )
    if program is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy chương trình đào tạo.")

    existing = db.scalars(
        select(ProgramFeeSchedule)
        .where(ProgramFeeSchedule.program_id == program_id)
        .with_for_update()
    ).all()
    for schedule in existing:
        starts_before_existing_ends = (
            schedule.valid_to is None or request.valid_from <= schedule.valid_to
        )
        existing_starts_before_ends = (
            request.valid_to is None or schedule.valid_from <= request.valid_to
        )
        if starts_before_existing_ends and existing_starts_before_ends:
            raise HTTPException(
                status_code=409,
                detail="Khoảng hiệu lực bị trùng với biểu phí đã khai báo.",
            )

    schedule = ProgramFeeSchedule(
        program_id=program_id,
        version=max((item.version for item in existing), default=0) + 1,
        amount_vnd=request.amount_vnd,
        valid_from=request.valid_from,
        valid_to=request.valid_to,
        created_by=current_user.id,
    )
    db.add(schedule)
    db.flush()
    _record_audit(
        db,
        current_user,
        "fee_schedule.created",
        details=(
            f"program_id={program_id};version={schedule.version};"
            f"amount_vnd={schedule.amount_vnd};valid_from={schedule.valid_from};"
            f"valid_to={schedule.valid_to}"
        ),
    )
    db.commit()
    db.refresh(schedule)
    return _program_fee_schedule_payload(schedule)


@router.post("/tuition-plans/batch", status_code=status.HTTP_201_CREATED)
def create_tuition_plans_batch(
    request: TuitionPlanBatchCreate,
    current_user: User = Depends(require_roles("admin", "accountant")),
    db: Session = Depends(get_db),
) -> dict:
    training_class = db.scalar(
        select(TrainingClass)
        .where(TrainingClass.id == request.class_id)
        .with_for_update()
    )
    if training_class is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy lớp học.")
    fee_schedule = db.scalar(
        select(ProgramFeeSchedule)
        .where(ProgramFeeSchedule.id == request.fee_schedule_id)
        .with_for_update()
    )
    if fee_schedule is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy biểu phí.")
    course_in_program = db.scalar(
        select(program_subjects.c.program_id).where(
            program_subjects.c.program_id == fee_schedule.program_id,
            program_subjects.c.subject_id == training_class.course_id,
        )
    )
    if course_in_program is None:
        raise HTTPException(
            status_code=422,
            detail="Môn học của lớp không thuộc chương trình trong biểu phí.",
        )
    if request.applied_on < fee_schedule.valid_from or (
        fee_schedule.valid_to is not None
        and request.applied_on > fee_schedule.valid_to
    ):
        raise HTTPException(
            status_code=422,
            detail="Ngày áp dụng nằm ngoài thời gian hiệu lực của biểu phí.",
        )

    installment_total = sum(
        (item.amount_vnd for item in request.installments), Decimal("0")
    )
    if installment_total != fee_schedule.amount_vnd:
        raise HTTPException(
            status_code=422,
            detail="Tổng các đợt phải bằng tổng học phí của biểu phí đã chọn.",
        )

    enrollments = db.scalars(
        select(ClassEnrollment)
        .where(
            ClassEnrollment.class_id == training_class.id,
            ClassEnrollment.enrollment_status == "active",
            or_(
                ClassEnrollment.effective_date.is_(None),
                ClassEnrollment.effective_date <= request.applied_on,
            ),
        )
        .order_by(ClassEnrollment.id)
        .with_for_update()
    ).all()
    existing_plans = (
        db.scalars(
            select(TuitionPlan).where(
                TuitionPlan.enrollment_id.in_([item.id for item in enrollments])
            )
        ).all()
        if enrollments
        else []
    )
    existing_by_enrollment = {
        plan.enrollment_id: plan for plan in existing_plans
    }
    if existing_plans:
        existing_installments = db.scalars(
            select(TuitionInstallment)
            .where(TuitionInstallment.plan_id.in_([plan.id for plan in existing_plans]))
            .order_by(
                TuitionInstallment.plan_id,
                TuitionInstallment.installment_number,
            )
        ).all()
        installments_by_plan: dict[int, list[TuitionInstallment]] = {}
        for item in existing_installments:
            installments_by_plan.setdefault(item.plan_id, []).append(item)
        requested_installment_values = [
            (item.amount_vnd, item.due_date) for item in request.installments
        ]
        for plan in existing_plans:
            persisted_values = [
                (item.amount_due_vnd, item.due_date)
                for item in installments_by_plan.get(plan.id, [])
            ]
            if (
                plan.fee_schedule_id != fee_schedule.id
                or plan.total_due_vnd != installment_total
                or plan.applied_on != request.applied_on
                or persisted_values != requested_installment_values
            ):
                raise HTTPException(
                    status_code=409,
                    detail=(
                        "Đã tồn tại kế hoạch khác cho một học viên trong lớp; "
                        "không thể ghi đè khi thử lại."
                    ),
                )

    created_plans = []
    skipped_enrollment_ids = []
    for enrollment in enrollments:
        if enrollment.id in existing_by_enrollment:
            skipped_enrollment_ids.append(enrollment.id)
            continue
        plan = TuitionPlan(
            enrollment_id=enrollment.id,
            fee_schedule_id=fee_schedule.id,
            total_due_vnd=installment_total,
            applied_on=request.applied_on,
            created_by=current_user.id,
        )
        db.add(plan)
        db.flush()
        db.add_all(
            [
                TuitionInstallment(
                    plan_id=plan.id,
                    installment_number=number,
                    amount_due_vnd=item.amount_vnd,
                    due_date=item.due_date,
                )
                for number, item in enumerate(request.installments, start=1)
            ]
        )
        created_plans.append(
            {"plan_id": plan.id, "enrollment_id": enrollment.id}
        )

    db.flush()
    _record_audit(
        db,
        current_user,
        "tuition.plan_batch_created",
        details=(
            f"class_id={training_class.id};fee_schedule_id={fee_schedule.id};"
            f"created={len(created_plans)};skipped={len(skipped_enrollment_ids)}"
        ),
    )
    try:
        db.commit()
    except IntegrityError as error:
        db.rollback()
        raise HTTPException(
            status_code=409,
            detail="Kế hoạch học phí được tạo đồng thời. Tải lại trước khi thử lại.",
        ) from error
    return {
        "class_id": training_class.id,
        "fee_schedule_id": fee_schedule.id,
        "created_count": len(created_plans),
        "skipped_count": len(skipped_enrollment_ids),
        "plans": created_plans,
        "skipped_enrollment_ids": skipped_enrollment_ids,
    }


def _installment_balance(db: Session, installment: TuitionInstallment) -> tuple[Decimal, Decimal]:
    paid = db.scalar(
        select(func.coalesce(func.sum(TuitionPayment.amount_vnd), 0)).where(
            TuitionPayment.installment_id == installment.id
        )
    ) or Decimal("0")
    reversed_amount = db.scalar(
        select(func.coalesce(func.sum(TuitionPaymentReversal.amount_vnd), 0))
        .join(
            TuitionPayment,
            TuitionPayment.id == TuitionPaymentReversal.payment_id,
        )
        .where(TuitionPayment.installment_id == installment.id)
    ) or Decimal("0")
    net_paid = Decimal(str(paid)) - Decimal(str(reversed_amount))
    gross_remaining = max(Decimal("0"), installment.amount_due_vnd - net_paid)
    approved_discount = db.scalar(
        select(func.coalesce(func.sum(TuitionDiscount.amount_vnd), 0)).where(
            TuitionDiscount.plan_id == installment.plan_id,
            TuitionDiscount.status == "approved",
        )
    ) or Decimal("0")
    later_installments = db.scalars(
        select(TuitionInstallment)
        .where(
            TuitionInstallment.plan_id == installment.plan_id,
            TuitionInstallment.installment_number > installment.installment_number,
        )
        .order_by(TuitionInstallment.installment_number.desc())
    ).all()
    discount_left = Decimal(str(approved_discount))
    for later in later_installments:
        later_paid, _ = _gross_installment_balance(db, later)
        later_open = max(Decimal("0"), later.amount_due_vnd - later_paid)
        discount_left = max(Decimal("0"), discount_left - later_open)
    allocated_discount = min(gross_remaining, discount_left)
    return net_paid, gross_remaining - allocated_discount


def _gross_installment_balance(
    db: Session, installment: TuitionInstallment
) -> tuple[Decimal, Decimal]:
    paid = db.scalar(
        select(func.coalesce(func.sum(TuitionPayment.amount_vnd), 0)).where(
            TuitionPayment.installment_id == installment.id
        )
    ) or Decimal("0")
    reversed_amount = db.scalar(
        select(func.coalesce(func.sum(TuitionPaymentReversal.amount_vnd), 0))
        .join(TuitionPayment, TuitionPayment.id == TuitionPaymentReversal.payment_id)
        .where(TuitionPayment.installment_id == installment.id)
    ) or Decimal("0")
    net_paid = Decimal(str(paid)) - Decimal(str(reversed_amount))
    return net_paid, max(Decimal("0"), installment.amount_due_vnd - net_paid)


def _installment_status(
    installment: TuitionInstallment, remaining: Decimal, paid: Decimal = Decimal("0")
) -> str:
    if remaining == 0:
        return "paid"
    if paid > 0:
        return "partial"
    if installment.due_date < datetime.now(HO_CHI_MINH).date():
        return "overdue"
    return "unpaid"


def _tuition_discount_payload(discount: TuitionDiscount) -> dict:
    return {
        "id": discount.id,
        "plan_id": discount.plan_id,
        "kind": discount.kind,
        "value": str(discount.value),
        "amount_vnd": str(discount.amount_vnd),
        "reason": discount.reason,
        "status": discount.status,
        "requested_by": discount.requested_by,
        "reviewed_by": discount.reviewed_by,
        "created_at": discount.created_at.isoformat() if discount.created_at else None,
        "reviewed_at": discount.reviewed_at.isoformat() if discount.reviewed_at else None,
    }


def _plan_open_balance(db: Session, plan_id: int) -> Decimal:
    installments = db.scalars(
        select(TuitionInstallment).where(TuitionInstallment.plan_id == plan_id)
    ).all()
    total = Decimal("0")
    for installment in installments:
        _, remaining = _installment_balance(db, installment)
        total += remaining
    return total


def _plan_discount_total(db: Session, plan_id: int) -> Decimal:
    return Decimal(str(db.scalar(
        select(func.coalesce(func.sum(TuitionDiscount.amount_vnd), 0)).where(
            TuitionDiscount.plan_id == plan_id,
            TuitionDiscount.status == "approved",
        )
    ) or Decimal("0")))


def _discount_request_amount(plan: TuitionPlan, kind: str, value: Decimal) -> Decimal:
    if kind == "fixed":
        return value.quantize(Decimal("1"))
    return (plan.total_due_vnd * value / Decimal("100")).quantize(
        Decimal("1"), rounding=ROUND_HALF_UP
    )


@router.post("/tuition-plans/{plan_id}/discounts", status_code=status.HTTP_201_CREATED)
def create_tuition_discount(
    plan_id: int,
    request: TuitionDiscountCreate,
    current_user: User = Depends(require_roles("admin", "accountant")),
    db: Session = Depends(get_db),
) -> dict:
    plan = db.scalar(select(TuitionPlan).where(TuitionPlan.id == plan_id).with_for_update())
    if plan is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy kế hoạch học phí.")
    amount = _discount_request_amount(plan, request.kind, request.value)
    outstanding = _plan_open_balance(db, plan.id)
    if amount <= 0 or amount > outstanding:
        raise HTTPException(status_code=409, detail="Mức giảm vượt nghĩa vụ học phí chưa thanh toán.")
    manager_approval_required = amount * Decimal("100") > plan.total_due_vnd * Decimal("30")
    discount = TuitionDiscount(
        plan_id=plan.id,
        kind=request.kind,
        value=request.value,
        amount_vnd=amount,
        reason=request.reason,
        status="pending" if manager_approval_required else "approved",
        requested_by=current_user.id,
        reviewed_by=None if manager_approval_required else current_user.id,
        reviewed_at=None if manager_approval_required else datetime.now(timezone.utc).replace(tzinfo=None),
    )
    db.add(discount)
    db.flush()
    _record_audit(
        db,
        current_user,
        "tuition.discount_requested",
        details=f"discount_id={discount.id};plan_id={plan.id};amount={amount};status={discount.status};reason={request.reason}",
    )
    db.commit()
    db.refresh(discount)
    return _tuition_discount_payload(discount)


@router.get("/tuition-plans/{plan_id}/discounts")
def list_tuition_discounts(
    plan_id: int,
    response: Response,
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=20, ge=1, le=100),
    _user: User = Depends(require_roles("admin", "accountant", "manager")),
    db: Session = Depends(get_db),
) -> list[dict]:
    if db.get(TuitionPlan, plan_id) is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy kế hoạch học phí.")
    query = select(TuitionDiscount).where(TuitionDiscount.plan_id == plan_id)
    total = db.scalar(select(func.count()).select_from(query.subquery())) or 0
    rows = db.scalars(query.order_by(TuitionDiscount.id)
        .offset((page - 1) * page_size).limit(page_size)).all()
    set_pagination_headers(response, total=total, page=page, page_size=page_size)
    return [_tuition_discount_payload(row) for row in rows]


def _review_tuition_discount(
    discount_id: int, decision: str, manager: User, db: Session
) -> dict:
    discount = db.scalar(
        select(TuitionDiscount).where(TuitionDiscount.id == discount_id).with_for_update()
    )
    if discount is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy yêu cầu giảm học phí.")
    if discount.status == decision:
        return _tuition_discount_payload(discount)
    if discount.status != "pending":
        raise HTTPException(status_code=409, detail="Yêu cầu giảm học phí đã được xử lý.")
    plan = db.scalar(select(TuitionPlan).where(TuitionPlan.id == discount.plan_id).with_for_update())
    if plan is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy kế hoạch học phí.")
    if decision == "approved" and discount.amount_vnd > _plan_open_balance(db, plan.id):
        raise HTTPException(status_code=409, detail="Số dư đã thay đổi; không thể áp dụng mức giảm này.")
    discount.status = decision
    discount.reviewed_by = manager.id
    discount.reviewed_at = datetime.now(timezone.utc).replace(tzinfo=None)
    _record_audit(
        db,
        manager,
        f"tuition.discount_{decision}",
        details=f"discount_id={discount.id};plan_id={plan.id};amount={discount.amount_vnd};reason={discount.reason}",
    )
    db.commit()
    db.refresh(discount)
    return _tuition_discount_payload(discount)


@router.post("/tuition-discounts/{discount_id}/approve")
def approve_tuition_discount(
    discount_id: int,
    manager: User = Depends(require_roles("manager")),
    db: Session = Depends(get_db),
) -> dict:
    return _review_tuition_discount(discount_id, "approved", manager, db)


@router.post("/tuition-discounts/{discount_id}/reject")
def reject_tuition_discount(
    discount_id: int,
    manager: User = Depends(require_roles("manager")),
    db: Session = Depends(get_db),
) -> dict:
    return _review_tuition_discount(discount_id, "rejected", manager, db)


def _tuition_payment_context(db: Session, payment_id: int):
    return db.execute(
        select(
            TuitionPayment,
            TuitionInstallment,
            TuitionPlan,
            ClassEnrollment,
            User,
            TrainingClass,
        )
        .join(
            TuitionInstallment,
            TuitionInstallment.id == TuitionPayment.installment_id,
        )
        .join(TuitionPlan, TuitionPlan.id == TuitionInstallment.plan_id)
        .join(ClassEnrollment, ClassEnrollment.id == TuitionPlan.enrollment_id)
        .join(User, User.id == ClassEnrollment.student_id)
        .join(TrainingClass, TrainingClass.id == ClassEnrollment.class_id)
        .where(TuitionPayment.id == payment_id)
    ).first()


def _assert_tuition_payment_access(current_user: User, student_id: int) -> None:
    roles = {role.code for role in current_user.roles}
    if roles.intersection(("admin", "accountant")):
        return
    if "student" in roles and current_user.id == student_id:
        return
    # Conceal payment existence from users outside its financial scope.
    raise HTTPException(status_code=404, detail="Không tìm thấy biên lai.")


def _tuition_payment_payload(db: Session, payment: TuitionPayment) -> dict:
    context = _tuition_payment_context(db, payment.id)
    if context is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy biên lai.")
    _, installment, plan, enrollment, student, training_class = context
    net_paid, remaining = _installment_balance(db, installment)
    reversal = db.scalar(
        select(TuitionPaymentReversal).where(
            TuitionPaymentReversal.payment_id == payment.id
        )
    )
    return {
        "id": payment.id,
        "receipt_code": payment.receipt_code,
        "installment_id": installment.id,
        "plan_id": plan.id,
        "student_id": student.id,
        "student_code": student.student_code or student.username,
        "student_name": student.full_name,
        "class_id": training_class.id,
        "class_code": training_class.code,
        "amount_vnd": str(payment.amount_vnd),
        "method": payment.method,
        "paid_on": payment.paid_on.isoformat(),
        "collected_by": payment.collected_by,
        "note": payment.note,
        "installment_amount_vnd": str(net_paid + remaining),
        "installment_paid_vnd": str(net_paid),
        "installment_remaining_vnd": str(remaining),
        "installment_status": _installment_status(installment, remaining, net_paid),
        "reversed": reversal is not None,
        "reversal_reason": reversal.reason if reversal else None,
        "created_at": payment.created_at.isoformat() if payment.created_at else None,
    }


@router.post(
    "/tuition-installments/{installment_id}/payments",
    status_code=status.HTTP_201_CREATED,
    response_model=None,
)
def record_tuition_payment(
    installment_id: int,
    request: TuitionPaymentCreate,
    current_user: User = Depends(require_roles("admin", "accountant")),
    db: Session = Depends(get_db),
) -> dict | JSONResponse:
    installment_plan_id = db.scalar(
        select(TuitionInstallment.plan_id).where(TuitionInstallment.id == installment_id)
    )
    if installment_plan_id is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy đợt học phí.")
    # Payments and discount approvals lock the plan first so neither can spend
    # the same remaining balance concurrently on databases with row locking.
    db.scalar(
        select(TuitionPlan.id).where(TuitionPlan.id == installment_plan_id).with_for_update()
    )
    installment = db.scalar(
        select(TuitionInstallment)
        .where(TuitionInstallment.id == installment_id)
        .with_for_update()
    )
    if installment is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy đợt học phí.")

    prior = db.scalar(
        select(TuitionPayment).where(
            TuitionPayment.installment_id == installment.id,
            TuitionPayment.idempotency_key == request.idempotency_key,
        )
    )
    if prior is not None:
        same_request = (
            prior.amount_vnd == request.amount_vnd
            and prior.method == request.method
            and prior.paid_on == request.paid_on
            and prior.note == request.note
        )
        if not same_request:
            raise HTTPException(
                status_code=409,
                detail="Idempotency key đã được dùng cho một giao dịch khác.",
            )
        return JSONResponse(
            status_code=200,
            content=_tuition_payment_payload(db, prior),
        )

    net_paid, remaining = _installment_balance(db, installment)
    if request.amount_vnd > remaining:
        raise HTTPException(
            status_code=409,
            detail="Số tiền thu vượt quá số dư còn phải thu của đợt.",
        )
    receipt_date = datetime.now(HO_CHI_MINH).strftime("%Y%m%d")
    payment = TuitionPayment(
        installment_id=installment.id,
        receipt_code=f"ICTU-{receipt_date}-{secrets.token_hex(5).upper()}",
        idempotency_key=request.idempotency_key,
        amount_vnd=request.amount_vnd,
        method=request.method,
        paid_on=request.paid_on,
        collected_by=current_user.id,
        note=request.note,
    )
    db.add(payment)
    db.flush()
    context = _tuition_payment_context(db, payment.id)
    student = context[4] if context else None
    _record_audit(
        db,
        current_user,
        "tuition.payment_recorded",
        target_user_id=student.id if student else None,
        details=(
            f"payment_id={payment.id};receipt={payment.receipt_code};"
            f"installment_id={installment.id};amount={payment.amount_vnd};"
            f"before_remaining={remaining};after_remaining={remaining - payment.amount_vnd}"
        ),
    )
    try:
        db.commit()
    except IntegrityError as error:
        db.rollback()
        raise HTTPException(
            status_code=409,
            detail="Giao dịch bị trùng hoặc được ghi đồng thời; tải lại số dư trước khi thử lại.",
        ) from error
    db.refresh(payment)
    return _tuition_payment_payload(db, payment)


@router.post(
    "/tuition-payments/{payment_id}/reverse",
    status_code=status.HTTP_201_CREATED,
    response_model=None,
)
def reverse_tuition_payment(
    payment_id: int,
    request: TuitionPaymentReversalCreate,
    current_user: User = Depends(require_roles("admin", "accountant")),
    db: Session = Depends(get_db),
) -> dict | JSONResponse:
    payment = db.scalar(
        select(TuitionPayment).where(TuitionPayment.id == payment_id).with_for_update()
    )
    if payment is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy giao dịch thu.")
    context = _tuition_payment_context(db, payment.id)
    if context is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy giao dịch thu.")
    installment = db.scalar(
        select(TuitionInstallment)
        .where(TuitionInstallment.id == payment.installment_id)
        .with_for_update()
    )
    if installment is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy đợt học phí.")

    prior = db.scalar(
        select(TuitionPaymentReversal).where(
            TuitionPaymentReversal.payment_id == payment.id,
            TuitionPaymentReversal.idempotency_key == request.idempotency_key
        )
    )
    if prior is not None:
        if prior.payment_id != payment.id or prior.reason != request.reason:
            raise HTTPException(
                status_code=409,
                detail="Idempotency key đã được dùng cho một giao dịch đảo khác.",
            )
        return JSONResponse(
            status_code=200,
            content={
                "id": prior.id,
                "payment_id": prior.payment_id,
                "amount_vnd": str(prior.amount_vnd),
                "reason": prior.reason,
                "created_at": prior.created_at.isoformat() if prior.created_at else None,
                "payment": _tuition_payment_payload(db, payment),
            },
        )
    if db.scalar(
        select(TuitionPaymentReversal.id).where(
            TuitionPaymentReversal.payment_id == payment.id
        )
    ) is not None:
        raise HTTPException(status_code=409, detail="Giao dịch đã được đảo trước đó.")

    reversal = TuitionPaymentReversal(
        payment_id=payment.id,
        idempotency_key=request.idempotency_key,
        amount_vnd=payment.amount_vnd,
        reason=request.reason,
        reversed_by=current_user.id,
    )
    db.add(reversal)
    db.flush()
    student = context[4]
    _record_audit(
        db,
        current_user,
        "tuition.payment_reversed",
        target_user_id=student.id,
        details=(
            f"payment_id={payment.id};reversal_id={reversal.id};"
            f"amount={reversal.amount_vnd};reason={reversal.reason}"
        ),
    )
    try:
        db.commit()
    except IntegrityError as error:
        db.rollback()
        raise HTTPException(
            status_code=409,
            detail="Giao dịch đảo được ghi đồng thời; tải lại trước khi thử lại.",
        ) from error
    db.refresh(reversal)
    return {
        "id": reversal.id,
        "payment_id": reversal.payment_id,
        "amount_vnd": str(reversal.amount_vnd),
        "reason": reversal.reason,
        "created_at": reversal.created_at.isoformat() if reversal.created_at else None,
        "payment": _tuition_payment_payload(db, payment),
    }


def _get_authorized_tuition_payment(
    payment_id: int, current_user: User, db: Session
) -> tuple[TuitionPayment, tuple]:
    context = _tuition_payment_context(db, payment_id)
    if context is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy biên lai.")
    payment = context[0]
    _assert_tuition_payment_access(current_user, context[4].id)
    return payment, context


@router.get("/tuition-payments")
def list_tuition_payments(
    response: Response,
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=20, ge=1, le=100),
    student_id: int | None = Query(default=None, ge=1),
    class_id: int | None = Query(default=None, ge=1),
    paid_from: date | None = None,
    paid_to: date | None = None,
    sort: str = Query(default="paid_on_desc", pattern="^(paid_on_asc|paid_on_desc)$"),
    current_user: User = Depends(require_roles("admin", "accountant", "student")),
    db: Session = Depends(get_db),
) -> dict:
    roles = {role.code for role in current_user.roles}
    query = (
        select(TuitionPayment)
        .join(TuitionInstallment, TuitionInstallment.id == TuitionPayment.installment_id)
        .join(TuitionPlan, TuitionPlan.id == TuitionInstallment.plan_id)
        .join(ClassEnrollment, ClassEnrollment.id == TuitionPlan.enrollment_id)
    )
    if "student" in roles and not roles.intersection(("admin", "accountant")):
        query = query.where(ClassEnrollment.student_id == current_user.id)
    else:
        if student_id is not None:
            query = query.where(ClassEnrollment.student_id == student_id)
        if class_id is not None:
            query = query.where(ClassEnrollment.class_id == class_id)
    if paid_from is not None:
        query = query.where(TuitionPayment.paid_on >= paid_from)
    if paid_to is not None:
        if paid_from is not None and paid_to < paid_from:
            raise HTTPException(status_code=422, detail="paid_to must not precede paid_from.")
        query = query.where(TuitionPayment.paid_on <= paid_to)
    total = db.scalar(select(func.count()).select_from(query.order_by(None).subquery())) or 0
    paid_order = TuitionPayment.paid_on.asc() if sort == "paid_on_asc" else TuitionPayment.paid_on.desc()
    payments = db.scalars(query.order_by(paid_order, TuitionPayment.id.asc())
        .offset((page - 1) * page_size).limit(page_size)).all()
    set_pagination_headers(response, total=total, page=page, page_size=page_size)
    return {"items": [_tuition_payment_payload(db, payment) for payment in payments],
        **page_metadata(total=total, page=page, page_size=page_size)}


@router.get("/tuition-payments/{payment_id}")
def get_tuition_payment(
    payment_id: int,
    current_user: User = Depends(require_roles("admin", "accountant", "student")),
    db: Session = Depends(get_db),
) -> dict:
    payment, _ = _get_authorized_tuition_payment(payment_id, current_user, db)
    return _tuition_payment_payload(db, payment)


@router.get("/tuition-payments/{payment_id}/receipt.pdf")
def download_tuition_receipt(
    payment_id: int,
    current_user: User = Depends(require_roles("admin", "accountant", "student")),
    db: Session = Depends(get_db),
) -> Response:
    payment, context = _get_authorized_tuition_payment(payment_id, current_user, db)
    _, installment, plan, enrollment, student, training_class = context
    payload = _tuition_payment_payload(db, payment)
    lines = [
        "ICTU CodeGym - Payment Receipt",
        f"Receipt: {payment.receipt_code}",
        f"Student: {student.full_name} ({student.student_code or student.username})",
        f"Class: {training_class.code}",
        f"Installment: {installment.installment_number}",
        f"Amount: {payment.amount_vnd} VND",
        f"Method: {payment.method}",
        f"Paid on: {payment.paid_on.isoformat()}",
        f"Balance remaining: {payload['installment_remaining_vnd']} VND",
        "Status: REVERSED" if payload["reversed"] else "Status: RECORDED",
    ]
    filename = f"{payment.receipt_code}.pdf"
    return Response(
        content=build_receipt_pdf(lines),
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/students/me/tuition")
def get_my_tuition(
    response: Response,
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=20, ge=1, le=100),
    current_user: User = Depends(require_roles("student")),
    db: Session = Depends(get_db),
) -> dict:
    """Return the signed-in learner's plans, balances, and payment history."""
    base = (
        select(TuitionPlan)
        .join(ClassEnrollment, ClassEnrollment.id == TuitionPlan.enrollment_id)
        .where(ClassEnrollment.student_id == current_user.id)
    )
    total = db.scalar(select(func.count()).select_from(base.subquery())) or 0
    plans = db.scalars(
        base
        .order_by(TuitionPlan.applied_on.desc(), TuitionPlan.id.desc())
        .offset((page - 1) * page_size).limit(page_size)
    ).all()
    result = []
    for plan in plans:
        enrollment = db.get(ClassEnrollment, plan.enrollment_id)
        training_class = db.get(TrainingClass, enrollment.class_id) if enrollment else None
        installments = db.scalars(
            select(TuitionInstallment)
            .where(TuitionInstallment.plan_id == plan.id)
            .order_by(TuitionInstallment.installment_number)
        ).all()
        installment_rows = []
        for installment in installments:
            net_paid, remaining = _installment_balance(db, installment)
            payments = db.scalars(
                select(TuitionPayment)
                .where(TuitionPayment.installment_id == installment.id)
                .order_by(TuitionPayment.paid_on, TuitionPayment.id)
            ).all()
            installment_rows.append({
                "id": installment.id,
                "number": installment.installment_number,
                "due_date": installment.due_date.isoformat(),
                "amount_due_vnd": str(net_paid + remaining),
                "paid_vnd": str(net_paid),
                "remaining_vnd": str(remaining),
                "status": _installment_status(installment, remaining, net_paid),
                "payments": [_tuition_payment_payload(db, payment) for payment in payments],
            })
        result.append({
            "plan_id": plan.id,
            "class_id": enrollment.class_id if enrollment else None,
            "class_code": training_class.code if training_class else None,
            "total_due_vnd": str(plan.total_due_vnd - _plan_discount_total(db, plan.id)),
            "discounts": [
                _tuition_discount_payload(row)
                for row in db.scalars(
                    select(TuitionDiscount)
                    .where(TuitionDiscount.plan_id == plan.id, TuitionDiscount.status == "approved")
                    .order_by(TuitionDiscount.id)
                ).all()
            ],
            "installments": installment_rows,
        })
    return {"student_id": current_user.id, "plans": result,
        **page_metadata(total=total, page=page, page_size=page_size)}


def _tuition_debt_rows(
    db: Session,
    *,
    class_id: int | None,
    student_id: int | None,
    status_filter: str | None,
) -> list[dict]:
    query = (
        select(TuitionInstallment, TuitionPlan, ClassEnrollment, User, TrainingClass,
               StudyProgram.id, StudyProgram.name)
        .join(TuitionPlan, TuitionPlan.id == TuitionInstallment.plan_id)
        .join(ClassEnrollment, ClassEnrollment.id == TuitionPlan.enrollment_id)
        .join(User, User.id == ClassEnrollment.student_id)
        .join(TrainingClass, TrainingClass.id == ClassEnrollment.class_id)
        .outerjoin(StudyProgram, StudyProgram.id == TrainingClass.program_id)
    )
    if class_id is not None:
        query = query.where(ClassEnrollment.class_id == class_id)
    if student_id is not None:
        query = query.where(ClassEnrollment.student_id == student_id)
    records = db.execute(query.order_by(TuitionInstallment.due_date, TuitionInstallment.id)).all()
    rows = []
    for installment, plan, enrollment, student, training_class, program_id, program_name in records:
        paid, remaining = _installment_balance(db, installment)
        state = _installment_status(installment, remaining, paid)
        overdue_days = (
            max(0, (datetime.now(HO_CHI_MINH).date() - installment.due_date).days)
            if state == "overdue"
            else 0
        )
        if status_filter is not None and state != status_filter:
            continue
        rows.append({
            "installment_id": installment.id,
            "plan_id": plan.id,
            "student_id": student.id,
            "student_name": student.full_name,
            "student_code": student.student_code,
            "class_id": training_class.id,
            "class_code": training_class.code,
            "class_name": training_class.name,
            "program_id": program_id,
            "program_name": program_name,
            "installment_number": installment.installment_number,
            "due_date": installment.due_date.isoformat(),
            "amount_due_vnd": str(paid + remaining),
            "paid_vnd": str(paid),
            "remaining_vnd": str(remaining),
            "status": state,
            "overdue_days": overdue_days,
        })
    return rows


def _tuition_debt_statement(*, class_id: int | None, student_id: int | None):
    paid_by_installment = select(
        TuitionPayment.installment_id.label("installment_id"),
        func.sum(TuitionPayment.amount_vnd).label("paid_vnd"),
    ).group_by(TuitionPayment.installment_id).subquery()
    reversed_by_installment = select(
        TuitionPayment.installment_id.label("installment_id"),
        func.sum(TuitionPaymentReversal.amount_vnd).label("reversed_vnd"),
    ).join(TuitionPayment, TuitionPayment.id == TuitionPaymentReversal.payment_id)
    reversed_by_installment = reversed_by_installment.group_by(
        TuitionPayment.installment_id
    ).subquery()
    discount_by_plan = select(
        TuitionDiscount.plan_id.label("plan_id"),
        func.sum(TuitionDiscount.amount_vnd).label("discount_vnd"),
    ).where(TuitionDiscount.status == "approved").group_by(
        TuitionDiscount.plan_id
    ).subquery()
    paid = (func.coalesce(paid_by_installment.c.paid_vnd, 0)
        - func.coalesce(reversed_by_installment.c.reversed_vnd, 0))
    gross_remaining = case(
        (TuitionInstallment.amount_due_vnd - paid > 0,
         TuitionInstallment.amount_due_vnd - paid),
        else_=0,
    )
    base = select(
        TuitionInstallment.id.label("installment_id"), TuitionInstallment.plan_id,
        TuitionInstallment.installment_number, TuitionInstallment.due_date,
        TuitionInstallment.amount_due_vnd, paid.label("paid_vnd"),
        gross_remaining.label("gross_remaining_vnd"),
        func.coalesce(discount_by_plan.c.discount_vnd, 0).label("discount_vnd"),
    ).join(TuitionPlan, TuitionPlan.id == TuitionInstallment.plan_id).outerjoin(
        paid_by_installment, paid_by_installment.c.installment_id == TuitionInstallment.id
    ).outerjoin(reversed_by_installment,
        reversed_by_installment.c.installment_id == TuitionInstallment.id
    ).outerjoin(discount_by_plan, discount_by_plan.c.plan_id == TuitionInstallment.plan_id)
    if class_id is not None:
        base = base.join(ClassEnrollment, ClassEnrollment.id == TuitionPlan.enrollment_id).where(
            ClassEnrollment.class_id == class_id
        )
    if student_id is not None:
        if class_id is None:
            base = base.join(ClassEnrollment, ClassEnrollment.id == TuitionPlan.enrollment_id)
        base = base.where(ClassEnrollment.student_id == student_id)
    base = base.subquery()
    later_open = func.coalesce(func.sum(base.c.gross_remaining_vnd).over(
        partition_by=base.c.plan_id, order_by=base.c.installment_number,
        rows=(1, None),
    ), 0)
    discount_available = base.c.discount_vnd - later_open
    discount_allocated = case(
        (discount_available <= 0, 0),
        (discount_available >= base.c.gross_remaining_vnd, base.c.gross_remaining_vnd),
        else_=discount_available,
    )
    remaining = case(
        (base.c.gross_remaining_vnd > discount_allocated,
         base.c.gross_remaining_vnd - discount_allocated),
        else_=0,
    )
    local_today = datetime.now(HO_CHI_MINH).date()
    state = case(
        (remaining <= 0, "paid"),
        (base.c.paid_vnd > 0, "partial"),
        (base.c.due_date < local_today, "overdue"),
        else_="unpaid",
    )
    return select(
        base.c.installment_id, base.c.plan_id, ClassEnrollment.student_id,
        User.full_name.label("student_name"), User.student_code,
        TrainingClass.id.label("class_id"), TrainingClass.code.label("class_code"),
        TrainingClass.name.label("class_name"), StudyProgram.id.label("program_id"),
        StudyProgram.name.label("program_name"), base.c.installment_number,
        base.c.due_date, base.c.paid_vnd, remaining.label("remaining_vnd"),
        state.label("status"),
    ).join(TuitionPlan, TuitionPlan.id == base.c.plan_id).join(
        ClassEnrollment, ClassEnrollment.id == TuitionPlan.enrollment_id
    ).join(User, User.id == ClassEnrollment.student_id).join(
        TrainingClass, TrainingClass.id == ClassEnrollment.class_id
    ).outerjoin(StudyProgram, StudyProgram.id == TrainingClass.program_id)


@router.get("/tuition-debts")
def list_tuition_debts(
    class_id: int | None = Query(default=None, ge=1),
    student_id: int | None = Query(default=None, ge=1),
    status_filter: str | None = Query(default=None, alias="status", pattern="^(unpaid|partial|overdue|paid)$"),
    sort: str = Query(default="due_date", pattern="^(due_date|remaining_desc)$"),
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=20, ge=1, le=100),
    _user: User = Depends(require_roles("admin", "accountant")),
    db: Session = Depends(get_db),
) -> dict:
    debts = _tuition_debt_statement(class_id=class_id, student_id=student_id).subquery()
    query = select(debts)
    if status_filter is not None:
        query = query.where(debts.c.status == status_filter)
    filtered_debts = query.subquery()
    aggregate = db.execute(select(
        func.count().label("total"),
        func.coalesce(func.sum(filtered_debts.c.paid_vnd + filtered_debts.c.remaining_vnd), 0).label("due"),
        func.coalesce(func.sum(filtered_debts.c.paid_vnd), 0).label("paid"),
        func.coalesce(func.sum(filtered_debts.c.remaining_vnd), 0).label("remaining"),
        func.coalesce(func.sum(case((filtered_debts.c.status == "overdue", filtered_debts.c.remaining_vnd),
            else_=0)), 0).label("overdue"),
    ).select_from(filtered_debts)).one()
    if sort == "remaining_desc":
        ordering = (debts.c.remaining_vnd.desc(), debts.c.due_date, debts.c.installment_id)
    else:
        ordering = (debts.c.due_date, debts.c.installment_id)
    page_rows = db.execute(query.order_by(*ordering)
        .offset((page - 1) * page_size).limit(page_size)).all()
    today = datetime.now(HO_CHI_MINH).date()
    items = [{
        "installment_id": row.installment_id,
        "plan_id": row.plan_id,
        "student_id": row.student_id,
        "student_name": row.student_name,
        "student_code": row.student_code,
        "class_id": row.class_id,
        "class_code": row.class_code,
        "class_name": row.class_name,
        "program_id": row.program_id,
        "program_name": row.program_name,
        "installment_number": row.installment_number,
        "due_date": row.due_date.isoformat(),
        "amount_due_vnd": str(Decimal(str(row.paid_vnd)) + Decimal(str(row.remaining_vnd))),
        "paid_vnd": str(row.paid_vnd),
        "remaining_vnd": str(row.remaining_vnd),
        "status": row.status,
        "overdue_days": max(0, (today - row.due_date).days) if row.status == "overdue" else 0,
    } for row in page_rows]
    return {
        **page_metadata(total=aggregate.total, page=page, page_size=page_size),
        "items": items,
        "total_due_vnd": str(aggregate.due),
        "total_paid_vnd": str(aggregate.paid),
        "total_remaining_vnd": str(aggregate.remaining),
        "total_overdue_vnd": str(aggregate.overdue),
    }


@router.get("/tuition-debts/export.xlsx")
def export_tuition_debts(
    class_id: int | None = Query(default=None, ge=1),
    student_id: int | None = Query(default=None, ge=1),
    status_filter: str | None = Query(default=None, alias="status", pattern="^(unpaid|partial|overdue|paid)$"),
    _user: User = Depends(require_roles("admin", "accountant")),
    db: Session = Depends(get_db),
) -> Response:
    rows = _tuition_debt_rows(db, class_id=class_id, student_id=student_id, status_filter=status_filter)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Tuition debts"
    sheet.append(["Student code", "Student", "Class", "Installment", "Due date", "Amount VND", "Paid VND", "Remaining VND", "Status", "Overdue days"])
    for row in rows:
        sheet.append([row["student_code"] or "", row["student_name"], row["class_code"], row["installment_number"], row["due_date"], row["amount_due_vnd"], row["paid_vnd"], row["remaining_vnd"], row["status"], row["overdue_days"]])
    sheet.append(["TOTAL", "", "", "", "", sum((Decimal(row["amount_due_vnd"]) for row in rows), Decimal("0")), sum((Decimal(row["paid_vnd"]) for row in rows), Decimal("0")), sum((Decimal(row["remaining_vnd"]) for row in rows), Decimal("0")), "", ""])
    buffer = BytesIO()
    workbook.save(buffer)
    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="tuition-debts.xlsx"'},
    )


def _tuition_revenue_events(db: Session, from_date: date | None, to_date: date | None) -> list[tuple[date, Decimal, str, int | None, str | None, str | None, int | None, str | None]]:
    events: list[tuple[date, Decimal, str, int | None, str | None, str | None, int | None, str | None]] = []
    payment_query = select(
        TuitionPayment.paid_on, TuitionPayment.amount_vnd,
        TrainingClass.id, TrainingClass.code, TrainingClass.name,
        StudyProgram.id, StudyProgram.name,
    ).join(TuitionInstallment, TuitionInstallment.id == TuitionPayment.installment_id)
    payment_query = payment_query.join(TuitionPlan, TuitionPlan.id == TuitionInstallment.plan_id)
    payment_query = payment_query.join(ClassEnrollment, ClassEnrollment.id == TuitionPlan.enrollment_id)
    payment_query = payment_query.join(TrainingClass, TrainingClass.id == ClassEnrollment.class_id)
    payment_query = payment_query.outerjoin(StudyProgram, StudyProgram.id == TrainingClass.program_id)
    payment_query = payment_query.order_by(TuitionPayment.paid_on)
    if from_date is not None:
        payment_query = payment_query.where(TuitionPayment.paid_on >= from_date)
    if to_date is not None:
        payment_query = payment_query.where(TuitionPayment.paid_on <= to_date)
    events.extend((paid_on, Decimal(amount), "payment", class_id, code, name, program_id, program_name)
                  for paid_on, amount, class_id, code, name, program_id, program_name in db.execute(payment_query))

    reversal_query = select(
        TuitionPaymentReversal.created_at, TuitionPaymentReversal.amount_vnd,
        TrainingClass.id, TrainingClass.code, TrainingClass.name,
        StudyProgram.id, StudyProgram.name,
    ).join(TuitionPayment, TuitionPayment.id == TuitionPaymentReversal.payment_id)
    reversal_query = reversal_query.join(TuitionInstallment, TuitionInstallment.id == TuitionPayment.installment_id)
    reversal_query = reversal_query.join(TuitionPlan, TuitionPlan.id == TuitionInstallment.plan_id)
    reversal_query = reversal_query.join(ClassEnrollment, ClassEnrollment.id == TuitionPlan.enrollment_id)
    reversal_query = reversal_query.join(TrainingClass, TrainingClass.id == ClassEnrollment.class_id)
    reversal_query = reversal_query.outerjoin(StudyProgram, StudyProgram.id == TrainingClass.program_id)
    reversal_query = reversal_query.order_by(TuitionPaymentReversal.created_at)
    if from_date is not None:
        reversal_query = reversal_query.where(TuitionPaymentReversal.created_at >= datetime.combine(from_date, datetime.min.time()))
    if to_date is not None:
        reversal_query = reversal_query.where(TuitionPaymentReversal.created_at < datetime.combine(to_date + timedelta(days=1), datetime.min.time()))
    events.extend((created_at.date(), -Decimal(amount), "reversal", class_id, code, name, program_id, program_name)
                  for created_at, amount, class_id, code, name, program_id, program_name in db.execute(reversal_query))
    return events


def _tuition_revenue_report(db: Session, from_date: date | None, to_date: date | None) -> dict:
    if from_date is not None and to_date is not None and from_date > to_date:
        raise HTTPException(status_code=422, detail="Khoảng ngày báo cáo không hợp lệ.")
    monthly: dict[str, Decimal] = {}
    by_class: dict[int, dict] = {}
    by_program: dict[int | None, dict] = {}
    payment_count = reversal_count = 0
    for event_date, amount, kind, class_id, class_code, class_name, program_id, program_name in _tuition_revenue_events(db, from_date, to_date):
        month = event_date.strftime("%Y-%m")
        monthly[month] = monthly.get(month, Decimal("0")) + amount
        row = by_class.setdefault(class_id, {"class_id": class_id, "class_code": class_code,
            "class_name": class_name, "net_revenue_vnd": Decimal("0"),
            "payment_count": 0, "reversal_count": 0,
            "total_due_vnd": Decimal("0"), "paid_vnd": Decimal("0"),
            "remaining_vnd": Decimal("0")})
        row["program_id"] = program_id
        row["program_name"] = program_name
        row["net_revenue_vnd"] += amount
        program_row = by_program.setdefault(program_id, {"program_id": program_id,
            "program_name": program_name or "Unassigned", "net_revenue_vnd": Decimal("0"),
            "payment_count": 0, "reversal_count": 0,
            "total_due_vnd": Decimal("0"), "paid_vnd": Decimal("0"),
            "remaining_vnd": Decimal("0")})
        program_row["net_revenue_vnd"] += amount
        if kind == "payment":
            payment_count += 1
            row["payment_count"] += 1
            program_row["payment_count"] += 1
        else:
            reversal_count += 1
            row["reversal_count"] += 1
            program_row["reversal_count"] += 1
    for debt in _tuition_debt_rows(db, class_id=None, student_id=None, status_filter=None):
        class_row = by_class.setdefault(debt["class_id"], {
            "class_id": debt["class_id"], "class_code": debt["class_code"],
            "class_name": debt.get("class_name"), "program_id": debt.get("program_id"),
            "program_name": debt.get("program_name"), "net_revenue_vnd": Decimal("0"),
            "payment_count": 0, "reversal_count": 0,
            "total_due_vnd": Decimal("0"), "paid_vnd": Decimal("0"),
            "remaining_vnd": Decimal("0"),
        })
        class_row["total_due_vnd"] += Decimal(debt["amount_due_vnd"])
        class_row["paid_vnd"] += Decimal(debt["paid_vnd"])
        class_row["remaining_vnd"] += Decimal(debt["remaining_vnd"])
        program_id = debt.get("program_id")
        program_row = by_program.setdefault(program_id, {
            "program_id": program_id, "program_name": debt.get("program_name") or "Unassigned",
            "net_revenue_vnd": Decimal("0"), "payment_count": 0, "reversal_count": 0,
            "total_due_vnd": Decimal("0"), "paid_vnd": Decimal("0"),
            "remaining_vnd": Decimal("0"),
        })
        program_row["total_due_vnd"] += Decimal(debt["amount_due_vnd"])
        program_row["paid_vnd"] += Decimal(debt["paid_vnd"])
        program_row["remaining_vnd"] += Decimal(debt["remaining_vnd"])
    ordered_months = sorted(monthly)
    return {
        "from_date": from_date.isoformat() if from_date else None,
        "to_date": to_date.isoformat() if to_date else None,
        "net_revenue_vnd": str(sum(monthly.values(), Decimal("0"))),
        "payment_count": payment_count,
        "reversal_count": reversal_count,
        "monthly": [{"month": month, "net_revenue_vnd": str(monthly[month])} for month in ordered_months],
        "by_class": [{**row, **{key: str(row[key]) for key in ("net_revenue_vnd", "total_due_vnd", "paid_vnd", "remaining_vnd")}}
                     for row in sorted(by_class.values(), key=lambda item: (item["class_code"] or "", item["class_id"] or 0))],
        "by_program": [{**row, **{key: str(row[key]) for key in ("net_revenue_vnd", "total_due_vnd", "paid_vnd", "remaining_vnd")}}
                       for row in sorted(by_program.values(), key=lambda item: (item["program_name"] or "", item["program_id"] or 0))],
    }


@router.get("/tuition-revenue")
def get_tuition_revenue(
    from_date: date | None = None,
    to_date: date | None = None,
    _user: User = Depends(require_roles("admin", "accountant")),
    db: Session = Depends(get_db),
) -> dict:
    return _tuition_revenue_report(db, from_date, to_date)


@router.get("/tuition-revenue/export.xlsx")
def export_tuition_revenue(
    from_date: date | None = None,
    to_date: date | None = None,
    _user: User = Depends(require_roles("admin", "accountant")),
    db: Session = Depends(get_db),
) -> Response:
    report = _tuition_revenue_report(db, from_date, to_date)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Tuition revenue"
    sheet.append(["Month", "Net revenue VND"])
    for row in report["monthly"]:
        sheet.append([row["month"], row["net_revenue_vnd"]])
    sheet.append(["Total", report["net_revenue_vnd"]])
    class_sheet = workbook.create_sheet("By class")
    class_sheet.append(["Class code", "Class name", "Program", "Net revenue VND", "Amount due VND", "Paid VND", "Remaining VND", "Payments", "Reversals"])
    for row in report["by_class"]:
        class_sheet.append([row["class_code"], row["class_name"], row["program_name"],
            row["net_revenue_vnd"], row["total_due_vnd"], row["paid_vnd"],
            row["remaining_vnd"], row["payment_count"], row["reversal_count"]])
    program_sheet = workbook.create_sheet("By program")
    program_sheet.append(["Program", "Net revenue VND", "Amount due VND", "Paid VND", "Remaining VND", "Payments", "Reversals"])
    for row in report["by_program"]:
        program_sheet.append([row["program_name"], row["net_revenue_vnd"],
            row["total_due_vnd"], row["paid_vnd"], row["remaining_vnd"],
            row["payment_count"], row["reversal_count"]])
    buffer = BytesIO()
    workbook.save(buffer)
    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="tuition-revenue.xlsx"'},
    )


def _study_program_catalog_payload(program: StudyProgram) -> dict:
    return {
        "id": program.id, "code": program.code, "name": program.name,
        "degree_level": program.degree_level, "total_credits": program.total_credits,
        "duration_semesters": program.duration_semesters, "description": program.description,
        "is_active": program.is_active,
        "standard_tuition_vnd": str(program.standard_tuition_vnd),
    }


@router.post("/study-programs", status_code=status.HTTP_201_CREATED)
def create_study_program(
    request: StudyProgramCreate,
    _manager: User = Depends(require_roles(*ADMIN_ROLES)),
    db: Session = Depends(get_db),
) -> dict:
    normalized_code = request.code.strip().upper()
    normalized_name = request.name.strip()
    if not normalized_name:
        raise HTTPException(status_code=422, detail="Program name is required.")
    if db.scalar(select(StudyProgram.id).where(or_(
        StudyProgram.code == normalized_code, StudyProgram.name == normalized_name
    ))):
        raise HTTPException(status_code=409, detail="Program code or name already exists.")
    fields = request.model_dump()
    fields["code"] = normalized_code
    fields["name"] = normalized_name
    program = StudyProgram(**fields)
    db.add(program)
    try:
        db.commit()
    except IntegrityError as error:
        db.rollback()
        raise HTTPException(status_code=409, detail="Program code or name already exists.") from error
    db.refresh(program)
    return _study_program_catalog_payload(program)


@router.patch("/study-programs/{program_id}")
def update_study_program(
    program_id: int,
    request: StudyProgramUpdate,
    _manager: User = Depends(require_roles(*ADMIN_ROLES)),
    db: Session = Depends(get_db),
) -> dict:
    program = db.get(StudyProgram, program_id)
    if program is None:
        raise HTTPException(status_code=404, detail="Study program not found.")
    updates = request.model_dump(exclude_unset=True)
    if "name" in updates:
        updates["name"] = updates["name"].strip()
        if not updates["name"]:
            raise HTTPException(status_code=422, detail="Program name is required.")
        duplicate = db.scalar(select(StudyProgram.id).where(
            StudyProgram.name == updates["name"], StudyProgram.id != program_id
        ))
        if duplicate:
            raise HTTPException(status_code=409, detail="Program name already exists.")
    for key, value in updates.items():
        setattr(program, key, value)
    try:
        db.commit()
    except IntegrityError as error:
        db.rollback()
        raise HTTPException(status_code=409, detail="Program update conflicts with existing data.") from error
    return _study_program_catalog_payload(program)


@router.delete("/study-programs/{program_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_study_program(
    program_id: int,
    _manager: User = Depends(require_roles(*ADMIN_ROLES)),
    db: Session = Depends(get_db),
) -> None:
    program = db.get(StudyProgram, program_id)
    if program is None:
        raise HTTPException(status_code=404, detail="Study program not found.")
    classes = db.scalars(select(TrainingClass).where(TrainingClass.program_id == program_id)).all()
    if any(training_class.status == "active" for training_class in classes):
        raise HTTPException(status_code=409, detail="Program has active classes; deactivate the program instead.")
    if classes or db.scalar(select(ProgramFeeSchedule.id).where(
        ProgramFeeSchedule.program_id == program_id
    ).limit(1)):
        raise HTTPException(status_code=409, detail="Program has retained class or fee history; deactivate it instead.")
    try:
        db.delete(program)
        db.commit()
    except IntegrityError as error:
        db.rollback()
        raise HTTPException(status_code=409, detail="Program is referenced by retained records; deactivate it instead.") from error


@router.get("/study-programs")
def list_study_programs(
    response: Response,
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=20, ge=1, le=100),
    course_page: int = Query(default=1, ge=1),
    course_page_size: int = Query(default=20, ge=1, le=100),
    course_q: str = Query(default="", max_length=150),
    course_sort: str = Query(default="sequence", pattern="^(sequence|code|name)$"),
    _user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
) -> list[dict]:
    base_query = select(StudyProgram)
    total = db.scalar(select(func.count()).select_from(base_query.subquery())) or 0
    programs = db.scalars(
        base_query
        .order_by(StudyProgram.name, StudyProgram.id)
        .offset((page - 1) * page_size).limit(page_size)
    ).all()
    set_pagination_headers(response, total=total, page=page, page_size=page_size)
    program_ids = [program.id for program in programs]
    course_order = {
        "sequence": (program_subjects.c.order_index, Course.id),
        "code": (Course.code, Course.id),
        "name": (Course.name, Course.id),
    }[course_sort]
    course_query = select(
        program_subjects.c.program_id.label("program_id"),
        Course.id.label("id"), Course.code.label("code"), Course.name.label("name"),
        program_subjects.c.order_index.label("sequence"),
        program_subjects.c.prerequisite_subject_id.label("prerequisite_course_id"),
        program_subjects.c.weight.label("weight"),
        Course.learning_outcomes.label("learning_outcomes"),
        func.count().over(partition_by=program_subjects.c.program_id).label("course_total"),
        func.row_number().over(partition_by=program_subjects.c.program_id,
            order_by=course_order).label("course_row"),
    ).join(program_subjects, program_subjects.c.subject_id == Course.id).where(
        program_subjects.c.program_id.in_(program_ids or [-1])
    )
    if course_q.strip():
        search = f"%{course_q.strip()}%"
        course_query = course_query.where(or_(Course.code.ilike(search), Course.name.ilike(search)))
    paged_courses = course_query.subquery()
    first_course = (course_page - 1) * course_page_size + 1
    last_course = course_page * course_page_size
    course_rows = db.execute(select(paged_courses).where(
        paged_courses.c.course_row.between(first_course, last_course)
    ).order_by(paged_courses.c.program_id, paged_courses.c.course_row)).all()
    courses_by_program: dict[int, list[dict]] = {program_id: [] for program_id in program_ids}
    course_totals: dict[int, int] = {program_id: 0 for program_id in program_ids}
    for row in course_rows:
        courses_by_program[row.program_id].append({
            "id": row.id, "code": row.code, "name": row.name,
            "sequence": row.sequence,
            "prerequisite_course_id": row.prerequisite_course_id,
            "weight": str(row.weight),
            "learning_outcomes": row.learning_outcomes,
        })
        course_totals[row.program_id] = row.course_total
    result = []
    for program in programs:
        result.append({
            "id": program.id,
            "code": program.code,
            "name": program.name,
            "degree_level": program.degree_level,
            "total_credits": program.total_credits,
            "duration_semesters": program.duration_semesters,
            "description": program.description,
            "is_active": program.is_active,
            "standard_tuition_vnd": str(program.standard_tuition_vnd),
            "courses": courses_by_program[program.id],
            "courses_pagination": page_metadata(
                total=course_totals[program.id], page=course_page,
                page_size=course_page_size,
            ),
        })
    return result


def _program_subject_payload(db: Session, program_id: int) -> list[dict]:
    rows = db.execute(
        select(
            Course.id,
            Course.code,
            Course.name,
            Course.learning_outcomes,
            program_subjects.c.order_index,
            program_subjects.c.prerequisite_subject_id,
            program_subjects.c.weight,
        )
        .join(program_subjects, program_subjects.c.subject_id == Course.id)
        .where(program_subjects.c.program_id == program_id)
        .order_by(program_subjects.c.order_index, Course.id)
    )
    return [
        {
            "subject_id": row.id,
            "id": row.id,
            "code": row.code,
            "name": row.name,
            "learning_outcomes": row.learning_outcomes,
            "order_index": row.order_index,
            "prerequisite_subject_id": row.prerequisite_subject_id,
            "weight": str(row.weight),
        }
        for row in rows
    ]


@router.get("/study-programs/{program_id}/subjects")
def get_study_program_subjects(
    program_id: int,
    response: Response,
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=20, ge=1, le=100),
    q: str = Query(default="", max_length=150),
    sort: str = Query(default="sequence", pattern="^(sequence|code|name)$"),
    _user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
) -> list[dict]:
    if db.get(StudyProgram, program_id) is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy chương trình đào tạo.")
    query = select(Course.id, Course.code, Course.name, Course.learning_outcomes,
        program_subjects.c.order_index, program_subjects.c.prerequisite_subject_id,
        program_subjects.c.weight
    ).join(program_subjects, program_subjects.c.subject_id == Course.id).where(
        program_subjects.c.program_id == program_id
    )
    if q.strip():
        search = f"%{q.strip()}%"
        query = query.where(or_(Course.code.ilike(search), Course.name.ilike(search)))
    total = db.scalar(select(func.count()).select_from(query.order_by(None).subquery())) or 0
    ordering = {
        "sequence": (program_subjects.c.order_index, Course.id),
        "code": (Course.code, Course.id),
        "name": (Course.name, Course.id),
    }[sort]
    rows = db.execute(query.order_by(*ordering)
        .offset((page - 1) * page_size).limit(page_size)).all()
    set_pagination_headers(response, total=total, page=page, page_size=page_size)
    return [{"subject_id": row.id, "id": row.id, "code": row.code, "name": row.name,
        "order_index": row.order_index, "prerequisite_subject_id": row.prerequisite_subject_id,
        "weight": str(row.weight), "learning_outcomes": row.learning_outcomes}
        for row in rows]


@router.post(
    "/study-programs/{program_id}/subjects",
    status_code=status.HTTP_201_CREATED,
)
def add_study_program_subject(
    program_id: int,
    request: StudyProgramSubjectCreate,
    _admin: User = Depends(require_roles(*ADMIN_ROLES)),
    db: Session = Depends(get_db),
) -> dict:
    if db.get(StudyProgram, program_id) is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy chương trình đào tạo.")
    if db.get(Course, request.subject_id) is None:
        raise HTTPException(status_code=422, detail="Môn học không tồn tại.")
    if db.scalar(
        select(program_subjects.c.subject_id).where(
            program_subjects.c.program_id == program_id,
            program_subjects.c.subject_id == request.subject_id,
        )
    ) is not None:
        raise HTTPException(status_code=409, detail="Môn học đã có trong chương trình.")
    next_order = db.scalar(
        select(func.max(program_subjects.c.order_index)).where(
            program_subjects.c.program_id == program_id
        )
    )
    db.execute(
        program_subjects.insert().values(
            program_id=program_id,
            subject_id=request.subject_id,
            order_index=(next_order or 0) + 1,
            prerequisite_subject_id=None,
            weight=request.weight,
        )
    )
    db.commit()
    return next(
        item
        for item in _program_subject_payload(db, program_id)
        if item["subject_id"] == request.subject_id
    )


@router.put("/study-programs/{program_id}/subjects")
def update_study_program_subjects(
    program_id: int,
    request: StudyProgramSubjectsUpdate,
    _admin: User = Depends(require_roles(*ADMIN_ROLES)),
    db: Session = Depends(get_db),
) -> dict:
    program = db.get(StudyProgram, program_id)
    if program is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy chương trình đào tạo.")
    subject_ids = [item.subject_id for item in request.subjects]
    if len(subject_ids) != len(set(subject_ids)):
        raise HTTPException(status_code=422, detail="Không thể gắn trùng môn học.")
    if subject_ids:
        found_ids = set(db.scalars(select(Course.id).where(Course.id.in_(subject_ids))).all())
        if found_ids != set(subject_ids):
            raise HTTPException(status_code=422, detail="Có môn học không tồn tại.")
    prerequisites = {
        item.subject_id: item.prerequisite_subject_id for item in request.subjects
    }
    positions = {subject_id: index for index, subject_id in enumerate(subject_ids)}
    for subject_id, prerequisite_id in prerequisites.items():
        if prerequisite_id == subject_id:
            raise HTTPException(
                status_code=422, detail="Môn học không thể là tiên quyết của chính nó."
            )
        if prerequisite_id is not None and prerequisite_id not in prerequisites:
            raise HTTPException(
                status_code=422,
                detail="Môn tiên quyết phải được gắn vào cùng chương trình đào tạo.",
            )
        if prerequisite_id is not None and positions[prerequisite_id] >= positions[subject_id]:
            raise HTTPException(
                status_code=422,
                detail="Môn tiên quyết phải đứng trước môn học trong chương trình.",
            )

    db.execute(
        program_subjects.delete().where(
            program_subjects.c.program_id == program_id
        )
    )
    rows = [
        {
            "program_id": program_id,
            "subject_id": item.subject_id,
            "order_index": index + 1,
            "prerequisite_subject_id": item.prerequisite_subject_id,
            "weight": item.weight,
        }
        for index, item in enumerate(request.subjects)
    ]
    if rows:
        db.execute(program_subjects.insert(), rows)
    db.commit()
    return {"program_id": program_id, "subject_ids": subject_ids}


@router.delete(
    "/study-programs/{program_id}/subjects/{subject_id}",
    status_code=status.HTTP_204_NO_CONTENT,
)
def remove_study_program_subject(
    program_id: int,
    subject_id: int,
    _admin: User = Depends(require_roles(*ADMIN_ROLES)),
    db: Session = Depends(get_db),
) -> Response:
    if db.get(StudyProgram, program_id) is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy chương trình đào tạo.")
    deleted = db.execute(
        program_subjects.delete().where(
            program_subjects.c.program_id == program_id,
            program_subjects.c.subject_id == subject_id,
        )
    )
    if not deleted.rowcount:
        raise HTTPException(status_code=404, detail="Môn học không thuộc chương trình này.")
    db.execute(
        program_subjects.update()
        .where(
            program_subjects.c.program_id == program_id,
            program_subjects.c.prerequisite_subject_id == subject_id,
        )
        .values(prerequisite_subject_id=None)
    )
    remaining = db.execute(
        select(
            program_subjects.c.subject_id,
            program_subjects.c.order_index,
        )
        .where(program_subjects.c.program_id == program_id)
        .order_by(program_subjects.c.order_index, program_subjects.c.subject_id)
    ).all()
    for index, row in enumerate(remaining, start=1):
        if row.order_index != index:
            db.execute(
                program_subjects.update()
                .where(
                    program_subjects.c.program_id == program_id,
                    program_subjects.c.subject_id == row.subject_id,
                )
                .values(order_index=index)
            )
    db.commit()
    return Response(status_code=status.HTTP_204_NO_CONTENT)


@router.put("/study-programs/{program_id}/courses")
def update_study_program_courses(
    program_id: int,
    request: StudyProgramCoursesUpdate,
    _admin: User = Depends(require_roles(*ADMIN_ROLES)),
    db: Session = Depends(get_db),
) -> dict:
    program = db.get(StudyProgram, program_id)
    if program is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy chương trình đào tạo.")
    course_ids = [item.course_id for item in request.courses]
    if len(course_ids) != len(set(course_ids)):
        raise HTTPException(status_code=422, detail="Không thể gắn trùng môn học.")
    if course_ids:
        found_ids = set(db.scalars(select(Course.id).where(Course.id.in_(course_ids))).all())
        if found_ids != set(course_ids):
            raise HTTPException(status_code=422, detail="Có môn học không tồn tại.")
    requested_prerequisites = {
        item.course_id: item.prerequisite_course_id for item in request.courses
    }
    course_positions = {course_id: index for index, course_id in enumerate(course_ids)}
    for course_id, prerequisite_id in requested_prerequisites.items():
        if prerequisite_id == course_id:
            raise HTTPException(
                status_code=422, detail="Môn học không thể là tiên quyết của chính nó."
            )
        if prerequisite_id is not None and prerequisite_id not in requested_prerequisites:
            raise HTTPException(
                status_code=422,
                detail="Môn tiên quyết phải được gắn vào cùng chương trình đào tạo.",
            )
        if prerequisite_id is not None and course_positions[prerequisite_id] >= course_positions[course_id]:
            raise HTTPException(
                status_code=422,
                detail="Môn tiên quyết phải đứng trước môn học trong chương trình.",
            )

    db.execute(
        program_subjects.delete().where(
            program_subjects.c.program_id == program_id
        )
    )
    rows = [
        {
            "program_id": program_id,
            "subject_id": item.course_id,
            "order_index": index + 1,
            "prerequisite_subject_id": item.prerequisite_course_id,
            "weight": item.weight,
        }
        for index, item in enumerate(request.courses)
    ]
    if rows:
        db.execute(program_subjects.insert(), rows)
    db.commit()
    return {"id": program_id, "course_ids": course_ids}


@router.post("/courses", status_code=status.HTTP_201_CREATED)
def create_course(
    request: CourseCreate,
    _admin: User = Depends(require_roles(*ADMIN_ROLES)),
    db: Session = Depends(get_db),
) -> dict:
    if db.scalar(select(Course.id).where(Course.code == request.code)):
        raise HTTPException(status_code=409, detail="Mã môn học đã tồn tại.")
    if request.prerequisite_id == 0:
        raise HTTPException(status_code=422, detail="Môn tiên quyết không hợp lệ.")
    course = Course(**request.model_dump())
    db.add(course)
    try:
        db.commit()
    except IntegrityError as error:
        db.rollback()
        raise HTTPException(
            status_code=422, detail="Môn tiên quyết không tồn tại."
        ) from error
    db.refresh(course)
    return _course_payload(course)


@router.patch("/courses/{course_id}")
def update_course(
    course_id: int,
    request: CourseUpdate,
    _admin: User = Depends(require_roles(*ADMIN_ROLES)),
    db: Session = Depends(get_db),
) -> dict:
    course = db.get(Course, course_id)
    if course is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy môn học.")
    updates = request.model_dump(exclude_unset=True)
    if updates.get("prerequisite_id") == course.id:
        raise HTTPException(status_code=422, detail="Môn học không thể là tiên quyết của chính nó.")
    if (
        "session_count" in updates
        and (db.scalar(
            select(func.max(CourseLesson.sequence)).where(
                CourseLesson.course_id == course.id
            )
        ) or 0) > updates["session_count"]
    ):
        raise HTTPException(
            status_code=409,
            detail="Không thể giảm số buổi xuống thấp hơn buổi đã khai báo trong syllabus.",
        )
    for field, value in updates.items():
        setattr(course, field, value)
    try:
        db.commit()
    except IntegrityError as error:
        db.rollback()
        raise HTTPException(status_code=422, detail="Dữ liệu môn tiên quyết không hợp lệ.") from error
    db.refresh(course)
    return _course_payload(course)


@router.delete("/courses/{course_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_course(
    course_id: int,
    _admin: User = Depends(require_roles(*ADMIN_ROLES)),
    db: Session = Depends(get_db),
) -> None:
    course = db.get(Course, course_id)
    if course is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy môn học.")
    if db.scalar(
        select(TrainingClass.id).where(TrainingClass.course_id == course_id).limit(1)
    ):
        raise HTTPException(
            status_code=409,
            detail="Không thể xóa môn học đã được xếp vào lớp. Hãy ngừng mở môn học.",
        )
    db.delete(course)
    db.commit()


def _visible_class_query(current_user: User):
    roles = {role.code for role in current_user.roles}
    query = select(TrainingClass).options(
        joinedload(TrainingClass.course),
        joinedload(TrainingClass.teacher),
        joinedload(TrainingClass.assistant),
        selectinload(TrainingClass.assistants),
        selectinload(TrainingClass.enrollments),
        selectinload(TrainingClass.sessions),
    )
    if roles.intersection(ADMIN_ROLES + ("admissions", "accountant")):
        return query
    if roles.intersection(("teacher", "ta")):
        return query.where(
            or_(
                TrainingClass.teacher_id == current_user.id,
                TrainingClass.assistant_id == current_user.id,
                TrainingClass.assistants.any(User.id == current_user.id),
            )
        )
    if "student" in roles:
        return query.join(ClassEnrollment).where(
            ClassEnrollment.student_id == current_user.id
        )
    return query.where(false())


@router.get("/classes")
def list_classes(
    response: Response,
    term: str | None = Query(default=None, max_length=50),
    q: str = Query(default="", max_length=150),
    class_status: str | None = Query(
        default=None,
        alias="status",
        pattern="^(preparing|active|completed|cancelled)$",
    ),
    teacher_id: int | None = Query(default=None),
    course_id: int | None = Query(default=None),
    program_id: int | None = Query(default=None),
    start_date: date | None = None,
    end_date: date | None = None,
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=20, ge=1, le=100),
    current_user: User = Depends(require_roles(*ALL_ROLES)),
    db: Session = Depends(get_db),
) -> list[dict]:
    if start_date is not None and end_date is not None and start_date > end_date:
        raise HTTPException(status_code=422, detail="start_date must not be after end_date.")
    query = _visible_class_query(current_user)
    if term:
        query = query.where(TrainingClass.term == term)
    if class_status:
        query = query.where(TrainingClass.status == class_status)
    if teacher_id:
        query = query.where(TrainingClass.teacher_id == teacher_id)
    if course_id:
        query = query.where(TrainingClass.course_id == course_id)
    if program_id:
        query = query.where(or_(
            TrainingClass.program_id == program_id,
            and_(TrainingClass.program_id.is_(None), TrainingClass.course_id.in_(
                select(program_subjects.c.subject_id).where(program_subjects.c.program_id == program_id)
            )),
        ))
    if start_date is not None:
        query = query.where(TrainingClass.sessions.any(AttendanceSession.session_date >= start_date))
    if end_date is not None:
        query = query.where(TrainingClass.sessions.any(AttendanceSession.session_date <= end_date))
    if q.strip():
        search = f"%{q.strip()}%"
        query = query.join(Course).where(
            or_(
                TrainingClass.code.ilike(search),
                TrainingClass.name.ilike(search),
                Course.code.ilike(search),
                Course.name.ilike(search),
            )
        )
    total = db.scalar(select(func.count()).select_from(
        query.with_only_columns(TrainingClass.id).order_by(None).distinct().subquery()
    )) or 0
    rows = db.scalars(query.order_by(TrainingClass.day_of_week, TrainingClass.start_time,
        TrainingClass.id).offset((page - 1) * page_size).limit(page_size)).unique().all()
    set_pagination_headers(response, total=total, page=page, page_size=page_size)
    return [_class_payload(item) for item in rows]


@router.post("/classes", status_code=status.HTTP_201_CREATED)
def create_class(
    request: TrainingClassCreate,
    _admin: User = Depends(require_roles(*ADMIN_ROLES)),
    db: Session = Depends(get_db),
) -> dict:
    if db.scalar(
        select(TrainingClass.id).where(TrainingClass.code == request.code)
    ):
        raise HTTPException(status_code=409, detail="Mã lớp học đã tồn tại.")
    course = db.get(Course, request.course_id)
    teacher = db.scalar(
        select(User)
        .options(selectinload(User.roles))
        .where(User.id == request.teacher_id)
    )
    if course is None:
        raise HTTPException(status_code=422, detail="Môn học không tồn tại.")
    if request.program_id is not None:
        program = db.scalar(
            select(StudyProgram)
            .options(selectinload(StudyProgram.courses))
            .where(StudyProgram.id == request.program_id)
        )
        if program is None or request.course_id not in {
            item.id for item in program.courses
        }:
            raise HTTPException(
                status_code=422,
                detail="Môn học không thuộc chương trình đào tạo đã chọn.",
            )
        if program is not None and not program.is_active:
            raise HTTPException(status_code=409, detail="Inactive study programs cannot accept new classes.")
    classroom = db.scalar(
        select(Classroom).where(Classroom.code == request.room.strip())
    )
    if classroom is None and db.scalar(select(Classroom.id).limit(1)):
        raise HTTPException(
            status_code=422, detail="Phòng học không có trong danh mục phòng."
        )
    if classroom is not None and request.capacity > classroom.capacity:
        raise HTTPException(
            status_code=422, detail="Sĩ số lớp vượt quá sức chứa của phòng học."
        )
    if teacher is None or "teacher" not in {role.code for role in teacher.roles}:
        raise HTTPException(status_code=422, detail="Giảng viên không hợp lệ.")
    assistant_ids = set(request.assistant_ids)
    if len(assistant_ids) != len(request.assistant_ids):
        raise HTTPException(status_code=422, detail="Danh sách trợ giảng bị trùng.")
    if request.assistant_id is not None:
        if request.assistant_id in assistant_ids:
            raise HTTPException(status_code=422, detail="Danh sách trợ giảng bị trùng.")
        assistant_ids.add(request.assistant_id)
    assistants = _resolve_assistants(db, assistant_ids, teacher.id)

    overlap = db.scalar(
        select(TrainingClass.id).where(
            TrainingClass.term == request.term,
            TrainingClass.day_of_week == request.day_of_week,
            TrainingClass.start_time < request.end_time,
            TrainingClass.end_time > request.start_time,
            or_(
                TrainingClass.teacher_id == request.teacher_id,
                TrainingClass.room == request.room,
                *([TrainingClass.assistant_id.in_(assistant_ids)] if assistant_ids else []),
                *(
                    [TrainingClass.assistants.any(User.id.in_(assistant_ids))]
                    if assistant_ids
                    else []
                ),
            ),
        ).limit(1)
    )
    if overlap:
        raise HTTPException(
            status_code=409,
            detail="Lịch bị trùng với lớp khác của giảng viên hoặc phòng học.",
        )
    training_class = TrainingClass(
        **request.model_dump(exclude={"assistant_ids", "assistant_id"}),
        assistant_id=assistants[0].id if assistants else None,
        assistants=assistants,
        program_curriculum_snapshot=_program_curriculum_snapshot(db, request.program_id),
    )
    db.add(training_class)
    db.commit()
    training_class = db.scalar(
        select(TrainingClass)
        .options(
            joinedload(TrainingClass.course),
            joinedload(TrainingClass.teacher),
            joinedload(TrainingClass.assistant),
            selectinload(TrainingClass.assistants),
            selectinload(TrainingClass.enrollments),
        )
        .where(TrainingClass.id == training_class.id)
    )
    return _class_payload(training_class)


def _resolve_assistants(
    db: Session, assistant_ids: set[int], teacher_id: int
) -> list[User]:
    if teacher_id in assistant_ids:
        raise HTTPException(
            status_code=422, detail="Giảng viên không thể đồng thời là trợ giảng."
        )
    if not assistant_ids:
        return []
    assistants = db.scalars(
        select(User)
        .options(selectinload(User.roles))
        .where(User.id.in_(assistant_ids))
    ).all()
    if len(assistants) != len(assistant_ids) or any(
        "ta" not in {role.code for role in assistant.roles}
        for assistant in assistants
    ):
        raise HTTPException(status_code=422, detail="Có trợ giảng không hợp lệ.")
    return assistants


def _assert_class_staff_access(current_user: User, training_class: TrainingClass) -> None:
    roles = {role.code for role in current_user.roles}
    if roles.intersection(ADMIN_ROLES):
        return
    if "teacher" in roles and training_class.teacher_id == current_user.id:
        return
    if "ta" in roles and (
        training_class.assistant_id == current_user.id
        or any(assistant.id == current_user.id for assistant in training_class.assistants)
    ):
        return
    raise HTTPException(status_code=403, detail="Bạn không được phân công vào lớp này.")


@router.get("/classes/{class_id}/students")
def list_class_students(
    class_id: int,
    response: Response,
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=20, ge=1, le=100),
    current_user: User = Depends(
        require_roles("admin", "manager", "admissions", "accountant", "teacher", "ta")
    ),
    db: Session = Depends(get_db),
) -> list[dict]:
    training_class = db.get(TrainingClass, class_id)
    if training_class is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy lớp học.")
    if {role.code for role in current_user.roles}.intersection(("teacher", "ta")):
        _assert_class_staff_access(current_user, training_class)
    base_query = select(ClassEnrollment).where(ClassEnrollment.class_id == class_id).join(
        User, ClassEnrollment.student_id == User.id
    )
    total = db.scalar(select(func.count()).select_from(
        base_query.with_only_columns(ClassEnrollment.id).order_by(None).subquery()
    )) or 0
    enrollments = db.scalars(
        select(ClassEnrollment)
        .options(
            joinedload(ClassEnrollment.student).selectinload(User.roles)
        )
        .where(ClassEnrollment.class_id == class_id)
        .join(User, ClassEnrollment.student_id == User.id)
        .order_by(User.full_name, User.id)
        .offset((page - 1) * page_size).limit(page_size)
    ).all()
    student_ids = [enrollment.student_id for enrollment in enrollments]
    attendance_records = db.execute(
        select(
            AttendanceRecord.student_id,
            AttendanceSession.session_date,
            AttendanceRecord.status,
        )
        .join(AttendanceSession, AttendanceRecord.session_id == AttendanceSession.id)
        .where(AttendanceSession.class_id == class_id,
            AttendanceRecord.student_id.in_(student_ids))
        .order_by(AttendanceSession.session_date.desc())
    ).all() if student_ids else []
    history_by_student: dict[int, list[dict[str, str]]] = {}
    for student_id, session_date, attendance_status in attendance_records:
        history_by_student.setdefault(student_id, []).append(
            {"session_date": session_date.isoformat(), "status": attendance_status}
        )
    set_pagination_headers(response, total=total, page=page, page_size=page_size)
    return [
        {
            **_user_payload(enrollment.student),
            "enrolled_at": enrollment.enrolled_at.isoformat() if enrollment.enrolled_at else None,
            "effective_date": enrollment.effective_date.isoformat() if enrollment.effective_date else None,
            "enrolled_by_id": enrollment.enrolled_by_id,
            "tuition_status": enrollment.tuition_status,
            "attendance_status": enrollment.attendance_status,
            "last_attendance_at": (
                enrollment.last_attendance_at.isoformat()
                if enrollment.last_attendance_at
                else None
            ),
            "attendance_history": history_by_student.get(enrollment.student_id, []),
        }
        for enrollment in enrollments
    ]


@router.get("/classes/{class_id}/students/export")
def export_class_students(
    class_id: int,
    current_user: User = Depends(
        require_roles("admin", "manager", "admissions", "accountant", "teacher", "ta")
    ),
    db: Session = Depends(get_db),
) -> Response:
    training_class = db.get(TrainingClass, class_id)
    if training_class is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy lớp học.")
    current_roles = {role.code for role in current_user.roles}
    if current_roles.intersection(("teacher", "ta")):
        _assert_class_staff_access(current_user, training_class)
    enrollments = db.scalars(
        select(ClassEnrollment)
        .join(User, ClassEnrollment.student_id == User.id)
        .options(joinedload(ClassEnrollment.student))
        .where(ClassEnrollment.class_id == class_id)
        .order_by(User.full_name)
    ).all()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Danh sách học viên"
    sheet.append(
        [
            "Mã học viên",
            "Họ và tên",
            "Email",
            "Số điện thoại",
            "Học phí",
            "Điểm danh",
        ]
    )
    for enrollment in enrollments:
        student = enrollment.student
        sheet.append(
            [
                student.student_code or student.username,
                student.full_name,
                student.email,
                student.phone or "",
                "Đã đóng" if enrollment.tuition_status == "paid" else "Chưa đóng",
                {
                    "not_recorded": "Chưa điểm danh",
                    "present": "Có mặt",
                    "absent": "Vắng",
                    "late": "Đi muộn",
                }[enrollment.attendance_status],
            ]
        )
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    output = BytesIO()
    workbook.save(output)
    filename = f"{training_class.code}-students.xlsx"
    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.patch("/classes/{class_id}/students/{student_id}/status")
def update_enrollment_status(
    class_id: int,
    student_id: int,
    request: EnrollmentStatusUpdate,
    current_user: User = Depends(
        require_roles("admin", "manager", "accountant", "teacher", "ta")
    ),
    db: Session = Depends(get_db),
) -> dict:
    training_class = db.get(TrainingClass, class_id)
    if training_class is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy lớp học.")
    current_roles = {role.code for role in current_user.roles}
    if current_roles.intersection(("teacher", "ta")):
        _assert_class_staff_access(current_user, training_class)
    if request.tuition_status is not None and not current_roles.intersection(
        ADMIN_ROLES + ("accountant",)
    ):
        raise HTTPException(
            status_code=403,
            detail="Chỉ quản trị viên, quản lý hoặc kế toán mới được cập nhật học phí.",
        )
    if request.attendance_status is not None and not current_roles.intersection(
        ADMIN_ROLES + ("teacher", "ta")
    ):
        raise HTTPException(
            status_code=403,
            detail="Chỉ giảng viên, trợ giảng hoặc quản lý mới được cập nhật điểm danh.",
        )
    enrollment = db.scalar(
        select(ClassEnrollment).where(
            ClassEnrollment.class_id == class_id,
            ClassEnrollment.student_id == student_id,
        )
    )
    if enrollment is None:
        raise HTTPException(status_code=404, detail="Học viên chưa đăng ký lớp này.")
    if (
        enrollment.enrollment_status != "active"
        and enrollment.effective_date is not None
        and enrollment.effective_date <= datetime.now(timezone.utc).date()
    ):
        raise HTTPException(
            status_code=409,
            detail="Không thể điểm danh học viên sau ngày hiệu lực biến động.",
        )
    if request.tuition_status is not None:
        previous_tuition_status = enrollment.tuition_status
        enrollment.tuition_status = request.tuition_status
        if previous_tuition_status != request.tuition_status:
            _record_audit(
                db,
                current_user,
                "enrollment.tuition_status_changed",
                target_user_id=student_id,
                details=(
                    f"class_id={class_id};"
                    f"before={previous_tuition_status};"
                    f"after={request.tuition_status}"
                ),
            )
    if request.attendance_status is not None:
        enrollment.attendance_status = request.attendance_status
        now = datetime.now(timezone.utc)
        enrollment.last_attendance_at = (
            None
            if request.attendance_status == "not_recorded"
            else now
        )
        attendance_session = db.scalar(
            select(AttendanceSession)
            .where(
                AttendanceSession.class_id == class_id,
                AttendanceSession.session_date == now.date(),
            )
            .limit(1)
        )
        if attendance_session is None:
            attendance_session = AttendanceSession(
                class_id=class_id,
                session_date=now.date(),
                start_time=training_class.start_time,
                end_time=training_class.end_time,
            )
            db.add(attendance_session)
            db.flush()
        attendance_record = db.scalar(
            select(AttendanceRecord).where(
                AttendanceRecord.session_id == attendance_session.id,
                AttendanceRecord.student_id == student_id,
            )
        )
        old_status = attendance_record.status if attendance_record else None
        if attendance_record is None:
            attendance_record = AttendanceRecord(
                session_id=attendance_session.id,
                student_id=student_id,
                marked_by=current_user.id,
                status=request.attendance_status,
                marked_at=now.replace(tzinfo=None),
            )
            db.add(attendance_record)
            db.flush()
        else:
            attendance_record.marked_by = current_user.id
            attendance_record.status = request.attendance_status
            attendance_record.marked_at = now.replace(tzinfo=None)
        if old_status != request.attendance_status:
            db.add(
                AttendanceChangeLog(
                    record_id=attendance_record.id,
                    changed_by=current_user.id,
                    old_status=old_status,
                    new_status=request.attendance_status,
                )
            )
    db.commit()
    attendance_history = []
    if request.attendance_status is not None:
        attendance_history = [
            {"session_date": session_date.isoformat(), "status": attendance_status}
            for session_date, attendance_status in db.execute(
                select(AttendanceSession.session_date, AttendanceRecord.status)
                .join(
                    AttendanceRecord,
                    AttendanceRecord.session_id == AttendanceSession.id,
                )
                .where(
                    AttendanceSession.class_id == class_id,
                    AttendanceRecord.student_id == student_id,
                )
                .order_by(AttendanceSession.session_date.desc())
                .limit(4)
            ).all()
        ]
    return {
        "student_id": student_id,
        "tuition_status": enrollment.tuition_status,
        "attendance_status": enrollment.attendance_status,
        "last_attendance_at": (
            enrollment.last_attendance_at.isoformat()
            if enrollment.last_attendance_at
            else None
        ),
        "attendance_history": attendance_history,
    }


@router.patch("/classes/{class_id}")
def update_class(
    class_id: int,
    request: TrainingClassUpdate,
    _admin: User = Depends(require_roles(*ADMIN_ROLES)),
    db: Session = Depends(get_db),
) -> dict:
    training_class = db.scalar(
        select(TrainingClass)
        .options(
            selectinload(TrainingClass.enrollments),
            selectinload(TrainingClass.assistants),
        )
        .where(TrainingClass.id == class_id)
    )
    if training_class is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy lớp học.")
    updates = request.model_dump(exclude_unset=True)
    if not updates:
        raise HTTPException(status_code=400, detail="Không có dữ liệu cần cập nhật.")
    requested_assistant_ids = updates.pop("assistant_ids", None)
    legacy_assistant_id = updates.pop("assistant_id", None) if "assistant_id" in updates else None
    for required_field in (
        "course_id",
        "teacher_id",
        "room",
        "term",
        "day_of_week",
        "start_time",
        "end_time",
        "capacity",
    ):
        if required_field in updates and updates[required_field] is None:
            raise HTTPException(
                status_code=422, detail=f"Trường {required_field} không được để trống."
            )

    course_id = updates.get("course_id", training_class.course_id)
    program_id = updates.get("program_id", training_class.program_id)
    teacher_id = updates.get("teacher_id", training_class.teacher_id)
    room = updates.get("room", training_class.room)
    term = updates.get("term", training_class.term)
    weekday = updates.get("day_of_week", training_class.day_of_week)
    start_time = updates.get("start_time", training_class.start_time)
    end_time = updates.get("end_time", training_class.end_time)
    capacity = updates.get("capacity", training_class.capacity)
    if end_time <= start_time:
        raise HTTPException(status_code=422, detail="Giờ kết thúc phải sau giờ bắt đầu.")
    if capacity < len(training_class.enrollments):
        raise HTTPException(
            status_code=409, detail="Sức chứa mới nhỏ hơn sĩ số hiện tại của lớp."
        )
    if db.get(Course, course_id) is None:
        raise HTTPException(status_code=422, detail="Môn học không tồn tại.")
    if program_id is not None and db.scalar(select(StudyProgram.id).join(
        program_subjects, program_subjects.c.program_id == StudyProgram.id
    ).where(StudyProgram.id == program_id, program_subjects.c.subject_id == course_id)) is None:
        raise HTTPException(status_code=422, detail="Class course must belong to its program.")
    classroom = db.scalar(select(Classroom).where(Classroom.code == room.strip()))
    if classroom is None and db.scalar(select(Classroom.id).limit(1)):
        raise HTTPException(
            status_code=422, detail="Phòng học không có trong danh mục phòng."
        )
    if classroom is not None and capacity > classroom.capacity:
        raise HTTPException(
            status_code=422, detail="Sĩ số lớp vượt quá sức chứa của phòng học."
        )
    if classroom is not None and len(training_class.enrollments) > classroom.capacity:
        raise HTTPException(
            status_code=409,
            detail="Sức chứa phòng mới nhỏ hơn sĩ số đã ghi danh.",
        )
    teacher = db.scalar(
        select(User).options(selectinload(User.roles)).where(User.id == teacher_id)
    )
    if teacher is None or "teacher" not in {role.code for role in teacher.roles}:
        raise HTTPException(status_code=422, detail="Giảng viên không hợp lệ.")
    current_assistant_ids = {
        assistant.id for assistant in training_class.assistants
    }
    if training_class.assistant_id is not None:
        current_assistant_ids.add(training_class.assistant_id)
    if requested_assistant_ids is not None:
        if len(set(requested_assistant_ids)) != len(requested_assistant_ids):
            raise HTTPException(status_code=422, detail="Danh sách trợ giảng bị trùng.")
        next_assistant_ids = set(requested_assistant_ids)
        if legacy_assistant_id is not None:
            next_assistant_ids.add(legacy_assistant_id)
    elif "assistant_id" in request.model_fields_set:
        next_assistant_ids = (
            {legacy_assistant_id} if legacy_assistant_id is not None else set()
        )
    else:
        next_assistant_ids = current_assistant_ids
    next_assistants = _resolve_assistants(db, next_assistant_ids, teacher_id)

    allowed_status_transitions = {
        "preparing": {"active", "cancelled"},
        "active": {"completed", "cancelled"},
        "completed": set(),
        "cancelled": set(),
    }
    next_status = updates.get("status", training_class.status)
    if next_status != training_class.status and next_status not in allowed_status_transitions[
        training_class.status
    ]:
        raise HTTPException(
            status_code=409,
            detail=f"Không thể chuyển vòng đời lớp từ {training_class.status} sang {next_status}.",
        )

    overlap = db.scalar(
        select(TrainingClass.id).where(
            TrainingClass.id != class_id,
            TrainingClass.term == term,
            TrainingClass.day_of_week == weekday,
            TrainingClass.start_time < end_time,
            TrainingClass.end_time > start_time,
            or_(
                TrainingClass.teacher_id == teacher_id,
                TrainingClass.room == room,
                *(
                    [TrainingClass.assistant_id.in_(next_assistant_ids)]
                    if next_assistant_ids
                    else []
                ),
                *(
                    [TrainingClass.assistants.any(User.id.in_(next_assistant_ids))]
                    if next_assistant_ids
                    else []
                ),
            ),
        ).limit(1)
    )
    if overlap:
        raise HTTPException(
            status_code=409,
            detail="Lịch bị trùng với lớp khác của giảng viên hoặc phòng học.",
        )

    student_ids = [enrollment.student_id for enrollment in training_class.enrollments]
    if student_ids:
        conflicting_enrollments = db.scalars(
            select(ClassEnrollment)
            .join(TrainingClass, ClassEnrollment.class_id == TrainingClass.id)
            .where(
                ClassEnrollment.student_id.in_(student_ids),
                TrainingClass.id != class_id,
                TrainingClass.term == term,
                TrainingClass.day_of_week == weekday,
                TrainingClass.start_time < end_time,
                TrainingClass.end_time > start_time,
            )
        ).all()
        if conflicting_enrollments:
            conflicts = sorted(
                {enrollment.student_id for enrollment in conflicting_enrollments}
            )
            raise HTTPException(
                status_code=409,
                detail=f"Lịch mới bị trùng với lớp khác của học viên: {conflicts}.",
            )

    previous_schedule = (
        training_class.term,
        training_class.day_of_week,
        training_class.start_time,
        training_class.end_time,
        training_class.room,
    )
    for field, value in updates.items():
        setattr(training_class, field, value)
    if "program_id" in updates:
        training_class.program_curriculum_snapshot = _program_curriculum_snapshot(db, program_id)
    training_class.assistants = next_assistants
    training_class.assistant_id = next_assistants[0].id if next_assistants else None
    next_schedule = (term, weekday, start_time, end_time, room)
    if previous_schedule != next_schedule:
        message = (
            f"Lịch lớp {training_class.code} đã thay đổi: {WEEKDAYS[weekday]}, "
            f"{start_time.strftime('%H:%M')}–{end_time.strftime('%H:%M')}, "
            f"phòng {room}, học kỳ {term}."
        )
        for student_id in student_ids:
            notification = Notification(
                user_id=student_id,
                class_id=class_id,
                title="Lịch học được cập nhật",
                message=message,
            )
            db.add(notification)
            enqueue_notification_email(db, notification, db.get(User, student_id))
    db.commit()
    updated = db.scalar(
        select(TrainingClass)
        .options(
            joinedload(TrainingClass.course),
            joinedload(TrainingClass.teacher),
            joinedload(TrainingClass.assistant),
            selectinload(TrainingClass.assistants),
            selectinload(TrainingClass.enrollments),
        )
        .where(TrainingClass.id == class_id)
    )
    return _class_payload(updated)


@router.post("/classes/{class_id}/students", status_code=status.HTTP_201_CREATED)
def enroll_students(
    class_id: int,
    request: EnrollmentRequest,
    _admin: User = Depends(require_roles(*ADMIN_ROLES, "admissions")),
    db: Session = Depends(get_db),
) -> dict:
    training_class = db.scalar(
        select(TrainingClass)
        .options(selectinload(TrainingClass.enrollments))
        .where(TrainingClass.id == class_id)
        .with_for_update()
    )
    if training_class is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy lớp học.")
    student_ids = set(request.student_ids)
    current_ids = {enrollment.student_id for enrollment in training_class.enrollments}
    new_ids = student_ids - current_ids
    if not new_ids:
        raise HTTPException(status_code=409, detail="Học viên đã có trong lớp.")
    if len(training_class.enrollments) + len(new_ids) > training_class.capacity:
        raise HTTPException(status_code=409, detail="Sĩ số lớp vượt quá sức chứa.")
    if training_class.status not in {"preparing", "active"}:
        raise HTTPException(
            status_code=409,
            detail="Không thể ghi danh vào lớp đã kết thúc hoặc bị hủy.",
        )
    classroom = db.scalar(
        select(Classroom).where(Classroom.code == training_class.room)
    )
    if (
        classroom is not None
        and len(training_class.enrollments) + len(new_ids) > classroom.capacity
    ):
        raise HTTPException(
            status_code=409,
            detail="Sĩ số ghi danh vượt quá sức chứa của phòng học.",
        )

    students = db.scalars(
        select(User).options(selectinload(User.roles)).where(User.id.in_(new_ids))
    ).all()
    student_by_id = {student.id: student for student in students}
    if len(student_by_id) != len(new_ids):
        raise HTTPException(status_code=422, detail="Có học viên không tồn tại.")
    if any(
        not student.is_active
        or student.is_locked
        or "student" not in {role.code for role in student.roles}
        for student in students
    ):
        raise HTTPException(
            status_code=422, detail="Chỉ thêm được tài khoản học viên đang hoạt động."
        )

    overlapping_classes = db.scalars(
        select(TrainingClass)
        .join(ClassEnrollment)
        .where(
            ClassEnrollment.student_id.in_(new_ids),
            TrainingClass.term == training_class.term,
            TrainingClass.day_of_week == training_class.day_of_week,
            TrainingClass.start_time < training_class.end_time,
            TrainingClass.end_time > training_class.start_time,
        )
    ).all()
    overlapping_class_ids = [item.id for item in overlapping_classes]
    conflicts = db.scalars(
        select(ClassEnrollment).where(
            ClassEnrollment.student_id.in_(new_ids),
            ClassEnrollment.class_id.in_(overlapping_class_ids or [-1]),
        )
    ).all()
    if conflicts:
        conflicted_ids = sorted({enrollment.student_id for enrollment in conflicts})
        raise HTTPException(
            status_code=409,
            detail=f"Lịch học bị trùng với lớp hiện tại của học viên: {conflicted_ids}.",
        )

    db.add_all(
        ClassEnrollment(class_id=class_id, student_id=student_id, enrolled_by_id=_admin.id)
        for student_id in sorted(new_ids)
    )
    db.commit()
    return {
        "class_id": class_id,
        "enrolled_count": len(new_ids),
        "student_count": len(current_ids) + len(new_ids),
    }


@router.delete("/classes/{class_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_class(
    class_id: int,
    _admin: User = Depends(require_roles(*ADMIN_ROLES)),
    db: Session = Depends(get_db),
) -> None:
    training_class = db.get(TrainingClass, class_id)
    if training_class is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy lớp học.")
    db.delete(training_class)
    db.commit()


@router.put("/me/class-reminder-preferences")
def update_class_reminder_preferences(
    request: ClassReminderPreferenceUpdate,
    student: User = Depends(require_roles("student")),
    db: Session = Depends(get_db),
) -> dict:
    profile = db.get(StudentProfile, student.id)
    if profile is None:
        profile = StudentProfile(user_id=student.id)
        db.add(profile)
    profile.class_reminder_opt_out = request.class_reminder_opt_out
    db.commit()
    return {"class_reminder_opt_out": profile.class_reminder_opt_out}


@router.get("/assignments")
def list_assignments(
    response: Response,
    q: str | None = Query(default=None, max_length=150),
    class_id: int | None = Query(default=None, ge=1),
    page: int | None = Query(default=None, ge=1),
    page_size: int | None = Query(default=None, ge=1, le=100),
    current_user: User = Depends(require_roles(*ALL_ROLES)),
    db: Session = Depends(get_db),
) -> list[dict] | dict:
    roles = {role.code for role in current_user.roles}
    query = (
        select(Assignment, TrainingClass)
        .join(TrainingClass, Assignment.class_id == TrainingClass.id)
        .options(joinedload(TrainingClass.course))
        .order_by(Assignment.due_at.is_(None), Assignment.due_at, Assignment.id.desc())
    )
    if roles.intersection(ADMIN_ROLES + ("admissions", "accountant")):
        pass
    elif roles.intersection(("teacher", "ta")):
        query = query.where(
            or_(
                TrainingClass.teacher_id == current_user.id,
                TrainingClass.assistant_id == current_user.id,
                TrainingClass.assistants.any(User.id == current_user.id),
            )
        )
    elif "student" in roles:
        legacy_enrollment = exists(
            select(1).where(
                ClassEnrollment.class_id == Assignment.class_id,
                ClassEnrollment.student_id == current_user.id,
            )
        )
        snapshot_recipient = exists(
            select(1).where(
                AssignmentRecipient.assignment_id == Assignment.id,
                AssignmentRecipient.student_id == current_user.id,
            )
        )
        query = query.where(
            Assignment.status == "published",
            or_(
                and_(Assignment.roster_snapshotted_at.is_(None), legacy_enrollment),
                and_(Assignment.roster_snapshotted_at.is_not(None), snapshot_recipient),
            ),
        )
    else:
        return []

    if q is not None and q.strip():
        search = f"%{q.strip()}%"
        query = query.where(
            or_(
                Assignment.title.ilike(search),
                TrainingClass.code.ilike(search),
                TrainingClass.name.ilike(search),
                TrainingClass.course.has(
                    or_(Course.code.ilike(search), Course.name.ilike(search))
                ),
            )
        )

    if class_id is not None:
        query = query.where(Assignment.class_id == class_id)

    paginated = q is not None or class_id is not None or page is not None or page_size is not None
    total = db.scalar(select(func.count()).select_from(query.order_by(None).subquery())) or 0
    actual_page = page or 1
    actual_page_size = page_size or 20
    query = query.offset((actual_page - 1) * actual_page_size).limit(actual_page_size)
    set_pagination_headers(response, total=total, page=actual_page, page_size=actual_page_size)

    assignment_rows = db.execute(query).all()
    items = [
        {
            "id": assignment.id,
            "class_id": training_class.id,
            "class_code": training_class.code,
            "class_name": training_class.name,
            "course_name": training_class.course.name,
            "title": assignment.title,
            "description": assignment.description,
            "due_at": assignment.due_at.isoformat() if assignment.due_at else None,
            "max_score": assignment.max_score,
            "status": assignment.status,
            "session_id": assignment.session_id,
            "created_at": assignment.created_at.isoformat(),
        }
        for assignment, training_class in assignment_rows
    ]
    if assignment_rows:
        assignment_ids = [assignment.id for assignment, _ in assignment_rows]
        files_by_assignment: dict[int, list[dict]] = {assignment_id: [] for assignment_id in assignment_ids}
        for file_record in db.scalars(
            select(PrivateAssignmentFile)
            .where(PrivateAssignmentFile.assignment_id.in_(assignment_ids))
            .order_by(PrivateAssignmentFile.assignment_id, PrivateAssignmentFile.id)
        ).all():
            files_by_assignment[file_record.assignment_id].append({
                "id": file_record.id,
                "filename": file_record.original_filename,
                "content_type": file_record.content_type,
                "size_bytes": file_record.size_bytes,
                "download_url": f"/api/assignment-files/{file_record.id}",
            })
        for item, (assignment, _) in zip(items, assignment_rows):
            item["attachments"] = files_by_assignment[assignment.id]
    is_student_only = "student" in roles and not roles.intersection(
        ("admin", "manager", "teacher", "ta")
    )
    if is_student_only and assignment_rows:
        assignment_ids = [assignment.id for assignment, _ in assignment_rows]
        latest_submissions: dict[int, AssignmentSubmission] = {}
        for submission in db.scalars(
            select(AssignmentSubmission)
            .where(
                AssignmentSubmission.assignment_id.in_(assignment_ids),
                AssignmentSubmission.student_id == current_user.id,
            )
            .order_by(
                AssignmentSubmission.assignment_id,
                AssignmentSubmission.version_number.desc(),
            )
        ).all():
            latest_submissions.setdefault(submission.assignment_id, submission)

        published_grades = db.scalars(
            select(AssignmentGradeVersion)
            .where(
                AssignmentGradeVersion.assignment_id.in_(assignment_ids),
                AssignmentGradeVersion.student_id == current_user.id,
                AssignmentGradeVersion.status == "published",
            )
            .order_by(
                AssignmentGradeVersion.assignment_id,
                AssignmentGradeVersion.version_number.desc(),
            )
        ).all()
        latest_published_grades: dict[int, AssignmentGradeVersion] = {}
        for grade in published_grades:
            latest_published_grades.setdefault(grade.assignment_id, grade)
        open_redo_assignment_ids = set(db.scalars(
            select(AssignmentReworkRequest.assignment_id).where(
                AssignmentReworkRequest.assignment_id.in_(assignment_ids),
                AssignmentReworkRequest.student_id == current_user.id,
                AssignmentReworkRequest.request_type == "redo",
                AssignmentReworkRequest.status == "open",
            )
        ).all())

        now = datetime.now(timezone.utc)
        for item, (assignment, _) in zip(items, assignment_rows):
            deadline = _submission_deadline(assignment)
            remaining_seconds = (
                int((deadline - now).total_seconds()) if deadline is not None else None
            )
            if remaining_seconds is not None:
                item["is_overdue"] = remaining_seconds < 0
                item["remaining_seconds"] = max(0, remaining_seconds)
            else:
                item["is_overdue"] = False
                item["remaining_seconds"] = None

            submission = latest_submissions.get(assignment.id)
            published_grade = latest_published_grades.get(assignment.id)
            if submission is None:
                item["submission_status"] = "not_submitted"
                item["submission_version"] = None
                item["submitted_at"] = None
                item["is_late"] = False
                item["late_hours"] = None
            else:
                item["submission_version"] = submission.version_number
                item["submitted_at"] = submission.submitted_at.replace(
                    tzinfo=timezone.utc
                ).isoformat()
                elapsed_late = bool(deadline is not None and submission.submitted_at.replace(
                    tzinfo=timezone.utc
                ) > deadline)
                item["is_late"] = submission.is_late or elapsed_late
                item["late_hours"] = submission.late_hours
                if item["late_hours"] is None and elapsed_late and deadline is not None:
                    item["late_hours"] = max(1, ceil((submission.submitted_at.replace(
                        tzinfo=timezone.utc
                    ) - deadline).total_seconds() / 3600))
                item["submission_status"] = (
                    "needs_redo"
                    if assignment.id in open_redo_assignment_ids
                    else "graded"
                    if published_grade is not None
                    and published_grade.submission_id == submission.id
                    else "submitted"
                )
    if not paginated:
        return items
    return {
        "items": items,
        "total": total,
        "page": actual_page,
        "page_size": actual_page_size,
        "pages": (total + actual_page_size - 1) // actual_page_size,
        "has_previous": actual_page > 1 and total > 0,
        "has_next": actual_page * actual_page_size < total,
    }


@router.post(
    "/classes/{class_id}/assignments", status_code=status.HTTP_201_CREATED
)
def create_assignment(
    class_id: int,
    request: AssignmentCreate,
    current_user: User = Depends(
        require_roles("admin", "manager", "teacher", "ta")
    ),
    db: Session = Depends(get_db),
) -> dict:
    training_class = db.scalar(
        select(TrainingClass)
        .options(joinedload(TrainingClass.course))
        .where(TrainingClass.id == class_id)
    )
    if training_class is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy lớp học.")
    _assert_class_staff_access(current_user, training_class)

    if request.session_id is not None:
        session = db.get(AttendanceSession, request.session_id)
        if session is None or session.class_id != class_id or session.is_cancelled:
            raise HTTPException(
                status_code=422,
                detail="Buổi học phải thuộc lớp được chọn.",
            )

    if request.status == "published" and request.due_at is not None:
        due_at_local = (
            request.due_at.replace(tzinfo=HO_CHI_MINH)
            if request.due_at.tzinfo is None
            else request.due_at.astimezone(HO_CHI_MINH)
        )
        if due_at_local < datetime.now(HO_CHI_MINH):
            raise HTTPException(
                status_code=422,
                detail="Hạn nộp không được ở trong quá khứ.",
            )

    requested_status = request.status
    assignment_data = request.model_dump(exclude={"status"})
    assignment = Assignment(
        class_id=class_id,
        created_by=current_user.id,
        status="draft",
        **assignment_data,
    )
    db.add(assignment)
    db.flush()
    if requested_status == "published":
        _publish_assignment(db, assignment, training_class)
    db.commit()
    db.refresh(assignment)
    return _assignment_payload(assignment, training_class)


@router.post("/assignments/{assignment_id}/files", status_code=201)
async def upload_assignment_instruction_file(
    assignment_id: int,
    request: Request,
    current_user: User = Depends(require_roles("admin", "manager", "teacher", "ta")),
    db: Session = Depends(get_db),
) -> dict:
    assignment = db.get(Assignment, assignment_id)
    if assignment is None:
        raise HTTPException(status_code=404, detail="Assignment not found.")
    training_class = db.get(TrainingClass, assignment.class_id)
    _assert_class_staff_access(current_user, training_class)
    filename = request.headers.get("x-file-name", "")
    content_type = request.headers.get("content-type", "application/octet-stream")
    content = bytearray()
    async for chunk in request.stream():
        content.extend(chunk)
        if len(content) > 20 * 1024 * 1024:
            raise HTTPException(status_code=413, detail="File exceeds the 20 MiB limit.")
    key, safe_name, size, digest = save_private_file(filename, content_type, bytes(content))
    record = PrivateAssignmentFile(
        assignment_id=assignment.id,
        original_filename=safe_name,
        storage_key=key,
        content_type=content_type.split(";")[0].strip().lower(),
        size_bytes=size,
        sha256=digest,
        uploaded_by=current_user.id,
    )
    try:
        db.add(record)
        db.commit()
    except Exception:
        db.rollback()
        private_file_path(key).unlink(missing_ok=True)
        raise
    db.refresh(record)
    return {
        "id": record.id,
        "assignment_id": assignment.id,
        "filename": record.original_filename,
        "content_type": record.content_type,
        "size_bytes": record.size_bytes,
        "sha256": record.sha256,
        "created_at": record.created_at.isoformat(),
        "download_url": f"/api/assignment-files/{record.id}",
    }


@router.get("/assignments/{assignment_id}/files")
def list_assignment_instruction_files(
    assignment_id: int,
    response: Response,
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=20, ge=1, le=100),
    current_user: User = Depends(require_roles(*ALL_ROLES)),
    db: Session = Depends(get_db),
) -> list[dict]:
    assignment = db.get(Assignment, assignment_id)
    if assignment is None:
        raise HTTPException(status_code=404, detail="Assignment not found.")
    roles = {role.code for role in current_user.roles}
    student_only = "student" in roles and not roles.intersection(("admin", "manager", "teacher", "ta"))
    if student_only:
        if assignment.status != "published" or not _student_is_assignment_recipient(db, assignment, current_user.id):
            raise HTTPException(status_code=404, detail="Assignment not found.")
    else:
        _assert_class_staff_access(current_user, db.get(TrainingClass, assignment.class_id))
    base = select(PrivateAssignmentFile).where(PrivateAssignmentFile.assignment_id == assignment.id)
    total = db.scalar(select(func.count()).select_from(base.subquery())) or 0
    records = db.scalars(
        base
        .order_by(PrivateAssignmentFile.id)
        .offset((page - 1) * page_size).limit(page_size)
    ).all()
    set_pagination_headers(response, total=total, page=page, page_size=page_size)
    return [{
        "id": record.id,
        "filename": record.original_filename,
        "content_type": record.content_type,
        "size_bytes": record.size_bytes,
        "sha256": record.sha256,
        "created_at": record.created_at.isoformat(),
        "download_url": f"/api/assignment-files/{record.id}",
    } for record in records]


def _assignment_payload(assignment: Assignment, training_class: TrainingClass) -> dict:
    return {
        "id": assignment.id,
        "class_id": training_class.id,
        "class_code": training_class.code,
        "class_name": training_class.name,
        "course_name": training_class.course.name,
        "title": assignment.title,
        "description": assignment.description,
        "due_at": assignment.due_at.isoformat() if assignment.due_at else None,
        "max_score": assignment.max_score,
        "status": assignment.status,
        "session_id": assignment.session_id,
        "grade_component_key": assignment.grade_component_key,
        "late_submission_policy": assignment.late_submission_policy,
        "created_at": assignment.created_at.isoformat(),
    }


def _publish_assignment(
    db: Session,
    assignment: Assignment,
    training_class: TrainingClass,
) -> None:
    if assignment.status == "published":
        return
    if assignment.due_at is not None:
        due_at_local = (
            assignment.due_at.replace(tzinfo=HO_CHI_MINH)
            if assignment.due_at.tzinfo is None
            else assignment.due_at.astimezone(HO_CHI_MINH)
        )
        if due_at_local < datetime.now(HO_CHI_MINH):
            raise HTTPException(status_code=422, detail="Hạn nộp không được ở trong quá khứ.")

    assignment.status = "published"
    assignment.roster_snapshotted_at = datetime.now(timezone.utc).replace(tzinfo=None)
    published_on = datetime.now(HO_CHI_MINH).date()
    class_enrollments = db.scalars(
        select(ClassEnrollment)
        .where(ClassEnrollment.class_id == assignment.class_id)
        .order_by(ClassEnrollment.id)
    ).all()
    recipients = [
        enrollment.student_id
        for enrollment in class_enrollments
        if (
            enrollment.enrollment_status == "active"
            and (
                enrollment.effective_date is None
                or enrollment.effective_date <= published_on
            )
        )
        or (
            enrollment.enrollment_status != "active"
            and enrollment.effective_date is not None
            and published_on < enrollment.effective_date
        )
    ]
    db.add_all(
        AssignmentRecipient(assignment_id=assignment.id, student_id=student_id)
        for student_id in recipients
    )
    for student_id in recipients:
        notification = Notification(
            user_id=student_id,
            class_id=training_class.id,
            title="Bài tập mới",
            message=f"{assignment.title} đã được giao cho lớp {training_class.code}.",
        )
        db.add(notification)
        recipient = db.get(User, student_id)
        enqueue_notification_email(db, notification, recipient)


@router.post("/assignments/{assignment_id}/publish")
def publish_assignment(
    assignment_id: int,
    current_user: User = Depends(require_roles("admin", "manager", "teacher", "ta")),
    db: Session = Depends(get_db),
) -> dict:
    assignment = db.get(Assignment, assignment_id)
    if assignment is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy bài tập.")
    training_class = db.scalar(
        select(TrainingClass)
        .options(joinedload(TrainingClass.course), selectinload(TrainingClass.assistants))
        .where(TrainingClass.id == assignment.class_id)
    )
    if training_class is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy lớp học.")
    _assert_class_staff_access(current_user, training_class)
    if assignment.status == "closed":
        raise HTTPException(status_code=409, detail="Bài tập đã đóng, không thể giao lại.")
    if assignment.status == "draft":
        _publish_assignment(db, assignment, training_class)
        db.commit()
        db.refresh(assignment)
    return _assignment_payload(assignment, training_class)


def _student_is_assignment_recipient(
    db: Session, assignment: Assignment, student_id: int
) -> bool:
    if assignment.roster_snapshotted_at is not None:
        return db.scalar(
            select(AssignmentRecipient.id).where(
                AssignmentRecipient.assignment_id == assignment.id,
                AssignmentRecipient.student_id == student_id,
            )
        ) is not None
    return db.scalar(
        select(ClassEnrollment.id).where(
            ClassEnrollment.class_id == assignment.class_id,
            ClassEnrollment.student_id == student_id,
        )
    ) is not None


def _submission_deadline(assignment: Assignment) -> datetime | None:
    if assignment.due_at is None:
        return None
    if assignment.due_at.tzinfo is None:
        # Existing assignment API treats naive due_at values as ICTU local time.
        local_due_at = assignment.due_at.replace(tzinfo=HO_CHI_MINH)
    else:
        local_due_at = assignment.due_at.astimezone(HO_CHI_MINH)
    return local_due_at.astimezone(timezone.utc)


def _submission_payload(submission: AssignmentSubmission) -> dict:
    return {
        "id": submission.id,
        "assignment_id": submission.assignment_id,
        "student_id": submission.student_id,
        "version_number": submission.version_number,
        "repository_url": submission.repository_url,
        "submitted_at": submission.submitted_at.replace(tzinfo=timezone.utc).isoformat(),
        "status": "late" if submission.is_late else "submitted",
        "is_late": submission.is_late,
        "late_hours": submission.late_hours,
    }


@router.post(
    "/assignments/{assignment_id}/submissions",
    status_code=status.HTTP_201_CREATED,
)
def create_assignment_submission(
    assignment_id: int,
    request: AssignmentSubmissionCreate,
    student: User = Depends(require_roles("student")),
    db: Session = Depends(get_db),
) -> dict:
    assignment = db.get(Assignment, assignment_id)
    if assignment is None or assignment.status != "published":
        raise HTTPException(status_code=404, detail="Không tìm thấy bài tập.")
    if not _student_is_assignment_recipient(db, assignment, student.id):
        raise HTTPException(status_code=403, detail="Bạn không có quyền nộp bài tập này.")

    latest_submission_id = db.scalar(
        select(AssignmentSubmission.id).where(
            AssignmentSubmission.assignment_id == assignment.id,
            AssignmentSubmission.student_id == student.id,
        ).order_by(AssignmentSubmission.version_number.desc()).limit(1)
    )
    published_grade = db.scalar(
        select(AssignmentGradeVersion.id).where(
            AssignmentGradeVersion.submission_id == latest_submission_id,
            AssignmentGradeVersion.status == "published",
        )
    ) if latest_submission_id is not None else None
    rework = db.scalar(
        select(AssignmentReworkRequest).where(
            AssignmentReworkRequest.assignment_id == assignment.id,
            AssignmentReworkRequest.student_id == student.id,
            AssignmentReworkRequest.status == "open",
        ).order_by(AssignmentReworkRequest.request_number.desc()).limit(1)
    )
    if published_grade is not None and rework is None:
        raise HTTPException(
            status_code=409,
            detail="Bài đã được công bố điểm; cần giảng viên mở lại trước khi nộp tiếp.",
        )

    now = datetime.now(timezone.utc)
    deadline = _submission_deadline(assignment)
    if rework is not None:
        if rework.request_type == "redo" and rework.due_at is not None:
            deadline = rework.due_at.replace(tzinfo=timezone.utc)
    is_late = deadline is not None and now > deadline
    if is_late:
        if assignment.late_submission_policy == "BLOCK_LATE":
            raise HTTPException(status_code=409, detail="This assignment does not accept late submissions.")
        if assignment.late_submission_policy != "ALLOW_LATE":
            raise HTTPException(status_code=409, detail="This assignment does not accept late submissions.")
        late_hours = max(1, ceil((now - deadline).total_seconds() / 3600))
    else:
        late_hours = None

    latest_version = db.scalar(
        select(func.max(AssignmentSubmission.version_number)).where(
            AssignmentSubmission.assignment_id == assignment.id,
            AssignmentSubmission.student_id == student.id,
        )
    ) or 0
    submission = AssignmentSubmission(
        assignment_id=assignment.id,
        student_id=student.id,
        version_number=latest_version + 1,
        repository_url=str(request.repository_url),
        submitted_at=now.replace(tzinfo=None),
        is_late=is_late,
        late_hours=late_hours,
    )
    db.add(submission)
    db.flush()
    if rework is not None:
        rework.status = "completed"
        rework.completed_submission_id = submission.id
    _record_audit(
        db,
        student,
        "assignment.submitted",
        target_user_id=student.id,
        details=(
            f"assignment_id={assignment.id};submission_id={submission.id};"
            f"version={submission.version_number}"
        ),
    )
    try:
        db.commit()
    except IntegrityError as error:
        db.rollback()
        raise HTTPException(
            status_code=409,
            detail="Có lượt nộp đồng thời. Vui lòng tải lại lịch sử và thử lại.",
        ) from error
    db.refresh(submission)
    return _submission_payload(submission)


@router.post("/assignments/{assignment_id}/students/{student_id}/redo-requests", status_code=201)
def request_assignment_redo(
    assignment_id: int,
    student_id: int,
    request: AssignmentReworkRequestCreate,
    current_user: User = Depends(require_roles("admin", "manager", "teacher", "ta")),
    db: Session = Depends(get_db),
) -> dict:
    assignment = db.get(Assignment, assignment_id)
    if assignment is None:
        raise HTTPException(status_code=404, detail="Assignment not found.")
    training_class = db.get(TrainingClass, assignment.class_id)
    _assert_class_staff_access(current_user, training_class)
    if not _student_is_assignment_recipient(db, assignment, student_id):
        raise HTTPException(status_code=404, detail="Student is not assigned this work.")
    latest = db.scalar(select(AssignmentSubmission).where(
        AssignmentSubmission.assignment_id == assignment_id,
        AssignmentSubmission.student_id == student_id,
    ).order_by(AssignmentSubmission.version_number.desc()).limit(1))
    if latest is None:
        raise HTTPException(status_code=409, detail="A submission is required before requesting rework.")
    grade = db.scalar(select(AssignmentGradeVersion.id).where(
        AssignmentGradeVersion.submission_id == latest.id,
        AssignmentGradeVersion.status == "published",
    ))
    if grade is None:
        raise HTTPException(status_code=409, detail="Rework can only be requested after grading is published.")
    open_request = db.scalar(select(AssignmentReworkRequest.id).where(
        AssignmentReworkRequest.assignment_id == assignment_id,
        AssignmentReworkRequest.student_id == student_id,
        AssignmentReworkRequest.status == "open",
    ))
    if open_request is not None:
        raise HTTPException(status_code=409, detail="An open rework request already exists.")
    due = request.due_at
    now = datetime.now(timezone.utc)
    due_utc = due.replace(tzinfo=HO_CHI_MINH).astimezone(timezone.utc) if due.tzinfo is None else due.astimezone(timezone.utc)
    if due_utc <= now:
        raise HTTPException(status_code=422, detail="The rework deadline must be in the future.")
    last_no = db.scalar(select(func.max(AssignmentReworkRequest.request_number)).where(
        AssignmentReworkRequest.assignment_id == assignment_id,
        AssignmentReworkRequest.student_id == student_id,
    )) or 0
    row = AssignmentReworkRequest(
        assignment_id=assignment_id, student_id=student_id, submission_id=latest.id,
        request_number=last_no + 1, request_type="redo", feedback=request.feedback,
        due_at=due_utc.replace(tzinfo=None), requested_by=current_user.id,
    )
    db.add(row)
    db.commit()
    db.refresh(row)
    return {"id": row.id, "assignment_id": row.assignment_id, "student_id": row.student_id,
            "submission_id": row.submission_id, "request_number": row.request_number,
        "request_type": row.request_type,
            "feedback": row.feedback, "due_at": row.due_at.replace(tzinfo=timezone.utc).isoformat(),
            "requested_by": row.requested_by, "requested_at": row.requested_at.isoformat(), "status": row.status}


@router.post("/assignments/{assignment_id}/students/{student_id}/reopen", status_code=201)
def reopen_assignment_submission(
    assignment_id: int,
    student_id: int,
    current_user: User = Depends(require_roles("admin", "manager", "teacher", "ta")),
    db: Session = Depends(get_db),
) -> dict:
    assignment = db.get(Assignment, assignment_id)
    if assignment is None:
        raise HTTPException(status_code=404, detail="Assignment not found.")
    if assignment.status != "published":
        raise HTTPException(status_code=409, detail="Only a published assignment can be reopened.")
    _assert_class_staff_access(current_user, db.get(TrainingClass, assignment.class_id))
    if not _student_is_assignment_recipient(db, assignment, student_id):
        raise HTTPException(status_code=404, detail="Student is not assigned this work.")
    latest = db.scalar(select(AssignmentSubmission).where(
        AssignmentSubmission.assignment_id == assignment_id,
        AssignmentSubmission.student_id == student_id,
    ).order_by(AssignmentSubmission.version_number.desc()).limit(1))
    if latest is None or db.scalar(select(AssignmentGradeVersion.id).where(
        AssignmentGradeVersion.submission_id == latest.id,
        AssignmentGradeVersion.status == "published",
    )) is None:
        raise HTTPException(status_code=409, detail="A published grade for the latest submission is required.")
    if db.scalar(select(AssignmentReworkRequest.id).where(
        AssignmentReworkRequest.assignment_id == assignment_id,
        AssignmentReworkRequest.student_id == student_id,
        AssignmentReworkRequest.status == "open",
    )) is not None:
        raise HTTPException(status_code=409, detail="An open redo or reopen request already exists.")
    last_no = db.scalar(select(func.max(AssignmentReworkRequest.request_number)).where(
        AssignmentReworkRequest.assignment_id == assignment_id,
        AssignmentReworkRequest.student_id == student_id,
    )) or 0
    row = AssignmentReworkRequest(assignment_id=assignment_id, student_id=student_id,
        submission_id=latest.id, request_number=last_no + 1, request_type="reopen",
        requested_by=current_user.id, status="open")
    db.add(row)
    db.commit()
    db.refresh(row)
    return {"id": row.id, "assignment_id": row.assignment_id, "student_id": row.student_id,
            "submission_id": row.submission_id, "request_number": row.request_number,
            "request_type": row.request_type, "requested_by": row.requested_by,
            "requested_at": row.requested_at.isoformat(), "status": row.status}


@router.get("/assignments/{assignment_id}/rework-requests")
def list_assignment_rework_requests(
    assignment_id: int,
    response: Response,
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=20, ge=1, le=100),
    current_user: User = Depends(require_roles("admin", "manager", "teacher", "ta", "student")),
    db: Session = Depends(get_db),
) -> list[dict]:
    assignment = db.get(Assignment, assignment_id)
    if assignment is None:
        raise HTTPException(status_code=404, detail="Assignment not found.")
    roles = {role.code for role in current_user.roles}
    query = select(AssignmentReworkRequest).where(AssignmentReworkRequest.assignment_id == assignment_id)
    if "student" in roles and not roles.intersection(("admin", "manager", "teacher", "ta")):
        if not _student_is_assignment_recipient(db, assignment, current_user.id):
            raise HTTPException(status_code=403, detail="Not assigned to this work.")
        query = query.where(AssignmentReworkRequest.student_id == current_user.id)
    else:
        _assert_class_staff_access(current_user, db.get(TrainingClass, assignment.class_id))
    total = db.scalar(select(func.count()).select_from(query.subquery())) or 0
    rows = db.scalars(query.order_by(AssignmentReworkRequest.student_id, AssignmentReworkRequest.request_number, AssignmentReworkRequest.id)
        .offset((page - 1) * page_size).limit(page_size)).all()
    set_pagination_headers(response, total=total, page=page, page_size=page_size)
    return [{"id": r.id, "student_id": r.student_id, "submission_id": r.submission_id,
             "request_number": r.request_number, "feedback": r.feedback,
             "request_type": r.request_type,
             "due_at": r.due_at.replace(tzinfo=timezone.utc).isoformat() if r.due_at else None,
             "requested_by": r.requested_by, "requested_at": r.requested_at.isoformat(),
             "status": r.status, "completed_submission_id": r.completed_submission_id} for r in rows]


@router.post("/assignments/{assignment_id}/submissions/{submission_id}/files", status_code=201)
async def upload_assignment_submission_file(
    assignment_id: int,
    submission_id: int,
    request: Request,
    current_user: User = Depends(require_roles("admin", "manager", "teacher", "ta", "student")),
    db: Session = Depends(get_db),
) -> dict:
    assignment = db.get(Assignment, assignment_id)
    submission = db.get(AssignmentSubmission, submission_id)
    if assignment is None or submission is None or submission.assignment_id != assignment_id:
        raise HTTPException(status_code=404, detail="Submission not found.")
    roles = {role.code for role in current_user.roles}
    student_only = "student" in roles and not roles.intersection(("admin", "manager", "teacher", "ta"))
    if student_only:
        if submission.student_id != current_user.id or not _student_is_assignment_recipient(db, assignment, current_user.id):
            raise HTTPException(status_code=404, detail="Submission not found.")
    else:
        _assert_class_staff_access(current_user, db.get(TrainingClass, assignment.class_id))
    filename = request.headers.get("x-file-name", "")
    content_type = request.headers.get("content-type", "application/octet-stream")
    chunks = bytearray()
    async for chunk in request.stream():
        chunks.extend(chunk)
        if len(chunks) > 20 * 1024 * 1024:
            raise HTTPException(status_code=413, detail="File exceeds the 20 MiB limit.")
    storage_key, safe_name, size_bytes, digest = save_private_file(filename, content_type, bytes(chunks))
    row = PrivateAssignmentFile(
        submission_id=submission.id,
        original_filename=safe_name,
        storage_key=storage_key,
        content_type=content_type.split(";")[0].strip().lower(),
        size_bytes=size_bytes,
        sha256=digest,
        uploaded_by=current_user.id,
    )
    try:
        db.add(row)
        db.commit()
    except Exception:
        db.rollback()
        private_file_path(storage_key).unlink(missing_ok=True)
        raise
    db.refresh(row)
    return {"id": row.id, "submission_id": row.submission_id, "filename": row.original_filename,
            "content_type": row.content_type, "size_bytes": row.size_bytes, "sha256": row.sha256,
            "created_at": row.created_at.isoformat(), "download_url": f"/api/assignment-files/{row.id}"}


@router.get("/assignments/{assignment_id}/submissions/{submission_id}/files")
def list_assignment_submission_files(
    assignment_id: int,
    submission_id: int,
    response: Response,
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=20, ge=1, le=100),
    current_user: User = Depends(require_roles("admin", "manager", "teacher", "ta", "student")),
    db: Session = Depends(get_db),
) -> list[dict]:
    assignment = db.get(Assignment, assignment_id)
    submission = db.get(AssignmentSubmission, submission_id)
    if assignment is None or submission is None or submission.assignment_id != assignment_id:
        raise HTTPException(status_code=404, detail="Submission not found.")
    roles = {role.code for role in current_user.roles}
    student_only = "student" in roles and not roles.intersection(("admin", "manager", "teacher", "ta"))
    if student_only:
        if submission.student_id != current_user.id or not _student_is_assignment_recipient(db, assignment, current_user.id):
            raise HTTPException(status_code=404, detail="Submission not found.")
    else:
        _assert_class_staff_access(current_user, db.get(TrainingClass, assignment.class_id))
    base = select(PrivateAssignmentFile).where(PrivateAssignmentFile.submission_id == submission_id)
    total = db.scalar(select(func.count()).select_from(base.subquery())) or 0
    files = db.scalars(base.order_by(PrivateAssignmentFile.id).offset((page - 1) * page_size).limit(page_size)).all()
    set_pagination_headers(response, total=total, page=page, page_size=page_size)
    return [{"id": f.id, "filename": f.original_filename, "content_type": f.content_type,
             "size_bytes": f.size_bytes, "sha256": f.sha256, "created_at": f.created_at.isoformat(),
             "download_url": f"/api/assignment-files/{f.id}"} for f in files]


@router.get("/assignment-files/{file_id}")
def download_assignment_file(
    file_id: int,
    current_user: User = Depends(require_roles("admin", "manager", "teacher", "ta", "student")),
    db: Session = Depends(get_db),
) -> FileResponse:
    record = db.get(PrivateAssignmentFile, file_id)
    if record is None:
        raise HTTPException(status_code=404, detail="File not found.")
    roles = {role.code for role in current_user.roles}
    student_only = "student" in roles and not roles.intersection(("admin", "manager", "teacher", "ta"))
    if record.assignment_id is not None:
        assignment = db.get(Assignment, record.assignment_id)
        if assignment is None:
            raise HTTPException(status_code=404, detail="File not found.")
        if student_only:
            if assignment.status != "published" or not _student_is_assignment_recipient(db, assignment, current_user.id):
                raise HTTPException(status_code=404, detail="File not found.")
        else:
            _assert_class_staff_access(current_user, db.get(TrainingClass, assignment.class_id))
    elif record.submission_id is not None:
        submission = db.get(AssignmentSubmission, record.submission_id)
        assignment = db.get(Assignment, submission.assignment_id)
        if student_only:
            if submission.student_id != current_user.id or not _student_is_assignment_recipient(db, assignment, current_user.id):
                raise HTTPException(status_code=404, detail="File not found.")
        else:
            _assert_class_staff_access(current_user, db.get(TrainingClass, assignment.class_id))
    else:
        session = db.get(AttendanceSession, record.session_id)
        training_class = db.get(TrainingClass, session.class_id)
        if student_only:
            enrolled = db.scalar(select(ClassEnrollment.id).where(
                ClassEnrollment.class_id == session.class_id,
                ClassEnrollment.student_id == current_user.id,
                ClassEnrollment.enrollment_status == "active",
            ))
            if enrolled is None:
                raise HTTPException(status_code=404, detail="File not found.")
            session_start = datetime.combine(session.session_date, session.start_time, tzinfo=HO_CHI_MINH)
            if session.is_cancelled or (not session.materials_open_early and datetime.now(HO_CHI_MINH) < session_start):
                raise HTTPException(status_code=403, detail="Materials are not available yet.")
        else:
            _assert_class_staff_access(current_user, training_class)
    return FileResponse(private_file_path(record.storage_key), media_type=record.content_type,
                        filename=record.original_filename, headers={"X-Content-Type-Options": "nosniff"})


@router.post("/sessions/{session_id}/materials", status_code=201)
async def upload_session_learning_material(
    session_id: int,
    request: Request,
    current_user: User = Depends(require_roles("admin", "manager", "teacher", "ta")),
    db: Session = Depends(get_db),
) -> dict:
    session = db.get(AttendanceSession, session_id)
    if session is None:
        raise HTTPException(status_code=404, detail="Session not found.")
    _assert_class_staff_access(current_user, db.get(TrainingClass, session.class_id))
    filename = request.headers.get("x-file-name", "")
    content_type = request.headers.get("content-type", "application/octet-stream")
    content = bytearray()
    async for chunk in request.stream():
        content.extend(chunk)
        if len(content) > 20 * 1024 * 1024:
            raise HTTPException(status_code=413, detail="File exceeds the 20 MiB limit.")
    key, safe_name, size, digest = save_private_file(filename, content_type, bytes(content))
    record = PrivateAssignmentFile(session_id=session_id, original_filename=safe_name,
        storage_key=key, content_type=content_type.split(";")[0].strip().lower(),
        size_bytes=size, sha256=digest, uploaded_by=current_user.id)
    try:
        db.add(record)
        db.commit()
    except Exception:
        db.rollback()
        private_file_path(key).unlink(missing_ok=True)
        raise
    db.refresh(record)
    return {"id": record.id, "session_id": session_id, "filename": record.original_filename,
            "content_type": record.content_type, "size_bytes": record.size_bytes,
            "sha256": record.sha256, "uploaded_by": record.uploaded_by,
            "created_at": record.created_at.isoformat(),
            "download_url": f"/api/assignment-files/{record.id}"}


@router.patch("/sessions/{session_id}/materials/visibility")
def update_session_materials_visibility(
    session_id: int,
    request: SessionMaterialsVisibilityUpdate,
    current_user: User = Depends(require_roles("admin", "manager", "teacher", "ta")),
    db: Session = Depends(get_db),
) -> dict:
    session = db.get(AttendanceSession, session_id)
    if session is None:
        raise HTTPException(status_code=404, detail="Session not found.")
    _assert_class_staff_access(current_user, db.get(TrainingClass, session.class_id))
    session.materials_open_early = request.materials_open_early
    db.commit()
    return {"session_id": session_id, "materials_open_early": session.materials_open_early}


@router.get("/sessions/{session_id}/materials")
def list_session_learning_materials(
    session_id: int,
    response: Response,
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=20, ge=1, le=100),
    current_user: User = Depends(require_roles("admin", "manager", "teacher", "ta", "student")),
    db: Session = Depends(get_db),
) -> list[dict]:
    session = db.get(AttendanceSession, session_id)
    if session is None:
        raise HTTPException(status_code=404, detail="Session not found.")
    roles = {role.code for role in current_user.roles}
    student_only = "student" in roles and not roles.intersection(("admin", "manager", "teacher", "ta"))
    if student_only:
        enrolled = db.scalar(select(ClassEnrollment.id).where(
            ClassEnrollment.class_id == session.class_id,
            ClassEnrollment.student_id == current_user.id,
            ClassEnrollment.enrollment_status == "active",
        ))
        if enrolled is None:
            raise HTTPException(status_code=404, detail="File not found.")
        session_start = datetime.combine(session.session_date, session.start_time, tzinfo=HO_CHI_MINH)
        if session.is_cancelled or (not session.materials_open_early and datetime.now(HO_CHI_MINH) < session_start):
            raise HTTPException(status_code=403, detail="Materials are not available yet.")
    else:
        _assert_class_staff_access(current_user, db.get(TrainingClass, session.class_id))
    base = select(PrivateAssignmentFile).where(PrivateAssignmentFile.session_id == session_id)
    total = db.scalar(select(func.count()).select_from(base.subquery())) or 0
    files = db.scalars(base.order_by(PrivateAssignmentFile.id).offset((page - 1) * page_size).limit(page_size)).all()
    set_pagination_headers(response, total=total, page=page, page_size=page_size)
    return [{"id": item.id, "filename": item.original_filename, "content_type": item.content_type,
             "size_bytes": item.size_bytes, "sha256": item.sha256,
             "uploaded_by": item.uploaded_by, "created_at": item.created_at.isoformat(),
             "download_url": f"/api/assignment-files/{item.id}"} for item in files]


@router.post("/sessions/{session_id}/materials/links", status_code=201)
def create_session_learning_link(
    session_id: int,
    request: SessionLearningLinkCreate,
    current_user: User = Depends(require_roles("admin", "manager", "teacher", "ta")),
    db: Session = Depends(get_db),
) -> dict:
    session = db.get(AttendanceSession, session_id)
    if session is None:
        raise HTTPException(status_code=404, detail="Session not found.")
    _assert_class_staff_access(current_user, db.get(TrainingClass, session.class_id))
    link = SessionLearningLink(session_id=session_id, title=request.title,
        url=str(request.url), created_by=current_user.id)
    db.add(link)
    db.commit()
    db.refresh(link)
    return {"id": link.id, "session_id": link.session_id, "title": link.title,
            "url": link.url, "created_by": link.created_by, "created_at": link.created_at.isoformat()}


@router.get("/sessions/{session_id}/materials/links")
def list_session_learning_links(
    session_id: int,
    response: Response,
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=20, ge=1, le=100),
    current_user: User = Depends(require_roles("admin", "manager", "teacher", "ta", "student")),
    db: Session = Depends(get_db),
) -> list[dict]:
    session = db.get(AttendanceSession, session_id)
    if session is None:
        raise HTTPException(status_code=404, detail="Session not found.")
    roles = {role.code for role in current_user.roles}
    student_only = "student" in roles and not roles.intersection(("admin", "manager", "teacher", "ta"))
    if student_only:
        enrolled = db.scalar(select(ClassEnrollment.id).where(
            ClassEnrollment.class_id == session.class_id,
            ClassEnrollment.student_id == current_user.id,
            ClassEnrollment.enrollment_status == "active",
        ))
        if enrolled is None:
            raise HTTPException(status_code=404, detail="Session not found.")
        session_start = datetime.combine(session.session_date, session.start_time, tzinfo=HO_CHI_MINH)
        if session.is_cancelled or (not session.materials_open_early and datetime.now(HO_CHI_MINH) < session_start):
            raise HTTPException(status_code=403, detail="Materials are not available yet.")
    else:
        _assert_class_staff_access(current_user, db.get(TrainingClass, session.class_id))
    base = select(SessionLearningLink).where(SessionLearningLink.session_id == session_id)
    total = db.scalar(select(func.count()).select_from(base.subquery())) or 0
    links = db.scalars(base.order_by(SessionLearningLink.id).offset((page - 1) * page_size).limit(page_size)).all()
    set_pagination_headers(response, total=total, page=page, page_size=page_size)
    return [{"id": link.id, "session_id": link.session_id, "title": link.title,
             "url": link.url, "created_by": link.created_by,
             "created_at": link.created_at.isoformat()} for link in links]


@router.get("/assignments/{assignment_id}/submissions")
def list_assignment_submissions(
    assignment_id: int,
    response: Response,
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=20, ge=1, le=100),
    submission_status: str = Query(
        default="all",
        alias="status",
        pattern="^(all|not_graded|graded|late|not_submitted)$",
    ),
    current_user: User = Depends(require_roles("admin", "manager", "teacher", "ta", "student")),
    db: Session = Depends(get_db),
) -> list[dict] | dict:
    assignment = db.get(Assignment, assignment_id)
    if assignment is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy bài tập.")
    roles = {role.code for role in current_user.roles}
    if "student" in roles and not roles.intersection(("admin", "manager", "teacher", "ta")):
        if (
            assignment.status != "published"
            or not _student_is_assignment_recipient(db, assignment, current_user.id)
        ):
            raise HTTPException(status_code=404, detail="Không tìm thấy bài tập.")
        query = select(AssignmentSubmission).where(
            AssignmentSubmission.assignment_id == assignment.id,
            AssignmentSubmission.student_id == current_user.id,
        )
        total = db.scalar(select(func.count()).select_from(query.subquery())) or 0
        submissions = db.scalars(query.order_by(AssignmentSubmission.version_number, AssignmentSubmission.id)
            .offset((page - 1) * page_size).limit(page_size)).all()
        set_pagination_headers(response, total=total, page=page, page_size=page_size)
        return [_submission_payload(item) for item in submissions]
    else:
        training_class = db.scalar(
            select(TrainingClass)
            .options(selectinload(TrainingClass.assistants))
            .where(TrainingClass.id == assignment.class_id)
        )
        if training_class is None:
            raise HTTPException(status_code=404, detail="Không tìm thấy lớp học.")
        _assert_class_staff_access(current_user, training_class)
        latest_versions = (
            select(AssignmentSubmission.student_id.label("student_id"),
                func.max(AssignmentSubmission.version_number).label("version_number"))
            .where(AssignmentSubmission.assignment_id == assignment.id)
            .group_by(AssignmentSubmission.student_id).subquery()
        )
        latest = aliased(AssignmentSubmission)
        roster_query = select(User.id, User.student_code, User.full_name, latest).distinct()
        if assignment.roster_snapshotted_at is not None:
            roster_query = roster_query.join(AssignmentRecipient, AssignmentRecipient.student_id == User.id).where(
                AssignmentRecipient.assignment_id == assignment.id)
        else:
            roster_query = roster_query.join(ClassEnrollment, ClassEnrollment.student_id == User.id).where(
                ClassEnrollment.class_id == assignment.class_id)
        roster_query = roster_query.outerjoin(
            latest_versions, latest_versions.c.student_id == User.id
        ).outerjoin(latest, and_(latest.assignment_id == assignment.id,
            latest.student_id == User.id, latest.version_number == latest_versions.c.version_number))
        deadline = _submission_deadline(assignment)
        is_late_expr = latest.is_late.is_(True)
        if deadline is not None:
            is_late_expr = or_(is_late_expr, latest.submitted_at > deadline.replace(tzinfo=None))
        has_published_grade = exists(select(AssignmentGradeVersion.id).where(
            AssignmentGradeVersion.submission_id == latest.id,
            AssignmentGradeVersion.status == "published"))
        status_expr = case(
            (latest.id.is_(None), "not_submitted"),
            (is_late_expr, "late"),
            (has_published_grade, "graded"),
            else_="not_graded",
        )
        if submission_status != "all":
            roster_query = roster_query.where(status_expr == submission_status)
        total = db.scalar(select(func.count()).select_from(roster_query.order_by(None).subquery())) or 0
        roster = db.execute(roster_query.order_by(User.full_name, User.id)
            .offset((page - 1) * page_size).limit(page_size)).all()
        page_student_ids = [row[0] for row in roster]
        submissions = db.scalars(
            select(AssignmentSubmission)
            .where(AssignmentSubmission.assignment_id == assignment.id,
                AssignmentSubmission.student_id.in_(page_student_ids or [-1]))
            .order_by(
                AssignmentSubmission.student_id,
                AssignmentSubmission.version_number,
            )
        ).all()
        by_student: dict[int, list[AssignmentSubmission]] = {}
        for item in submissions:
            by_student.setdefault(item.student_id, []).append(item)
        published_grades = db.scalars(
            select(AssignmentGradeVersion)
            .where(
                AssignmentGradeVersion.assignment_id == assignment.id,
                AssignmentGradeVersion.status == "published",
                AssignmentGradeVersion.student_id.in_(page_student_ids or [-1]),
            )
            .order_by(
                AssignmentGradeVersion.student_id,
                AssignmentGradeVersion.version_number.desc(),
            )
        ).all()
        grades_by_student: dict[int, AssignmentGradeVersion] = {}
        for grade in published_grades:
            grades_by_student.setdefault(grade.student_id, grade)
        roster_items = []
        for student_id, student_code, full_name, latest_row in roster:
            versions = by_student.get(student_id, [])
            latest_submission = versions[-1] if versions else None
            latest_grade = grades_by_student.get(student_id)
            calculated_late = bool(
                latest_submission is not None
                and deadline is not None
                and latest_submission.submitted_at.replace(tzinfo=timezone.utc) > deadline
            )
            is_late = bool(latest_submission and (latest_submission.is_late or calculated_late))
            late_hours = (
                latest_submission.late_hours
                if latest_submission and latest_submission.late_hours is not None
                else max(1, ceil((latest_submission.submitted_at.replace(tzinfo=timezone.utc) - deadline).total_seconds() / 3600))
                if latest_submission and deadline and is_late
                else None
            )
            grade_is_current = bool(
                latest_submission is not None
                and latest_grade is not None
                and latest_grade.submission_id == latest_submission.id
            )
            derived_status = "not_submitted" if latest_submission is None else "late" if is_late else "graded" if grade_is_current else "not_graded"
            roster_items.append({
                "student_id": student_id,
                "student_code": student_code,
                "full_name": full_name,
                "status": derived_status,
                "is_late": is_late,
                "late_hours": late_hours,
                "version_count": len(versions),
                "latest_submission": (
                    _submission_payload(latest_submission)
                    if latest_submission is not None
                    else None
                ),
                "published_grade": (
                    _grade_payload(latest_grade, include_internal=False)
                    if grade_is_current and latest_grade is not None
                    else None
                ),
            })
        set_pagination_headers(response, total=total, page=page, page_size=page_size)
        return roster_items


@router.get("/assignments/{assignment_id}/submissions/progress")
def get_assignment_submission_progress(
    assignment_id: int,
    current_user: User = Depends(require_roles("admin", "manager", "teacher", "ta")),
    db: Session = Depends(get_db),
) -> dict:
    assignment = db.get(Assignment, assignment_id)
    if assignment is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy bài tập.")
    _assert_assignment_staff_scope(db, current_user, assignment)

    if assignment.roster_snapshotted_at is not None:
        recipient_ids = select(AssignmentRecipient.student_id).where(
            AssignmentRecipient.assignment_id == assignment.id
        )
        total = db.scalar(
            select(func.count()).select_from(AssignmentRecipient).where(
                AssignmentRecipient.assignment_id == assignment.id
            )
        ) or 0
    else:
        recipient_ids = select(ClassEnrollment.student_id).where(
            ClassEnrollment.class_id == assignment.class_id
        )
        total = db.scalar(
            select(func.count()).select_from(ClassEnrollment).where(
                ClassEnrollment.class_id == assignment.class_id
            )
        ) or 0

    submitted = db.scalar(
        select(func.count(func.distinct(AssignmentSubmission.student_id))).where(
            AssignmentSubmission.assignment_id == assignment.id,
            AssignmentSubmission.student_id.in_(recipient_ids),
        )
    ) or 0
    graded = db.scalar(
        select(func.count(func.distinct(AssignmentGradeVersion.student_id))).where(
            AssignmentGradeVersion.assignment_id == assignment.id,
            AssignmentGradeVersion.status == "published",
            AssignmentGradeVersion.student_id.in_(recipient_ids),
        )
    ) or 0
    return {
        "assignment_id": assignment.id,
        "total_recipients": total,
        "submitted_count": submitted,
        "not_submitted_count": max(0, total - submitted),
        "graded_count": graded,
        "not_graded_count": max(0, submitted - graded),
        "grading_progress_percent": (
            round(graded / total * 100, 1) if total else None
        ),
    }


@router.get("/assignments/{assignment_id}/grading-queue")
def get_assignment_grading_queue(
    assignment_id: int,
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=20, ge=1, le=100),
    current_user: User = Depends(require_roles("admin", "manager", "teacher", "ta")),
    db: Session = Depends(get_db),
) -> dict:
    assignment = db.get(Assignment, assignment_id)
    if assignment is None or assignment.status != "published":
        raise HTTPException(status_code=404, detail="Không tìm thấy bài tập.")
    _assert_assignment_staff_scope(db, current_user, assignment)

    latest_versions = (
        select(
            AssignmentSubmission.student_id.label("student_id"),
            func.max(AssignmentSubmission.version_number).label("version_number"),
        )
        .where(AssignmentSubmission.assignment_id == assignment.id)
        .group_by(AssignmentSubmission.student_id)
        .subquery()
    )
    already_published = exists(
        select(AssignmentGradeVersion.id).where(
            AssignmentGradeVersion.submission_id == AssignmentSubmission.id,
            AssignmentGradeVersion.status == "published",
        )
    )
    if assignment.roster_snapshotted_at is not None:
        recipient_scope = exists(
            select(AssignmentRecipient.id).where(
                AssignmentRecipient.assignment_id == assignment.id,
                AssignmentRecipient.student_id == AssignmentSubmission.student_id,
            )
        )
    else:
        recipient_scope = exists(
            select(ClassEnrollment.id).where(
                ClassEnrollment.class_id == assignment.class_id,
                ClassEnrollment.student_id == AssignmentSubmission.student_id,
            )
        )
    queue_query = (
        select(AssignmentSubmission, User)
        .join(
            latest_versions,
            and_(
                latest_versions.c.student_id == AssignmentSubmission.student_id,
                latest_versions.c.version_number == AssignmentSubmission.version_number,
            ),
        )
        .join(User, User.id == AssignmentSubmission.student_id)
        .where(
            AssignmentSubmission.assignment_id == assignment.id,
            recipient_scope,
            ~already_published,
        )
    )
    total_ungraded = db.scalar(select(func.count()).select_from(queue_query.subquery())) or 0
    rows = db.execute(
        queue_query.order_by(User.full_name, User.id, AssignmentSubmission.id)
        .offset((page - 1) * page_size).limit(page_size)
    ).all()
    submission_ids = [submission.id for submission, _ in rows]
    files_by_submission: dict[int, list[dict]] = {submission_id: [] for submission_id in submission_ids}
    if submission_ids:
        for file_record in db.scalars(
            select(PrivateAssignmentFile)
            .where(PrivateAssignmentFile.submission_id.in_(submission_ids))
            .order_by(PrivateAssignmentFile.submission_id, PrivateAssignmentFile.id)
        ).all():
            files_by_submission[file_record.submission_id].append({
                "id": file_record.id,
                "filename": file_record.original_filename,
                "content_type": file_record.content_type,
                "size_bytes": file_record.size_bytes,
                "download_url": f"/api/assignment-files/{file_record.id}",
            })
    items = [
        {
            "submission_id": submission.id,
            "student_id": student.id,
            "student_code": student.student_code or student.username,
            "student_name": student.full_name,
            "version_number": submission.version_number,
            "repository_url": submission.repository_url,
            "files": files_by_submission[submission.id],
            "submitted_at": submission.submitted_at.replace(
                tzinfo=timezone.utc
            ).isoformat(),
            "grade_path": (
                f"/api/assignments/{assignment.id}/students/{student.id}/grades"
            ),
        }
        for submission, student in rows
    ]
    return {
        "assignment_id": assignment.id,
        "items": items,
        "next_submission_id": items[0]["submission_id"] if items else None,
        "total_ungraded": total_ungraded,
        **page_metadata(total=total_ungraded, page=page, page_size=page_size),
    }


def _assert_assignment_staff_scope(
    db: Session, current_user: User, assignment: Assignment
) -> TrainingClass:
    training_class = db.scalar(
        select(TrainingClass)
        .options(selectinload(TrainingClass.assistants))
        .where(TrainingClass.id == assignment.class_id)
    )
    if training_class is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy lớp học.")
    _assert_class_staff_access(current_user, training_class)
    return training_class


def _rubric_template_payload(template: RubricTemplate) -> dict:
    return {
        "id": template.id,
        "name": template.name,
        "description": template.description,
        "version": template.version,
        "criteria": template.criteria,
    }


def _assert_rubric_template_access(template: RubricTemplate, current_user: User) -> None:
    roles = {role.code for role in current_user.roles}
    if template.created_by != current_user.id and "admin" not in roles:
        raise HTTPException(status_code=404, detail="Không tìm thấy rubric.")


@router.post("/rubric-templates", status_code=status.HTTP_201_CREATED)
def create_rubric_template(
    request: RubricTemplateCreate,
    current_user: User = Depends(require_roles("admin", "manager", "teacher", "ta")),
    db: Session = Depends(get_db),
) -> dict:
    criteria = [
        {
            "id": index,
            "title": item.title,
            "description": item.description,
            "weight": str(item.weight),
        }
        for index, item in enumerate(request.criteria, start=1)
    ]
    template = RubricTemplate(
        name=request.name,
        description=request.description,
        criteria=criteria,
        created_by=current_user.id,
    )
    db.add(template)
    db.flush()
    _record_audit(
        db,
        current_user,
        "rubric.template_created",
        details=f"template_id={template.id};before=none;after=version=1;criteria_count={len(criteria)}",
    )
    db.commit()
    db.refresh(template)
    return _rubric_template_payload(template)


@router.get("/rubric-templates")
def list_rubric_templates(
    response: Response,
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=20, ge=1, le=100),
    current_user: User = Depends(require_roles("admin", "manager", "teacher", "ta")),
    db: Session = Depends(get_db),
) -> list[dict]:
    roles = {role.code for role in current_user.roles}
    query = select(RubricTemplate)
    if "admin" not in roles:
        query = query.where(RubricTemplate.created_by == current_user.id)
    total = db.scalar(select(func.count()).select_from(query.order_by(None).subquery())) or 0
    templates = db.scalars(query.order_by(RubricTemplate.name, RubricTemplate.id)
        .offset((page - 1) * page_size).limit(page_size)).all()
    set_pagination_headers(response, total=total, page=page, page_size=page_size)
    return [_rubric_template_payload(template) for template in templates]


@router.put("/rubric-templates/{template_id}")
def update_rubric_template(
    template_id: int,
    request: RubricTemplateCreate,
    current_user: User = Depends(require_roles("admin", "manager", "teacher", "ta")),
    db: Session = Depends(get_db),
) -> dict:
    template = db.get(RubricTemplate, template_id)
    if template is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy rubric.")
    _assert_rubric_template_access(template, current_user)
    previous_version = template.version
    template.version += 1
    template.name = request.name
    template.description = request.description
    template.criteria = [
        {
            "id": index,
            "title": item.title,
            "description": item.description,
            "weight": str(item.weight),
        }
        for index, item in enumerate(request.criteria, start=1)
    ]
    _record_audit(
        db,
        current_user,
        "rubric.template_updated",
        details=(
            f"template_id={template.id};before=version={previous_version};"
            f"after=version={template.version};criteria_count={len(template.criteria)}"
        ),
    )
    db.commit()
    db.refresh(template)
    return _rubric_template_payload(template)


@router.put("/assignments/{assignment_id}/rubric")
def snapshot_assignment_rubric(
    assignment_id: int,
    template_id: int = Query(ge=1),
    current_user: User = Depends(require_roles("admin", "manager", "teacher", "ta")),
    db: Session = Depends(get_db),
) -> dict:
    assignment = db.get(Assignment, assignment_id)
    if assignment is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy bài tập.")
    _assert_assignment_staff_scope(db, current_user, assignment)
    if assignment.status != "draft" or assignment.rubric_snapshot is not None:
        raise HTTPException(status_code=409, detail="Rubric chỉ được gắn một lần khi bài tập còn nháp.")
    template = db.get(RubricTemplate, template_id)
    if template is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy rubric.")
    _assert_rubric_template_access(template, current_user)
    assignment.rubric_snapshot = {
        "template_id": template.id,
        "template_name": template.name,
        "template_version": template.version,
        "criteria": [dict(item) for item in template.criteria],
    }
    db.commit()
    return assignment.rubric_snapshot


@router.get("/assignments/{assignment_id}/rubric")
def get_assignment_rubric(
    assignment_id: int,
    current_user: User = Depends(require_roles("admin", "manager", "teacher", "ta")),
    db: Session = Depends(get_db),
) -> dict | None:
    assignment = db.get(Assignment, assignment_id)
    if assignment is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy bài tập.")
    _assert_assignment_staff_scope(db, current_user, assignment)
    return assignment.rubric_snapshot


def _grade_payload(grade: AssignmentGradeVersion, *, include_internal: bool) -> dict:
    payload = {
        "id": grade.id,
        "assignment_id": grade.assignment_id,
        "submission_id": grade.submission_id,
        "student_id": grade.student_id,
        "version_number": grade.version_number,
        "score": grade.score,
        "feedback": grade.feedback,
        "status": grade.status,
        "created_at": grade.created_at.isoformat(),
        "published_at": grade.published_at.isoformat() if grade.published_at else None,
    }
    if include_internal:
        payload["internal_comment"] = grade.internal_comment
        payload["rubric_scores"] = grade.rubric_scores
    elif grade.status == "published":
        payload["rubric_scores"] = grade.rubric_scores
    return payload


@router.post(
    "/assignments/{assignment_id}/students/{student_id}/grades",
    status_code=status.HTTP_201_CREATED,
)
def create_assignment_grade_draft(
    assignment_id: int,
    student_id: int,
    request: AssignmentGradeCreate,
    current_user: User = Depends(require_roles("admin", "manager", "teacher", "ta")),
    db: Session = Depends(get_db),
) -> dict:
    assignment = db.get(Assignment, assignment_id)
    if assignment is None or assignment.status != "published":
        raise HTTPException(status_code=404, detail="Không tìm thấy bài tập.")
    _assert_assignment_staff_scope(db, current_user, assignment)
    if assignment.max_score is None:
        raise HTTPException(status_code=409, detail="Bài tập chưa có điểm tối đa.")
    rubric_snapshot = assignment.rubric_snapshot
    rubric_scores: list[dict] | None = None
    grade_score = request.score
    if rubric_snapshot:
        if request.rubric_scores is None:
            raise HTTPException(status_code=422, detail="Bài tập yêu cầu chấm đủ tiêu chí rubric.")
        try:
            parsed_scores = [RubricScoreInput.model_validate(item) for item in request.rubric_scores]
        except Exception as error:
            raise HTTPException(status_code=422, detail="Điểm tiêu chí rubric không hợp lệ.") from error
        expected_ids = {item["id"] for item in rubric_snapshot["criteria"]}
        provided_ids = [item.criterion_id for item in parsed_scores]
        if set(provided_ids) != expected_ids or len(provided_ids) != len(expected_ids):
            raise HTTPException(status_code=422, detail="Cần chấm mỗi tiêu chí rubric đúng một lần.")
        by_id = {item.criterion_id: item for item in parsed_scores}
        weighted_percent = Decimal("0")
        rubric_scores = []
        for criterion in rubric_snapshot["criteria"]:
            result = by_id[criterion["id"]]
            weight = Decimal(criterion["weight"])
            weighted_percent += weight * result.score_percent / Decimal("100")
            rubric_scores.append({
                "criterion_id": criterion["id"],
                "title": criterion["title"],
                "weight": criterion["weight"],
                "score_percent": str(result.score_percent),
                "comment": result.comment,
            })
        computed_score = (assignment.max_score * weighted_percent / Decimal("100")).quantize(
            Decimal("0.01"), rounding=ROUND_HALF_UP
        )
        if grade_score is not None and grade_score != computed_score:
            raise HTTPException(status_code=422, detail="Điểm tổng không khớp với rubric.")
        grade_score = computed_score
    elif grade_score is None:
        raise HTTPException(status_code=422, detail="Cần nhập điểm tổng khi bài tập chưa gắn rubric.")
    if grade_score is None or grade_score > assignment.max_score:
        raise HTTPException(status_code=422, detail="Điểm vượt quá điểm tối đa của bài tập.")
    submission = db.get(AssignmentSubmission, request.submission_id)
    if (
        submission is None
        or submission.assignment_id != assignment.id
        or submission.student_id != student_id
        or not _student_is_assignment_recipient(db, assignment, student_id)
    ):
        raise HTTPException(status_code=404, detail="Không tìm thấy bài nộp.")

    latest = db.scalar(
        select(AssignmentGradeVersion)
        .where(
            AssignmentGradeVersion.assignment_id == assignment.id,
            AssignmentGradeVersion.student_id == student_id,
            AssignmentGradeVersion.status == "published",
        )
        .order_by(AssignmentGradeVersion.version_number.desc())
    )
    version_number = (
        db.scalar(
            select(func.max(AssignmentGradeVersion.version_number)).where(
                AssignmentGradeVersion.assignment_id == assignment.id,
                AssignmentGradeVersion.student_id == student_id,
            )
        )
        or 0
    ) + 1
    grade = AssignmentGradeVersion(
        assignment_id=assignment.id,
        submission_id=submission.id,
        student_id=student_id,
        version_number=version_number,
        score=grade_score,
        rubric_scores=rubric_scores,
        feedback=request.feedback,
        internal_comment=request.internal_comment,
        status="draft",
        created_by=current_user.id,
    )
    db.add(grade)
    db.flush()
    before_score = str(latest.score) if latest is not None else "none"
    _record_audit(
        db,
        current_user,
        "grade.draft_saved",
        target_user_id=student_id,
        details=(
            f"assignment_id={assignment.id};submission_id={submission.id};"
            f"before={before_score};after={grade_score};version={version_number}"
        ),
    )
    db.commit()
    db.refresh(grade)
    return _grade_payload(grade, include_internal=True)


@router.get("/assignments/{assignment_id}/grades")
def list_assignment_grades(
    assignment_id: int,
    response: Response,
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=20, ge=1, le=100),
    current_user: User = Depends(require_roles("admin", "manager", "teacher", "ta", "student")),
    db: Session = Depends(get_db),
) -> list[dict]:
    assignment = db.get(Assignment, assignment_id)
    if assignment is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy bài tập.")
    roles = {role.code for role in current_user.roles}
    student_only = "student" in roles and not roles.intersection(("admin", "manager", "teacher", "ta"))
    query = select(AssignmentGradeVersion).where(
        AssignmentGradeVersion.assignment_id == assignment.id
    )
    if student_only:
        if (
            assignment.status != "published"
            or not _student_is_assignment_recipient(db, assignment, current_user.id)
        ):
            raise HTTPException(status_code=404, detail="Không tìm thấy bài tập.")
        query = query.where(
            AssignmentGradeVersion.student_id == current_user.id,
            AssignmentGradeVersion.status == "published",
        )
    else:
        _assert_assignment_staff_scope(db, current_user, assignment)
    total = db.scalar(select(func.count()).select_from(query.subquery())) or 0
    grades = db.scalars(
        query.order_by(
            AssignmentGradeVersion.student_id,
            AssignmentGradeVersion.version_number,
            AssignmentGradeVersion.id,
        ).offset((page - 1) * page_size).limit(page_size)
    ).all()
    set_pagination_headers(response, total=total, page=page, page_size=page_size)
    return [_grade_payload(item, include_internal=not student_only) for item in grades]


@router.post("/assignments/{assignment_id}/grades/{grade_id}/publish")
def publish_assignment_grade(
    assignment_id: int,
    grade_id: int,
    current_user: User = Depends(require_roles("admin", "manager", "teacher", "ta")),
    db: Session = Depends(get_db),
) -> dict:
    assignment = db.get(Assignment, assignment_id)
    grade = db.get(AssignmentGradeVersion, grade_id)
    if assignment is None or grade is None or grade.assignment_id != assignment_id:
        raise HTTPException(status_code=404, detail="Không tìm thấy bản chấm.")
    training_class = _assert_assignment_staff_scope(db, current_user, assignment)
    if grade.status != "draft":
        raise HTTPException(status_code=409, detail="Bản chấm này đã được công bố.")
    latest_version = db.scalar(
        select(func.max(AssignmentGradeVersion.version_number)).where(
            AssignmentGradeVersion.assignment_id == assignment.id,
            AssignmentGradeVersion.student_id == grade.student_id,
        )
    )
    if grade.version_number != latest_version:
        raise HTTPException(status_code=409, detail="Chỉ có thể công bố bản chấm mới nhất.")
    grade.status = "published"
    grade.published_at = datetime.now(timezone.utc).replace(tzinfo=None)
    _record_audit(
        db,
        current_user,
        "grade.published",
        target_user_id=grade.student_id,
        details=(
            f"assignment_id={assignment.id};submission_id={grade.submission_id};"
            f"before=draft;after=published;score={grade.score};version={grade.version_number}"
        ),
    )
    notification = Notification(
        user_id=grade.student_id,
        class_id=training_class.id,
        title="Điểm bài tập đã được công bố",
        message=f"Điểm bài tập {assignment.title} đã sẵn sàng.",
    )
    db.add(notification)
    enqueue_notification_email(db, notification, db.get(User, grade.student_id))
    db.commit()
    db.refresh(grade)
    return _grade_payload(grade, include_internal=True)


@router.post("/announcements", status_code=status.HTTP_201_CREATED)
def create_class_announcement(
    request: AnnouncementCreate,
    current_user: User = Depends(require_roles("admin", "manager")),
    db: Session = Depends(get_db),
) -> dict:
    classes = db.scalars(select(TrainingClass).where(TrainingClass.id.in_(request.class_ids))).all()
    if len(classes) != len(request.class_ids):
        raise HTTPException(status_code=404, detail="One or more selected classes were not found.")
    announcement = Announcement(title=request.title, body=request.body, created_by=current_user.id)
    db.add(announcement)
    db.flush()
    db.execute(announcement_classes.insert(), [
        {"announcement_id": announcement.id, "class_id": class_id}
        for class_id in request.class_ids
    ])
    eligible_enrollments = db.execute(
        select(ClassEnrollment, User)
        .join(User, User.id == ClassEnrollment.student_id)
        .where(
            ClassEnrollment.class_id.in_(request.class_ids),
            ClassEnrollment.enrollment_status == "active",
            or_(ClassEnrollment.effective_date.is_(None), ClassEnrollment.effective_date <= date.today()),
            User.is_active.is_(True),
            User.is_locked.is_(False),
        )
    ).all()
    recipient_by_id = {student.id: student for _, student in eligible_enrollments}
    notification_class_id = request.class_ids[0] if len(request.class_ids) == 1 else None
    for recipient in recipient_by_id.values():
        notification = Notification(
            user_id=recipient.id,
            class_id=notification_class_id,
            announcement_id=announcement.id,
            title=request.title,
            message=request.body,
            navigation_url=f"/announcements/{announcement.id}",
        )
        db.add(notification)
        db.flush()
        db.add(NotificationOutbox(
            notification_id=notification.id,
            dedupe_key=f"announcement:{announcement.id}:user:{recipient.id}",
            recipient_email=recipient.email,
            status="pending",
            attempt_count=0,
        ))
    _record_audit(db, current_user, "announcement.created",
        details=f"announcement_id={announcement.id};class_count={len(classes)};recipient_count={len(recipient_by_id)}")
    db.commit()
    db.refresh(announcement)
    return {"id": announcement.id, "title": announcement.title, "body": announcement.body,
            "class_ids": request.class_ids, "recipient_count": len(recipient_by_id),
            "created_by": announcement.created_by, "created_at": announcement.created_at.isoformat()}


@router.get("/announcements/{announcement_id}/read-count")
def get_announcement_read_count(
    announcement_id: int,
    current_user: User = Depends(require_roles("admin", "manager")),
    db: Session = Depends(get_db),
) -> dict:
    announcement = db.get(Announcement, announcement_id)
    if announcement is None:
        raise HTTPException(status_code=404, detail="Announcement not found.")
    total = db.scalar(select(func.count(Notification.id)).where(Notification.announcement_id == announcement_id)) or 0
    read = db.scalar(select(func.count(Notification.id)).where(
        Notification.announcement_id == announcement_id, Notification.read_at.is_not(None)
    )) or 0
    return {"announcement_id": announcement_id, "recipient_count": total,
            "read_count": read, "unread_count": total - read}


@router.get("/notifications")
def list_notifications(
    limit: int = Query(default=20, ge=1, le=100),
    page: int | None = Query(default=None, ge=1),
    page_size: int | None = Query(default=None, ge=1, le=100),
    cursor_created_at: str | None = Query(default=None, max_length=40),
    cursor_id: int | None = Query(default=None, ge=1),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
) -> dict:
    if (cursor_created_at is None) != (cursor_id is None):
        raise HTTPException(
            status_code=422,
            detail="cursor_created_at và cursor_id phải được truyền cùng nhau.",
        )
    query = select(Notification).where(Notification.user_id == current_user.id)
    page_mode = page is not None or page_size is not None
    if page_mode and cursor_id is not None:
        raise HTTPException(status_code=422, detail="Không kết hợp page với cursor.")
    total_notifications = db.scalar(select(func.count(Notification.id)).where(
        Notification.user_id == current_user.id
    )) or 0
    if cursor_created_at is not None and cursor_id is not None:
        try:
            cursor_time = datetime.fromisoformat(cursor_created_at.replace("Z", "+00:00"))
        except ValueError as exc:
            raise HTTPException(status_code=422, detail="Cursor thời gian không hợp lệ.") from exc
        if cursor_time.tzinfo is not None:
            cursor_time = cursor_time.astimezone(timezone.utc).replace(tzinfo=None)
        query = query.where(
            or_(
                Notification.created_at < cursor_time,
                (Notification.created_at == cursor_time) & (Notification.id < cursor_id),
            )
        )
    actual_page = page or 1
    actual_page_size = page_size or limit
    if page_mode:
        fetched = db.scalars(query.order_by(Notification.created_at.desc(), Notification.id.desc())
            .offset((actual_page - 1) * actual_page_size).limit(actual_page_size)).all()
        notifications = fetched
        has_more = actual_page * actual_page_size < total_notifications
    else:
        fetched = db.scalars(
            query.order_by(Notification.created_at.desc(), Notification.id.desc()).limit(limit + 1)
        ).all()
        has_more = len(fetched) > limit
        notifications = fetched[:limit]
    unread_count = db.scalar(
        select(func.count(Notification.id)).where(
            Notification.user_id == current_user.id,
            Notification.read_at.is_(None),
        )
    ) or 0
    return {
        "unread_count": unread_count,
        "total": total_notifications,
        "page": actual_page if page_mode else None,
        "page_size": actual_page_size if page_mode else limit,
        "pages": (total_notifications + actual_page_size - 1) // actual_page_size,
        "has_previous": page_mode and actual_page > 1 and total_notifications > 0,
        "has_next": has_more,
        "has_more": has_more,
        "next_cursor": (
            {
                "created_at": notifications[-1].created_at.isoformat(),
                "id": notifications[-1].id,
            }
            if has_more and notifications
            else None
        ),
        "items": [
            {
                "id": item.id,
                "class_id": item.class_id,
                "lead_id": item.lead_id,
                "announcement_id": item.announcement_id,
                "title": item.title,
                "message": item.message,
                "navigation_url": item.navigation_url,
                "scheduled_for": (
                    f"{item.scheduled_for.isoformat()}Z" if item.scheduled_for else None
                ),
                "is_read": item.read_at is not None,
                "created_at": item.created_at.isoformat(),
            }
            for item in notifications
        ],
    }


@router.post("/notifications/read-all")
def mark_all_notifications_read(
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
) -> dict:
    result = db.execute(
        update(Notification)
        .where(
            Notification.user_id == current_user.id,
            Notification.read_at.is_(None),
        )
        .values(read_at=datetime.now(timezone.utc))
    )
    db.commit()
    return {"updated_count": result.rowcount or 0, "unread_count": 0}


@router.post("/notifications/{notification_id}/read")
def mark_notification_read(
    notification_id: int,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
) -> dict:
    notification = db.scalar(
        select(Notification).where(
            Notification.id == notification_id,
            Notification.user_id == current_user.id,
        )
    )
    if notification is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy thông báo.")
    if notification.read_at is None:
        notification.read_at = datetime.now(timezone.utc)
        db.commit()
    return {"id": notification.id, "is_read": True}
